from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path(__file__).parent
text=(root/'Omniarch-Synopsis-to-Chapter-8-Benchmark.md').read_text()
d=Document(); sec=d.sections[0]; sec.top_margin=sec.bottom_margin=Inches(.72); sec.left_margin=sec.right_margin=Inches(.78)
styles=d.styles
for name in ['Normal','Body Text']:
 styles[name].font.name='Calibri';styles[name].font.size=Pt(10.5);styles[name].paragraph_format.space_after=Pt(6)
for name in ['Title','Heading 1','Heading 2','Heading 3']:
 styles[name].font.name='Calibri'; styles[name].font.color.rgb=RGBColor.from_string('163B4B')
styles['Heading 1'].paragraph_format.space_before=Pt(18)
styles['Heading 2'].paragraph_format.space_before=Pt(14)

def put(p,t):
 # Preserve link destinations in the editable copy.
 t=re.sub(r'\[([^]]+)\]\(([^)]+)\)',r'\1 (\2)',t)
 for x in re.split(r'(\*\*.*?\*\*|`[^`]*`)',t):
  r=p.add_run(x[2:-2] if x.startswith('**') else x[1:-1] if x.startswith('`') else x)
  if x.startswith('**'):r.bold=True
  if x.startswith('`'):r.font.name='Consolas';r.font.size=Pt(9)
lines=text.splitlines();i=0
while i<len(lines):
 l=lines[i].strip()
 if not l or l=='---':i+=1;continue
 if l.startswith('|'):
  rows=[]
  while i<len(lines) and lines[i].strip().startswith('|'):
   row=[x.strip() for x in lines[i].strip().strip('|').split('|')]
   if not all(re.fullmatch(r'[:\- ]+',x) for x in row):rows.append(row)
   i+=1
  t=d.add_table(rows=0,cols=len(rows[0]));t.style='Light Shading Accent 1'
  for row in rows:
   cells=t.add_row().cells
   for c,v in zip(cells,row):put(c.paragraphs[0],v)
  continue
 if l.startswith('# '):d.add_heading(l[2:],0)
 elif l.startswith('## '):d.add_heading(l[3:],1)
 elif l.startswith('### '):d.add_heading(l[4:],2)
 elif l.startswith('> '):put(d.add_paragraph(style='Quote'),l[2:])
 elif re.match(r'^\d+\. ',l):put(d.add_paragraph(style='List Number'),re.sub(r'^\d+\. ','',l))
 elif l.startswith('- '):put(d.add_paragraph(style='List Bullet'),l[2:])
 else:put(d.add_paragraph(),l)
 i+=1
header=sec.header.paragraphs[0];header.text='EDITORIAL BENCHMARK  /  THE OMNIARCH  /  SYNOPSIS–CHAPTER 8';header.style='Caption'
f=sec.footer.paragraphs[0];f.alignment=2
r=f.add_run('12 October 2026  •  ');r.font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
d.core_properties.title='The Omniarch — Synopsis to Chapter 8: Benchmark Editorial Assessment'
d.core_properties.subject='Evidence-led translation assessment and revision priorities'
d.save(root/'Omniarch-Synopsis-to-Chapter-8-Benchmark.docx')
print('Created Word report')
