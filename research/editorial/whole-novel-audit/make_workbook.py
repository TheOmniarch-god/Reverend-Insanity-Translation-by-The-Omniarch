from pathlib import Path
import json
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.table import Table,TableStyleInfo
p=Path(__file__).parent
coverage=json.loads((p/'coverage.json').read_text());hits=json.loads((p/'hits.json').read_text())+json.loads((p/'supplementary-hits.json').read_text());summary=json.loads((p/'summary.json').read_text());selected=json.loads((p/'selected-evidence.json').read_text())
wb=Workbook();s=wb.active;s.title='Read me'
for row in [ ['THE OMNIARCH | WHOLE-NOVEL TOPICAL AUDIT',''],['Date','13 October 2026'],['Scope','Aperture percentages, essence accounting, and regional travel'],['Narrative chapters screened',2334],['Nonempty body paragraphs screened',213892],['Characters screened',7275173],['Primary candidate paragraphs',959],['Additional context-screen candidates',225],['Total candidate paragraphs inspected',1184],['Status','Automated screening complete; all candidate paragraph texts inspected; selective contextual verification.'],['NOT completed','A line-by-line reading and certification of all 2,334 chapters.'],['Important','Candidate excerpts include false positives. Presence in the evidence inventory is not a finding of error.'],['Numbering','Corpus order follows the supplied EPUB spine. Volume, printed title and filename identify the source.'],['Source SHA-256',summary['source_sha256']],['Changes applied','No source or public-site changes. Article revised to reflect stronger located evidence.'],['Authority','Suggested numerical emendations remain subject to approval.'] ]:s.append(row)
counts=Counter(h['file'] for h in hits)
s=wb.create_sheet('Chapter coverage');s.append(['Corpus order','Volume','Printed chapter title','EPUB file','Paragraphs','Characters','Candidate paragraphs','Screening status','Full chapter close-read certification'])
for x in coverage:s.append([x['corpus_chapter'],x['volume'],x['title'],x['file'],x['paragraph_count'],x['characters'],counts[x['file']],'Automated topical screen completed','Not certified by this register'])
s=wb.create_sheet('Candidate excerpts');s.append(['Corpus order','Volume','Printed chapter title','EPUB file','Body paragraph','Search categories','Chinese passage','Preceding paragraph','Following paragraph','Assessment scope'])
for x in sorted(hits,key=lambda a:(a['corpus_chapter'],a['paragraph'])):s.append([x['corpus_chapter'],x['volume'],x['title'],x['file'],x['paragraph'],', '.join(x['categories']),x['text'],'\n'.join(x['before']),'\n'.join(x['after']),'Candidate text inspected; not automatically a confirmed finding'])
s=wb.create_sheet('Selected evidence');s.append(['Evidence ID','Corpus order','Volume','Printed title','EPUB file','Source paragraphs','Chinese text'])
for x in selected:s.append([x['id'],x['corpus_chapter'],x['volume'],x['title'],x['file'],', '.join(str(y['paragraph']) for y in x['paragraphs']),'\n\n'.join(f"p{y['paragraph']}: {y['text']}" for y in x['paragraphs'])])
findings=[
 ['A01','Volume-based capacity','V1 Ch10','Direct source statement','Keep established capacity figures','No geometric conversion of all percentages'],
 ['A02','44 minus 8 equals 36','V1 Ch80','Explicit arithmetic','Use capacity-point accounting here','Do not interpret the 8 as 8% of personal maximum'],
 ['A03','90 minus 70 equals 20','V2 Ch100','Explicit arithmetic','Preserve additive scale','Do not replace remaining 20 with 27'],
 ['A04','Late 20-plus reserve permits two 10-point uses','V5 Ch481','Explicit example in dream episode','Preserve local activation cost','Not a universal Gu cost'],
 ['A05','Less than half of peak reserve','V3 Ch67','Explicit personal denominator','Retain personal-reserve reference','Do not convert automatically to 50% of full aperture'],
 ['A06','60% transfer, twice','V2 Ch31','10 to 6 to 3.6 is consistent','Keep rounding-compatible wording','Efficiency is not capacity'],
 ['A07','Two 90% apertures','V2 Ch178; V3 Ch99','Separate stores and sharing','Retain per-aperture accounting','180% total is not one overfilled aperture'],
 ['A08','Height and volume','V1 Ch6/10/92; V2 Ch98','Source conflates exact measures in places','Local wording harmonisation proposed','Preserve aptitude numbers; log emendation'],
 ['A09','First Revolution stage equivalence','V1 Ch80','Explicit 1:2:4:8','Article corrected','Does not settle every cross-revolution comparison'],
 ['N01','45% versus repeated 44%','V1 Ch43 versus Ch34/50/80','Strong inconsistency candidate','44% proposed, not approved','No corrected witness authenticated'],
 ['N02','42% versus later 43%','V1 Ch159/162/167 versus Ch198','Conflict in checked readings','Linked-sequence repair pending','If retaining 42, later 53/63 must be addressed'],
 ['N03','92% plus two points but 93%','V2 Ch102/105','Baseline ambiguity','Hold automatic edit','Could be cumulative from earlier 91'],
 ['N04','Grade C / 38% / three attacks','V1 Ch22','Previously confirmed inconsistency','Exact repair pending','Grade-label change affects staging'],
 ['G01','Decades of migration; crane transport','V1 Ch192/199','Migration and transport explicit','Preserve','Exact itinerary and speed not established'],
 ['G02','Mortals cross barriers more easily','V4 Ch61','Explicit comparison','Use proper cultivation context','Tea-time is barrier crossing, not continental journey'],
 ['G03','Beast kings versus desolate beasts','V5 Ch24','Explicit distinction','Do not apply desolate restrictions to ordinary beast kings','Does not prove mount speed'],
 ['G04','Years of flight versus a day/night with formations','V4 Ch307','Explicit long-route workaround','Preserve travel mechanics','No need to invent ordinary flight speed'],
 ['G05','Earth Hole versus Star Gate','V3 Ch54/57','Different regional reach','Preserve individual limits','Do not assign all transport to all travellers'],
 ['G06','Barrier disappearance','V6 Ch2','Explicit era change','Audit routes chronologically','Early and late conditions differ'],
 ['G07','Central extent and neighbouring boundaries','V2 Ch135; V6 Ch253','Source magnitude and adjacency explicit','Hold global rescaling','Extent is not each trip distance'],
 ['G08','8,000-plus-km river and region coverage','V2 Ch1 versus V1 Ch29','Scale tension; shape/basin unspecified','Further route/geography work','No unique replacement length established'],
 ['G09','Vast region, nearby destination','V4 Ch88','Explicit distinction','Preserve local-distance explanation','Area alone does not determine journey length']]
s=wb.create_sheet('Decision register');s.append(['ID','Issue','Source location','Evidence status','Recommended disposition','Limit / dependency'])
for row in findings:s.append(row)
for s in wb:
 s.freeze_panes='A2';s.sheet_view.showGridLines=False
 for c in s[1]:c.fill=PatternFill('solid',fgColor='163B4B');c.font=Font(bold=True,color='FFFFFF');c.alignment=Alignment(wrap_text=True,vertical='top')
 for row in s.iter_rows(min_row=2):
  for c in row:c.alignment=Alignment(wrap_text=True,vertical='top');c.font=Font(name='Calibri',size=11)
 if s.title!='Read me':
  s.auto_filter.ref=s.dimensions
  for col in s.columns:s.column_dimensions[col[0].column_letter].width=24
  s.column_dimensions['A'].width=15
 else:s.column_dimensions['A'].width=40;s.column_dimensions['B'].width=100
 if s.title=='Candidate excerpts':
  for col in ['G','H','I']:s.column_dimensions[col].width=70
  s.column_dimensions['C'].width=40;s.column_dimensions['J'].width=45
 if s.title=='Selected evidence':s.column_dimensions['G'].width=115
 if s.title=='Decision register':
  for col in ['B','C','D','E','F']:s.column_dimensions[col].width=43
 for i in range(2,s.max_row+1):s.row_dimensions[i].height=45 if s.title not in ['Candidate excerpts','Selected evidence'] else 95
 s.row_dimensions[1].height=32
wb.save(p/'Whole-Novel-Evidence-and-Coverage.xlsx')
print('Workbook:',len(coverage),'coverage rows;',len(hits),'candidate rows;',len(findings),'decision rows')
