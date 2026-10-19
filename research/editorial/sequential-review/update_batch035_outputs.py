import json
from pathlib import Path
from openpyxl import load_workbook
r = Path(__file__).parent
WB = r/'Sequential-Reading-and-Proposals.xlsx'
wb = load_workbook(WB)
led = {v['id']: v for v in json.loads((r/'proposal-ledger.json').read_text())}
src = json.loads((r/'batch-035-source.json').read_text())
ids_new = [f"SR-{i}" for i in range(419, 443)]
pl = wb['Proposal ledger']
present = {pl.cell(rr,1).value for rr in range(2, pl.max_row+1)}
added_ledger = 0
for sid in ids_new:
    if sid in present: continue
    e = led[sid]; loc = e['locations'][0]
    pl.append([sid, e['category'], f"Ch{loc['chapter']} p{loc['paragraph']}", e['recommended_action'], e['rationale'],
               e['confidence'], e['status'], e['dependencies'], e['publication_status'], 'No', 'No', e.get('witness_035', '')])
    added_ledger += 1
cc = wb['Chapter completion']
have = {cc.cell(rr,1).value for rr in range(2, cc.max_row+1)}
added_ch = 0
for c in src:
    if c['corpus_chapter'] in have: continue
    cc.append([c['corpus_chapter'], c['file'], c['title'], len(c['paragraphs']), 'Full sequential read']); added_ch += 1
pr = wb['Paragraph register']
have_p = {pr.cell(rr,1).value for rr in range(2, pr.max_row+1)}
pmap = {}
for sid in ids_new:
    loc = led[sid]['locations'][0]; pmap.setdefault((loc['chapter'], loc['paragraph']), []).append(sid)
added_p = 0
for c in src:
    if c['corpus_chapter'] in have_p: continue
    for i, t in enumerate(c['paragraphs'], 1):
        pids = ', '.join(pmap.get((c['corpus_chapter'], i), [])) or None
        pr.append([c['corpus_chapter'], i, 'Read; not certified error-free', pids, None, None, t]); added_p += 1
rm = wb['Read me']
upd = {'Completed': '1–685 of 2334', 'Next unread': '686', 'Paragraphs': '61680', 'Characters': '2171266',
       'Proposals': f"{len(led)} total; {len(led)-5} active (held or query); 4 withdrawn; 1 locally resolved"}
for rr in range(1, rm.max_row+1):
    k = rm.cell(rr,1).value
    if k in upd: rm.cell(rr,2).value = upd[k]
wb.save(WB)
# reading progress
pj = r/'reading-progress.json'
p = json.loads(pj.read_text())
done = list(p['sequential_chapters_completed'])
for n in range(666, 686):
    if n not in done: done.append(n)
p['sequential_chapters_completed'] = done
p['paragraphs_read'] = 61680
p['paragraph_characters_read'] = 2171266
p['next_chapter'] = 686
p['remaining_chapters'] = 1649
p['batch_report'] = 'Batch-035-Chapters-666-685.md (latest).'
p['batches'] = 35
p['chapters'] = 685
pj.write_text(json.dumps(p, ensure_ascii=False, indent=2))
print('ledger', added_ledger, 'chapters', added_ch, 'paragraphs', added_p, 'ledger rows', pl.max_row-1)
