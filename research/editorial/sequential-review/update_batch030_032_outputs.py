from pathlib import Path
import json, re
from openpyxl import load_workbook
from docx import Document
from docx.shared import Pt
r = Path(__file__).parent
srcs = [json.loads((r/f'batch-0{n}-source.json').read_text()) for n in (30, 31, 32)]
src = [c for s in srcs for c in s]
led = json.loads((r/'proposal-ledger.json').read_text())
assert len(led) == 380 and led[-1]['id'] == 'SR-380'
HOLD = 'HOLD: full2334-chapter reading and global reconciliation outstanding'
prog = json.loads((r/'reading-progress.json').read_text())
N = sum(len(c['paragraphs']) for c in src); CH = sum(len(t) for c in src for t in c['paragraphs'])
assert N == 6052 and CH == 207293, (N, CH)
assert prog['paragraphs_read'] == 56533 and prog['paragraph_characters_read'] == 1982128
assert [c['corpus_chapter'] for c in src] == list(range(566, 626))
pmap = {}
for v in led:
    if int(v['id'][3:]) >= 338:
        for l in v['locations']:
            pmap.setdefault((l['chapter'], l['paragraph']), []).append(v['id'])

wb = load_workbook(r/'Sequential-Reading-and-Proposals.xlsx')
def drop(ws, pred):
    for idx in range(ws.max_row, 1, -1):
        if pred(ws.cell(idx, 1).value): ws.delete_rows(idx)
isnum = lambda v: v is not None and str(v).isdigit()
drop(wb['Chapter completion'], lambda v: isnum(v) and int(v) >= 566)
drop(wb['Paragraph register'], lambda v: isnum(v) and int(v) >= 566)
drop(wb['Proposal ledger'], lambda v: v is not None and str(v).startswith('SR-') and int(str(v)[3:]) >= 338)

ws = wb['Chapter completion']
for c in src:
    ws.append([c['corpus_chapter'], c['file'], c['title'], len(c['paragraphs']), 'Full sequential read'])
ws = wb['Paragraph register']
for c in src:
    for i, t in enumerate(c['paragraphs'], 1):
        ids = pmap.get((c['corpus_chapter'], i))
        ws.append([c['corpus_chapter'], i, 'Read; not certified error-free', ', '.join(ids) if ids else 'None', 'None', 'None', t])
ws = wb['Proposal ledger']
for v in led:
    n = int(v['id'][3:])
    if n >= 338:
        refs = '; '.join(f"Ch{l['chapter']} p{l['paragraph']}" for l in v['locations'])
        batch = 'Batch 030' if n <= 352 else ('Batch 031' if n <= 365 else 'Batch 032')
        ws.append([v['id'], v['category'], refs, v['recommended_action'], v['rationale'], v['confidence'], v['status'],
                   v['dependencies'], HOLD, 'No', 'No', f'{batch} sequential reading; source-internal; witness not yet collated'])

withdrawn = sum(1 for v in led if v['status'].startswith('Withdrawn'))
resolved = sum(1 for v in led if v['status'].startswith('Resolved locally'))
vals = {'Completed': '1–625 of 2334', 'Next unread': '626', 'Paragraphs': str(56533), 'Characters': str(1982128),
        'Proposals': f"{len(led)} total; {len(led)-withdrawn-resolved} active (held or query); {withdrawn} withdrawn; {resolved} locally resolved"}
for row in wb['Read me'].iter_rows(min_row=1, max_col=2):
    if row[0].value in vals: row[1].value = vals[row[0].value]
wb.save(r/'Sequential-Reading-and-Proposals.xlsx')
print('workbook saved', N, CH, vals['Proposals'])

def build_docx(md_name, out_name):
    md = (r/md_name).read_text(encoding='utf-8').splitlines()
    doc = Document()
    st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
    table = None; in_table = False
    clean = lambda s: re.sub(r'\*\*(.+?)\*\*', r'\1', s)
    for line in md:
        if not line.strip() or line.strip() == '---': in_table = False; continue
        if line.startswith('|'):
            if set(line.replace('|', '').strip()) <= set('-: '): continue
            cells = [clean(c.strip()) for c in line.strip('|').split('|')]
            if not in_table:
                table = doc.add_table(rows=0, cols=len(cells)); table.style = 'Light Grid Accent 1'; in_table = True
            row = table.add_row().cells
            for k, cv in enumerate(cells): row[k].text = cv
            continue
        in_table = False
        if line.startswith('### '): doc.add_heading(line[4:], 2)
        elif line.startswith('## '): doc.add_heading(line[3:], 1)
        elif line.startswith('# '): doc.add_heading(line[2:], 0)
        elif line.startswith('- '): doc.add_paragraph(clean(line[2:]), style='List Bullet')
        else: doc.add_paragraph(clean(line))
    doc.save(r/out_name)
    return len(doc.tables)
for m, o in [('Batch-030-Chapters-566-585.md', 'Batch-030-Chapters-566-585.docx'),
             ('Batch-031-Chapters-586-605.md', 'Batch-031-Chapters-586-605.docx'),
             ('Batch-032-Chapters-606-625.md', 'Batch-032-Chapters-606-625.docx')]:
    print('docx', o, 'tables', build_docx(m, o))
