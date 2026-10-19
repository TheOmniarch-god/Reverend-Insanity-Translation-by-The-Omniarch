from pathlib import Path
import json, re
from openpyxl import load_workbook
from docx import Document
from docx.shared import Pt
r = Path(__file__).parent
src = json.loads((r/'batch-028-source.json').read_text())
led = json.loads((r/'proposal-ledger.json').read_text())
HOLD = 'HOLD: full2334-chapter reading and global reconciliation outstanding'
assert len(led) == 306 and led[-1]['id'] == 'SR-306'
pmap = {}
for v in led:
    if int(v['id'][3:]) >= 286:
        for l in v['locations']:
            pmap.setdefault((l['chapter'], l['paragraph']), []).append(v['id'])
wb = load_workbook(r/'Sequential-Reading-and-Proposals.xlsx')
def drop(ws, pred):
    for idx in range(ws.max_row, 1, -1):
        if pred(ws.cell(idx, 1).value): ws.delete_rows(idx)
# idempotent: remove any Batch 028 rows left by an earlier run
drop(wb['Chapter completion'], lambda v: v is not None and str(v).isdigit() and int(v) >= 526)
drop(wb['Paragraph register'], lambda v: v is not None and str(v).isdigit() and int(v) >= 526)
drop(wb['Proposal ledger'], lambda v: v is not None and str(v).startswith('SR-') and int(str(v)[3:]) >= 286)
assert not any(x[0] and str(x[0]).isdigit() and int(x[0]) >= 526 for x in wb['Chapter completion'].iter_rows(min_row=2, values_only=True))
assert not any(x[0] and str(x[0]).isdigit() and int(x[0]) >= 526 for x in wb['Paragraph register'].iter_rows(min_row=2, values_only=True))
ws = wb['Chapter completion']
for c in src:
    ws.append([c['corpus_chapter'], c['file'], c['title'], len(c['paragraphs']), 'Full sequential read'])
ws = wb['Paragraph register']
for c in src:
    for i, t in enumerate(c['paragraphs'], 1):
        ids = pmap.get((c['corpus_chapter'], i))
        ws.append([c['corpus_chapter'], i, 'Read; not certified error-free', ', '.join(ids) if ids else 'None', 'None', 'None', t])
ws = wb['Proposal ledger']
for v in led:
    if int(v['id'][3:]) >= 286:
        refs = '; '.join(f"Ch{l['chapter']} p{l['paragraph']}" for l in v['locations'])
        ws.append([v['id'], v['category'], refs, v['recommended_action'], v['rationale'], v['confidence'], v['status'], v['dependencies'], HOLD, 'No', 'No', 'Batch 028 sequential reading; source-internal; witness collated'])
vals = {'Completed': '1–545 of 2334', 'Next unread': '546', 'Paragraphs': '48681', 'Characters': '1711344',
        'Proposals': '306 total; 301 active (held or query); 4 withdrawn; 1 locally resolved'}
for row in wb['Read me'].iter_rows(min_row=1, max_col=2):
    k = row[0].value
    if k in vals: row[1].value = vals[k]
wb.save(r/'Sequential-Reading-and-Proposals.xlsx')
print('workbook saved')
md = (r/'Batch-028-Chapters-526-545.md').read_text(encoding='utf-8').splitlines()
doc = Document()
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
doc._t = None; doc._t_done = True
for line in md:
    if not line.strip() or line.strip() == '---': continue
    if line.startswith(('# ','## ','### ')): doc._t_done = True
    if line.startswith('# '): doc.add_heading(line[2:], 0)
    elif line.startswith('## '): doc.add_heading(line[3:], 1)
    elif line.startswith('### '): doc.add_heading(line[4:], 2)
    elif line.startswith('- '): doc.add_paragraph(re.sub(r'\*\*(.+?)\*\*', r'\1', line[2:]), style='List Bullet')
    elif line.startswith('|'):
        if set(line.replace('|', '').strip()) <= set('-: '): continue
        cells = [c.strip() for c in line.strip('|').split('|')]
        if doc._t is None or doc._t_done:
            doc._t = doc.add_table(rows=0, cols=len(cells)); doc._t.style = 'Light Grid Accent 1'; doc._t_done = False
        row = doc._t.add_row().cells
        for k, cv in enumerate(cells): row[k].text = re.sub(r'\*\*(.+?)\*\*', r'\1', cv)
    else:
        doc._t_done = True
        doc.add_paragraph(re.sub(r'\*\*(.+?)\*\*', r'\1', line))
doc.save(r/'Batch-028-Chapters-526-545.docx')
print('docx saved', len(doc.paragraphs), len(doc.tables))
