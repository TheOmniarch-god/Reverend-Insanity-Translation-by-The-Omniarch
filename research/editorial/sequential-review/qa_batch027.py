from pathlib import Path
import json, zipfile, re
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from docx import Document
r = Path(__file__).parent
load = lambda n: json.loads((r/n).read_text())
b = load('batch-027-source.json'); led = load('proposal-ledger.json'); p = load('reading-progress.json')
z = zipfile.ZipFile(r.parents[1]/'benchmark/sources/Gu-Zhen-Ren.epub')

# source vs EPUB
meta = load('../whole-novel-audit/coverage.json') if (r.parent/'whole-novel-audit/coverage.json').exists() else None
for c in b:
    actual = [n.get_text('', strip=True) for n in BeautifulSoup(z.read(c['file']), 'html.parser').find_all('p') if n.get_text('', strip=True)]
    assert c['paragraphs'] == actual, c['corpus_chapter']
assert [c['corpus_chapter'] for c in b] == list(range(506, 526))
N = sum(len(c['paragraphs']) for c in b); CH = sum(len(t) for c in b for t in c['paragraphs'])
assert N == 1696 and CH == 63354

# numbered source exact
src = (r/'Batch-027-Numbered-Chinese-Source.md').read_text(encoding='utf-8')
cur = None; found = []
for line in src.splitlines():
    m = re.match(r'## Chapter (\d+) ', line)
    if m: cur = int(m[1])
    m = re.match(r'(\d+)\.(\d+) (.*)', line)
    if m:
        assert int(m[1]) == cur
        assert m[3] == b[cur-506]['paragraphs'][int(m[2])-1]; found.append((cur, int(m[2])))
assert found == [(c['corpus_chapter'], i) for c in b for i in range(1, len(c['paragraphs'])+1)]
assert len(found) == 1696

# ledger
ids = [v['id'] for v in led]
assert len(led) == 285 and len(set(ids)) == 285 and ids[-1] == 'SR-285'
assert all(not v['approved'] and not v['applied'] for v in led)
for v in led:
    if int(v['id'][3:]) >= 276:
        for l in v['locations']:
            assert 1 <= l['paragraph'] <= len(b[l['chapter']-506]['paragraphs']) if 506 <= l['chapter'] <= 525 else True

# progress
assert p['next_chapter'] == 526 and p['sequential_chapters_completed'] == list(range(1, 526))
assert p['paragraphs_read'] == 45224 + N and p['paragraph_characters_read'] == 1582593 + CH
assert p['remaining_chapters'] == 1809

# workbook
w = load_workbook(r/'Sequential-Reading-and-Proposals.xlsx', read_only=True, data_only=True)
cc = [row for row in w['Chapter completion'].iter_rows(min_row=2, values_only=True) if row[0] and str(row[0]).isdigit() and int(row[0]) >= 506]
assert len(cc) == 20
pr = [row for row in w['Paragraph register'].iter_rows(min_row=2, values_only=True) if row[0] and str(row[0]).isdigit() and int(row[0]) >= 506]
assert len(pr) == N and all(row[6] == b[row[0]-506]['paragraphs'][row[1]-1] for row in pr)
pl = [row for row in w['Proposal ledger'].iter_rows(min_row=2, values_only=True) if row[0] in [f'SR-{i}' for i in range(276, 286)]]
assert len(pl) == 10
rm = dict(w['Read me'].iter_rows(min_row=2, max_col=2, values_only=True))
assert rm['Next unread'] == '526' and rm['Completed'] == '1–525 of 2334'

# docx
d = Document(r/'Batch-027-Chapters-506-525.docx')
txt = '\n'.join(x.text for x in d.paragraphs) + '\n' + '\n'.join(c.text for t in d.tables for row in t.rows for c in row.cells)
assert 'SR-276' in txt and 'SR-285' in txt and len(d.tables) == 2

# report hygiene
rep = (r/'Batch-027-Chapters-506-525.md').read_text(encoding='utf-8')
assert '—' not in rep
for line in rep.splitlines():
    if line.startswith('> '): assert line[2:] in {t for c in b for t in c['paragraphs']}

qa = {'status': 'passed', 'epub_paragraphs_rechecked': N, 'numbered_source_exact': True, 'numbered_source_lines': N,
      'workbook_chapter_rows': 20, 'workbook_paragraph_rows': N, 'workbook_paragraph_text_exact': True,
      'workbook_ledger_rows_added': 10, 'workbook_read_me_next_unread': 526, 'docx_proposal_ids_present': True,
      'docx_table_count': 2, 'ledger_total': 285, 'approved_or_applied': 0, 'witness_hashes_verified': 0,
      'report_blockquotes_exact': True, 'em_dashes_in_report': 0, 'progress_next': 526,
      'progress_paragraphs': 45224 + N, 'progress_characters': 1582593 + CH}
(r/'batch-027-qa.json').write_text(json.dumps(qa, ensure_ascii=False, indent=2))
print(json.dumps(qa, ensure_ascii=False, indent=2))
