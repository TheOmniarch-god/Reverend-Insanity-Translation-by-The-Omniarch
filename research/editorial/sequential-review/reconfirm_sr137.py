from pathlib import Path
import json
from docx import Document
from openpyxl import load_workbook
r=Path(__file__).parent
p=r/'proposal-ledger.json';ledger=json.loads(p.read_text());x=next(x for x in ledger if x['id']=='SR-137')
x.update(category='Source-name consistency query, not a translation error',recommended_action='Retain Hail Gu for 冰雹蛊 at312p97. Suspend the proposed 冰雹蛊→冰爆蛊 emendation pending corroboration.',rationale='冰雹 correctly means hail. The EPUB reads 冰雹蛊 at312p97.299p25 defines 冰爆蛊 as detonating ice, and313p40 names 冰爆蛊. This supports a source-name consistency query, not proof that Hail is mistranslated or that the Chinese requires correction.',confidence='High confidence in the lexical translation; source emendation unconfirmed.',status='Provisional source-name query; replacement suspended; no approval',dependencies='Reconfirm against independent Chinese textual witnesses and other Gu-name occurrences before recommending an emendation. Functional similarity alone is insufficient.',latest_update='User challenged Hail finding; dictionary and local-source recheck completed. Earlier very-high-confidence replacement downgraded.')
p.write_text(json.dumps(ledger,ensure_ascii=False,indent=2))
notice='**Fact-check correction, SR-137:** “Hail Gu” correctly translates the source 冰雹蛊. Retain it. The proposed replacement with 冰爆蛊 is suspended; the difference from nearby ice-detonation passages is an unresolved source-name query, not a demonstrated translation error. Earlier “very high” confidence is withdrawn.'
for p in r.rglob('*.md'):
 if '016' not in p.name:continue
 s=p.read_text()
 if any(t in s for t in ['冰雹','Hail']):
  s=s.replace('| SR137 | Hail Gu named for an ice-detonation effect | 冰雹蛊 → 冰爆蛊 | Very high |','| SR137 | Source-name consistency query, not a translation error | Retain Hail Gu; replacement suspended | Emendation unconfirmed |')
  s=s.replace('Preferred **冰雹蛊 → 冰爆蛊** at312p97.','**Retain 冰雹蛊 / Hail Gu at312p97; replacement suspended pending corroboration.**')
  s=s.replace('IceExplosionmisnamed冰雹 at31297,correct冰爆 based31340/314effect.','312p97 reads冰雹 (Hail); nearby冰爆 (Ice Explosion) creates an unresolved source-name query, not a proved error.')
  if p.name=='Batch-016-Continuity.md':
   s=s.replace('**recommended_action:** Preferred312p97 冰雹蛊→冰爆蛊.','**recommended_action:** '+x['recommended_action']).replace('Hail is the mismatched technical name.','Hail correctly translates the source; source-name identity remains unconfirmed.').replace('Very high functional and adjacent-name support.',x['confidence'])
  if notice not in s:s=notice+'\n\n'+s
  p.write_text(s)
p=r/'Batch-016-Chapters-301-315.docx';d=Document(p)
d.paragraphs[0].insert_paragraph_before(notice.replace('**',''))
for paragraph in d.paragraphs:
 if 'Preferred 冰雹蛊 → 冰爆蛊' in paragraph.text:paragraph.text=paragraph.text.replace('Preferred 冰雹蛊 → 冰爆蛊','Replacement suspended: retain 冰雹蛊 / Hail Gu')
for table in d.tables:
 for row in table.rows:
  if any('SR137' in c.text or 'SR-137' in c.text for c in row.cells):
   for c,text in zip(row.cells,['SR137','Source-name consistency query, not a translation error','Retain Hail Gu; replacement suspended','Emendation unconfirmed']):c.text=text
d.save(p)
p=r/'Sequential-Reading-and-Proposals.xlsx';w=load_workbook(p);s=w['Proposal ledger']
for row in s:
 if row[0].value=='SR-137':
  for col,key in [(2,'category'),(4,'recommended_action'),(5,'rationale'),(6,'confidence'),(7,'status'),(8,'dependencies'),(12,'latest_update')]:s.cell(row[0].row,col,x[key])
w.save(p)
(r/'SR-137-Fact-Check.md').write_text('''# SR-137 fact-check and correction

'''+notice+'''

## Rechecked primary evidence
- Ch299 p25: 冰爆蛊是三转蛊，能引爆冰块，形成瞬间的强大杀伤力。
- Ch312 p97: “除非我动用冰雹蛊，引爆冰锥，才可压过丹火。但是如此一来，我的真元消耗就多了。这场战斗，才刚刚开始……”
- Ch313 p40: 同时，她暗暗催动冰爆蛊。

## Dictionary corroboration
- https://zdic.net/hans/冰雹 : 冰雹 = hail / hailstone; falling ice, often associated with thunderstorms.
- https://dictionary.cambridge.org/us/dictionary/english-chinese-simplified/hail : hail = 雹，冰雹.

## Limits and next verification requirements
Dictionary meaning is verified. An independent Chinese witness for this exact passage has not yet been verified. The original EPUB is unchanged. No translation correction is warranted merely because the name describes hail while the narrated effect involves detonation. The existing local contrast remains worth investigating without assuming an error.

For further review: check lexical meanings, reread full context and counterevidence, distinguish faithful translation from suspected source inconsistency, and label unverified findings as queries. Reconfirm earlier high-confidence factual proposals before elevating them to recommendations. This targeted recheck is not a claim that all 139 proposals have been independently revalidated.

Local report, workbook, ledger and Batch016 Markdown notes updated. The earlier Chapter315 remote snapshot predates this correction; no updated remote backup has been claimed.
''')
p=r/'CURRENT-STATUS.md';p.write_text('## Subsequent SR-137 fact-check correction\nHail correctly translates 冰雹. Retain Hail Gu; emendation suspended pending corroboration. Local exports updated. Remote Chapter315 snapshot predates this correction. See SR-137-Fact-Check.md.\n\n'+p.read_text())
print('Corrected SR-137 in local ledger, workbook, report and notes. Source unchanged.')
