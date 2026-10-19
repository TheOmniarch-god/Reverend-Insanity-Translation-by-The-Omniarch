from pathlib import Path
import json, zipfile, re
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from docx import Document
r = Path(__file__).parent
load = lambda n: json.loads((r/n).read_text())
b = load('batch-029-source.json'); led = load('proposal-ledger.json'); p = load('reading-progress.json')
z = zipfile.ZipFile(r.parents[1]/'benchmark/sources/Gu-Zhen-Ren.epub')

# source vs EPUB
for c in b:
    actual = [n.get_text('', strip=True) for n in BeautifulSoup(z.read(c['file']), 'html.parser').find_all('p') if n.get_text('', strip=True)]
    assert c['paragraphs'] == actual, c['corpus_chapter']
assert [c['corpus_chapter'] for c in b] == list(range(546, 566))
N = sum(len(c['paragraphs']) for c in b); CH = sum(len(t) for c in b for t in c['paragraphs'])
assert N == 1800 and CH == 63491

# numbered source exact
src = (r/'Batch-029-Numbered-Chinese-Source.md').read_text(encoding='utf-8')
cur = None; found = []
for line in src.splitlines():
    m = re.match(r'## Chapter (\d+) ', line)
    if m: cur = int(m[1])
    m = re.match(r'(\d+)\.(\d+) (.*)', line)
    if m:
        assert int(m[1]) == cur
        assert m[3] == b[cur-546]['paragraphs'][int(m[2])-1]; found.append((cur, int(m[2])))
assert found == [(c['corpus_chapter'], i) for c in b for i in range(1, len(c['paragraphs'])+1)]

# ledger
ids = [v['id'] for v in led]
assert len(led) == 337 and len(set(ids)) == 337 and ids[-1] == 'SR-337'
assert all(not v['approved'] and not v['applied'] for v in led)
new = [v for v in led if int(v['id'][3:]) >= 307]
assert len(new) == 31
for v in new:
    for l in v['locations']:
        assert 1 <= l['chapter'] <= 565
        if 546 <= l['chapter'] <= 565:
            assert 1 <= l['paragraph'] <= len(b[l['chapter']-546]['paragraphs']), (v['id'], l)
    assert v['witness_collation'] and v['status'].startswith(('Held repair', 'Provisional')), v['id']
held = sum(1 for v in led if int(v['id'][3:]) >= 286 and v['status'].startswith('Held repair'))
prov = sum(1 for v in led if int(v['id'][3:]) >= 286 and v['status'].startswith('Provisional'))
assert held == 30 and prov == 22, (held, prov)
assert sum(1 for v in new if v['status'].startswith('Held repair')) == 18
witness_supported = [v['id'] for v in new if 'Witness-supported' in v['confidence'] or 'Witness supports' in v['witness_collation']]
assert set(witness_supported) >= {'SR-315','SR-318','SR-322','SR-326'}, witness_supported
assert [v['confidence'] for v in led if v['id']=='SR-299'][0].startswith('Moderate (0.6)')

# progress
assert p['next_chapter'] == 566 and p['sequential_chapters_completed'] == list(range(1, 566))
assert p['paragraphs_read'] == 50481 and p['paragraph_characters_read'] == 1774835
assert p['remaining_chapters'] == 1769

# workbook
w = load_workbook(r/'Sequential-Reading-and-Proposals.xlsx', read_only=True, data_only=True)
cc = [row for row in w['Chapter completion'].iter_rows(min_row=2, values_only=True) if row[0] and str(row[0]).isdigit() and int(row[0]) >= 526]
assert len(cc) == 40
pr = [row for row in w['Paragraph register'].iter_rows(min_row=2, values_only=True) if row[0] and str(row[0]).isdigit() and int(row[0]) >= 526]
allsrc = load('batch-028-source.json') + b
assert len(pr) == 1761 + N and all(row[6] == allsrc[row[0]-526]['paragraphs'][row[1]-1] for row in pr)
pl = [row for row in w['Proposal ledger'].iter_rows(min_row=2, values_only=True) if row[0] and str(row[0]).startswith('SR-') and int(str(row[0])[3:]) >= 286]
assert len(pl) == 52, len(pl)
rm = dict(w['Read me'].iter_rows(min_row=2, max_col=2, values_only=True))
assert rm['Next unread'] == '566' and rm['Completed'] == '1–565 of 2334' and rm['Paragraphs'] == '50481'

# docx
for name, n_ids in [('Batch-029-Chapters-546-565.docx', ('SR-307', 'SR-337')), ('Batch-028-Chapters-526-545.docx', ('SR-286', 'SR-306'))]:
    d = Document(r/name)
    txt = '\n'.join(x.text for x in d.paragraphs) + '\n' + '\n'.join(c.text for t in d.tables for row in t.rows for c in row.cells)
    assert n_ids[0] in txt and n_ids[1] in txt and len(d.tables) == 2, name

# report hygiene
for name in ['Batch-029-Chapters-546-565.md', 'Batch-028-Chapters-526-545.md']:
    rep = (r/name).read_text(encoding='utf-8')
    assert '—' not in rep, name
    assert not any(line.startswith('> ') for line in rep.splitlines()), name

qa = {'status': 'passed', 'epub_paragraphs_rechecked': N, 'numbered_source_exact': True, 'numbered_source_lines': N,
      'workbook_chapter_rows_batch_029_range': 40, 'workbook_paragraph_rows_526_565': 1761 + N, 'workbook_paragraph_text_exact': True,
      'workbook_ledger_rows_286_337': 52, 'workbook_read_me_next_unread': 566, 'docx_built': ['Batch-028-Chapters-526-545.docx', 'Batch-029-Chapters-546-565.docx'],
      'docx_table_count_each': 2, 'ledger_total': 337, 'held_repair_candidates_286_plus': held, 'provisional_286_plus': prov,
      'batch_029_new_ledger': 31, 'batch_029_held': 18, 'batch_029_queries': 13, 'approved_or_applied': 0,
      'witness_collated_chapters': 20, 'witness_supported_proposals_029': sorted(witness_supported),
      'sr299_confidence': 'Moderate (0.6)', 'em_dashes_in_reports': 0, 'progress_next': 566, 'progress_paragraphs': 50481, 'progress_characters': 1774835}
(r/'batch-029-qa.json').write_text(json.dumps(qa, ensure_ascii=False, indent=2))
print(json.dumps(qa, ensure_ascii=False, indent=2))
