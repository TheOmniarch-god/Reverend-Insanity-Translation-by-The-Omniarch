from pathlib import Path
import json,hashlib
from docx import Document
from openpyxl import load_workbook
r=Path(__file__).parent
p=r/'proposal-ledger.json';ledger=json.loads(p.read_text());updates={
'SR-134':dict(recommended_action='Retain as a temporal-scope query. Test whether 这些年 scopes only frugal city residence or also eye-treatment saving. 这段时间 remains conditional if a shared scope is intended; no global date repair selected.',confidence='Moderate-to-high local temporal-scope concern; exact repair unconfirmed; broader clock unresolved.'),
'SR-135':dict(recommended_action='Retain93% and seven points remaining. Keep 两分→一分 only as a conditional candidate if92% is confirmed as the intended baseline; investigate retrospective91% baseline first.',confidence='High confidence in quoted figures and arithmetic; emendation unconfirmed because the retrospective baseline is not explicit.',status='Provisional baseline query; numerical replacement not selected; no approval'),
'SR-136':dict(recommended_action='Retain strong source-office inconsistency finding; 家老 / clan elder is a conjectural candidate, not an attested replacement in the examined witness.',confidence='High local office conflict; moderate confidence in exact replacement.'),
'SR-137':dict(category='Chinese textual variant: 冰雹蛊 versus 冰爆蛊; not an English mistranslation',recommended_action='Retain Hail Gu for EPUB 冰雹蛊. Record externally attested 冰爆蛊 / Ice Explosion Gu variant at the same sentence. Replacement suspended pending source selection and corroboration.',rationale='Hail correctly translates冰雹. Hetushu Ch312 reads冰爆 in the same sentence, consistent with local299/313 ice detonation. This establishes a variant, not authorial priority or an English mistranslation.',confidence='High lexical equivalence and variant existence; authorial reading unestablished.',dependencies='External witness provenance and textual relationship unknown; source selection and further corroboration required.'),
'SR-138':dict(recommended_action='Retain figures. Treat four monthly bouts as phase-specific practice unless permanence is proved. Reconcile carried-in wins and eighteen-month anchor; no replacement interval selected.',confidence='High confidence in located constraints; inconsistency unproved. Monthly allocations are not automatically full elapsed months.')}
for x in ledger:
 if x['id'] not in updates:continue
 x.update(updates[x['id']]);x['latest_update']='Second-pass fact-check: see Batch-016-Fact-Check-Addendum.md; prior confidence/action superseded. All substantive changes remain held.'
 x['fact_check_addendum']='Batch-016-Fact-Check-Addendum.md'
 assert x['approved'] is False and x['applied'] is False
p.write_text(json.dumps(ledger,ensure_ascii=False,indent=2))
w=load_workbook(r/'Sequential-Reading-and-Proposals.xlsx');s=w['Proposal ledger'];byid={x['id']:x for x in ledger}
for row in s:
 if row[0].value not in updates:continue
 x=byid[row[0].value]
 for col,key in [(2,'category'),(4,'recommended_action'),(5,'rationale'),(6,'confidence'),(7,'status'),(8,'dependencies'),(12,'latest_update')]:s.cell(row[0].row,col,x[key])
w.save(r/'Sequential-Reading-and-Proposals.xlsx')
notice='SECOND-PASS FACT-CHECK: SR-134 through SR-138 are superseded by Batch-016-Fact-Check-Addendum.md/.docx. This includes a verified external 冰爆蛊 variant, an unresolved aptitude baseline, and separate confidence for identifying an issue and selecting its repair. Original discussion below is retained as review history, not current authority.'
for p in [r/'Batch-016-Chapters-301-315.md',r/'jotters/Batch-016-Continuity.md',r/'jotters/batch-016-working.md',r/'SR-137-Fact-Check.md']:
 if p.exists() and notice not in p.read_text():p.write_text(notice+'\n\n'+p.read_text())
p=r/'Batch-016-Chapters-301-315.docx';d=Document(p)
if notice not in '\n'.join(p.text for p in d.paragraphs):d.paragraphs[0].insert_paragraph_before(notice)
d.save(p)
d=Document()
for line in (r/'Batch-016-Fact-Check-Addendum.md').read_text().splitlines():
 if line.startswith('# '):d.add_heading(line[2:],0)
 elif line.startswith('## '):d.add_heading(line[3:],1)
 elif line.strip():d.add_paragraph(line.lstrip('> ').replace('**',''))
d.save(r/'Batch-016-Fact-Check-Addendum.docx')
p=r/'CURRENT-STATUS.md';p.write_text('## Batch016 second-pass completed for SR134–138\nSee Batch-016-Fact-Check-Addendum.md/.docx for controlling decisions. External Chinese Hail/Ice Explosion variant verified; authorial priority unknown. Aptitude baseline query downgraded; office contradiction retained with conjectural repair. No new sequential coverage; next316. Remote snapshot not updated.\n\n'+p.read_text())
sha=hashlib.sha256(Path('/home/user/benchmark/sources/Gu-Zhen-Ren.epub').read_bytes()).hexdigest();assert sha=='1ef1c29b6bddd7994f56696e2d091363aa8bbcee4889439465a9e2c51437dd5d'
check=load_workbook(r/'Sequential-Reading-and-Proposals.xlsx',read_only=True);assert check['Proposal ledger'].max_row==140
assert len(ledger)==139 and all(not x.get('approved') and not x.get('applied') for x in ledger)
(r/'fact-check-016/qa.json').write_text(json.dumps({'proposals_rechecked':list(updates),'proposal_count':139,'source_epub_sha256':sha,'substantive_edits_applied':False,'sequential_progress_advanced':False,'external_witnesses':5,'remote_uploaded':False},indent=2))
print('Fact-check addendum and updated ledger/workbook saved; source unchanged; next316.')
