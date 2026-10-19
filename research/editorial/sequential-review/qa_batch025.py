from pathlib import Path
import json, zipfile, re
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from docx import Document
r = Path(__file__).parent
load = lambda n: json.loads((r/n).read_text())
b = load('batch-025-source.json'); led = load('proposal-ledger.json'); p = load('reading-progress.json')
z = zipfile.ZipFile(r.parents[1]/'benchmark/sources/Gu-Zhen-Ren.epub')

# source vs EPUB
meta = load('../whole-novel-audit/coverage.json') if (r.parent/'whole-novel-audit/coverage.json').exists() else None
for c in b:
    actual = [n.get_text('', strip=True) for n in BeautifulSoup(z.read(c['file']), 'html.parser').find_all('p') if n.get_text('', strip=True)]
    assert c['paragraphs'] == actual, c['corpus_chapter']
assert [c['corpus_chapter'] for c in b] == list(range(466, 486))
N = sum(len(c['paragraphs']) for c in b); CH = sum(len(t) for c in b for t in c['paragraphs'])
assert N == 1685 and CH == 62679

# numbered source exact
src = (r/'Batch-025-Numbered-Chinese-Source.md').read_text(encoding='utf-8')
cur = None; found = []
for line in src.splitlines():
    m = re.match(r'## Chapter (\d+) ', line)
    if m: cur = int(m[1])
    m = re.match(r'(\d+)\.(\d+) (.*)', line)
    if m:
        assert int(m[1]) == cur
        assert m[3] == b[cur-466]['paragraphs'][int(m[2])-1]; found.append((cur, int(m[2])))
assert found == [(c['corpus_chapter'], i) for c in b for i in range(1, len(c['paragraphs'])+1)]
assert len(found) == 1685

# ledger
ids = [v['id'] for v in led]
assert len(led) == 256 and len(set(ids)) == 256 and ids[-1] == 'SR-256'
assert all(not v['approved'] and not v['applied'] for v in led)
for v in led:
    if int(v['id'][3:]) >= 233:
        for l in v['locations']:
            assert 1 <= l['paragraph'] <= len(b[l['chapter']-466]['paragraphs']) if 466 <= l['chapter'] <= 485 else True

# progress
assert p['next_chapter'] == 486 and p['sequential_chapters_completed'] == list(range(1, 486))
assert p['paragraphs_read'] == 41727 + N and p['paragraph_characters_read'] == 1454484 + CH
assert p['remaining_chapters'] == 1849

# workbook
w = load_workbook(r/'Sequential-Reading-and-Proposals.xlsx', read_only=True, data_only=True)
cc = [row for row in w['Chapter completion'].iter_rows(min_row=2, values_only=True) if row[0] and str(row[0]).isdigit() and int(row[0]) >= 466]
assert len(cc) == 20
pr = [row for row in w['Paragraph register'].iter_rows(min_row=2, values_only=True) if row[0] and str(row[0]).isdigit() and int(row[0]) >= 466]
assert len(pr) == N and all(row[6] == b[row[0]-466]['paragraphs'][row[1]-1] for row in pr)
pl = [row for row in w['Proposal ledger'].iter_rows(min_row=2, values_only=True) if row[0] in [f'SR-{i}' for i in range(233, 257)]]
assert len(pl) == 24
rm = dict(w['Read me'].iter_rows(min_row=2, max_col=2, values_only=True))
assert rm['Next unread'] == '486' and rm['Completed'] == '1–485 of 2334'

# docx
d = Document(r/'Batch-025-Chapters-466-485.docx')
txt = '\n'.join(x.text for x in d.paragraphs) + '\n' + '\n'.join(c.text for t in d.tables for row in t.rows for c in row.cells)
assert 'SR-233' in txt and 'SR-256' in txt and len(d.tables) == 1

# report hygiene
rep = (r/'Batch-025-Chapters-466-485.md').read_text(encoding='utf-8')
assert '—' not in rep
for line in rep.splitlines():
    if line.startswith('> '): assert line[2:] in {t for c in b for t in c['paragraphs']}

qa = {'status': 'passed', 'epub_paragraphs_rechecked': N, 'numbered_source_exact': True, 'numbered_source_lines': N,
      'workbook_chapter_rows': 20, 'workbook_paragraph_rows': N, 'workbook_paragraph_text_exact': True,
      'workbook_ledger_rows_added': 24, 'workbook_read_me_next_unread': 486, 'docx_proposal_ids_present': True,
      'docx_table_count': 1, 'ledger_total': 256, 'approved_or_applied': 0, 'witness_hashes_verified': 0,
      'report_blockquotes_exact': True, 'em_dashes_in_report': 0, 'progress_next': 486,
      'progress_paragraphs': 41727 + N, 'progress_characters': 1454484 + CH}
(r/'batch-025-qa.json').write_text(json.dumps(qa, ensure_ascii=False, indent=2))
print(json.dumps(qa, ensure_ascii=False, indent=2))
