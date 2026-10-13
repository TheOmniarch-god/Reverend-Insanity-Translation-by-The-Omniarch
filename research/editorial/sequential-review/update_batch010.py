from pathlib import Path
import json, shutil
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
r=Path('/home/user/editorial/sequential-review')
def load(n):return json.loads((r/n).read_text())
def save(n,v):(r/n).write_text(json.dumps(v,ensure_ascii=False,indent=2))
for n in ['reading-progress.json','proposal-ledger.json','chronology-register.json','resource-register.json','working-standards-register.json']:
 d=r/'history'/('before-batch010-'+n)
 if not d.exists():shutil.copy2(r/n,d)
cs=sum([load(f'batch-{i:03d}-source.json') for i in range(1,11)],[]);by={c['corpus_chapter']:c for c in cs};batch=cs[200:220]
def src(c,p):return dict(chapter=c,paragraph=p,file=by[c]['file'],text=by[c]['paragraphs'][p-1])
def ss(rr):return [src(*z) for z in rr]
def loc(rr):return [{k:v for k,v in z.items() if k!='text'} for z in ss(rr)]
def refs(z):return '; '.join(f"Ch{x['chapter']} p{x['paragraph']}" for x in z)
def reader(c,p,t):
 if (c==151 and (31<=p<=67 or p in [98,99])) or (c,p) in [(201,22),(203,84)]:return '[Restricted passage omitted from reader-facing copy. Original numbering retained; original EPUB unchanged. See earlier non-graphic editorial discussion.]'
 return t
ledger=load('history/before-batch010-proposal-ledger.json');new=[]
def pro(i,cat,rr,action,why,conf,dep):
 v=dict(id=f'SR-{i:03d}',category=cat,locations=loc(rr),recommended_action=action,rationale=why,confidence=conf,status='Active provisional proposal; no approval',dependencies=dep,approved=False,applied=False,publication_status='HOLD: full 2334-chapter reading and global reconciliation outstanding');ledger.append(v);new.append(v);return v
pro(95,'Tie Xue Leng inventory scope',[(184,4),(184,20),(184,48),(185,8),(185,16),(185,22),(185,55),(193,21),(202,42)],'Conditional candidate202p42: replace 不过七只左右 with 也不过数种主要蛊虫 if this is an approximate comparison of principal types. If an exact personal inventory is intended, enumerate the holdings and specify the reference time instead.','At least nine named types are identifiable: Heavenly Sound, Righteousness, Empathy, Iron Hand, Oil Dragon, Fire Dragon, Mountain Puppet, Demon-Suppressing Chains and Talisman. Empathy itself was deployed in dozens. Types, specimens, deployed Gu and the current carried inventory are not identical counts.','Moderate scope query; not a certified seven→nine substitution.','Check later inventory and source witnesses. 七只左右 is approximate; a temporary carried subset may be intended. Do not collapse all deployed copies to one specimen.')
pro(96,'Tie family relationship',[(182,22),(211,57),(211,58),(211,63)],'Conditional preference:182叔叔→舅舅 and211侄女→外甥女, retaining211’s repeated maternal-uncle address. In English, leave uncle/niece neutral pending genealogy.','182 identifies the clan chief as paternal uncle;211 twice uses maternal uncle but also says paternal niece. The father is recalled as a close friend. That helps context but does not prove whether the men were related.','High inconsistent kinship terminology; moderate choice of relationship.','Later Tie genealogy may favour the opposite repair: keep paternal uncle and amend舅舅. Do not invent a marriage or sibling relationship.')
pro(97,'Bai Mo Ting rank',[(214,56),(214,57),(215,31),(217,45),(217,46),(217,47)],'If no intervening promotion is established, qualify217p46: 其中三位是三转，百陌亭则是二转. Do not automatically raise the entire earlier scouting party toRevolutionThree.','214 gives oneRevolutionThree leader and threeRevolutionTwo companions.215 places the elder’s nephew in that party;217 names the nephew BaiMoTing among four youths all said to beRevolutionThree. A promotion during the short interval is possible but not narrated.','Moderate-to-high continuity query; conditional repair.','Confirm identity, interval, later ranks and promotion evidence before choosing which passage to change.')
pro(98,'inheritance master title',[(214,28),(214,43),(214,45)],'Preferred214p28 肉骨上师→骨肉上师, if later inheritance material confirms that title.','The same account attributes the inheritance to肉骨上师 then骨肉上师. 肉白骨 is a separate healingGu name and must not be mechanically reordered.','High local title variation; preferred candidate not final.','Later signature, inscription and source variants. Keep the separateGu name untouched.')
pro(99,'cultivation forecast scope',[(214,45),(214,50),(214,51)],'If the pair’s advantage chiefly comes from receiving aRevolutionThree partner’s essence while below her rank, consider214p51 在三转以后→在三转以前. Until the actual mechanism is read, preserve the source and label the claim as FangYuan’s forecast.','The stated after-RevolutionThree acceleration sits beside an emphasis on the early period. This is not yet a demonstrated mechanical contradiction: reciprocal cultivation may still help at higher ranks.','Low; deliberately conditional mechanics query.','Read acquisition and operation of BoneFleshReunion before deciding. No importing later plot knowledge or replacing the source on expectation alone.')
pro(100,'hunt calendar',[(216,70),(217,47),(217,77),(218,1),(218,8),(218,18),(219,47),(219,59),(219,71)],'Provisional nonnumerical option216p70: 从今日起连续七天→从今日起的这些天. Alternatively retain seven if later evidence establishes that the private contest continues beyond the official hunt. Do not invent an extension announcement.','Explicit day labels yield huntday1,2,3,5; the challenge starts nextmorning(day6), its second day isday7, thirdday isday8, and the mountain outing follows(day9). This exceeds the announcedseven days if both contests share one duration.','High counted sequence; moderate inconsistency because contest scopes may differ.','Later closing ceremony, private-contest rules and source witnesses. Do not change七→九/十 without evidence for the intended total.')
pro(101,'Tie pursuit casualty count',[(214,66),(214,82),(214,88),(214,98),(220,41),(220,49),(220,51),(220,52),(220,54)],'Preferred220p51 其他的几人也遭殃，被炸死数人→其他三人也遭殃，一人被当场炸死. Retain the two initial wounded, one later death, and TieDaoKu’s survival.','Nine set out; two die to crocodiles and three to the pheasant, leaving four. At the blast AoTian dies, two are wounded, and one of those later dies. Only one additional immediate fatality fits the established four; 数人 after separating AoTian overstates that remainder.','High local count; strong minimal candidate.','Check casualties and whether any replacement pursuers joined. None has been narrated through220.')
pro(102,'Chen Cui Hua reconnaissance loot',[(210,77),(210,78),(211,110),(211,115),(211,117),(211,118),(211,127),(211,129)],'Retain the five named liveGu as the confirmed recoveries; keep the unidentified reconnaissanceGu as an unresolved disposal query. If later evidence confirms it died with the damaged head, add only that source-supported explanation.','The350-pace sightGu is explicit but absent from the recovered list. Bookworm is explicitly dead; the consumed ThirdWatchGu is an effect, not recoverableloot. Neither automatically accounts for the unnamed sightGu.','High inventory gap; low confidence in any explanatory addition.','Possible unlisted loss, bodily location, identity overlap or incomplete enumeration. No invented sixth live acquisition or automatic destruction claim.')
fixes=[
(201,2,'在血罩内白凝冰携手','在血罩内与白凝冰携手'),(201,3,'在江山漂流','在江上漂流'),(201,5,'也至于受伤','也不至于受伤'),(201,12,'递给这些东西分作两份','将这些东西分作两份'),(201,40,'变不会男儿身','变不回男儿身'),(201,48,'水整掉','水拧掉'),(201,51,'他嚼得','她嚼得'),(201,58,'要去木柴','要去取木柴'),(201,60,'就在这是','就在这时'),(201,67,'火升好','火生好'),
(202,1,'芊芊玉手','纤纤玉手'),(202,16,'用资源和蛊虫来千方百计地来','用资源和蛊虫千方百计地'),(202,20,'真元海地','真元海底'),(202,67,'目光抖寒','目光陡寒'),
(203,5,'不同戴天','不共戴天'),(203,48,'沙硕','沙砾'),(203,94,'缓了一口回来','缓了一口气回来'),
(204,10,'多少次被当面族长、家老夸赞','多少次被族长、家老当面夸赞'),(204,15,'偏僻入里','鞭辟入里'),(204,69,'却捂住不断流淌的血液','却捂不住不断流淌的血液'),(204,91,'不顾上愤怒','顾不上愤怒'),
(205,19,'野外流魄','野外流落'),(206,18,'一夜修行','一夜休息'),(206,25,'再催出一记血月蛊','再催出一记血色月刃'),(206,53,'无以伦比','无与伦比'),
(207,12,'四肢粗壮的腿','四条粗壮的腿'),(207,25,'后辈','后背'),(207,27,'不敢迎接','不敢硬接'),(207,59,'一张纸巨大的蜘蛛网','一张张巨大的蜘蛛网'),
(208,21,'残羹剩炙','残羹剩炙'),(208,60,'力气，却剿除','力气，去剿除'),(208,67,'楞了楞','愣了愣'),(208,74,'在加上高速飞行','再加上高速飞行'),(208,97,'向随水中深处','向水中深处'),
(209,12,'而是当然独处时','而是当人独处时'),(209,44,'目无表情','面无表情'),
(210,4,'作鬼','作怪'),(210,12,'外门风大雨大','外面风大雨大'),(210,17,'别惹这么麻烦','别惹这种麻烦'),(210,18,'楞了一下','愣了一下'),(210,32,'尽量得夺得','尽量地夺得'),(210,40,'极可是原先','极可能原先'),
(211,7,'背后黑的蝠翅','背后黑色的蝠翅'),(211,14,'存息葬玉蛊','存息玉葬蛊'),(211,25,'眉间一直垂到腰间','眉尖一直垂到腰间'),(211,28,'天壤之距','天壤之别'),(211,37,'存息葬玉蛊','存息玉葬蛊'),(211,59,'脸上愣住','愣住'),(211,100,'怕不起来','爬不起来'),(211,100,'淌外来','淌出来'),(211,105,'正要动手','正要动身'),(211,107,'他利用山猪死亡','她利用山猪死亡'),(211,124,'收刮出来','搜刮出来'),(211,151,'撇了地上的尸体一眼','瞥了地上的尸体一眼'),
(212,14,'重生一来','重生以来'),(212,26,'比酒虫还有珍贵','比酒虫还要珍贵'),(212,34,'他浑身上下','她浑身上下'),(212,37,'炼成了他的本命蛊','炼成了她的本命蛊'),(212,38,'他一边嘴中抱怨着','她一边嘴中抱怨着'),(212,59,'月谷蛊','月光蛊'),
(213,33,'最好不够','最好不过'),(213,52,'枝牙','枝芽'),(214,28,'肉骨上师','肉骨上师'),(214,98,'在加上其他两位','再加上其他两位'),
(216,17,'微楞','微愣'),(217,31,'半柱香','半炷香'),(217,51,'然而目光紧紧地盯住','然后目光紧紧地盯住'),(217,70,'带到酒宴','待到酒宴'),(217,85,'一对胞胎兄妹','一对双胞胎兄妹'),
(218,27,'说吧。”。','说吧。”'),(218,44,'接收挑战','接受挑战'),(218,46,'白战猎','百战猎'),(218,66,'其热融融','其乐融融'),(218,75,'骨兽体制强硬','骨兽体质强硬'),
(219,6,'如果操作得很','如果操作得好'),(219,22,'品性不断','品性不端'),(219,34,'成功再望','成功在望'),(219,38,'世间拖得越久','时间拖得越久'),(219,61,'砰然一动','怦然一动'),(219,66,'保驾护的隐蛊师','保驾护航的隐身蛊师'),(219,67,'楞了一下','愣了一下'),(219,84,'白莲有意','百莲有意'),(219,85,'白莲故意','百莲故意'),
(220,15,'这在这一刹那间','就在这一刹那间'),(220,16,'他们俩个少年','他们两个少年'),(220,34,'楞了一下','愣了一下'),(220,35,'没有涂到','没有吐到'),(220,74,'头昏脑涨','头昏脑胀'),(220,88,'他们一定从青茅山','他们一定是从青茅山')]
fixes=[z for z in fixes if z[2]!=z[3]]
for c,p,a,b in fixes:assert a in src(c,p)['text'],(c,p,a)
v=pro(103,'located copy candidates',sorted(set((c,p) for c,p,a,b in fixes)),'Apply only individually reviewed substitutions after approval. Current-scene pronoun candidates are local consistency choices, not a blanket policy for Bai’s earlier life or identity.','Obvious wrong characters, omissions, names and a few contextual diction choices. Bigger inventory, genealogy and calendar questions have separate records.','High for clear copy slips; moderate for interpretive wording.','Source variants and contextual review. Preserve quoted originals; do not globally replace白 with百 or change historical male references.')
moderate={(201,48),(202,16),(203,48),(206,18),(206,25),(207,27),(210,4),(210,17),(211,59),(211,105),(217,51),(218,75),(219,66)}
v['copy_items']=[dict(chapter=c,paragraph=p,original=a,replacement=b,confidence='Moderate contextual candidate' if (c,p) in moderate else 'High local copy candidate',status='HOLD; proposed only') for c,p,a,b in fixes]
updates={'SR-001':'202p10 independently confirms final90% aptitude.212p82 restores current reserves to90; no change to final90 proposed.', 'SR-007':'211 ThirdWatch is3x personal time, not+300%=4x and not3x aptitude.212 fine-control use exhausts then refills90points; do not treat it as a full moonblade count.', 'SR-060':'209p27 explicitly says Bai previously used IceMuscle; persistent defence does not itself prove current ownership or its original consumption.', 'SR-081':'209p5 again explicitly roots Ear for stronger hearing; no unconditional correction of190 moving300 or169200 is established.', 'SR-082':'215–219 female Bai chief and217p90 inheritance by children reinforce need for clan-specific patrimony/office scope, not a universal male-ruler prohibition.', 'SR-085':'202 final90 corroborated;212 midstage uses90reserves.198’s43/53/63 chain remains a source discrepancy, not permission to lower90.', 'SR-090':'No new StoneAperture mechanism adjudication; postreversal native recovery is observed normally.', 'SR-091':'216–219 now introduces a separate seven-day Bai hunt with an internal day-count querySR100. Do not use it to revise the earlier three-clan tournament.'}
for v in ledger:
 if v['id'] in updates:v['batch010_update']=updates[v['id']]
events=load('history/before-batch010-chronology-register.json')['entries'];resources=load('history/before-batch010-resource-register.json')['entries'];ne=[];nr=[]
def ev(timeline,kind,event,constraint,rr,status='Explicit source anchors; interpretation remains revisable'):
 v=dict(id=f'T{len(events)+1:03d}',timeline=timeline,kind=kind,event=event,constraint=constraint,evidence_status=status,dependency='Whole-novel reconciliation outstanding',sources=ss(rr));events.append(v);ne.append(v)
def res(kind,entry,calc,rr,status='Held evidence; no manuscript edit'):
 v=dict(id=f'R{len(resources)+1:03d}',kind=kind,entry=entry,calculation=calc,status=status,sources=ss(rr));resources.append(v);nr.append(v)
ev('Revised current life','five-day anchor','Shore rest and Tie search','201 states five days since rebirth;202 Tie death five/six days ago. Search days/nights occur within that interval, not an extra trip to add.',[(201,1),(201,46),(202,80),(202,83),(202,86)])
ev('Revised current life','combat duration','Crocodile battle','Over halfshichen before fatigue; another quarterhour before retreat; kingfight overhalfshichen. Do not halve时辰.',[(203,82),(204,22),(204,69)])
ev('Revised current life','overnight sequence','Tree camp, stone camp and illness','Eveningdeparture; tree sleep; late-morning wake/noon departure; nextnightrockcamp; illnessafterdays and sixdayrecovery.',[(205,25),(205,78),(205,89),(206,1),(206,23),(206,65),(206,79),(206,87),(206,89)])
ev('Revised current life','approximate interval','Wilderness detour','207 morethanhalfmonth ellipsis; five further scoutingdays; nocturnalpheasant; morningbutchery.209 halfmonthrecap mayoverlap; no exact summed calendar certified.',[(207,2),(207,55),(208,1),(208,6),(208,22),(209,36)])
ev('Revised current life','evidence age versus elapsed','Chen’s trail','Fire estimatedhalfmonthold isnot newelapsedtime. Woundrestseveraldays;restartday2fire/day3fight;laterdayslead tocave.',[(209,8),(209,51),(209,54),(209,56),(209,67),(209,74)])
ev('Recorded revised battle / later sect scene','replay and forecast','Sky Crane’s report','Opening211 is recording, not newlivingFounder. Returntraveltimeunquantified; SkyCranehasdaysleft and promises20yearresult.',[(211,9),(211,11),(211,14),(211,16),(211,17),(211,37),(211,40),(211,47)])
ev('Parallel Tie clan scene','return and genealogy','Tie Ruo Nan at home','Returned severaldays, nineteenthrequest. No synchronised datedarrival or established travel speed. Uncle terminology unresolvedSR096.',[(211,49),(211,57),(211,58)])
ev('External clock versus accelerated personal clock','threefold time','Chen Cui Hua’s consumed Third Watch','Oneexternalday≈threepersonal; recovery andpoison3x. Lessoneexternalyear≈2–3personalyears. Plannedthreedaywait abandonedafteroneexternalday.',[(210,79),(211,77),(211,129),(211,133),(211,135),(211,139),(211,140),(211,142)])
ev('Revised current life','cultivation interval','Fang Yuan reaches RevolutionOne middle','Morethan20days afterChenkill. FourFlavour held but noordinaryLiquorWorm; no reverse-refinement yet.',[(212,8),(212,11),(212,16),(212,18),(212,21)])
ev('Revised current life','preparation and onwardtravel','Five pits and wine','Initialday plus nextthreedays preparation; onlytwopitsdetonated. Five days afterwine, mountaininview. No claimallwine fitsTushita.',[(213,73),(213,76),(213,79),(213,80),(213,82),(213,84),(214,7)])
ev('Original500-year life / current scouting','future recollection','Bai migration and twins','Migration tenyearslater;twins18atdiscovery, notexplicitlysameyear. RevolutionFive nearcenturylater. Currentpresenceisscouting,notfullmigration.',[(214,19),(214,26),(214,29),(215,34),(215,41)])
ev('Revised current life','hunt calendar','Seven-day announcement versus later days','Day1bear/Relic;day2rhino;day3stones;day5banquet;day6challenge;day7secondchallenge;day8third;day9mountain. Possible private-extension scope unconfirmed.',[(216,70),(217,47),(217,77),(218,1),(218,8),(218,18),(219,47),(219,59),(219,71)],'Held calendar querySR100')
ev('Revised current life / casualty flashback','pursuit losses','Tie team nine to one','Initial9, crocodileskill2,pheasantkills3=>4. BlastAoTian+oneotherimmediate;twoinjured,oneexpireslater=>DaoKualone.220数人 conflictsremainder.',[(214,66),(214,82),(214,88),(214,98),(220,41),(220,49),(220,51),(220,52),(220,54)])
ev('Revised current life','mountain approach','Escorted path versus solo estimate','Soloestimated>halfyear;Bai escortclearsduring3shichen≈6h, dozenscasualties,threeeldersinjured. Differentforcesnotdistancecontradiction.',[(219,12),(219,77),(219,79),(219,80),(219,84)])
ev('Revised current life','endpoint / surveillance','Hostage scene revealed','Chiefswitchesoffsmokeafterapparentlyemptycave;questioningTieDaoKu;reopensseeingtwohostages. Guinheritancenotyetacquiredby220.',[(220,4),(220,6),(220,28),(220,97),(220,98)])
res('Inventory checkpoint','Twelve Gu owned by Fang Yuan','Lotus,Cicada,FourFlavour,HiddenScale,Plunder,Yang,Skull,BloodMoon,Ear,Tushita,Canopy,Centipede. Borrowed-outGu count as owned.',[(202,26),(202,27),(202,33),(202,34),(202,35),(202,36),(202,37),(202,38),(202,39),(202,40)])
res('Unenumerated loot / losses','Massacre loot and ordinary Liquor Worm','201injuredvoyagelootstarves.212 retrospectivelythirdordinaryWorm fromYaoLe destroyedwhenTiearrived, alongwithotherloot. These aredifferent lossoccasions.',[(201,2),(201,3),(212,15),(212,17),(212,18),(212,20)])
res('Loans / ownership','Revocable use permissions','Canopy+Centipede lent201,Moon203,Lotus204 thenrecalled205. Earrotationplanned210 butBainevercutsEar211. HiddenScaleactuallylent211.',[(201,83),(201,87),(201,88),(203,91),(204,33),(205,16),(210,50),(211,80),(211,86),(211,96)])
res('Food / capacity','Forecasts and replenishment','Wine6months,Plunder5months,Earnearyear. Crocmeatadded205 then2083monthsfeed;213wine2yearforecastwithsurplusnotallstored. Notelapsedtime.',[(202,46),(202,49),(202,51),(202,53),(205,23),(205,24),(208,23),(213,33),(213,84),(213,85)])
res('Acquisitions / physical limits','Back Carapace and Crocodile Strength','TwoRevolution2refined204. BaiassignedStrength205, fullonecrocpowerby213. FYdoubleboarbodycannotyetacceptmorewithoutreinforcement.',[(204,75),(204,79),(204,81),(204,83),(204,86),(205,12),(205,16),(205,17),(207,19),(213,75)])
res('Unobtained targets','Lava crocodile Gu','ThreeidentifiedRevolution3Gu;armourdestroyed,otherstwoabsent208. AshGu neveracquired. PoisonFrog alsoonlytarget.',[(207,37),(207,60),(208,12),(208,14),(208,19)])
res('Acquisitions / consumed effect','Chen Cui Hua loot','Five liveGu: RiceBag,IronThorn,MotherPotato,HairyFoot,JumpingGrass. Bookwormdead. ThirdWatchconsumedbeforemeeting, notsixthliveGu. SightGu350paceunaccounted.',[(210,77),(211,115),(211,117),(211,118),(211,129)])
res('Ownership transfer','Bai’s vital Iron Thorn','212 explicitlyrequested/received/refinedasBai vitalGu; unlike revocableCanopyloan. FYrecallsBloodMoonforfinemanipulation.',[(212,36),(212,37),(212,51),(212,53),(212,55)])
res('Consumable production','Mother and potatoes','Revolution3motherproducesRevolution2seeds;lessthanthirdusable,individuallyrefinedplantedandgrown. Fivepits,twoexplode,threeunusednotautomaticallyremoved.',[(213,48),(213,52),(213,53),(213,54),(213,55),(213,58),(213,80),(213,82),(220,60)])
res('Cultivation and food','FourFlavour retained / refinement options','WorksRevolution2notRevolution1;reverseoptionsyieldtwoWorms,oneWorm,butterflyoruncertainGu. Noneperformedthrough220. FineBloodMooncontrolusesandrefills90points.',[(212,16),(212,21),(212,24),(212,55),(212,60),(212,82)])
res('Acquisition / consumption','Bronze Relic','Bearrewardpromised216,delivered217p36,consumed217p37: RevolutionOne middle→upper. NotRevolutionTwo.',[(216,79),(217,36),(217,37)])
res('Healing / loan / cash','Clear Heat and Bai gifts','Revolution2detoxGu acquiredhuntday2, lentBaiLian day3night;returnnotnarrated220. Unquantifiedstonesday3. Doesnotprovewoundhealingcapability.',[(217,79),(217,80),(217,81),(218,1),(218,4),(218,5),(218,65)])
res('Gu loss after owner death','Tie rescue Gu','Mask+handrescueRuoNan thenBloodFrenzy contaminationdestroysboth. No additionalpresentGu heldbyher inferred.',[(202,80),(202,81),(202,82)])
np=sum(len(c['paragraphs']) for c in batch);nc=sum(sum(map(len,c['paragraphs'])) for c in batch);tp=sum(len(c['paragraphs']) for c in cs);tc=sum(sum(map(len,c['paragraphs'])) for c in cs)
save('proposal-ledger.json',ledger);save('chronology-register.json',dict(scope='Sequential through220; separate external/personal clocks, recordings, original and revised histories',entries=events));save('resource-register.json',dict(scope='Sequential through220; holdings, loans, consumed effects and missing items distinguished',entries=resources))
progress=load('history/before-batch010-reading-progress.json');progress.update(sequential_chapters_completed=list(range(1,221)),paragraphs_read=tp,paragraph_characters_read=tc,next_chapter=221,remaining_chapters=2114,batch_report='Batch-010-Chapters-201-220.md',scope='Full sequential Chinese-source reading through220; no fresh bilingual/live-site audit',whole_novel_chronology_verified=False)
progress['batches'].append(dict(id='010',chapters=[201,220],paragraphs=np,characters=nc,reading_segments='Ten complete untruncated two-chapter outputs'))
progress['chapters'] += [dict(chapter=c['corpus_chapter'],file=c['file'],title=c['title'],paragraphs=len(c['paragraphs']),first_paragraph=1,last_paragraph=len(c['paragraphs'])) for c in batch]
save('reading-progress.json',progress)
save('batch-010-reading-log.json',dict(status='Complete',endpoint=220,next_unread=221,segments=[dict(chapters=[a,a+1],status='Full paragraph sequences printed and read; untruncated',paragraphs=sum(len(by[n]['paragraphs']) for n in [a,a+1])) for a in range(201,221,2)],note='Actual sequential reading; no invented timestamps. Reader copy marks201p22 and203p84 omissions, originalEPUB unchanged.'))
text=['# Numbered Chinese source: Chapters201–220','Original wording retained except explicit reader-facing omission markers at201p22 and203p84 for sexualised teenage-body descriptions. Numbering and originalEPUB unchanged. Corpus201–220 = VolumeII Chapters2–21.']
for c in batch:text += [f"\n## Corpus chapter {c['corpus_chapter']}: {c['title']}",f"Volume: {c['volume']} | EPUB: {c['file']}"]+[f"{i}. {reader(c['corpus_chapter'],i,t)}" for i,t in enumerate(c['paragraphs'],1)]
(r/'Batch-010-Numbered-Chinese-Source.md').write_text('\n\n'.join(text))
for name,vs in [('Chronology',ne),('Resources',nr)]:
 out=[f'# Batch010 {name} jotter','All evidence provisional; no source edits.']
 for v in vs:
  out += [f"## {v['id']}: {v.get('event',v.get('entry'))}",v.get('constraint',v.get('calculation')),v.get('evidence_status',v.get('status'))]
  if 'timeline' in v:out.append('Timeline: '+v['timeline'])
  for z in v['sources']:out += [f"Ch{z['chapter']} p{z['paragraph']}",'> '+reader(z['chapter'],z['paragraph'],z['text'])]
 (r/f'jotters/Batch-010-{name}.md').write_text('\n\n'.join(out))
out=['# Batch010 Continuity and proposed corrections','Every proposal remains unapproved and unapplied.']
for v in new:
 out += [f"## {v['id']}: {v['category']}",f"**Action:** {v['recommended_action']}",f"**Rationale:** {v['rationale']}",f"**Confidence:** {v['confidence']}",f"**Dependencies:** {v['dependencies']}",f"**Status:** {v['publication_status']}"]
 if 'copy_items' in v:
  for z in v['copy_items']:out += [f"### Ch{z['chapter']} p{z['paragraph']}: {z['original']} → {z['replacement']}",z['confidence'], '> '+src(z['chapter'],z['paragraph'])['text']]
 else:
  for z in v['locations']:out += [f"Ch{z['chapter']} p{z['paragraph']}",'> '+src(z['chapter'],z['paragraph'])['text']]
(r/'jotters/Batch-010-Continuity.md').write_text('\n\n'.join(out))
(r/'jotters/Batch-010-Preserve.md').write_text('''# Batch010 preservation cautions

- The twelve-Gu count202 is correct when borrowed-out Gu remain in the owner’s inventory.
- Ordinary Liquor Worm destroyed before escape; FourFlavour remains but cannot serveRevolutionOne. Do not turn205’s wording into loss ofFourFlavour.
- ThirdWatchGu is alreadyconsumed; its threefold personalclock is not extra recoverableloot or a tripled aptitude.
- Chen’s plannedthree-day wait and actualone-daywait are not contradictory. Poison progresses faster forher.
- The Ear loan/ear-cutting is proposed then cancelled; no new ear injury has occurred.
- IronThorn becomesBai’s own vitalGu; Canopy/Centipede/Moon use remainspermission-based. ClearHeat is lent, not destroyed.
- The surviving IceMuscle effect does not solve when the originalGu leftinventory.
- 211’s founder fight is a recording, not another resurrection or thirdrebirth.
- Six-dayillness pause, halfmonthellipsis, five-dayscouting,20-plus-daysafterkill andfive-daysafterwine must retain individualanchors; do not add overlapping recaps.
- Ten-yearfuturemigration and twins’age18 atinheritance are not explicitly the same date. Currentage8 is not established.
- Gu taxonomy’s six practical functions versus seven includingcultivation are different scopes.
- Bai-clanfemalechief does not byitself settle anotherclan’spatrimony rules.
- FangYuan’s alias FangZheng, claimedclansurvivors, apparentromance and opponents’guesses are intentionaldeceptions, not name/castmistakes.
- 220 ends with twochildren aliveas hostages. InheritanceGu notyetobtained and escape notcompleted.
''')
stand=load('history/before-batch010-working-standards-register.json');stand['coverage_through']=220
su={'WS-01':'211 opens with a recording.215–220 aliases, invented survivors and staged courtship require speaker/provenance labels. The recording’s fall of the CraneKing is not explicit proof of death.', 'WS-02':'202 confirms90% final aptitude;211 gives3x personal time, not+300% or3xcapacity.212 temporary essence depletion/refill uses the same90%maximum.', 'WS-03':'212 RevolutionOne initial→middle bycultivation,217 middle→upper byRelic. BoneFlesh forecast214 awaits actualmechanics.', 'WS-04':'202 twelve includesborrowed-outGu.211 Earloan cancelled;212 IronThorn explicitlytransferred/vital.218 ClearHeatloan notsale. ThirdWatch tattoo isconsumedeffect.', 'WS-05':'Track external and3xpersonalclocks separately. Recordings, estimatedtraceage and20yearforecasts are notpresentelapsedtime. Seven-dayhunt vsday9sequenceheldSR100.', 'WS-06':'Feedingforecasts6months/5months/year, then3monthsmeat/2yearswine are stock horizons, nottimejumps. SurpluswinecannotallfitTushita.', 'WS-07':'211 citywalls千里 and214 mountain千丈 are sourcefantasygeometry, notpermissiontoshrinkallgeography. Escortedvs solo route estimates usedifferentresources.', 'WS-08':'Futuremigrationin10years andtwinsinheritanceat18 do notestablishpresentage8. Maternal/paternalTieuncle distinctionsremainopen.', 'WS-09':'Localcopycandidates locatedindividually. 百clan vs白in names iscontextual; no globalreplace. Readercopy201p22/203p84 usesexplicitomissionmarkers.', 'WS-10':'Actualcoverage220/2334;2114remain. No approved/appliedchanges or globalcertification.'}
for v in stand['entries']:v['batch010_update']=su[v['id']]
save('working-standards-register.json',stand)
out=['# Working editorial standards','## Evidence checkpoint: Chapters1–220','No story-wide factual standard is whole-novel-confirmed. Standards guide review, not manuscript approval.']
for v in stand['entries']:out += [f"## {v['id']}: {v['topic']}",f"**Status:** {v['status']}",f"**Rule:** {v['rule']}",f"**Earlier evidence:** {v['evidence']}",f"**Latest qualification:** {v['batch010_update']}",f"**Reopen when:** {v['reopen']}"]
(r/'WORKING-STANDARDS.md').write_text('\n\n'.join(out))
print('REGISTERS',np,nc,tp,tc,len(ledger),len(fixes),len(events),len(resources))
def quote(c,p):return f"**Ch{c} p{p}:**\n\n> {src(c,p)['text']}"
report=f'''# Sequential close-reading review: Chapters 201–220
## Batch 010 | The Omniarch

**Actual coverage: Chapters 1–220 of 2,334. Next unread: 221. Remaining: 2,114.**

Twenty chapters were read in ten complete, untruncated two-chapter sequences: **{np:,} paragraphs / {nc:,} source characters**. Cumulative coverage is **{tp:,} paragraphs / {tc:,} characters**. These are actual reading totals, not counts inferred from extraction or topical screening.

The ledger now holds **103 proposal records: 102 active and held, one withdrawn**. Nine records are new, including SR-103’s **{len(fixes)} individually located copy candidates**. All remain unapproved and unapplied. This is Chinese-source review, not fresh bilingual or live-site verification.

## 1. The twelve-Gu inventory is correct

Chapter 202 supplies a useful explicit checkpoint:

{quote(202,40)}

Counting the named holdings gives exactly twelve:

| Gu | Ownership and use at the checkpoint |
|---|---|
|Heavenly Essence Treasure Lotus|Fang Yuan’s; later lent during the crocodile fight and recalled|
|Spring Autumn Cicada|Vital Gu, physically fixed in the aperture; weak after rebirth|
|Four Flavours Liquor Worm|Still held, but its refinement effect serves Revolution Two, not current Revolution One|
|Hidden Scale Gu|Held; later lent for the ambush|
|Plunder Gu|Held; later used on the dead cultivator’s aperture|
|Yang Gu|The retained half of the matched Yin–Yang pair|
|Blood Skull Gu|Held; no new aptitude increase in this batch|
|Blood Moon Gu|Held; subsequently loaned and recalled|
|Earth Hearing Flesh Ear Grass|Still Fang Yuan’s ear; the proposed exchange is cancelled|
|Tushita Flower|Stores supplies; not an unlimited storehouse|
|Sky Canopy Gu|Owned by Fang Yuan, borrowed by Bai|
|Chainsaw Golden Centipede|Owned by Fang Yuan, borrowed by Bai|

Borrowing does not subtract a Gu from its owner’s inventory. Chapter 201 explicitly states that permission can be revoked by the owner. By contrast, Chapter 212 explicitly gives Iron Thorn Bramble Gu to Bai and makes it Bai’s vital Gu. Those are different transactions.

### Two previously unclear losses are now explained

Chapter 201 says the extra loot collected during the massacre was injured and subsequently died from lack of food during the river journey. This explains why unenumerated emergency spoils should not be added indiscriminately to the twelve.

Chapter 212 separately identifies a third ordinary Liquor Worm, taken from Yao Le, and says Fang Yuan deliberately destroyed it when the Tie investigators arrived:

{quote(212,18)}

This is an earlier disposal, not one of the later starvation deaths. It also resolves the apparent problem with Chapter 205 saying he lacks a Liquor Worm: he still possesses the Four Flavours version, but it cannot provide the required Revolution One effect. Reverse refinement is discussed, not performed.

## 2. Personal time now runs at a different rate

Chen Cui Hua’s Third Watch Gu was consumed before the encounter. The remaining mark represents its effect; it is not an additional live Gu recovered from her body.

{quote(211,133)}

{quote(211,135)}

The text makes the ratio clear:

| Quantity | External clock | Chen’s personal clock |
|---|---|---|
|One ordinary day|1 day|3 days of personal time|
|Cultivation and native recovery|Ordinary reference rate|Threefold temporal opportunity/rate in the stated account|
|Time-dependent poisoning|One day of exposure|Three days’ progression|
|Remaining natural lifespan|Ordinary reference duration|About one-third as much external time, under this mechanism|

**Translate the rate as “three times as fast,” not “increased by 300%,” which would imply four times the original rate.** Do not triple her aptitude or aperture capacity.

Her plan to wait three days and her departure after only one day are not contradictory: the text says the worsening poison makes her abandon her plan. Less than one external year since receiving the inheritance can correspond to roughly two or three personal years. That helps explain her cultivation without inventing extra external years.

The working time standard is therefore being expanded, not discarded: units, elapsed intervals, viewpoints and personal time-flow all need separate labels.

## 3. Wilderness travel is much longer than the initial estimate

The new chronology register preserves the individual anchors rather than forcing a single exact journey total:

- Chapter 201: five days since the rebirth, continuing the shore landing from200.
- Chapters 203–204: prolonged crocodile fighting, including periods measured in 时辰. These are not silently converted to single hours.
- Chapters 205–206: the tree-camp night, a later stone-camp night, subsequent illness and **six days of recovery**.
- Chapter 207: a more-than-half-month transition, then **five days of scouting** around the lava crocodile.
- Chapter 208: the pheasant’s night of predation, followed by morning scavenging.
- Chapter 209: additional wound recovery, explicitly numbered early tracking days, and further days before reaching the cave.
- Chapter 212: **more than twenty days after Chen’s death**, Fang Yuan reaches Revolution One middle stage.
- Chapter 213: pit construction extends beyond the first afternoon, with five pits completed over the stated preparation sequence.
- Chapter 214: **five days after obtaining the monkey wine**, White Bone Mountain comes into view.

Some recaps can overlap, especially the half-month language around207–209. No exact cumulative day count is certified. Chapter 216’s statement that news of Qing Mao’s destruction has spread over months is not automatically a contradiction.

An estimated fifteen-day-old campfire is the **age of evidence**, not fifteen new days added to the travellers’ clock. Feeding reserves for six months, three months or two years are likewise projections, not time jumps.

## 4. The seven-day hunt has a real scope question

The announced duration is explicit:

{quote(216,70)}

The later sequence maps as follows:

| Announced hunt day | Event |
|---|---|
|1|Bear hunt; Bronze Relic delivered and consumed|
|2|Rhino hunt; Clear Heat Gu awarded|
|3|Primeval-stone gift; Clear Heat borrowed by Bai Lian that evening|
|5|Banquet questioning about settlement|
|6|Next-morning confrontation and the private team contest’s first day|
|7|“Another day” of the team contest|
|8|The private contest’s explicitly named third day|
|9|Following-day journey onto White Bone Mountain|

**SR-100, held:** if the entire sequence belongs to the same seven-day event, the announcement and subsequent labels conflict. A conservative candidate removes the exact seven-day commitment while retaining the original reading in the register.

However, a private contest continuing after the official hunt could explain the sequence. That extension has not been explicitly established by220. I am therefore not changing seven to nine or ten, nor inventing an announcement that extends the hunt.

## 5. Tie-family findings: a count, a relationship and a recording

### Nine pursuers become one survivor

The pursuit begins with Tie Ao Tian and eight companions. Two die to crocodiles; three more die to the pheasant. Four remain before the explosion.

Chapter 220 identifies Ao Tian’s immediate death and two initial wounded, one of whom dies shortly afterwards. This permits **one other immediate fatality**, leaving Tie Dao Ku as the sole survivor.

{quote(220,51)}

**SR-101:** the phrase saying several of the other members were killed overstates the remainder if read literally. The preferred local candidate is: “The other three were also caught in the blast; one was killed immediately.” No replacement members have been narrated.

### Uncle versus uncle is not neutral in Chinese

Chapter 182 uses 叔叔, while211 twice uses 舅舅 and also uses 侄女. The paternal and maternal relationship terms disagree.

**SR-096:** a provisional preference follows the repeated maternal-uncle address, changing182’s term and211’s niece term accordingly. But later genealogy must decide this. English “uncle” and “niece” can remain neutral in the meantime. The father’s description as a close friend does not by itself exclude a family relationship.

### The founder’s reappearance is a recording

Chapter 211 opens with the founder’s last fight, then explicitly reveals that it is being replayed on a wall. It is neither another rebirth nor a newly living founder.

The Crane King is alive at the start of that recorded fight and subsequently falls. Falling is not an explicit death statement. The precise relationship between that bird and the later giant mount remains open; the earlier character guess about the king’s death is not upgraded to fact.

Sky Crane’s remaining days of life, his purchased Soul Incubating Flea, and his promise of success within twenty years are respectively a current condition, an acquisition and a forecast. They are not twenty elapsed years.

### Tie Xue Leng’s “about seven Gu”

**SR-095** queries the scope of202’s estimate. At least nine named types occur in the earlier battle and tracking account, including Empathy Gu, of which multiple copies were deployed. An approximate count of principal types, a carried inventory at one time and all deployed specimens are different things. I have not silently changed seven to nine.

## 6. New Gu, missing Gu and healing limits

| Development | Confirmed outcome |
|---|---|
|Crocodile king defeated|Back Carapace and Crocodile Strength, both Revolution Two, captured and refined|
|Crocodile Strength given to Bai|A training process, not instant full strength; one crocodile’s strength is complete by213|
|Fang Yuan’s strength|Still two boars; further growth requires bodily reinforcement to avoid damage|
|Lava crocodile defeated by the pheasant|Flame Armour destroyed; the other two Gu missing. Ash Gu **not acquired**|
|Chen Cui Hua killed|Five live Gu recovered; Bookworm explicitly dead|
|Third Watch Gu|Previously consumed; no new acquisition|
|Bronze Relic awarded|Actually delivered and consumed217; Revolution One middle becomes upper stage|
|Clear Heat Gu awarded|Revolution Two detoxification Gu acquired217; lent to Bai Lian218; return not narrated through220|

The five confirmed live recoveries from Chen are Rice Bag Grass, Iron Thorn Bramble, the potato-producing mother Gu, Hairy Foot Gu and Jumping Grass. **SR-102** preserves an inventory gap: her350-pace reconnaissance Gu was described earlier but is not identified among these recoveries. It might have been lost or destroyed, but the text has not supplied that explanation. The dead Bookworm and consumed Third Watch Gu must not automatically be used to fill the gap.

Clear Heat ends the absolute absence of any healing-category Gu. Its stated function is detoxification. It does **not** establish that the party now possesses general wound regeneration, and while it is borrowed out, ownership is not the same as immediate availability.

Chapter 209 confirms that the effect of Bai’s earlier Ice Muscle use still protects the body. That persistent effect does not settle whether the original Gu remained in the disputed Chapter140 inventory.

## 7. Revolution and inheritance wording still need later evidence

### Bai Mo Ting: SR-097

The first scouting group consists of one Revolution Three elder and three Revolution Two companions. The elder’s nephew is present. The next banquet names that nephew, Bai Mo Ting, among four young cultivators all described as Revolution Three.

An unmentioned promotion is possible, but not established. The conditional proposal qualifies the later blanket rank statement rather than upgrading the entire earlier party. Identity and later rank recaps still need checking.

### The inheritance master’s title: SR-098

Chapter 214 alternates 肉骨上师 and 骨肉上师 for the same master. The latter is the preferred local candidate, pending inscriptions or later naming evidence. The separate healing Gu 肉白骨 must not be reordered with it.

### Bone Flesh Reunion: SR-099

The forecast of accelerated cultivation “after Revolution Three” sits beside an emphasis on the early stages and the benefit of Bai’s strength. A possible “before Revolution Three” reading is logged at **low confidence only**. The actual mechanism has not yet been read. A reciprocal cultivation method may still offer higher-rank benefits, so expectation is not enough to emend the source.

The inheritance has not been acquired by220. Its forecasted contents are not present inventory.

## 8. Preserve the deliberate deceptions and distinct scopes

- Fang Yuan calling himself **Fang Zheng** is intentional. Do not correct his dialogue to his real name.
- The surviving Gu Yue delegation, its approaching elders and its supposed spring information are fabricated leverage.
- The apparent romantic interest, rivalry, honesty during the hunt and worry about the mountain are performances. The Bai observers’ conclusions do not become narrator-confirmed facts.
- The female Bai clan chief reinforces the need to distinguish clan-specific inheritance/property rules from a universal prohibition on female leadership. It does not alone resolve SR-082.
- The future migration in ten years and the twins’ inheritance discovery at eighteen are not explicitly the same event-date. Their present age is not thereby proved to be eight.
- Six practical Gu functions and a seven-category list including cultivation are different classifications, not automatically inconsistent arithmetic.
- The half-year estimate for two fugitives to reach the inheritance and the Bai party’s six-hour escorted approach use different forces and routes. They do not justify shrinking the mountain.

## 9. Copy candidates and source preservation

SR-103 records **{len(fixes)} candidates** with exact paragraph locations, original text, proposed wording and confidence. Examples include:

| Location | Source | Candidate |
|---|---|---|
|201p3|江山漂流|江上漂流|
|204p15|偏僻入里|鞭辟入里|
|211pp14,37|存息葬玉蛊|存息玉葬蛊|
|212p59|月谷蛊|月光蛊|
|218p46|白战猎|百战猎|
|219pp84–85|白莲|百莲|

These are not global substitutions. 白 remains correct in White Bone Mountain and Bai Ning Bing. Local narrative pronoun candidates do not authorise rewriting historical male references or imposing a blanket identity policy.

The numbered reading copy retains source wording and paragraph numbering, with **explicit omission markers at201p22 and203p84** for sexualised teenage-body descriptions. The original EPUB and internal source extraction are unchanged. These omissions do not affect reading-completion claims or the numerical/chronological evidence quoted here.

## 10. Endpoint and saved checkpoint

At the end of220, the Bai chief reopens the surveillance display and sees Fang Yuan holding the two children as hostages. The children are alive at this endpoint. No inheritance Gu has yet been recovered, and no escape has been completed.

Fang Yuan is **Revolution One upper stage with90% aptitude**. Bai remains **Revolution Three peak stage**, with one crocodile’s strength and Iron Thorn Bramble as a newly established vital Gu. Loans, retained ownership and uncertain returns remain separately tracked.

Saved: this report, numbered Chinese reading copy, four jotters, full working notes, ten-segment reading log, **103 proposals, {len(events)} chronology entries, {len(resources)} resource entries**, the updated workbook and ten working standards.

**Next actual unread:221.** All edits remain held. Whole-novel consistency is not certified. No source, site or repository changes, uploads or deletions occurred. Remote research backup remains unresolved.
'''
(r/'Batch-010-Chapters-201-220.md').write_text(report)
# Workbook: cumulative evidence, not merely this batch's summary.
w=Workbook();s=w.active;s.title='Read me'
for row in [('Sequential reading','Status'),('Completed','1–220 of2334'),('Next unread',221),('Paragraphs',tp),('Characters',tc),('Proposals','103total;102activeheld;1withdrawn'),('New copy candidates',len(fixes)),('Whole-novel certified','No'),('Approved/applied changes','None'),('Scope','Chinese-source review, not a fresh live-site audit'),('Timeline','Overwritten and revised histories must not be counted twice'),('Source note','Original EPUB unchanged; Ch151,201p22,203p84 restricted excerpts retain omission markers')]:s.append(row)
s=w.create_sheet('Chapter completion');s.append(['Chapter','File','Title','Paragraphs','Status'])
for c in cs:s.append([c['corpus_chapter'],c['file'],c['title'],len(c['paragraphs']),'Full sequential read'])
def ids(vs,k,c,p):return ', '.join(v['id'] for v in vs if any(z['chapter']==c and z['paragraph']==p for z in v[k]))
s=w.create_sheet('Paragraph register');s.append(['Chapter','Paragraph','Status','Proposal IDs','Chronology IDs','Resource IDs','Chinese reading copy'])
for c in cs:
 for p,t in enumerate(c['paragraphs'],1):
  n=c['corpus_chapter'];s.append([n,p,'Read; not certified error-free',ids(ledger,'locations',n,p),ids(events,'sources',n,p),ids(resources,'sources',n,p),reader(n,p,t)])
s=w.create_sheet('Proposal ledger');s.append(['ID','Category','References','Action','Rationale','Confidence','Status','Dependencies','Publication','Approved','Applied','Latest update'])
for v in ledger:
 latest=next((v[f'batch{i:03d}_update'] for i in range(10,0,-1) if f'batch{i:03d}_update' in v),'')
 s.append([v['id'],v['category'],refs(v['locations']),v['recommended_action'],v.get('rationale','See original report'),v.get('confidence','See originating report'),v['status'],v['dependencies'],v['publication_status'],'No','No',latest])
s=w.create_sheet('Copy items');s.append(['Proposal','Chapter','Paragraph','Original','Candidate','Confidence','Status'])
for v in ledger:
 for z in v.get('copy_items',[]):s.append([v['id'],z['chapter'],z['paragraph'],z['original'],z['replacement'],z.get('confidence','See originating proposal'),z['status']])
s=w.create_sheet('Chronology');s.append(['ID','Timeline','Kind','Event','Constraint','Evidence status','Dependencies','References','Exact source'])
for v in events:s.append([v['id'],v['timeline'],v['kind'],v['event'],v['constraint'],v['evidence_status'],v['dependency'],refs(v['sources']),'\n'.join(reader(z['chapter'],z['paragraph'],z['text']) for z in v['sources'])])
s=w.create_sheet('Resources');s.append(['ID','Kind','Entry','Calculation','Status','References','Exact source'])
for v in resources:s.append([v['id'],v['kind'],v['entry'],v['calculation'],v['status'],refs(v['sources']),'\n'.join(reader(z['chapter'],z['paragraph'],z['text']) for z in v['sources'])])
s=w.create_sheet('Working standards');s.append(['ID','Topic','Status','Rule','Evidence','Reopening trigger','Latest qualification'])
for v in stand['entries']:s.append([v['id'],v['topic'],v['status'],v['rule'],v['evidence'],v['reopen'],v['batch010_update']])
for s in w:
 s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
 for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='163B4B')
 for row in s.iter_rows(min_row=2):
  for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
 for col in s.columns:s.column_dimensions[col[0].column_letter].width=min(65,max(14,max(len(str(c.value or '')) for c in col[:30])+2))
w.save(r/'Sequential-Reading-and-Proposals.xlsx')
code=(r.parent/'make_whole_novel_audit_docx.py').read_text().replace("root=Path(__file__).parent","root=Path('/home/user/editorial/sequential-review')").replace('Whole-Novel-Aperture-and-Travel-Audit','Batch-010-Chapters-201-220').replace('WHOLE-NOVEL SOURCE AUDIT','SEQUENTIAL REVIEW / CHAPTERS201–220').replace('The Omniarch: Aperture and Travel Source Audit','The Omniarch: Sequential Review, Chapters201–220').replace('Full-corpus topical screening with selective contextual verification','Twenty complete chapters; personal clocks and inventory scope distinguished')
exec(compile(code,'docx-export','exec'),{})
(r/'CURRENT-STATUS.md').write_text(f'''# Sequential review checkpoint

Completed **1–220 of2334**. Next unread **221**. Remaining **2114**. Suggested next batch221–240, adaptive to density.

Batch010: {np}paragraphs/{nc}characters. Cumulative: {tp}paragraphs/{tc}characters.103proposalrecords:102activeheld,1withdrawn. SR103contains{len(fixes)}copycandidates.160chronology entries,97resources. No approved/applied changes. Whole-novel chronology NOTcertified.

Latest deliverable: Batch-010-Chapters-201-220.md/.docx. Numbered source, four010jotters, batch-010-working.md, readinglog, registers, workbook and workingstandards updated. Readercopy marks201p22/203p84 omissions;sourceEPUBunchanged. Corpus201–220=VolIICh2–21.

Carry forward:
-202inventory12 EXACT: Lotus,Cicada,FourFlavour,HiddenScale,Plunder,Yang,Skull,Moon,Ear,Tushita,Canopy,Centipede. Borrowed-outGuowned.201massacrelootstarvation separatefrom212retrospectivedestructionofYaoLeLiquorWorm+otherlootwhenTiearrived. FourFlavourRevolution2only stillheld;notreverserefined.
-211ThirdWatchalreadyconsumed:3xpersonalclock,cultivation/recovery/poison/lifeexternaltime1/3. Not3xcapacityor+300%. Chenplans3daywaitbutleavesafter1externalday. NoEarcutting/installation:rotationcancelled. HiddenScaleactuallyloaned.
-SR095Tieabout7Gu versus9namedtypes/manydeployedcopies scopequery. SR096182paternaluncle vs211maternaluncle/paternalniece unresolved. SR097BaiMoTingRevolution2scoutvsRevolution3banquet,unreportedpromotionpossible. SR098214Master肉骨vs骨肉preferredlatterconditional. SR099214afterRevolution3forecastpossiblebeforeLOWconfidenceawaitmechanism. SR100seven-dayhunt versusday9mountain,privateextensionpossiblebutunconfirmed. SR101220blastcasualties4present=>AoTian+1immediatedead,2wounded1laterdead,DaoKualone. SR102Chen350pacesightGuunlocatedinloot,don'tinventloss. SR10387copyitems.
-211Founderfightrecordingnotrebirth;CraneKingfallsnotexplicitdeath;SkyCraneSoulIncubatingFleaalreadybought,20yearsuccessforecastnotelapsed. TieRuoNanhome,5/6daydeathanchor202,homescentimingnotexact. Citywallsthousandli/tower100zhangsourcefantasynotreduced.
-Travel:206sixdayillness,207>halfmonthellipsis+5scoutingdays,209woundpause+traildays,212>20daysafterChendeath,213fivepitsinitialday+3more,2145daysafterwine. Overlaprecapsnotalladditive;216monthssinceQingMaoplausible. Twoof5pitsexplode;remainingtrapsnotnecessarilyremoved,220causalityconfirmed.
-204BackCarapace/CrocStrengthacquired,205BaiassignedStrength/Lotusrecalled;Bai1crocstrengthcomplete213. FYdoubleboarcan’taddwithoutboneGu. AshGuNOTacquired;LavaArmourdestroyed,otherwildGuabsent. 2115liveGu: RiceBag,IronThorn,MotherPotato,HairyFoot,JumpingGrass;Bookwormdead,ThirdWatchconsumed.212IronThornBaiowned/vital. IceMuscleeffect209persistencesupportnotSR060disposalproof.
-212FYRevolution1middle,217Relicdelivered/consumed→upper.90aptitudeconfirmed202/212;SR08542/43chainunchangedheldfinal90. StoneApertureSR090stillunresolved.
-217ClearHeatRevolution2detox acquired,218lentBaiLianreturnNOTnarrated. Noassumptiongeneralwoundheal. Giftstonesunquantified. LoanedCanopy/Centipede/Moon stillFYowned. BoneFleshReunion/YadeBone/etcNOTacquired220.
-215aliasFangZheng/clanresidualpartyfabricationintentional. FemalebaiChief no universalpatrimonyproofSR082. Twins18informerlife/10yearmigrationnotcurrent8proved. Huntd1bearRelic,d2rhinoHeat,d3stonesloan,d5banquet,d6challenge1,d7challenge2,d8challenge3,d9mountain.219three时辰≈6hescortedprogressvs6monthssolo differentforces.
-220endpointFYholdsbothchildrenaliveashostages,chiefseesafterturningsmokeoffandon. Noescape,noinheritanceacquisitionyet.Next221.

No source/site/repository edits, uploads or deletions. Remote backup unresolved. Earliergeneratorrerunsunsafe;010rerunsfrompre010history. AlloriginalEPUBparagraphswerefullyread;readeromissionsnotunreadgaps.
''')
assert len(cs)==220 and len(ledger)==103 and len(new)==9
assert (np,nc,tp,tc)==(1934,65084,18787,671746)
assert all(not v['approved'] and not v['applied'] for v in ledger)
assert all(z['text']==src(z['chapter'],z['paragraph'])['text'] for v in events+resources for z in v['sources'])
texts={t for c in cs for t in c['paragraphs']}
for line in report.splitlines():
 if line.startswith('> '):assert line[2:] in texts
print('QA PASSED',dict(chapters=220,paragraphs=tp,characters=tc,proposals=len(ledger),new_records=len(new),copy_items=len(fixes),chronology=len(events),resources=len(resources),next=221))
