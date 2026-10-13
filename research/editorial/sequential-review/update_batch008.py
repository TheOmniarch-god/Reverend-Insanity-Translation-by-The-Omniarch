from pathlib import Path
import json,shutil,re
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
r=Path('/home/user/editorial/sequential-review')
def load(n):return json.loads((r/n).read_text())
def save(n,v):(r/n).write_text(json.dumps(v,ensure_ascii=False,indent=2))
for n in ['reading-progress.json','proposal-ledger.json','chronology-register.json','resource-register.json']:
 d=r/'history'/('before-batch008-'+n)
 if not d.exists():shutil.copy2(r/n,d)
chapters=sum([load(f'batch-{i:03d}-source.json') for i in range(1,9)],[]);by={c['corpus_chapter']:c for c in chapters};batch=chapters[160:180]
def src(c,p):return dict(chapter=c,paragraph=p,file=by[c]['file'],text=by[c]['paragraphs'][p-1])
def ss(rr):return [src(*x) for x in rr]
def loc(rr):return [{k:v for k,v in z.items() if k!='text'} for z in ss(rr)]
def refs(z):return '; '.join(f"Ch{x['chapter']} p{x['paragraph']}" for x in z)
def reader(c,p,t):
 if c==151 and (31<=p<=67 or p in [98,99]):return '[Passage omitted from the reader-facing copy: sexualized material involving a minor. Original paragraph numbering retained; original EPUB unchanged. See the non-graphic editorial discussion.]'
 return t
ledger=load('history/before-batch008-proposal-ledger.json');new=[]
updates={'SR-001':'Ch162p56 and167p78 retain42% capacity after159’s explicit loss.173pp15–16 independently show aptitude changing99%→TenExtreme100%. No universal fixed-for-life aptitude assumption.', 'SR-007':'167less-than3percentage points→42refill in8–9minutes;168forty points remaining follows expenditure, not a new40% capacity.', 'SR-037':'Ch175p62 explicitly identifies the formerhunter/demonicGuMaster asWangDa, reinforcing the160p62 correction.', 'SR-047':'Still withdrawn. No Gregorian mapping reinstated.162late-summer wording remains another source calendar anchor, not proof of the earlier categorical diagnosis.', 'SR-050':'180is a recording of the original awakening, not a new ceremony that can repair the113YaoLe timing query.', 'SR-053':'No final former-life destruction date established in161–180. Do not replace the original numerical conflict with a forecast of current events.', 'SR-067':'180p68 again associates the original awakening withage15; this is retrospective/hypothetical, not a claim thatFangYuan is currently15.'}
for v in ledger:
 if v['id'] in updates:v['batch008_update']=updates[v['id']]
def prop(i,cat,rr,action,reason,conf,dep):
 v=dict(id=f'SR-{i:03d}',category=cat,locations=loc(rr),recommended_action=action,rationale=reason,confidence=conf,status='Active provisional proposal; no approval',dependencies=dep,approved=False,applied=False,publication_status='HOLD: full2334-chapter reading and global reconciliation outstanding');ledger.append(v);new.append(v);return v
prop(74,'wolf nomenclature',[(163,77),(164,5),(164,21)],'Ch164p5 雷光头狼→雷冠头狼.','The continuing assault targets the same ThunderCrownWolfKing, consistently named before and after this line.','High.','Local substitution only; retain other legitimate lightning-related names.')
prop(75,'Gu nomenclature',[(167,64),(167,68),(167,70),(168,46),(168,57)],'Ch167p70 雷壳蛊→雷盾蛊.','The capturedGu is namedThunderShield immediately before and on its next use. No secondGu capture occurs.','High.','Do not confuse shell-like appearance with a differentGu species.')
prop(76,'elder identity',[(155,24),(165,71),(171,13)],'Preferred Ch165p71 药钟→赤钟, if the retained official is the already-established ChiZhong.','The head of medicine isChiZhong;165leavesYaoZhong andGeYan behind;171listsChiZhong andGeYan among surviving hosts. Strong context supports a name slip.','High contextual candidate, not independently witnessed.','A distinct elderYaoZhong is not formally excluded; check later cast references before applying.')
prop(77,'clan branch identity',[(172,62),(172,70),(172,72),(172,77),(172,83)],'Ch172p70 赤脉大院→漠脉大院; p77赤脉正统→漠脉正统.','The scene isMoChen’s courtyard and the succession crisis of theMo branch, not theChi branch.','High.','No globalMo/Chi replacement. PreserveChi references elsewhere.')
prop(78,'speaker accusation / twin identity',[(175,62),(175,67),(175,72),(176,1),(176,2)],'Ch175p67 方正滥杀无辜→方源滥杀无辜.','TieRuoNan saysFangZheng was mistaken for the killer, yet one clause names him as the killer. The same account andFangYuan’s subsequent admission resolve the identity.','High.','Repair this clause only; the subsequentFangZheng victim references are correct.')
prop(79,'dialogue attribution',[(178,50),(178,51),(178,52)],'Ch178p51 铁若男语气沧桑而又悲凉→铁血冷语气沧桑而又悲凉.','The speaker addresses “my child”; the next line answers “Yes, Father.” It isTieXueLeng’s response toTieRuoNan.','High.','Do not change the speech content or the daughter’s contrasting position.')
prop(80,'attack distance / staging',[(168,20),(168,34),(168,37),(168,38),(168,42)],'Conditional minimal candidate at168p42: 跨越百里距离→穿过浓烟. English: tore through the dense smoke. Preserve the separate stated hundred-li smoke extent pending global comparison.','The strike is described as crossing100li immediately afterFangYuan begins escaping the same battlefield. No corresponding travel interval is supplied. The text does not establish that he has moved100li away.','High staging query; moderate repair.','Do not assert thatRevolutionFour lightning cannot have a100li range. Do not globally shrinkli, change smoke dimensions automatically, or invent a flight speed.')
prop(81,'reconnaissance range',[(128,31),(129,33),(129,35),(138,62),(169,56),(169,61),(169,63)],'If169p56 intends the established fully rooted maximum, 两百步→三百步. If it instead gives an effective working range, label that limitation explicitly rather than presenting it as an unexplained new maximum.','Earlier maximum is300paces withgroundcontact; moving unrooted wasbelowhalf atRevolutionTwo.169gives200before describing a separate rooting action. CurrentRevolutionThree cultivation and combat conditions prevent a categorical numerical repair.','High need to distinguish modes; moderate-to-low replacement certainty.','LaterGu upgrade/essence scaling, root contact and environmental effects. Do not impose300 as an unconditional range in every scene.')
prop(82,'succession rule scope',[(152,75),(172,76),(172,77),(172,83)],'Investigate whether172’s male-line rule concerns hereditary clan property, not all offices or political succession. If the rule is specificallyMo practice, candidate: 漠脉父系家业历来传男不传女. Otherwise retain the source and add a clearly provisional explanatory note.','YaoLe is explicitly being raised as theYao-branch successor;172broadly states propertypassesonlytomales. Offices, independently established property and inheritedpatrimony may differ, but that distinction is not yet fully proved.','Moderate scope tension; low confidence in any universal rewrite.','Full-novel inheritance rules andYao-branch history. Do not erase femaleleadership or invent a prohibition applying to everyclan.')
prop(83,'expedition clock / translation safeguard',[(162,3),(162,18),(162,73),(162,82),(163,67),(163,73)],'Translate163p73 as “It was dawn,” not “the next morning.” If later reconciliation requires an explicit day boundary, clarify that a further day has elapsed rather than silently halving时辰. No change to the current Chinese is necessary on present evidence.','Three时辰≈6hours, six-plus时辰>12hours, plus the30minute deadend and unquantified retracing already exceed one ordinary overnight. The source does not explicitly say nextmorning.','High unit distinction; actual return-day unresolved.','Preserve小时 versus时辰. Additionalreturntime could place this at a laterdawn; do not invent an exact36-hour journey.')
fixes=[(161,2,'不再此列','不在此列'),(161,6,'下人的感激','小人的感激'),(161,8,'发生是的','发生似的'),(161,14,'站起来身','站起身'),(161,28,'出于组织高层','处于组织高层'),(161,29,'心甘如怡','心甘如饴'),(161,31,'被偷被盗白抢','被偷被盗被抢'),(161,95,'物质榜','物资榜'),(162,20,'迈入进入其中','迈入其中'),(164,14,'猛地一口气','猛地吸一口气'),(164,107,'倒时就有','到时就有'),(165,71,'顺着顺着','顺着'),(166,4,'终身后跃','纵身后跃'),(166,29,'面临惊异之色','面露惊异之色'),(166,44,'灌注到着锯齿','灌注到这锯齿'),(166,78,'向抛飞','向后抛飞'),(166,96,'黑风狂舞','黑发狂舞'),(167,9,'风暴卷席而来','风暴席卷而来'),(167,61,'电流催敌','电流摧敌'),(167,65,'能成为一张','能形成一张'),(168,6,'强制镇定','强自镇定'),(168,29,'眨眼睛','眨眼间'),(168,74,'侦察氛围','侦察范围'),(168,93,'不属于人的智慧','不输于人的智慧'),(169,15,'却和他的心意','却合他的心意'),(170,10,'拘泥与','拘泥于'),(170,14,'另一家势必被沦为','另一家势必沦为'),(170,14,'修养生息','休养生息'),(171,16,'密室中修养','密室中休养'),(171,24,'那里修养','那里休养'),(171,57,'想要我要解释','想要我解释'),(172,18,'剥削家老','剥夺家老'),(172,32,'方源可以比敌','方源可以匹敌'),(172,39,'语出惊人死不休','语不惊人死不休'),(172,53,'这一次也漠脉也','这一次和漠脉也'),(173,25,'站起来身来','站起身来'),(173,55,'双人皆','二人皆'),(174,25,'任何的留念','任何的留恋'),(174,28,'落入二阶','落入二转'),(174,46,'开馆的家奴','开棺的家奴'),(174,54,'他身前','他生前'),(174,58,'所以的服饰','所有的服饰'),(174,58,'齐根而段','齐根而断'),(174,58,'陷阱所制','陷阱所致'),(175,31,'漠尘家可','漠尘家老可'),(175,37,'不在受他掌控','不再受他掌控'),(175,38,'这是以药脉','这是药脉'),(175,41,'将个人阵营','将各人阵营'),(175,62,'机源巧合','机缘巧合'),(175,62,'亲身哥哥','亲生哥哥'),(175,65,'亲身弟弟','亲生弟弟'),(175,80,'弟弟的职责','弟弟的指责'),(176,7,'倾述','倾诉'),(176,26,'道歉道；','道歉道：'),(177,2,'才孔洞中','从孔洞中'),(177,13,'但靠锯齿','单靠锯齿'),(177,15,'全方外','全方位'),(178,11,'有夹杂着','又夹杂着'),(179,1,'不妨他突然','不防他突然'),(179,5,'就从迎了上去','就迎了上去'),(179,43,'芊芊手指','纤纤手指'),(179,64,'你你要','你要'),(180,63,'传入众人耳帘','传入众人耳畔'),(180,73,'直觉得自己','只觉得自己')]
for c,p,a,b in fixes:assert a in src(c,p)['text'],(c,p,a)
v=prop(84,'source-copy bundle',[(c,p) for c,p,a,b in fixes],'Apply only the individually located substitutions in copy_items after approval.','Miskeys, duplicatecharacters and malformedphrases; localcopyrepairs, notblanketwordreplacement.','High contextual confidence overall; candidates remainheld.','Independent witnesses notchecked; aphorism wording and near-synonyms remaineditorialcandidates rather than authenticatedcorrections.')
v['copy_items']=[dict(chapter=c,paragraph=p,original=a,replacement=b,status='Held; unapproved/unapplied') for c,p,a,b in fixes]
save('proposal-ledger.json',ledger)
tdoc=load('history/before-batch008-chronology-register.json');events=tdoc['entries'];te=len(events)
def event(timeline,kind,title,constraint,status,rr,dep='Full-corpus reconciliation outstanding.'):
 events.append(dict(id=f'T{len(events)+1:03d}',timeline=timeline,kind=kind,event=title,constraint=constraint,evidence_status=status,dependency=dep,sources=ss(rr)))
event('current life','banquet / residence','Gift collection and elder housing','Samebanquetends;newclan-allottedhousealreadyavailable;>10daysuntilnextcavevisit.','Notoldrental;carryingservantsnotnewlyownedstaff.',[(161,13),(161,73),(161,74),(161,75),(162,2)])
event('current life / prior-life forecast','late-summer anchor','ThunderCrown schedule','Departurelate-summernight;oldlifekingarrived3daysearlier;actualarrivalatreturn.','Explicitbranchschedulechange; noexactAugustday inferred.',[(149,41),(162,3),(162,4),(162,8),(163,73),(163,77)])
event('current life','hours and day boundary','Long cave expedition','3shichen+6-plusshichen+30minutes+retrace/otheractivity;returnatdawn.','Cannotcallthisimmediatelynextmorningwithoutresolving>18.5hours ofstatedoutwardwork.',[(162,18),(162,73),(162,82),(163,67),(163,73)],'SR-083; no inventedexactday.')
event('historical reconstruction / clan tradition','founder and inheritance','Near-millennium history','163FangYuanreconstructsfounder/FlowerWine/Lotus;165clanbelievesfounderdeadnear1000years,lastcrisis200yearsago.','Reconstruction/traditionnotindependentlyverifiedhistoricaltruth.',[(163,1),(163,2),(163,7),(163,16),(165,5),(165,6)])
event('current life','battle clocks and casualties','ThunderCrown and chain-array','15minbattle→sixeldersdead→anotherkilled;chainarray5–6min;reinforcementsalterparticipantcount.','No fixedneartenparticipantdenominator;ritualdurationnotquantified.',[(164,66),(164,72),(164,78),(164,93),(165,41)])
event('current life','pursuit and replenishment','Valley battle','30minpursuit;less3points→42in8–9min;fiveeldersdeadobservedafterinterval.','FiveobservedmayincludeearlierBai elderdeath;donotaddsixwithoutproof.',[(165,78),(166,65),(167,78),(167,83),(167,84),(167,85)])
event('current life / corrected identification','wolf versus lycan','Apparent fourth king','167suspectsfourth;168identifiesLycanandmistakenfrontleginjuryassumption.','No confirmedfourthThunderCrown. Xiongdestructionassumptionlaterreversed175.',[(167,109),(168,11),(168,14),(168,18),(175,10),(175,11)])
event('forecast corrected by event','investigator arrival','TieXueLeng earlierthanexpected','113year-after-nextforecast;170actualarrival;172explicitlycallsitunexpectedlyearly.','Predictionnotbindingelapsedclock; do notsilentlyrepairforecast.',[(113,60),(170,42),(172,30)])
event('current life','scope of survivors','Battlefield elders and remaining clan','170eldersatfielddead;171survivingtable<10,MoCheninjured2weeksearlier,ChiLianconfirmeddead.','Localbattlefieldscope; notallclaneldersdead.',[(170,55),(171,13),(171,14),(171,16),(171,17)])
event('current life / prior-life comparison','self-injury and false alibis','YaoJi and MoChen','YaoJiself-amputationmatchesformerlife;MoChenclaimsdruggedFYforlovingMoYan.','MoChenstoryfalse;176MoYandislikeexplicit. Punishmentnotfinalised.',[(171,54),(171,55),(172,18),(172,46),(172,57),(176,71)])
event('current life','aptitude and physiology','Bai99% becomesTenExtreme','173past99GradeA→laterTenExtreme;178realrightarmregeneration.','Distinctfromearlierfracturefeint/actualamputation; nobirthdateinferred.',[(173,15),(173,16),(178,75),(178,76)])
event('current life','investigation clocks','Nextday inquiry and village search','174nextmorningaftertour;halfdayfailedvillagesearch;175laterJiangHeevidence.','Failedsearchnotallleadsclosed; novicehypothesesremainhypotheses.',[(174,36),(174,66),(175,62),(175,72)])
event('current life / forecasts','main tide versus residual wolves','Wolf crisis ends gradually','Mainassaultdefeated,remainingmonthsminorwolves;year-enddisappearanceforecast.','Do not equateendofworstcrisiswithlastwolfvanishingorcontradict>=1yeardurationautomatically.',[(171,22),(171,23),(172,8),(124,38)])
event('current life','Gu healing clocks','PigIron treatment','Halfshichen≈1hour;remaining5–6scratcheshealweeks;untreatedteethwouldtake>halfyear.','SpecificGuhealing,notFYbodyhealingorLightningWingsrepair.',[(177,8),(177,11),(177,12),(177,13)])
event('previous life / present pressure','Cicada','Advancement duration and aperture danger','400-plusyearswithin500-yearformerlife toRevolution6;currentCicadarecoveryaccelerates.','Do notadd400+500 asdisjointlifespans; noboundeddeadlinegivennumerically.',[(177,22),(177,26),(177,28),(177,36)])
event('current life / recorded past','tournament and replay','Dayaftertomorrow / originalawakening','179tournamentdayaftertomorrow;180sameeveningbargainandrecordingoforiginalceremony.','No newawakening; tournamentNOTyetheld;Tie’slast-linehypothesisnotyetrevealed.',[(179,70),(179,79),(180,2),(180,31),(180,47),(180,91)])
tdoc.update(scope='Sequential chronology throughChapter180; sourceclaims ranked by provenance and evidence, full-novelchronologynotcertified.',entries=events);save('chronology-register.json',tdoc)
rdoc=load('history/before-batch008-resource-register.json');resources=rdoc['entries'];re0=len(resources)
def res(kind,entry,calc,status,rr):resources.append(dict(id=f'R{len(resources)+1:03d}',kind=kind,entry=entry,calculation=calc,status=status,sources=ss(rr)))
res('banquet gifts / storage','Near10000stonesreceived;Tushita30000maximum,15000workingallocationwithothergoods.','10000generalminimumtravelreserveisguidance,notthirdactualreceipt.','Newloan3000alreadyearlier; noexactrunningbalanceclaimed.',[(161,13),(161,26),(161,82),(161,87),(161,89)])
res('capturedGu','ChainsawGoldenCentipedeactuallyrefined162.','Wildwillreplaced,losescommandoverordinarycentipedes.','Captureandcontrolconfirmed; predictednewwildking10+yearsisnotelapsedtime.',[(162,65),(162,66),(162,69),(162,70)])
res('Lotus prospective acquisition','Revolution3Lotusobserved8.5leaves;need9;500+trialstones.','Estimated50000+required,10000+funds,approx40000gap.','Approximationnotexactmismatch; nine-daymaturationgenericrecipeknowledge; NOTownedby180.',[(163,31),(163,34),(163,50),(163,52),(163,54),(180,15)])
res('Lotus rank progression','Springscost1atRevolution3,then7,9,11atRevolution4/5/6.','Incrementalstages;notallone-stagecost.','FangYuanhaspartialknowledge,nofullrecipe;Revolution6ordinal6inrankingnotcultivation6thstage.',[(163,18),(163,21),(163,24),(163,32)])
res('BloodGu swarm','Hundredsinitial→dozensfoldgrowth→tensofthousands.','DoublingreproductionfromGuMasterblood;beastbloodfeedwithoutreproduction.','Rangescanoverlap;donotclaimanexactexponentialgenerationcount.',[(165,11),(165,25),(165,29),(165,34),(165,36)])
res('loot and loss','ThunderShieldcaught167,dies168;LightningWingscriticallyinjured,Canopyhurt.','Shortownershipwindow; noduplicateShieldorautomaticwingdeath.','OnlyShieldlootconfirmed;otherescaped8–9Gu notautomaticallyowned.',[(167,68),(167,71),(167,91),(168,57),(170,25)])
res('essence and range','42capacity;below3remaining→42refill8–9min;40afterattack.','40notnewcapacity;ElectricEye50normal/30smoke;ear200queriedagainst300rooted.','FlyingwingsnotconcealedbyRevolution2HiddenScale;terrainhidesfromEye.',[(165,74),(167,75),(167,76),(167,78),(167,83),(168,69),(169,51),(169,56)])
res('snake status','WhiteFormImmortalSnakeRevolution5residesinBaiaperture.','Notrefined,notinheritanceapproval;automaticprotectivelimits.','173leader’srecognitionguesscorrected178. Newrightarmfromphysique,notSnakehealing.',[(173,78),(173,84),(178,76),(178,77),(178,81),(178,82)])
res('marriage negotiations','100000+10rareRevolution3minimumGuasked;30%reduction;40000advance.','Cashwouldbe70000ifdiscountappliestothatprice;Guquantitydiscountunspecified.','Notfinalmarriageorfullsettlement;40000actuallyreceivedby180,notadditionaltoagreedtotal.',[(174,20),(176,58),(176,59),(176,61),(180,6)])
res('consumables and repair','PigIronreceivedandconsumed;GoldenCentipederestoredmostly.','Revolution2consumable;healinggrassnotdelivered.','Do notcountusedPigIronasretainedinventory. Own-bodyhealinggapremains.',[(177,1),(177,8),(177,9),(177,11),(180,6),(180,7)])
res('evidence-generatingGu','DivinePointerRevolution3consumed179;MemoryGucrushed180toprojectrecording.','InferenceGuusesknownclues,notomniscientanswer. Replaypast≠presentevent.','ChiChengcheatcorroborated;LiquorWormblackmailoriginstillfalseinference.',[(179,34),(179,44),(179,48),(179,55),(180,43),(180,47),(180,60)])
res('funds versus carryingcapacity','40000advancearrivedininstalments;earlier10000+held.','AggregatefundsmayexceedTushita30000absolute/15000practicalcapacity.','No statementyetthatallcashisinsideflower. Carryforwardphysicaltransportquery,notcontradictionestablishedat180.',[(161,82),(163,54),(180,6)])
rdoc.update(scope='Resources throughChapter180; acquisitions,loans,forecasts,storageandcontrolseparatelylabelled.',entries=resources);save('resource-register.json',rdoc)
np=sum(len(c['paragraphs']) for c in batch);nc=sum(sum(map(len,c['paragraphs'])) for c in batch);tp=sum(len(c['paragraphs']) for c in chapters);tc=sum(sum(map(len,c['paragraphs'])) for c in chapters)
p=load('history/before-batch008-reading-progress.json');p.update(sequential_chapters_completed=list(range(1,181)),paragraphs_read=tp,paragraph_characters_read=tc,next_chapter=181,remaining_chapters=2154,batch_report='Batch-008-Chapters-161-180.md',scope='Full sequential Chinese-source reading throughChapter180. EPUB counts include paratext; notfreshbilingual/live-siteaudit.',whole_novel_chronology_verified=False,chapters=[dict(chapter=c['corpus_chapter'],file=c['file'],title=c['title'],paragraphs=len(c['paragraphs']),first_paragraph=1,last_paragraph=len(c['paragraphs'])) for c in chapters]);p['batches'].append(dict(id='008',chapters=[161,180],paragraphs=np,characters=nc,reading_segments='Ten complete untruncated two-chapter outputs'));save('reading-progress.json',p)
save('batch-008-reading-log.json',dict(status='Complete',endpoint=180,next_unread=181,segments=[dict(chapters=[c,c+1],status='Full paragraph sequences printed and read; untruncated',paragraphs=sum(len(by[k]['paragraphs']) for k in [c,c+1])) for c in range(161,181,2)],note='Extractionalone notcompletion; no inventedtimestamps.'))
lines=['# Numbered Chinese source: Chapters161–180','Unemended EPUB paragraphs, numbered to match the review. No source substitutions applied.','']
for c in batch:
 lines += [f"## Chapter {c['corpus_chapter']}: {c['title']}",f"Source: `{c['file']}`",'']
 for i,t in enumerate(c['paragraphs'],1):lines += [f'**{i:03d}.** {t}','']
(r/'Batch-008-Numbered-Chinese-Source.md').write_text('\n'.join(lines))
for name,vs in [('Chronology',events[te:]),('Resources',resources[re0:])]:
 lines=[f'# Batch008: {name} jotter','Full sequential read throughChapter180; proposalsheld.','']
 for v in vs:
  lines += [f"## {v['id']}: {v.get('event',v['kind'])}"]
  for k in ['timeline','constraint','evidence_status','dependency','entry','calculation','status']:
   if k in v:lines += [f"**{k}:** {v[k]}"]
  for s in v['sources']:lines += [f"\n**Ch{s['chapter']} p{s['paragraph']}**",'> '+reader(s['chapter'],s['paragraph'],s['text']),'']
 (r/'jotters'/f'Batch-008-{name}.md').write_text('\n'.join(lines))
lines=['# Batch008: Continuity and proposals',f'11newrecords;SR-084contains{len(fixes)}individualcopyitems. Allheld.','']
for v in new:
 lines += [f"## {v['id']}: {v['category']}"]
 for k in ['recommended_action','rationale','confidence','dependencies','publication_status']:lines += [f"**{k}:** {v[k]}"]
 for s in v['locations']:lines += [f"\n**Ch{s['chapter']} p{s['paragraph']}**",'> '+reader(s['chapter'],s['paragraph'],src(s['chapter'],s['paragraph'])['text'])]
 for z in v.get('copy_items',[]):lines += [f"- Ch{z['chapter']} p{z['paragraph']}: {z['original']} → {z['replacement']} (held)"]
 lines += ['']
lines += ['## Earlier records updated']
for k,v in updates.items():lines += [f'**{k}:** {v}','']
(r/'jotters'/'Batch-008-Continuity.md').write_text('\n'.join(lines))
(r/'jotters'/'Batch-008-Preserve.md').write_text('''# Batch008: Preserve, do not repair

- TheXiongclan’s apparent destruction is reversedby175. Earlierjudgments arebeliefs,notproofthatsurvivorsresurrected.
- TheapparentfourthThunderCrownWolf isidentifiedasLightningLycan. Itsforelegshapeisnotestablishedbattleinjury.
- TieXueLengarrivesearly;172explicitlyacknowledgesit. Earlierforecastdoesnotbecomeafixedappointment.
- Themainwolfcrisiscanendwhileminorwolvesremainformonths. Year-endforecastisnotalreadyelapsedduration.
- Battlefieldeldersdie170;otherssurviveatthe171banquet. Donottranslateawaythatlocalscope.
- MoChen’sMoYan-romance/druggingalibiisfalse. Herdislikeisexplicit176. ProposedremovalofFY’selderstatusisnotacompletedremoval.
- Bai99%→TenExtreme100%isexplicit;notallGradeAmustbecome100. FY42%remainscurrent,prior44%remainsvalid.
- WhiteFormSnakeinhabitationisnotrefinementornarrativeconfirmationofheirapproval.178explainswhyitentered. Actualrightarmregenerationisphysique-driven.
- Publichistory179isfalsifiedas180confirms. Do notharmoniseitwiththeearlierrecordingoftheFlowerWineconfrontation.
- DivinePointerworksfromknowncluesandcanproducefallibleinferences. CorrectidentificationofChiChengcheatingdoesnotprovetheguessedLiquorWormorigin.
- Sixstonesbought,fiveopenedonthespotmatchesCh40/41/43. TheLiquorWormfindisanalibi,notanactualsecondgamblingprize.
- FangZheng’sdrunkennessdescription177ishismistakeninterpretation. Itdoesnotmakeearlier买醉necessarilymeaninsensibledrunkenness.
- 3shichenand6-plusshichenareabout6and12-plushours,not3and6.163onlysaysdawn,nottheimmediatelynextmorning.
- Tushita’s30000absoluteand15000practicalstoneallocationsarenotcontradictory. CashownedisnotautomaticallycashpackedinsidetheGu.
- Lotus8.5leaves,9required,500-plusspent,50000-plusestimatedneed:roundfiguresarenotexactaccountingidentities. NOTacquiredby180.
- ThunderShieldacquired167dies168. LightningWingsisbadlyinjured,notdeclareddead. PigIronspecificallyhealsGoldenCentipede,notallGu.
- Contractaskingprice,30%concession,40000advanceandfullsettlementaredifferent.40kreceivedby180;healinggrassandmarriagenotcompleted.
- GoldenCentipede’snear2muncontractedlengthfitsgeneric1-plusmetrewording;1.5mcontractedisanexplicitmodechange.
- Electricwolves’weakhearing178doesnotbyitselfcontradicttheirearlierweaksmell.Bothcanbetrue.
- Chapter180’sawakeningisarecording;noadditionalpresent-dayceremony. FinalhypothesisandLotusattemptremainunfinished.
''')
standards=[
 dict(id='WS-01',topic='Evidence and viewpoint',status='Established review method; not manuscript approval',rule='Label narrated action, testimony, inference, official record, recording, forecast and previous-life memory separately. A later explicit correction can overturn an earlier belief.',evidence='168/175: Xiong destruction belief reversed.173/178: snake recognition inferred then denied.179/180: official history exposed as falsified.',reopen='An apparent narrator statement may itself be a copy error; examine wording and provenance rather than apply a mechanical hierarchy.'),
 dict(id='WS-02',topic='Aperture quantities',status='Strong working standard in covered mortal-aperture scenes',rule='Distinguish capacity, current reserves, essence quality and recovery rate. Use percentage points when the source measures changes against aperture capacity.',evidence='159:44→42 aptitude loss.162/167:42 remains the maximum.167/168: refill then40 remaining is expenditure, not another aptitude change.',reopen='Later cultivation systems, changed denominators, special transformations or explicit exceptions. No universal claim about every percentage in the novel.'),
 dict(id='WS-03',topic='Revolution, stage and aptitude',status='Strong working distinction; mechanics remain revisable',rule='Revolution, within-rank stage, refined essence quality and aptitude are separate. Aptitude is not necessarily fixed for life.',evidence='145:upper cultivation withpeak-quality essence.135/140:Bai suppression without permanent demotion.173:99%GradeA becomesTenExtreme100%.',reopen='A new method can change aptitude or rank. Do not pre-emptively change earlier valid values or give allGradeA charactersTenExtreme status.'),
 dict(id='WS-04',topic='Gu possession and control',status='Strong working standard',rule='Distinguish possession, refinement, borrowed permission, autonomous protection, recipe consumption, destruction and injury. Record acquisition and disposal events.',evidence='138:Bai heldPlunder unrefined.157:permission permits borrowedGu use.162:refinement removes wild insect command.173/178:snake residence is not refinement or heir approval.',reopen='LaterGu with special control rules. Do not invent consumption or disposal merely to force an inventory total.'),
 dict(id='WS-05',topic='Time units and interval scope',status='Strong working unit distinction; calendar unresolved',rule='Keep时辰 and小时 distinct; conventionally one时辰 is abouttwohours. Separate elapsed intervals, inclusive day labels, retrospective scenes and forecasts. Do not add an unexpressed next-day boundary.',evidence='61:1时辰 traversal versus30-minute expressions elsewhere.162/163:>18.5hours of outward activity before an unspecified dawn return.172:Tie arrival explicitly early.',reopen='Explicit in-world clock redefinition, tightly dated later evidence or source variants. Month/season mapping remains open.'),
 dict(id='WS-06',topic='Resources and storage',status='Established accounting method; balances locally incomplete',rule='Separate asking price, estimate, accepted contract, payment, debt, repayment, stock and carrying capacity. Approximate figures stay approximate.',evidence='161:near10000gifts,30000absolute storage versus15000mixed allocation.163:50000-plus estimate.176/180:40000advance later actually received.',reopen='Later receipts, expenditures, discounts, extra storage or explicit transport details. No invented running balance.'),
 dict(id='WS-07',topic='Ranges and physical staging',status='Conditional factual standards only',rule='Preserve the source units and identify attack/reconnaissance mode and environment. Do not globally shrink geography or silently equatepaces andmetres.',evidence='129:300paces requiresgroundcontact.169:Eye50normal/30smoke;Ear200 needs clarification.168:100li lightning path is an unresolved staging query.',reopen='Gu upgrades, cultivation effects, terrain, source witnesses or explicit distance/time anchors. Current disputed numbers remain held.'),
 dict(id='WS-08',topic='Ages and calendars',status='Open global reconciliation',rule='Separate cohort age, nominal/NewYear age, birthday age, academic year and elapsed year. Do not invent birthdays or impose a calendar.',evidence='159:FangYuan17.180:recorded awakening at15, not present age.116/145:conflictingJuly season labels. YaoLe age querySR-067 remains conditional.',reopen='Any later explicit date, birthday, age convention, cohort rule or calendar definition.'),
 dict(id='WS-09',topic='Copy corrections and textual provenance',status='Local confidence can be high; all proposals still held',rule='Exact local name/speaker slips may be strongly supported without settling global lore. Separate author-attributed notes, compiler speculation and narrative. Preserve original quotations and locations.',evidence='175:FangZheng/FangYuan clause.178:father’s speech misattributed to daughter.151:appended commentary outside narrative.',reopen='Authenticated source variant, deliberate wordplay or later evidence. EnglishGu labels in reports are working identifiers, not an approved global glossary.'),
 dict(id='WS-10',topic='Completion and approval gates',status='Standing workflow rule',rule='Stabilising a working standard does not end sequential reading. Whole-novel confirmation requires review through2334 and reconciliation of outstanding conflicts. No publication change is authorised merely by confidence or completed reading.',evidence='User requested full1–2334 coverage and reaffirmed thatonevolume is insufficient. Currentactual endpoint180; oldtopicalscreening is not sequential completion.',reopen='Only an explicit user change of scope or approval changes the work gate. Later contradictions may reopen working standards.')]
# Normalise reader-facing spacing without altering source quotations.
replacements={'RevolutionThree':'Revolution Three','RevolutionFour':'Revolution Four','GradeA':'Grade A','TenExtreme':'Ten Extreme','FangYuan':'Fang Yuan','BaiNingBing':'Bai Ning Bing','ChiCheng':'Chi Cheng','LiquorWorm':'Liquor Worm','DivinePointer':'Divine Pointer','YaoLe':'Yao Le','TieXueLeng':'Tie Xue Leng','PigIron':'Pig Iron','Plunder unrefined':'Plunder unrefined','NewYear':'New Year','GuMaster':'Gu Master','whole-novel-confirmed':'whole-novel-confirmed','twohours':'two hours','groundcontact':'ground contact','RevolutionTwo':'Revolution Two','onevolume':'one volume','full2334reading':'full 2334 reading','oldtopicalscreening':'old topical screening','Currentactual':'Current actual','allGradeA':'all Grade A','laterGu':'later Gu','specialcontrol':'special control','andmetres':'and metres','everyGu':'every Gu','Snake':'Snake'}
def readable(t):
 for a,b in replacements.items():t=t.replace(a,b)
 return re.sub(r'(?<=[A-Za-z])(?=\d)|(?<=\d)(?=[A-Za-z])',' ',t)
for v in standards:
 for k in ['topic','status','rule','evidence','reopen']:v[k]=readable(v[k])
save('working-standards-register.json',dict(coverage_through=180,whole_novel_confirmed=False,approval='No manuscript edits approved or applied',entries=standards))
lines=['# Working editorial standards','## Evidence checkpoint: Chapters1–180','These standards organise the review; they do not authorise manuscript changes or replace the remaining sequential reading. **No story-wide factual standard is labelled whole-novel-confirmed at this checkpoint.**','']
for v in standards:
 lines += [f"## {v['id']}: {v['topic']}",f"**Status:** {v['status']}",f"**Working rule:** {v['rule']}",f"**Evidence:** {v['evidence']}",f"**Reopening trigger:** {v['reopen']}",'']
lines += ['## Current non-negotiable distinctions','- Well-supported local evidence is not whole-novel certification.','- A source correction is not a translation preference, and neither is an approved adaptation.','- A current-life event, a previous-life memory and a recorded past event do not share one clock.','- Actual reading coverage is180/2334, with2154chapters remaining.']
(r/'WORKING-STANDARDS.md').write_text('\n'.join(lines))
report=f'''# Sequential close-reading review: Chapters161–180
## Batch008 | The Omniarch

**Actual completed coverage: Chapters1–180 of2,334. Next unread:181. Remaining:2,154.**

Twenty complete chapters were read in ten untruncated two-chapter sequences: **{np:,} paragraphs / {nc:,} source characters**. Cumulative coverage is **{tp:,} paragraphs / {tc:,} characters**. Extraction and earlier topical searches are not counted as sequential reading.

The ledger now contains **84 proposal records:83 active and held, one withdrawn**. Eleven records are new; SR-084 contains{len(fixes)} separately located copy-edit candidates. Nothing is approved or applied. This remains Chinese-source review, not a fresh bilingual or live-site audit.

## 1. We can stabilise working standards without stopping the whole-novel review

This batch supports a distinction between:

- **Established review methods:** separating claims from events, keeping precise source locations, distinguishing proposals from approvals, and not turning estimates into payments.
- **Strong working story mechanics:** the42% capacity after the explicit loss, distinct rank/stage/essence quality, and possession versus refinement.
- **Unresolved global rules:** calendar and seasonal mapping, some ages, range scaling, inheritance law and historical chronology.

The new **WORKING-STANDARDS.md** and matching JSON register record ten standards with evidence and reopening triggers. None is presented as whole-novel factual certification. A well-supported standard can guide the current work while later evidence remains capable of changing it.

Completing one volume would not fulfil the requested2334-chapter review. Nor will finishing the reading automatically authorise applying the proposals.

## 2. Three apparent facts are explicitly reversed or qualified

### The Xiong clan was not simply annihilated

The newly encountered wolf king leads the characters to infer that Xiong Village has fallen. Chapter175 then supplies survivors and an explanation:

> {src(175,10)['text']}

> {src(175,11)['text']}

Preserve the earlier inference as an inference. Do not change every earlier “destroyed” statement into a contradiction or treat the survivors as a resurrection.

The same sequence first suggests a fourth Thunder Crown Wolf. Chapter168 identifies a Lightning Lycan instead. Its short forelegs are its anatomy, not a confirmed injury merely because an observer initially interpreted them that way.

### The white snake has not accepted Bai Ning Bing as an heir

Chapter173’s clan leader infers recognition when the snake enters Bai’s aperture. Chapter178 explains the actual relationship:

> {src(178,81)['text']}

> {src(178,82)['text']}

It is attracted to his physique, lives in his aperture, and protects him imperfectly. It is not refined, and its arrival does not confirm inheritance approval. This strongly supports the possession/control distinction in the working standards.

Bai’s right arm also genuinely regrows in178 through his developed physique. This is distinct from the earlier false fracture, the later real amputation, and any Gu sacrifice. Frost Demon Gu has not thereby returned.

### The official history is intentionally false

Chapter179’s public history says Flower Wine Wanderer was killed on the spot. Chapter180 explicitly distinguishes public history from secret history:

> {src(180,39)['text']}

The reader has already seen evidence inconsistent with the public version. That is the point of the scene, not a defect to harmonise. Likewise, Fang Yuan’s earlier reconstruction of the Lotus’s history is not a newly authenticated record of every past event.

## 3. The investigator arrives early, not on a fixed promised date

Chapter113 forecast the year after next. The current arrival happens sooner, and172 acknowledges it:

> {src(172,30)['text']}

This is stronger evidence for retaining the earlier statement as a forecast than for mechanically changing its wording. Exact scheduling, prediction and actual arrival must remain separate.

The wolf tide also has phases. The most dangerous assault ends, but171–172 still expect smaller groups for months and their near-disappearance toward year’s end. “The crisis has been weathered” is not the same as “every wolf has gone.” It does not by itself refute an earlier forecast of a year-long tide.

Chapter170’s “all elders” death statement is local to those remaining on that battlefield. The surviving elders at171’s banquet do not contradict it. Chi Lian’s death is now confirmed; Mo Chen avoided that pursuit after being injured two weeks earlier.

## 4. Expedition clock: do not silently turn时辰 into hours

Chapter162 supplies three separate durations:

| Activity | Source duration | Conventional equivalent |
|---|---|---|
|Reopening the monkey route|三个时辰|About6hours|
|Exploring the dead-end branch|半个小时|30minutes|
|Exploring the longer passage|六个多时辰|More than12hours|

These alone exceed18.5hours, before retracing, Lotus inspection and the final return. Chapter163 says:

> {src(163,73)['text']}

**SR-083:** translate this as “It was dawn,” not automatically “the next morning.” The Chinese does not supply that next-day claim. A later dawn and an unquantified return interval can accommodate the stated durations; the exact elapsed total remains open.

Do not repair the apparent overnight pressure by halving时辰, shrinking the passage, or inventing a precisely36-hour expedition.

## 5. Two range questions remain conditional

### Hundred-li lightning path

Chapter168 describes a hundred-li smoke extent and then a lightning strike crossing a hundred-li distance toward Fang Yuan. Yet he has only just begun escaping the same battlefield.

**SR-080, held:** “tore through the dense smoke” is a minimal candidate for the strike’s path if later staging confirms a nearby target. It is not a claim that RevolutionFour lightning cannot possess a hundred-li range. The smoke extent is a separate assertion and is not automatically reduced.

### Earth Hearing Flesh Ear Grass:200 versus300paces

Earlier passages give a300-pace rooted maximum and a smaller unrooted range. Chapter169 says:

> {src(169,56)['text']}

A later action in the same scene separately roots the grass. CurrentRevolutionThree cultivation and effective working conditions complicate the comparison.

**SR-081, held:**300 is a candidate only if the passage means the previously stated fully rooted maximum. If200 is a practical operating range, clarify that condition instead. I am not treating300 as unconditional in every mode.

Bai’s Electric Eye is much clearer here:50paces normally,30in the smoke. Smoke reduces its vision; it does not make it see through solid terrain.

## 6. Resource accounting and the unfinished Lotus acquisition

### Cash and carrying capacity are different

The banquet yields **nearly10000stones**, not exactly10000. Tushita Flower can hold30000stones if used for that purpose; Fang Yuan budgets15000when reserving space for food and other goods. These are not contradictory capacities.

He then estimates that restoring the Lotus’s missing half-leaf will require50000-plus stones. More than500have only produced slight growth; he has10000-plus and describes a roughly40000gap. Those rounded amounts do not establish a subtraction error.

By180, the40000advance has actually arrived in instalments. His aggregate funds may exceed the flower’s absolute capacity, but the text has **not yet said all those funds are packed inside it**. This is a transport question to carry forward, not a proved contradiction at the endpoint.

### Marriage bargaining is not a completed marriage

The initial demand is100000stones plus ten rare Gu of at leastRevolutionThree. The later30% concession suggests70000for the cash component if that is its scope; it does not establish a precisely discounted seven-Gu package. The40000is an advance, not an extra full payment on top.

Pig Iron Gu is delivered and consumed to repair the centipede. The requested healing grass is not delivered. No completed marriage is narrated.

### Acquisition and loss ledger

| Item | Confirmed status through180 |
|---|---|
|Chainsaw Golden Centipede|Captured and refined in162; mostly repaired with Pig Iron in177|
|Thunder Shield Gu|Captured in167; destroyed in168|
|Lightning Wings Gu|Badly injured; still unusable at170, not declared dead|
|Sky Canopy Gu|Injured, not destroyed|
|White Form Immortal Snake|In Bai’s aperture but unrefined; not Fang Yuan’s property|
|Heavenly Essence Treasure Lotus|Observed and partially fed; **not acquired** by180|
|Healing grass requested from Mo Chen|Still undelivered|

Pig Iron heals the centipede specifically. It does not establish that Fang Yuan’s wounds or Lightning Wings have also been healed. The remaining five or six shallow centipede marks are expected to heal over weeks; without the treatment, replacement of all its damaged teeth would take more than half a year.

## 7. Aptitude and evidence-generating Gu

Chapter162 and167 preserve Fang Yuan’s42% capacity. The later40points are current reserves after expenditure, not another aptitude reduction.

Bai’s history is independently important:

> {src(173,16)['text']}

His original99% GradeA developed into the TenExtreme physique. This supports a **mutable aptitude** model where a stated mechanism changes it, not a global assumption that all aptitude is fixed or that all GradeA must reach100%.

The investigation also demonstrates fallible evidence tools. Divine Pointer works from the planter’s known clues; it is not an omniscient oracle. Chi Cheng’s cheating is corroborated, but the resulting claim that Chi Lian supplied the original Liquor Worm remains an incorrect hypothesis.

The six stones bought and five opened on the spot match Chapters40–43. There is no count repair to make. The supposed Liquor Worm discovery is Fang Yuan’s alibi, not an actual second Gu found by gambling.

Fang Zheng’s later account of drunkenness is his mistaken interpretation. It does not retroactively require translating the opening买醉 as drinking oneself insensible.

## 8. Strong local corrections, still held

| Record | Location | Candidate |
|---|---|---|
|SR-074|164p5|雷光头狼→雷冠头狼|
|SR-075|167p70|雷壳蛊→雷盾蛊|
|SR-076|165p71|药钟→赤钟, strongly supported by the retained-official pairing and171’s survivors|
|SR-077|172pp70,77|赤脉→漠脉 in theMo courtyard andMo succession discussion|
|SR-078|175p67|方正滥杀→方源滥杀|
|SR-079|178p51|Father’s speech attributed to铁血冷, not铁若男|

The Continuity jotter gives the exact passages, rationale, confidence, dependencies and publication status. SR-084 adds{len(fixes)} individually listed copy candidates, not a global search-and-replace instruction.

SR-082 remains a lower-confidence scope query:172’s male-line inheritance language needs reconciliation with Yao Le’s stated role as a future branch successor. Hereditary property, independently established property, office and political leadership may have different rules. That distinction must be demonstrated, not invented to smooth over the passages.

## 9. Endpoint and saved checkpoint

At180, the clan leader has secured a delay so Fang Yuan can participate in the tournament scheduled for the day after tomorrow. The tournament has not happened. The awakening on the wall is a **recording of the original ceremony**, not a new present event. Tie Ruo Nan’s final hypothesis remains unrevealed at this endpoint.

Fang Yuan has received the40000advance and leaves to act on his Lotus plan. He has not yet picked or refined it.

Saved: numbered Chinese source, four jotters, reading log,84-record proposal ledger, **{len(events)} chronology entries**, **{len(resources)} resource entries**, updated workbook and the new working-standards register.

**Next actual unread:181. Next planned batch:181–200.** No source, site or repository edits, uploads or deletions occurred. Remote research backup remains unresolved. We continue beyond any locally stabilised standard or single-volume milestone.
'''
report=readable(report)
for a,b in {'RevolutionFour':'Revolution Four','RevolutionThree':'Revolution Three','GradeA':'Grade A','nextday':'next day'}.items():report=report.replace(a,b)
(r/'Batch-008-Chapters-161-180.md').write_text(report)
w=Workbook();s=w.active;s.title='Read me'
for row in [('Sequential review','Status'),('Completed','1–180 of2334'),('EPUB paragraphs',tp),('Source characters',tc),('Next unread',181),('Proposal records','84total;83activeheld;1withdrawn;noneapproved/applied'),('New copy bundle',f'SR-084contains{len(fixes)}individualitems'),('Standards','10workingstandards; nonewhole-novel-confirmed'),('Scope','Chinese-source review; notfreshbilingual/liveEnglishaudit'),('Reader copy','EarlierCh151restrictedpassagesretainexplicitomissionmarkers;originalEPUBunchanged'),('Chronology entries',len(events)),('Resource entries',len(resources)),('Global consistency','Not yet certified')]:s.append(row)
s=w.create_sheet('Chapter completion');s.append(['Chapter','File','Title','Paragraphs','Status'])
for c in chapters:s.append([c['corpus_chapter'],c['file'],c['title'],len(c['paragraphs']),'Full sequential read;151includes3paratextparagraphs' if c['corpus_chapter']==151 else 'Full sequential read'])
def ids(vs,k,c,p):return ', '.join(v['id'] for v in vs if any(z['chapter']==c and z['paragraph']==p for z in v[k]))
s=w.create_sheet('Paragraph register');s.append(['Chapter','Paragraph','Status','Proposal IDs','Chronology IDs','Resource IDs','Chinese reading copy'])
for c in chapters:
 for p,t in enumerate(c['paragraphs'],1):
  n=c['corpus_chapter'];s.append([n,p,'Read; not certified error-free',ids(ledger,'locations',n,p),ids(events,'sources',n,p),ids(resources,'sources',n,p),reader(n,p,t)])
s=w.create_sheet('Proposal ledger');s.append(['ID','Category','References','Action','Rationale','Confidence','Status','Dependencies','Publication','Approved','Applied','Latest update'])
for v in ledger:
 latest=next((v[f'batch{i:03d}_update'] for i in range(8,0,-1) if f'batch{i:03d}_update' in v),'')
 s.append([v['id'],v['category'],refs(v['locations']),v['recommended_action'],v.get('rationale','See originating report'),v.get('confidence','See status'),v['status'],v['dependencies'],v['publication_status'],'No','No',latest])
s=w.create_sheet('Copy items');s.append(['Proposal','Chapter','Paragraph','Original','Candidate','Status'])
for v in ledger:
 for z in v.get('copy_items',[]):s.append([v['id'],z['chapter'],z['paragraph'],z['original'],z['replacement'],z['status']])
s=w.create_sheet('Chronology');s.append(['ID','Timeline','Kind','Event','Constraint','Evidence status','Dependencies','References','Exact permitted source'])
for v in events:s.append([v['id'],v['timeline'],v['kind'],v['event'],v['constraint'],v['evidence_status'],v['dependency'],refs(v['sources']),'\n'.join(reader(z['chapter'],z['paragraph'],z['text']) for z in v['sources'])])
s=w.create_sheet('Resources');s.append(['ID','Kind','Entry','Calculation','Status','References','Exact permitted source'])
for v in resources:s.append([v['id'],v['kind'],v['entry'],v['calculation'],v['status'],refs(v['sources']),'\n'.join(reader(z['chapter'],z['paragraph'],z['text']) for z in v['sources'])])
s=w.create_sheet('Working standards');s.append(['ID','Topic','Status','Rule','Evidence','Reopening trigger'])
for v in standards:s.append([v['id'],v['topic'],v['status'],v['rule'],v['evidence'],v['reopen']])
for s in w:
 s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
 for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='163B4B')
 for row in s.iter_rows(min_row=2):
  for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
 for col in s.columns:s.column_dimensions[col[0].column_letter].width=min(65,max(14,max(len(str(c.value or '')) for c in col[:30])+2))
w.save(r/'Sequential-Reading-and-Proposals.xlsx')
code=(r.parent/'make_whole_novel_audit_docx.py').read_text().replace("root=Path(__file__).parent","root=Path('/home/user/editorial/sequential-review')").replace('Whole-Novel-Aperture-and-Travel-Audit','Batch-008-Chapters-161-180').replace('WHOLE-NOVEL SOURCE AUDIT','SEQUENTIAL REVIEW / CHAPTERS161–180').replace('The Omniarch: Aperture and Travel Source Audit','The Omniarch: Sequential Review, Chapters161–180').replace('Full-corpus topical screening with selective contextual verification','Twenty complete chapters; working standards remain revisable')
exec(compile(code,'docx-export','exec'),{})
(r/'CURRENT-STATUS.md').write_text(f'''# Sequential review checkpoint

Completed **Chapters1–180**. Next unread **181**. Next planned batch **181–200**. Remaining **2154**.

Batch008: {np}paragraphs/{nc}characters. Cumulative: {tp}paragraphs/{tc}characters.84proposalrecords:83activeheld,1withdrawn. SR-084contains{len(fixes)}copyitems. No approved/appliedchanges. Whole-novel chronology NOTcertified.

Latest main deliverable: Batch-008-Chapters-161-180.md/.docx. Actual numbered Chinese source: Batch-008-Numbered-Chinese-Source.md. Four008jotters, readinglog, registers and workbook updated.

New WORKING-STANDARDS.md + working-standards-register.json:10standards, each withstatus/evidence/reopeningtrigger. User explicitly reaffirmed no stopping afteronevolume; strongworkingstandards do notreplace full2334reading or authoriseemendations.

Key carry-forwards:
- Xiongdestruction was inference, corrected175survivors/groupinvisibility. ApparentfourthThunderCrown isLightningLycan. No resurrection.
-17399GradeA→100TenExtreme explicitlydynamic;FY42confirmed162/167,remaining40afterspendingnotnewcapacity.
- WhiteFormSnake173inhabitationNOTinheritanceapprovalorrefinement;178explains. BaiRIGHTarmgenuinelyregeneratedphysique. FrostDemonnotrestored.
-179publichistoryfalse,180secret/public distinctionexplicit. DivinePointerfallibleclue-basedhypothesis,notomniscience. Sixstones/fiveopenedvalidated40–43.
-162expedition3shichen+6plusshichen+30min+return exceeds18.5h;163atdawnnotnecessarilyimmediatenextmorning. SR083translationguard,actualdayunresolved.
- SR080168hundred-li strikepath conditionalstagingquery,noglobalshrink. SR081169Ear200vs300rootedconditional;ElectricEye50normal/30smokeexplicit.
- SR074wolfname164;075Shieldname167;076165YaoZhong→ChiZhongcandidate;077172Mo/Chi;078175FangZheng→FangYuan;079178speakerTieXueLeng. SR082maleinheritancescopeunresolved.
- Near10000gifts161;storage30000absolute/15000practical;Lotus8.5leaves,500+trial,50000+estimate,40000advanceactuallyreceived180. Cashownednotautomaticallyallpacked. LotusNOTacquired,healinggrassNOTdelivered,marriageNOTcompleted.
- GoldenCentipedecaptured162/repaired177;PigIronconsumed. ThunderShieldacquired167/destroyed168. LightningWingsbadlyinjured/notdead;repairnotestablished. No exactinventorygrandtotalbecauseearlierYaoLelootunnumbered.
- Tiearrivalearlyexplicit172;113forecastnotfixeddate. Wolfmaincrisisendswhilemonthsminorwolvesremain. Battlefieldeldersdead170notallclaneldersdead;ChiLiandead171,MoChenRevolution2concealed.
-180recordedawakeningnotpresentceremony;competitiondayaftertomorrowNOTheld;lastlineTiehypothesisNOTrevealed. Continue181withoutimportinglaterknowledge.

No source/site/repository edits, uploads or deletions. Remote researchbackup unresolved. OriginalEPUBunchanged; cumulativeworkbook retains earlierCh151restrictedpassageomissionmarkers. Do notrerun earliergenerators unchanged. Topical2334screening isNOTsequentialcompletion.
''')
assert len(chapters)==180 and len(ledger)==84 and len(new)==11
assert np==1861 and nc==62666 and tp==15028 and tc==543457
assert all(not v['approved'] and not v['applied'] for v in ledger)
assert all(z['text']==src(z['chapter'],z['paragraph'])['text'] for v in events+resources for z in v['sources'])
assert '—' not in report
texts={t for c in chapters for t in c['paragraphs']}
for line in report.splitlines():
 if line.startswith('> '):assert line[2:] in texts,line
print(dict(batch_paragraphs=np,batch_characters=nc,total_paragraphs=tp,total_characters=tc,proposal_records=len(ledger),active=83,withdrawn=1,copy_items=len(fixes),chronology=len(events),resources=len(resources),standards=len(standards),next=181))
