import json
from pathlib import Path
from openpyxl import load_workbook
from pathlib import Path as P
r = Path('/home/user/editorial/sequential-review')
led = json.loads((r/'proposal-ledger.json').read_text())
WD = ('Withdrawn: 508p29 and 517p21 refer to 方源 (real identity, five-level), not to the 常山阴 persona. '
      'Earlier text (431p24, 493p59, 495p37, 499p24) and 515p69, 525p82 all treat the 常山阴 persona as four-level. No discrepancy.')
for v in led:
    if v['id'] == 'SR-282':
        v['status'] = WD
        v['confidence'] = 'Withdrawn.'
        v['recommended_action'] = 'Withdrawn. No change.'
        v['rationale'] = 'The five-level statements describe 方源 under his own name. The four-level statements describe the 常山阴 disguise.'
        v['dependencies'] = 'Resolved against earlier batches and the 508–517 statements. Register note: 常山阴 is the persona; 方源 is the real identity.'
        v['approved'] = False; v['applied'] = False
(r/'proposal-ledger.json').write_text(json.dumps(led, ensure_ascii=False, indent=1))
wd = sum(1 for v in led if v.get('status','').startswith('Withdrawn'))
rs = sum(1 for v in led if v.get('status','').startswith('Resolved'))
print('withdrawn', wd, 'resolved', rs, 'total', len(led), 'active', len(led)-wd-rs)

# workbook: SR-282 row and Read me label
wb = load_workbook(r/'Sequential-Reading-and-Proposals.xlsx')
ws = wb['Proposal ledger']
hdr = [c.value for c in ws[1]]; ci = {h:i for i,h in enumerate(hdr)}
for row in ws.iter_rows(min_row=2):
    if row[0].value == 'SR-282':
        row[ci['Confidence']].value = 'Withdrawn.'
        row[ci['Action']].value = 'Withdrawn. No change.'
        row[ci['Status']].value = WD
        row[ci['Dependencies']].value = 'Resolved against earlier batches and the 508–517 statements. Register note: 常山阴 is the persona; 方源 is the real identity.'
        print('SR-282 row updated')
rm = wb['Read me']
for row in rm.iter_rows(min_row=1, max_col=2):
    if row[0].value == 'Proposals':
        row[1].value = f'285 total; {len(led)-wd-rs} active (held or query); {wd} withdrawn; {rs} locally resolved'
        print('Read me Proposals ->', row[1].value)
wb.save(r/'Sequential-Reading-and-Proposals.xlsx')
