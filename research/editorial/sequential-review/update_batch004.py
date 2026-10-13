from pathlib import Path
import json, shutil
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
r=Path('/home/user/editorial/sequential-review')
def load(n):return json.loads((r/n).read_text())
def save(n,v):(r/n).write_text(json.dumps(v,ensure_ascii=False,indent=2))
chapters=sum([load(f'batch-{i:03d}-source.json') for i in range(1,5)],[])
by={c['corpus_chapter']:c for c in chapters};batch=[by[i] for i in range(51,81)]
def src(c,p):return dict(chapter=c,paragraph=p,file=by[c]['file'],text=by[c]['paragraphs'][p-1])
def sources(refs):return [src(*x) for x in refs]
def loc(refs):return [{k:v for k,v in s.items() if k!='text'} for s in sources(refs)]
def refs(ss):return '; '.join(f"Ch{s['chapter']} p{s['paragraph']}" for s in ss)
for n in ['reading-progress.json','proposal-ledger.json','chronology-register.json','resource-register.json']:
 dst=r/'history'/('before-batch004-'+n)
 if not dst.exists():shutil.copy2(r/n,dst)
ledger=load('proposal-ledger.json');ledger=[v for v in ledger if int(v['id'][3:])<=28]
for v in ledger:
 if v['id']=='SR-024':
  for s in loc([(62,5),(62,83),(71,56),(73,35),(73,37)]):
   if s not in v['locations']:v['locations'].append(s)
  v['batch004_update']='Further Mo Chen→Mo Bei occurrences in Ch62 pp5,83; Ch71 p56; Ch73 pp35,37. Ch63 p13 explicitly identifies Mo Bei as the beetle recipient. Preserve elder Mo Chen in Ch77–78.'
 if v['id']=='SR-017':v['batch004_update']='Ch60 supplies >10-day wait and sixth excavation day; do not fold these later delays into the earlier Ch46 recording duration. Ch72 half-year and Ch80 seasonal labels add separate dependencies.'
 if v['id']=='SR-018':v['batch004_update']='Ch56 pp11,24 supports next-noon visit→same-evening dispute and rainy departure. Investigation supplies no exact new murder/departure interval; retain held linked repair.'
 if v['id']=='SR-022':v['batch004_update']='Ch80 p39 explicitly gives full44% volume; supports local45→44 candidate without recalculating aptitude globally.'
new=[]
def proposal(i,category,where,action,rationale,confidence,deps):
 v=dict(id=f'SR-{i:03d}',category=category,locations=loc(where),recommended_action=action,rationale=rationale,confidence=confidence,status='Provisional active proposal; not approved',dependencies=deps,approved=False,applied=False,publication_status='HOLD: full2334-chapter reading and global reconciliation outstanding');ledger.append(v);new.append(v)
proposal(29,'ranking',[(53,22),(53,24),(53,26),(53,27),(54,22),(62,87)],'If overall ranking is intended, change Chi Cheng third→fourth at all three linked mentions. Alternatively specify third among the remaining candidates.','Fang Yuan first, Mo Bei second, Fang Zheng next, Chi Cheng next. Ch62 explicitly makes Chi Cheng fourth in selection order.','High for overall sequence; medium for repair because candidate-subset usage is possible.','Resolve whether Ch53/54 excludes Fang Yuan; preserve correct officer appointments.')
proposal(30,'accounting',[(54,55),(54,58)],'Retain68 pending reconciliation. If56 paying peers remain, use67 for levy-only receipts or70 including own3-stone stipend, and label the accounting scope.','53 ordinary peers×1 +2 deputies×3 +1 head×8=67. Former59 fits56+3, but new68 does not match either consistent scope.','High conditional arithmetic; medium defect assessment because a pupil is speaking.','Confirm later cohort headcounts and receipts; do not manufacture a missing payer or expense.')
proposal(31,'witness time',[(44,1),(55,67),(56,11)],'If harmonising the witness rather than preserving mistaken testimony, change 那天上午→那天中午; English: that noon.','Direct Ch44 narration and Fang Yuan’s Ch56 account place the visit at noon; witness says morning.','High local mismatch; medium emendation.','Character testimony need not be accurate; distinguish this from murder-date repair SR-018.')
proposal(32,'chronology',[(7,4),(7,47),(72,2),(72,26)],'Working repair: 在之前的半年里→在过去这几个月里; English: Over the past few months. Retain mid-year examination.','Ch7 is April and seven days after awakening; Ch72 is lateJune. The training retrospect calls this the past half-year. A mid-year exam need not follow six full months of attendance.','High tension in short chronology; moderate preferred wording.','Reconcile SR-002/SR-017 and all later academy dates; do not shift April to another month simply to save this sentence.')
proposal(33,'appointment schedule',[(26,14),(53,28),(53,95)],'If early appointment was intended, revise 三个月后→届时 and explain selection once the leading candidates qualify; otherwise preserve as an unannounced schedule change.','Ch26 announces appointments after3months, but the linked extortion/caravan/breakthrough sequence reaches appointment much sooner.','Moderate continuity query; low-to-moderate repair confidence.','Whole academy calendar and possibility of changed policy; no invented two-month hiatus.')
proposal(34,'scene time',[(66,49),(67,16)],'If the household scene is continuous with the afternoon arrival, 今天早上你就多睡一会儿嘛→今天你就多休息一会儿嘛. English: You should rest a little more today.','The killing is explicitly afternoon; the immediately preceding family dialogue refers to more sleep this morning. An earlier scene setup remains possible.','Moderate query, conditional repair.','Determine whether scene is an implicit earlier cut; do not move the killing to morning or insert a day.')
proposal(35,'numerical wording',[(62,40),(62,62)],'持续时间是蛮力天牛蛊的一倍→持续时间是蛮力天牛蛊的两倍. English: lasts twice as long.','The paragraph explicitly compares10breaths with5; standalone 是…一倍 states equality rather than the intended doubled duration.','High local confidence.','Apply only to this explicit ratio, not every倍 construction.')
proposal(36,'source identity',[(66,14),(66,28),(66,74)],'另外两个蛊师→另外两个猎手; English: the other two hunters.','These are the mortal companions, repeatedly identified as hunters, not Gu Masters.','High contextual confidence.','Source witness validation; not a claim that liveEnglish currently contains the error.')
proposal(37,'source names',[(74,38),(76,17)],'Change 王二→王大 at Ch74p38 and Ch76p17.','The elder son who vanished and the surviving attacker are Wang Da; Wang Er is the younger hunter already killed.','High contextual confidence.','Do not change deliberate Fang Yuan/Fang Zheng mistaken identity in the same sequence.')
proposal(38,'Gu nomenclature',[(75,75),(75,82),(75,83),(76,23),(76,36),(76,76),(77,68)],'Working standard: 爱别离, with one consistent English Gu name. Harmonise earlier 爱生离 only after later usage is checked.','One poison Gu, one owner and one continuing fight use both forms; Ch76–77 settles locally on 爱别离.','High identity confidence; medium preferred canonical form.','Later appearances, recipes and independent witnesses; do not invent a second poison Gu.')
save('proposal-ledger.json',ledger)
tdoc=load('chronology-register.json');events=tdoc['entries'][:41]
def event(timeline,kind,title,constraint,status,where,dep='Full-corpus reconciliation remains outstanding.'):
 events.append(dict(id=f'T{len(events)+1:03d}',timeline=timeline,kind=kind,event=title,constraint=constraint,evidence_status=status,dependency=dep,sources=sources(where)))
event('current life','relative interval and ranking','Breakthroughs and stipend day','Mo Bei second; Fang Zheng3hours later; fifth-day stipend and appointments.','Events narrated; Chi Cheng third label queried underSR-029.',[(53,21),(53,22),(53,24),(53,26),(53,28)])
event('current life','witness testimony','Wine-shop visit time','Narrated noon versus witness morning; following-evening dispute.','Separate conflicting testimony from actual narration and later deceptive testimony.',[(44,1),(55,67),(56,11),(56,24)],'SR-031 and SR-018.')
event('current life','investigation window','Surveillance and footprint evidence','Surveillance follows advancement; footprintGu shows last30,000steps, not every past movement.','Insufficient to establish a real murder-time alibi; political protection and limited evidence must not be rewritten into omniscience.',[(57,20),(57,28),(57,32),(57,42)])
event('Ren Zu myth','physical age','Renewed youth','LifespanGu restores Ren Zu to20; subsequent transactions exchange life phases.','Mythic physical-age account, not current-life elapsed calendar.',[(58,8),(58,26),(58,31)])
event('clan history','approximate duration','Clan foundation','Nearly1000years; 百年基业 need not mean exactly100.','Preserve approximate/idiomatic descriptions rather than forcing100=1000.',[(55,45),(58,60)])
event('current life','planned event','Fang Zheng private instruction','Leader plans to begin tomorrow night; Ch61 confirms instruction underway.','Plan subsequently corroborated, not a second start when publicised inCh78.',[(58,54),(61,19),(78,32),(78,34)])
event('forecast / changed history','prediction','Tie investigator and Jia conflict','Tie arrival predicted at least2–3years away; Jia competition may be advanced by interference.','Fang Yuan forecast, not fixed future event date.',[(59,44),(59,45),(59,46)])
event('current life','relative intervals','Return to cave and excavation','Revelation over1month earlier; wait>10days after Jia departure before return; current sixth digging day.','Broad intervals can coexist. Six consecutive days span5days from first to sixth; do not double-count.',[(60,9),(60,12),(60,17)])
event('current life','local clock and travel','Cave exploration and second selection','Midnight return→nextmorning lesson→evening test→nextmorning selection→sameevening inheritance. Ropework1hour; cautious300m traversal2hours.','Explicit narrated sequence; 时辰 and 小时 distinguished.',[(60,70),(61,2),(61,22),(61,42),(61,62),(62,1),(63,7),(63,44)])
event('current life','scene clock','Afternoon hunting and morning dialogue','Ch66 afternoon; Ch67 daughter references this morning.','Conditional scene-continuity query, not proven day change.',[(66,49),(67,16)],'SR-034.')
event('Wang family past','reported death / revised evidence','Wang Da and wife','Supposed death3yearsago; mother died>10yearsago. Later narrative confirms Wang Da alive and wife murdered to refineGu.','Rumours must not be treated as verified deaths; narration supersedes the hearsay and Fang Yuan’s tentative motive.',[(69,8),(69,13),(74,56),(74,57),(75,84),(77,74),(78,63)])
event('prior life / current life','rebirth comparison','Cultivation and Bookworm','Bookworm accompanied priorlife60years. Classmates formerly advanced faster; present combat practice diverted their cultivation time.','Explicit branch difference, not duplicated current elapsed time.',[(71,49),(71,50),(71,54)])
event('current life','calendar anchors','June and the exam','MidJune→lateJune→exam3dayslater; training retrospect sayspast half-year.','June anchors firm in text; half-year elapsed-duration query held.',[(71,1),(72,2),(72,26),(72,27),(73,16)],'SR-032 and opening chronology.')
event('current life','cache / clock / travel','Exam-day sequence','Tusks cached nightbeforelast; treehouse500–600m away reached inhalf刻; detour1hour; arrives1hourlate.','Approximate clock relationships; no speed repair required.',[(73,68),(73,74),(73,80),(77,17),(77,48)])
event('current life','cooldown and age','Wang Da attack','ShadowGu3hour cooldown; Wang Da dies during fight, aged35 per later report.','Cooldown is prospective, not a3hour battle duration.',[(76,14),(76,89),(77,68)])
event('current life','retrospect and report delay','After assassination','Death confirmed to Fang Yuan3dayslater; Fang Zheng coma7days/nights; another十几天 before Fang Yuan/Chi Cheng upper-stage advancement.','Death-date distinct from knowledge-date; coma compatible with third-day survival report.',[(78,58),(78,72),(78,73),(80,53)])
event('current life','season / retrospective interval','Autumn transition','Summer→autumn; previousweek45steps; fullboarstrength3daysago; later deepautumn and assassination>2monthsago.','Seasonal watch: lateJune/earlyJuly plus2+months roughlySeptember. No calendar system or exact day established.',[(79,1),(79,6),(79,11),(80,51),(80,52)],'Keep 深秋 and 两个多月 until later exact anchors justify a linked repair; do not assume Gregorian dates.')
event('forecast / current life','cultivation deadline','Peak stage','Fang Zheng alreadypeak; Fang Yuan stillupper and predictspeak in10–15days.','Distinguish refinedpeak essence from actualpeak cultivation; prediction not completion.',[(80,50),(80,56),(80,57)])
tdoc.update(scope='Sequential evidence throughChapter80; global chronology not certified.',entries=events);save('chronology-register.json',tdoc)
rdoc=load('resource-register.json');resources=rdoc['entries'][:18]
def resource(kind,entry,calculation,status,where):resources.append(dict(id=f'R{len(resources)+1:03d}',kind=kind,entry=entry,calculation=calculation,status=status,sources=sources(where)))
resource('reward and stipend','30-stone breakthrough prize received;10-stone head stipend declined.','Do not book rejected stipend as income.','Actual reward separated from proposed office benefit.',[(53,39),(53,41),(53,43)])
resource('levy receipts','Newrates1/3/8; pupil states former59/current68.','56payingpeers implies53+6+8=67; plusown3stipend=70.','SR-030 conditional accounting query; no cumulative cash correction.',[(54,55),(54,58)])
resource('food and living costs','Four costlyGu, >2stonesdailyfeeding, ≥5totaldaily; hundreds held.','Cicada excluded from costlyfeeding subset; fiveowned inCh70.','Approximate ongoingexpenses, not fully itemised account. Proposed100+jars order has unquoted deposit.',[(64,8),(64,60),(64,61),(64,62),(64,64),(70,13)])
resource('feeding versus activation','WhiteBoar adultpig every5days; daily useabout15minutes.','Feeding interval≠activation interval.','No dailywholepig requirement.',[(64,71),(70,62),(70,83)])
resource('essence balance','Ch70 hunt leaves23high-stagepoints; cost stated20.','44−20=24 only if huntbeganfull and20isexact; neither independently established.','Do not silentlycorrect23→24 or20→21.',[(70,72),(70,73)])
resource('exam tusks and rewards','23actualboars→46fresh tusks; separateoldercacheexistsbutretrieval/submissionisnotnarrated; latepenalty4tusks; chiefannounces30stones.','Conditional freshcount afterpenalty42; totalcache/submittedcountnotgiven.10stonespertusk reward vs20tusks/stone market.','No invented finalcashbalance; Fang Yuan’s10+personalclaim isdeliberatefalsehood.',[(72,27),(72,34),(73,66),(77,48),(77,55),(78,38)])
resource('Wang Da essence','Remaining50→40→30→20→5 throughassassination.','Locally decreasingreserve, not maximumcapacity50.','Ch75 likelyaptituderange isindividualinference, not newgradeboundaries.',[(75,8),(75,53),(75,68),(75,71),(76,20),(76,49)])
resource('food weight and stock','JadeSkin2liang/10days; eachmonkeyeye≥1liang; sixthGu acquired.','Pair ofeyes≥2liang coversonefeed.','Retain sourceweightunit; no unjustified gramconversion. Six includesCicada.',[(79,34),(79,41),(79,70)])
resource('stage potency / capacity','Full44volume−8peakpoints=36remaining;8peak=64initial.','Stagepotency1:2:4:8; Moonlightcost10/5/2.5/1.25.44/64=68.75% ofhypotheticalinitial-stagebudget.','Exact localarithmetic confirmed; Ch80p49 mid-actiondepletion isloosewording, not a32capacity claim.',[(80,39),(80,40),(80,46),(80,47),(80,48),(80,49)])
rdoc.update(scope='Resources throughChapter80; no certified full cumulativecashbalance.',entries=resources);save('resource-register.json',rdoc)
np=sum(len(c['paragraphs']) for c in batch);nc=sum(sum(map(len,c['paragraphs'])) for c in batch);tp=sum(len(c['paragraphs']) for c in chapters);tc=sum(sum(map(len,c['paragraphs'])) for c in chapters)
progress=load('reading-progress.json');progress.update(sequential_chapters_completed=list(range(1,81)),paragraphs_read=tp,paragraph_characters_read=tc,next_chapter=81,remaining_chapters=2254,batch_report='Batch-004-Chapters-51-80.md',scope='Full sequential Chinese source reading throughChapter80. Not a fresh liveEnglish/bilingual audit.',whole_novel_chronology_verified=False,chapters=[dict(chapter=c['corpus_chapter'],file=c['file'],title=c['title'],paragraphs=len(c['paragraphs']),first_paragraph=1,last_paragraph=len(c['paragraphs'])) for c in chapters]);progress['batches']=[x for x in progress['batches'] if x['id']!='004']+[dict(id='004',chapters=[51,80],paragraphs=np,characters=nc,reading_segments='Fifteen full untruncated two-chapter outputs')];save('reading-progress.json',progress)
save('batch-004-reading-log.json',dict(status='Complete',endpoint=80,next_unread=81,segments=[dict(chapters=[a,a+1],status='Full paragraph sequences printed and read; output untruncated',paragraphs=sum(len(by[k]['paragraphs']) for k in [a,a+1])) for a in range(51,81,2)],note='No invented reading timestamps. Extraction did not count as completion.'))
lines=['# Numbered Chinese source: Chapters 51–80','Unemended nonempty EPUB body paragraphs. Paragraph numbers match this review.','']
for c in batch:
 lines += [f"## Chapter {c['corpus_chapter']}: {c['title']}",f"Source: `{c['file']}`",'']
 for i,t in enumerate(c['paragraphs'],1):lines += [f'**{i:03d}.** {t}','']
(r/'Batch-004-Numbered-Chinese-Source.md').write_text('\n'.join(lines))
# Clean evidence jotters.
for name,vs in [('Chronology',events[41:]),('Resources',resources[18:])]:
 lines=[f'# Batch 004: {name} jotter','Complete source read through Chapter 80. All changes held.','']
 for v in vs:
  lines += [f"## {v['id']}: {v.get('event',v['kind'])}"]
  for k in ['timeline','kind','constraint','evidence_status','dependency','entry','calculation','status']:
   if k in v:lines += [f"**{k.replace('_',' ').title()}:** {v[k]}"]
  for s in v['sources']:lines += [f"\n**Ch{s['chapter']} p{s['paragraph']}**",'> '+s['text'],'']
 (r/'jotters'/f'Batch-004-{name}.md').write_text('\n'.join(lines))
lines=['# Batch 004: Continuity and active proposals','All proposals unapproved and unapplied.','']
for v in new:
 lines += [f"## {v['id']}: {v['category']}",f"**Action:** {v['recommended_action']}",f"**Rationale:** {v['rationale']}",f"**Confidence:** {v['confidence']}",f"**Dependencies:** {v['dependencies']}",f"**Status:** {v['publication_status']}",'']
 for s in v['locations']:lines += [f"**Ch{s['chapter']} p{s['paragraph']}**",'> '+src(s['chapter'],s['paragraph'])['text'],'']
lines += ['## SR-024 extension',next(v['batch004_update'] for v in ledger if v['id']=='SR-024'),'','## Additional watch items','- Ch68p28 削蛊: likely source corruption, but 削骨 versus an intended flesh-cutting comparison is not settled. Preserve as a textual query; do not silently redesign blade mechanics.','- 幽影随行蛊 / 幽影虫: likely full name versus short form. Prefer one English name if identity is confirmed globally; no need to invent another Gu or forcibly expand every Chinese abbreviation.','- Ch75p26 categorical claim that breakthroughs during battle are impossible: retain for now and cross-check later ordinary breakthroughs versus exceptional means.','- Ch80 deep-autumn label: preserve the text while collecting exact subsequent dates; a descriptive season is weaker evidence than a numbered month.','']
(r/'jotters'/'Batch-004-Continuity.md').write_text('\n'.join(lines))
preserve='''# Batch 004: Preserve, do not repair

- Ch51–52 guards survive. Earlier assumptions of death are corrected by the scene itself.
- Fang Yuan's provocation story, gambling origin for the Liquor Worm, and denial of murder are deliberate lies.
- Bamboo Gentleman is received unrefined and suppressed by the Cicada. This is not ordinary successful refinement by a liar.
- The elders' surveillance is not a genuine murder-time alibi. Preserve the distinction between evidence and politically convenient conclusions.
- Ch55's toad buyer is Jia Jin Sheng: Ch42pp58–63 directly confirms him. A working-memory suspicion that Jia Fu bought it was wrong and has been withdrawn. No new correction belongs here.
- Ch58 Ren Zu's physical youth and clan-history approximations belong in separate registers.
- Ch64 four feeding-cost Gu and Ch70 five owned Gu are compatible because Cicada is excluded from the expense subset.
- Ch65 independent wilderness travel and hunting near patrol-cleared villages are different risk settings.
- Ch66 slow-motion depiction is perception and narrative technique, not a supernatural time-flow change.
- Fang Yuan's promises to spare the hunters do not become binding narration. His lies are intentional.
- Wang Da's supposed death and his wife's supposed suicide are rumours. Later evidence revises them.
- The attacker mistakes Fang Zheng for Fang Yuan. Do not replace every Fang Yuan occurrence before the reveal.
- The academy guard named Wang Da is not established as the hunter family's Wang Da.
- Jiang Ya can be called 弟弟 and then more precisely 表弟. This is not proof of contradictory kinship.
- Fang Yuan's exam claim of personally collecting only ten-plus tusks is deliberately false. Actual23boars yield46fresh tusks before penalties. The earlier cache is not explicitly retrieved or submitted after the escape detour.
- The chief's public announcement of private teaching does not restart teaching that already began secretly.
- White Boar feeding every5days is compatible with daily activation; strength persists after acquisition.
- One-boar strength does not mean the boulder weighs oneboar. Rolling, lifting, traction and leverage differ.
- The three-hour Shadow Gu recovery never elapses on the page before its owner's death.
- Fang Zheng's seven-day coma does not conflict with Fang Yuan learning of Wang Da's death after three days.
- Ch80 confirms44% aperture volume and44−8=36, with stage potency1:2:4:8. No global aptitude recalculation is warranted.
- Fang Yuan's peak-quality refined essence does not make him peak-stage yet. His10–15day forecast remains a forecast.
'''
(r/'jotters'/'Batch-004-Preserve.md').write_text(preserve)
# Report.
report=f'''# Sequential close-reading review: Chapters 51–80
## Batch 004 | The Omniarch

**Completed:** Chapters 1–80 of 2,334. **Next unread:** Chapter 81. **Remaining:** 2,254 chapters.

This batch contains **{np:,} source paragraphs and {nc:,} paragraph characters**, read in fifteen complete, untruncated two-chapter sequences. Cumulative coverage is **{tp:,} paragraphs and {tc:,} characters**. Extraction alone was not counted. This is Chinese-source review, not a fresh bilingual or live-site audit.

**Publication status: HOLD.** All {len(ledger)} proposals remain unapproved and unapplied. The whole novel has not yet been read or temporally reconciled. No source, site or repository changes were made.

## 1. The strongest new chronology issue

Chapter 7 places the lesson in April, seven days after the awakening ceremony. Chapter 72 explicitly reaches late June, but its teacher describes the preceding training as a half-year:

> {src(72,26)['text']}

The existing short sequence does not supply six months of academy training. The working proposal is **“Over the past few months”**, not moving the April anchor or renaming the mid-year exam. A mid-year examination can occur in June even if the pupils entered partway through the year. This remains dependent on the opening chronology and all later recaps.

A second schedule issue predates this batch: Chapter 26 promised appointment of officers after three months, while Chapter 53 makes those appointments much earlier in the linked narrative. A changed schedule is possible, but no change is announced. The held repair would make appointment contingent on the qualifying results rather than a mandatory three-month wait. Do not invent an unseen multi-month gap.

## 2. The newly established sequence

| Section | Textual timing | Editorial treatment |
|---|---|---|
| Ch51–53 | Breakthrough morning; intervening breakthroughs; fifth-day stipend | Preserve sequence; query Chi Cheng's ordinal |
| Ch54–59 | Extortion, investigation and Jia Fu's departure that evening/night | Treat false testimony as testimony, not objective chronology |
| Ch60 | Wait over ten days after Jia Fu leaves; now sixth digging day; revelation over a month earlier | Broad intervals can coexist; avoid double-counting excavation days |
| Ch61–63 | Next morning lesson; evening exploration; next morning Gu selection; same evening inheritance | Ropework takes one hour and cautious traversal two hours |
| Ch64–70 | Later noon, hunting and household confrontation, later night hunt | Exact intervening day count not supplied everywhere |
| Ch71–73 | Mid-June, late June, then three days to the examination | Strong calendar anchors |
| Ch73–77 | Examination day, assassination, one-hour detour and late return | One afternoon; attacker's cooldown is not elapsed time |
| Ch78 | Death news three days later; another ten-plus days to upper stage | News date is not death date |
| Ch79–80 | Summer becomes autumn; later deep autumn, over two months after attack | Seasonal wording still needs later calendar reconciliation |

**The autumn watch item:** late June or early July plus two-plus months points roughly toward September, while Chapter 80 says 深秋. The novel has not established enough exact calendar detail to justify forcing this to three or four months. Retain both readings in the register. Do not silently impose Gregorian seasonal boundaries.

Fang Zheng's seven-day, seven-night coma is narrated retrospectively in Chapter 80. It does not conflict with Fang Yuan receiving confirmation of Wang Da's death three days after the attack. Similarly, the investigator's predicted arrival in two or three years is Fang Yuan's forecast, not a fixed appointment.

## 3. New active proposals

The table gives actions, not approvals. Full quotations, rationale and confidence are in the continuity jotter and cumulative ledger.

| ID | Location | Working action | Confidence / dependency |
|---|---|---|---|
'''
for v in new:report+=f"| {v['id']} | {refs(v['locations'])} | {v['recommended_action']} | {v['confidence']} {v['dependencies']} |\n"
report+='''
### Existing proposals updated

- **SR-024, Mo Chen/Mo Bei:** add Ch62pp5,83; Ch71p56; Ch73pp35,37. These are the pupil Mo Bei, not his grandfather. Ch63p13 independently confirms the pupil's Gu choice. Elder Mo Chen in the council must remain unchanged.
- **SR-017, recording duration:** Chapter60's later wait and excavation cannot be moved backward to lengthen the recording before Chapter46. They do not solve that earlier issue.
- **SR-018, caravan murder night:** Chapter56 reinforces the noon visit and following-evening dispute, and confirms rain on departure day. It does not supply a decisive new date resolving every “last night” reference.
- **SR-022, 45 versus44:** Chapter80 explicitly confirms44% full aperture volume. This supports the existing local proposal, not a global recalculation of aptitudes.

## 4. Resources: exact arithmetic versus missing balances

### New levy: conditional discrepancy

Chapter54 sets ordinary pupils at1stone, deputies at3each, and the head at8. With the previously established56 paying peers, that is:

**53 × 1 + 2 × 3 + 8 = 67 stones.**

Adding Fang Yuan's own ordinary stipend gives70. The quoted old total59 fits56+3, but the new stated68 fits neither consistent scope. It is a pupil's statement, so later headcounts and receipts must be checked before calling this an authoritative account or selecting67 versus70.

### Do not manufacture a full cash balance

The30-stone breakthrough reward is received; the10-stone head stipend is refused. The proposed hundred-plus jars of wine have no quoted deposit yet. Chapter64 gives ongoing food costs above2stones a day and total daily expenditure of at least5, but not a complete running account.

Fang Yuan kills23boars on examination day, giving46fresh tusks. A separate earlier cache exists, but its retrieval and submission are not narrated after the escape detour. Do not automatically add it to the bag handed over. Late arrival carries a four-tusk penalty. The final submitted count is not specified. His claim to have personally hunted only ten-plus tusks is intentional deception, not a replacement for the narrated23kills. The chief also announces a30-stone award. These facts do not establish an exact ending cash balance.

### Essence and feeding

Chapter70 leaves23high-stage essence points and describes the hunt as costing20. A full44-point starting reserve is not explicitly established. Do not automatically change23to24 or invent a missing1-point expense.

Chapter80 is different: the text explicitly supplies **44 − 8 = 36**. Its stage potency is **1:2:4:8**, so Moonlight activation costs **10 / 5 / 2.5 / 1.25** aperture-capacity points. Eight peak-stage points correspond to64initial-stage points. The correct equations need no repair.

White Boar Gu eats once per five days but can be activated daily for about fifteen minutes. Jade Skin Gu needs two liang of jade every ten days; two monkey eyes, each at least one liang, supply a feeding. Keep feeding intervals, activation intervals and permanent strength gains distinct.

## 5. Important non-errors and one withdrawn suspicion

**The toad buyer is Jia Jin Sheng.** I rechecked Chapter42pp58–63 before creating a proposal against Chapter55's witness. The transaction explicitly names Jia Jin Sheng. The working-memory suspicion that Jia Fu was the buyer was wrong. No correction is proposed.

Other distinctions that must survive editing:

- The guards in Chapters51–52 are wounded, not killed and resurrected.
- Fang Yuan lies about the Liquor Worm's origin and the murder. The Bamboo Gentleman test is deliberately defeated through Cicada suppression.
- Surveillance begun after his breakthrough is not, on its own, a real alibi for an earlier murder. Do not backdate surveillance to make the elders' conclusion valid.
- Wang Da's supposed death and his wife's suicide are rumours later overturned. His attack on the wrong twin is intentional narrative misidentification, not a blanket Fang Yuan→Fang Zheng correction.
- Four Gu discussed as feeding expenses and five Gu actually owned are compatible when Cicada is outside the expense subset.
- Jiang Ya being called 弟弟 and later 表弟 does not establish contradictory kinship. A live-English “brother” would need its own contextual check.
- The public announcement of Fang Zheng's private instruction does not mean his already ongoing lessons only begin then.
- Prior-life cultivation rates are explicitly altered by rebirth and training choices. Do not synchronise both branches mechanically.
- RevolutionThree independent wilderness travel is not the same claim as safe hunting on patrolled village outskirts.

## 6. Resumable files

- `Batch-004-Numbered-Chinese-Source.md`: actual complete Chinese text, with stable paragraph numbers.
- `jotters/Batch-004-Chronology.md`: chronology entries with exact quotations and evidence status.
- `jotters/Batch-004-Resources.md`: costs, essence, feeding and conditional calculations.
- `jotters/Batch-004-Continuity.md`: active proposals, rationale, confidence and dependencies.
- `jotters/Batch-004-Preserve.md`: safeguards against false corrections.
- `Sequential-Reading-and-Proposals.xlsx`: cumulative paragraph-level source and registers through80.
- `reading-progress.json` and `CURRENT-STATUS.md`: resumption checkpoint.

**Next planned reporting batch: Chapters81–110**, still read in smaller complete segments. No reading beyond80 is claimed by this report.
'''
(r/'Batch-004-Chapters-51-80.md').write_text(report)
# Workbook retains exact paragraph evidence and cross references.
w=Workbook();s=w.active;s.title='Read me'
for row in [('Sequential review','Status'),('Completed','1–80 of2334'),('Paragraphs',tp),('Source characters',tc),('Next unread',81),('Proposals',f'{len(ledger)} held; none approved/applied'),('Scope','Chinese-source review, not fresh liveEnglish validation'),('Global chronology','Not certified')]:s.append(row)
s=w.create_sheet('Chapter completion');s.append(['Chapter','File','Title','Paragraphs','Status'])
for c in chapters:s.append([c['corpus_chapter'],c['file'],c['title'],len(c['paragraphs']),'Full sequential read'])
s=w.create_sheet('Paragraph register');s.append(['Chapter','Paragraph','Status','Proposal IDs','Chronology IDs','Resource IDs','Chinese source'])
def ids(vs,k,c,p):return ', '.join(v['id'] for v in vs if any(z['chapter']==c and z['paragraph']==p for z in v[k]))
for c in chapters:
 for p,t in enumerate(c['paragraphs'],1):
  n=c['corpus_chapter'];s.append([n,p,'Read, not certified error-free',ids(ledger,'locations',n,p),ids(events,'sources',n,p),ids(resources,'sources',n,p),t])
s=w.create_sheet('Held proposals');s.append(['ID','Category','Locations','Working action','Rationale','Confidence / status','Dependencies','Publication','Approved','Applied','Latest update'])
for v in ledger:s.append([v['id'],v['category'],refs(v['locations']),v['recommended_action'],v.get('rationale','See earlier batch report and evidence'),v.get('confidence',v['status']),v['dependencies'],v['publication_status'],'No','No',v.get('batch004_update','')])
s=w.create_sheet('Chronology');s.append(['ID','Timeline','Kind','Event','Constraint','Evidence status','Dependencies','References','Exact source'])
for v in events:s.append([v['id'],v['timeline'],v['kind'],v['event'],v['constraint'],v['evidence_status'],v['dependency'],refs(v['sources']),'\n'.join(z['text'] for z in v['sources'])])
s=w.create_sheet('Resources');s.append(['ID','Kind','Entry','Calculation','Status','References','Exact source'])
for v in resources:s.append([v['id'],v['kind'],v['entry'],v['calculation'],v['status'],refs(v['sources']),'\n'.join(z['text'] for z in v['sources'])])
for s in w:
 s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
 for cell in s[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='163B4B')
 for row in s.iter_rows(min_row=2):
  for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
 for col in s.columns:s.column_dimensions[col[0].column_letter].width=min(65,max(14,max(len(str(x.value or '')) for x in col[:30])+2))
w.save(r/'Sequential-Reading-and-Proposals.xlsx')
code=(r.parent/'make_whole_novel_audit_docx.py').read_text().replace("root=Path(__file__).parent","root=Path('/home/user/editorial/sequential-review')").replace('Whole-Novel-Aperture-and-Travel-Audit','Batch-004-Chapters-51-80').replace('WHOLE-NOVEL SOURCE AUDIT','SEQUENTIAL REVIEW / CHAPTERS51–80').replace('The Omniarch: Aperture and Travel Source Audit','The Omniarch: Sequential Review, Chapters51–80').replace('Full-corpus topical screening with selective contextual verification','Thirty complete source chapters; all emendations held')
exec(compile(code,'docx-export','exec'),{})
(r/'CURRENT-STATUS.md').write_text(f'''# Sequential review checkpoint

Actual full sequential reading completed: **Chapters1–80**. Next unread: **81**. Next planned batch: **81–110**.

Batch004: {np}paragraphs, {nc}characters. Cumulative: {tp}paragraphs, {tc}characters. Remaining2254chapters. All{len(ledger)}proposals held, none approved/applied. Whole-novel chronology not verified.

Latest report: Batch-004-Chapters-51-80.md/.docx. Actual source: Batch-004-Numbered-Chinese-Source.md. Four clean jotters and cumulative workbook updated. No source/site/repository changes. No remote backup created. Do not rerun earlier batch generators, which would reset coverage.

Critical carry-forwards: SR-032 half-year training versus April/June; SR-033 three-month appointment announcement versus earlyappointment; SR-029 third/fourth ranking; SR-030 levy67/70 versus68. Ch80 deepautumn/>2months remains seasonalwatch, not approvedrepair. Ch70 startingreserve unknown, noautomatic23→24.

Correction: Chapter55 toadbuyer isJiaJinSheng, confirmedCh42pp58–63; priorworking-memoryJiaFu suspicionwithdrawn. Ch23 three-stoneadvance remains viableexplanationof20cash; noautomatic20→23. Deliberatelies, falsealibiclaims, mistakenidentity, rumours, forecasts and changedrebirthbranches remain separatelylabelled.
''')
assert len(chapters)==80 and len(ledger)==38 and all(not v['approved'] and not v['applied'] for v in ledger)
assert '—' not in report
print(dict(batch_paragraphs=np,batch_characters=nc,total_paragraphs=tp,total_characters=tc,proposals=len(ledger),chronology=len(events),resources=len(resources),next=81))
