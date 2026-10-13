from pathlib import Path
import json, shutil
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
r=Path(__file__).parent
load=lambda name:json.loads((r/name).read_text())
save=lambda name,x:(r/name).write_text(json.dumps(x,ensure_ascii=False,indent=2))
old=load('reading-progress.json'); batch=load('batch-002-source.json'); allchap=load('batch-001-source.json')+batch
by={x['corpus_chapter']:x for x in allchap}
h=r/'history';h.mkdir(exist_ok=True)
for name in ['reading-progress.json','proposal-ledger.json','Batch-001-Chapters-1-10.md','Batch-001-Chapters-1-10.docx']:
 if not (h/('before-batch002-'+name)).exists():shutil.copy2(r/name,h/('before-batch002-'+name))
ledger=load('proposal-ledger.json')
for v in ledger:
 v['publication_status']='HOLD: complete-corpus reading and global reconciliation outstanding'
 v['approved']=False;v['applied']=False
 if v['id']=='SR-002':
  v['batch002_update']='Ch11–20 supports short chronology and three-day rain interval; does not establish whole-novel consistency.'
  v['dependencies']='Continue collecting later chronological references through Chapter 2334; no release after a local batch.'
 if v['id']=='SR-003':
  v['batch002_update']='Ch18 p17 recalls aunt/uncle provision shortly after parental deaths; p19 confirms factual recollection. Supports continuity of care but gives no numerical fostering age.'
  if not any(l['chapter']==18 for l in v['locations']):v['locations'] += [dict(chapter=18,paragraph=p,file=by[18]['file']) for p in [17,19]]
new=[
 ('SR-012','resource accounting',[(18,84),(20,29)],'Investigate twelve spent versus six cumulative. Candidate if twelve survives: at least twelve in Ch20; alternative is six in Ch18 if later accounts support it. No direction finalised.','Definite local numerical conflict on the stated common basis; repair direction unresolved','Later balances, first-night cost, scope of this evening, and independent witnesses'),
 ('SR-013','source copy',[(12,19),(12,26)],'小儿 → 小二 in waiter context','High contextual confidence','Retain already-correct English; no independent witness verified'),
 ('SR-014','source copy',[(15,39)],'巨大的力量用来 → 巨大的力量涌来','Contextual source-copy candidate','Convey force striking him; no independent witness verified'),
 ('SR-015','source copy',[(16,9)],'能方才 → 可能方才; preserve inference rather than certainty','Contextual source-copy candidate','No independent witness verified'),
 ('SR-016','source copy',[(17,26)],'自己都不舍不得 → 自己都舍不得','High contextual confidence','Remove duplicated negative; retain unwillingness to spend')]
for id,kind,locs,action,status,dep in new:
 v=dict(id=id,category=kind,locations=[dict(chapter=c,paragraph=p,file=by[c]['file']) for c,p in locs],recommended_action=action,status=status,dependencies=dep,approved=False,applied=False,publication_status='HOLD: complete-corpus reading and global reconciliation outstanding')
 ledger=[x for x in ledger if x['id']!=id]+[v]
save('proposal-ledger.json',ledger)
progress={k:v for k,v in old.items() if k not in ['chapters']}
progress.update(sequential_chapters_completed=list(range(1,21)),paragraphs_read=sum(len(x['paragraphs']) for x in allchap),paragraph_characters_read=sum(sum(map(len,x['paragraphs'])) for x in allchap),next_chapter=21,remaining_chapters=2314,batch_report='Batch-002-Chapters-11-20.md',publication_policy='All emendations held for full 2334-chapter review and global reconciliation.',whole_novel_chronology_verified=False,scope='Full sequential Chinese source reading through Chapter 20. Not a fresh bilingual/live-site audit. Topical excerpts beyond Chapter 20 do not count as sequential completion.',batches=[dict(id='001',chapters=[1,10],paragraphs=715,characters=26627),dict(id='002',chapters=[11,20],paragraphs=750,characters=27937)],chapters=[dict(chapter=x['corpus_chapter'],file=x['file'],title=x['title'],paragraphs=len(x['paragraphs']),first_paragraph=1,last_paragraph=len(x['paragraphs'])) for x in allchap])
save('reading-progress.json',progress)
# No automated selection is used to claim reading: these sources were printed and read in full before this register was built.
events=[]
def event(id,era,kind,locs,description,constraint,certainty,dependency=''):
 events.append(dict(id=id,timeline=era,kind=kind,event=description,constraint=constraint,evidence_status=certainty,dependency=dependency,sources=[dict(chapter=c,paragraph=p,file=by[c]['file'],text=by[c]['paragraphs'][p-1]) for c,p in locs]))
event('T001','current life','anchor',[(1,78),(1,81),(1,84)],'Rebirth on eve of ceremony','Ceremony tomorrow; reversal stated as five hundred years','Explicit relative anchor; centuries not converted into exact age')
event('T002','current life','age',[(4,3),(9,4)],'Present age of brothers','Fifteen','Explicit present-age evidence','Do not infer birthday')
event('T003','shared childhood','retrospective',[(2,49),(2,50),(9,16),(18,17),(18,19)],'Orphaning and fostering','Orphaned three; aunt claims fostered from five for ten years; later childhood recollection connects aunt/uncle care with parental deaths','Conflicting/partly underdetermined start; SR-003 held','Later childhood references and witnesses')
event('T004','current life','elapsed interval',[(7,1),(7,4),(7,52)],'Week-later lesson and Gu selection','Ceremony seven days earlier; lesson leads to Gu selection','Explicit interval and instruction')
event('T005','current life','calendar',[(7,47)],'Arrival of fourth month','April/fourth month wording; no exact date','Explicit month, unspecified calendar-day position')
event('T006','current life','conflict',[(8,48),(8,49),(8,58)],'Opening month/week/ten-plus-days conflict','Competing elapsed descriptions in continuous sequence','SR-002 held','Whole-novel chronology')
event('T007','current life','elapsed interval',[(10,5),(10,6),(11,55),(13,26),(14,51)],'Three-day rain ends before discovery','Three days and nights; Ch11 approaching evening after rain','Multiple explicit local anchors','Analytical position about ceremony +10 days, not a canonical date')
event('T008','current life','activity duration',[(12,1),(12,33),(13,29),(16,55)],'Search and wine purchases','More than week elapsed; seven active/purchase days recalled','Different measured activities; do not equate automatically','Rain interruption may explain distinction; preserve wording')
event('T009','previous life','relative future event',[(13,9),(13,10),(16,39)],'Other Gu Master discovers legacy','A month later in previous life; exact referent not restated','Previous-life relative anchor','Do not transplant date to current-life discovery')
event('T010','historical clan account','retrospective narrative',[(13,1),(13,4),(13,7),(13,8)],'Received story of fourth leader and Flower Wine Monk','More than three hundred years ago; account explicitly questioned','Public story, not final omniscient history','Preserve reveal in Ch15–16')
event('T011','historical recorded encounter','threat',[(15,11),(16,2)],'Poison death threatened','Seven days if not treated','Conditional threat, not elapsed duration','Fourth leader dies during encounter')
event('T012','historical recorded encounter','proposed schedule',[(15,32)],'Weekly tribute demanded','30 jin petals and 3000 stones each week','Demand only; no payment established')
event('T013','current life','playback',[(16,4),(16,5)],'Historical recording loops','Second and third playback','One recorded historical event, multiple present replays')
event('T014','current life','scene continuity',[(16,53),(16,61),(17,8),(17,26)],'Return to inn from legacy','Same night as inn confrontation and wine purchase','Explicit same-evening recall')
event('T015','current life','overnight interval',[(17,57),(17,58),(17,60),(18,2),(18,30)],'Dawn and breakfast confrontation','Half-night initial refinement; Shen Cui visit yesterday','Explicit next-morning link')
event('T016','current life','forecast',[(17,69)],'Liquor Worm projected refinement time','Not within a month under current conditions; stronger-Gu suppression offered as exception','Conditional forecast, not elapsed month','Ch19–20 changes conditions')
event('T017','current life','future eligibility',[(18,72),(18,73),(18,74)],'Inheritance eligibility','About one year later, age sixteen, minimum Revolution1 middle-stage cultivation','Rule and family plans, not elapsed year','Track eventual actual inheritance event')
event('T018','current life','elapsed interval',[(18,79),(18,80),(18,84)],'Further refinement until backlash','Two further days; two days/nights with two hours rest per day','Explicit coarse interval; exact midnight boundary not given','Approximate ceremony +13-day region only')
event('T019','previous life','project duration',[(19,31)],'Preparation of Spring Autumn Cicada','Thirty years of preparation','Subperiod of longer life; not total age')
event('T020','cross-life transition','mechanism',[(19,33),(19,34),(19,36),(19,56),(19,57)],'Cicada survives rebirth and sleeps','Five-hundred-year reversal; bond persists, detection prevented by weakness','Stated explanation at this point in story','Track later revisions without retroactively changing viewpoint')
event('T021','current life','short interval',[(20,10),(20,26),(20,38),(20,40)],'Assisted refinement and academy visit','Liquor Worm moments; Moonlight <5 minutes; quarter-hour transition; night','Nested successive intervals, not a shared deadline')
event('T022','current life','resource over time',[(18,84),(20,29)],'Refinement expenditure conflict','Twelve in specified period versus six recalled cumulatively','SR-012 unresolved','Subsequent balances and meaning of this evening')
save('chronology-register.json',dict(scope='Sequential source evidence through Chapter 20; chronology unresolved; analytic day labels are not canon.',entries=events))
# Exact source copy for this batch.
lines=['# Numbered Chinese source: Volume 1, Chapters 11–20','Unemended EPUB paragraphs. Numbers match Batch 002 references.','']
for x in batch:
 lines += [f"## Chapter {x['corpus_chapter']}: {x['title']}",f"Source: `{x['file']}`",'']
 for i,t in enumerate(x['paragraphs'],1):lines += [f'**{i:03d}.** {t}','']
(r/'Batch-002-Numbered-Chinese-Source.md').write_text('\n'.join(lines))
# Add an explicit superseding hold to the earlier report, preserving an archive.
p=r/'Batch-001-Chapters-1-10.md';s=p.read_text()
banner='**SUPERSEDING STATUS: All proposals are on publication hold pending the complete 2,334-chapter reading and global reconciliation. Earlier local readiness assessments below do not authorise implementation. See Batch 002 and the cumulative ledger.**\n\n'
if banner not in s:s=s.replace('## Reading status\n\n','## Reading status\n\n'+banner);p.write_text(s)
# Cumulative workbook: raw evidence stays distinct from editorial conclusions.
w=Workbook();a=w.active;a.title='Read me'
for row in [ ['Sequential review','Status'],['Completed chapters','1–20 of 2334'],['Read paragraphs',1465],['Source characters',54564],['Next chapter',21],['Publication policy','HOLD all emendations pending full-corpus review and global reconciliation'],['Scope','Chinese source only; no fresh bilingual/public-site validation'],['Per-paragraph status','Read does not mean certified error-free; no proposal recorded is not a correctness guarantee'],['Later topical excerpts','Do not count as sequentially completed chapters'],['Previous batch report','Historical local assessments superseded by global hold'] ]:a.append(row)
a=w.create_sheet('Chapter completion');a.append(['Chapter','Source file','Title','Paragraphs','Status'])
for x in allchap:a.append([x['corpus_chapter'],x['file'],x['title'],len(x['paragraphs']),'Full source sequence read'])
b=w.create_sheet('Paragraph register');b.append(['Chapter','Paragraph','Status','Proposal IDs','Chronology IDs','Chinese source'])
for x in allchap:
 for i,t in enumerate(x['paragraphs'],1):
  ids=[v['id'] for v in ledger if any(l['chapter']==x['corpus_chapter'] and l['paragraph']==i for l in v['locations'])]
  tids=[v['id'] for v in events if any(l['chapter']==x['corpus_chapter'] and l['paragraph']==i for l in v['sources'])]
  b.append([x['corpus_chapter'],i,'Read; proposal recorded' if ids else 'Read; no proposal recorded in this batch',', '.join(ids),', '.join(tids),t])
c=w.create_sheet('Active proposals');c.append(['ID','Category','Locations','Recommended action','Evidence status','Dependencies','Publication status','Approved','Applied'])
for v in ledger:c.append([v['id'],v['category'],'; '.join(f"Ch{l['chapter']} p{l['paragraph']}" for l in v['locations']),v['recommended_action'],v['status'],v['dependencies'],v['publication_status'],'No','No'])
d=w.create_sheet('Chronology');d.append(['ID','Timeline','Kind','Event','Constraint','Evidence status','Dependencies','References','Exact source'])
for v in events:d.append([v['id'],v['timeline'],v['kind'],v['event'],v['constraint'],v['evidence_status'],v['dependency'],'; '.join(f"Ch{l['chapter']} p{l['paragraph']}" for l in v['sources']),'\n'.join(f"Ch{l['chapter']} p{l['paragraph']}: {l['text']}" for l in v['sources'])])
for sh in w:
 sh.freeze_panes='A2';sh.auto_filter.ref=sh.dimensions
 for cell in sh[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='163B4B')
 for row in sh.iter_rows(min_row=2):
  for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
 for col in sh.columns:sh.column_dimensions[col[0].column_letter].width=min(65,max(14,max(len(str(x.value or '')) for x in col[:30])+2))
w.save(r/'Sequential-Reading-and-Proposals.xlsx')
# Established Word renderer, retargeted to each report.
base=(r.parent/'make_whole_novel_audit_docx.py').read_text()
for stem,scope in [('Batch-001-Chapters-1-10','1–10'),('Batch-002-Chapters-11-20','11–20')]:
 code=base.replace("root=Path(__file__).parent", "root=Path('/home/user/editorial/sequential-review')")
 code=code.replace('Whole-Novel-Aperture-and-Travel-Audit',stem).replace('WHOLE-NOVEL SOURCE AUDIT','SEQUENTIAL SOURCE REVIEW / CHAPTERS '+scope).replace('The Omniarch: Aperture and Travel Source Audit','The Omniarch: Sequential Review, Chapters '+scope).replace('Full-corpus topical screening with selective contextual verification','Complete source reading of named chapters; all proposals held for global reconciliation')
 exec(compile(code,'docx-export','exec'), {})
# Structural and coverage checks.
assert [x['corpus_chapter'] for x in allchap]==list(range(1,21))
assert progress['paragraphs_read']==1465 and progress['paragraph_characters_read']==54564
assert all(not x['approved'] and not x['applied'] for x in ledger)
assert len({x['id'] for x in ledger})==len(ledger)==16
assert all('—' not in (r/f).read_text() for f in ['Batch-001-Chapters-1-10.md','Batch-002-Chapters-11-20.md'])
print('Verified:',len(allchap),'chapters;',1465,'paragraphs;',len(events),'chronology entries;',len(ledger),'held proposals. Next: 21.')
