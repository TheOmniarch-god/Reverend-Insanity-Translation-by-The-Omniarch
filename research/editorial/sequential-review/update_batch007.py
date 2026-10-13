from pathlib import Path
import json,shutil,re
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
r=Path('/home/user/editorial/sequential-review')
def load(n):return json.loads((r/n).read_text())
def save(n,v):(r/n).write_text(json.dumps(v,ensure_ascii=False,indent=2))
for n in ['reading-progress.json','proposal-ledger.json','chronology-register.json','resource-register.json']:
 d=r/'history'/('before-batch007-'+n)
 if not d.exists():shutil.copy2(r/n,d)
chapters=sum([load(f'batch-{i:03d}-source.json') for i in range(1,8)],[]);by={c['corpus_chapter']:c for c in chapters};batch=chapters[140:160]
def src(c,p):return dict(chapter=c,paragraph=p,file=by[c]['file'],text=by[c]['paragraphs'][p-1])
def ss(rr):return [src(*x) for x in rr]
def loc(rr):return [{k:v for k,v in z.items() if k!='text'} for z in ss(rr)]
def refs(z):return '; '.join(f"Ch{x['chapter']} p{x['paragraph']}" for x in z)
def reader(c,p,t):
 if c==151 and (31<=p<=67 or p in [98,99]):return '[Passage omitted from the reader-facing copy: sexualized material involving a minor. Original paragraph numbering retained; original EPUB unchanged. See the non-graphic editorial discussion.]'
 return t
ledger=load('history/before-batch007-proposal-ledger.json');new=[]
updates={
'SR-001':'Ch159p41 explicitly changes44%→42% because of the breakthrough method. Preserve44% before that consequence and42% afterwards; this is not a retroactive aperture repair.',
'SR-007':'Ch152p48 consumes38percentage points;159p41 loss is2percentage points,44→42. Ch156p60 MoonShadow suppression uses capacity points, not a relative percentage of available essence.',
'SR-021':'159p41 remainsGradeC at42%; no blanket recalculation of earlier aptitude scenes.',
'SR-037':'Add Ch160p62: 王二→王大. The poisonGu belonged to the elder brother who became a demonic GuMaster, not the younger mortal hunter.',
'SR-038':'Ch160p61 again uses爱别离, strengthening the proposed consistent name.',
'SR-054':'Ch145p10 explicitly callsJuly hot summer, strengthening the internal conflict with116Julyearlyautumn. Separate fromSR-065 month regression. Hold any final choice of seasonal wording.',
'SR-060':'Still unresolved. Ch155’s firstRevolution3Gu language also excludes a consumedGu unless qualified; do not use that as proof of IceMuscle’s disposal.'}
for v in ledger:
 if v['id'] in updates:v['batch007_update']=updates[v['id']]
 if v['id']=='SR-037':v['locations']+=loc([(160,62)])
def prop(i,cat,rr,action,reason,conf,dep):
 v=dict(id=f'SR-{i:03d}',category=cat,locations=loc(rr),recommended_action=action,rationale=reason,confidence=conf,status='Active provisional proposal; no approval',dependencies=dep,approved=False,applied=False,publication_status='HOLD: full2334-chapter reading and global reconciliation outstanding');ledger.append(v);new.append(v);return v
prop(62,'character identity',[(141,67),(142,75),(142,76)],'Ch142p76 白凝冰所化的巨大树精→古月青书所化的巨大树精.','QingShu is the transformed tree spirit attacking intruders; BaiNingBing is its prisoner. The same paragraph’s explanation names QingShu.','High.','Local source-copy repair; independent witness not checked.')
prop(63,'destroyed Gu / narrative tense',[(139,24),(140,23),(140,25),(143,35),(143,36),(143,37),(143,38)],'Frame143pp35–36 explicitly as retrospective: 已经自爆的霜妖蛊，本是白凝冰身上最有价值的蛊虫… Then143p37: 如今价值最高的则是蓝鸟冰棺蛊… English: The destroyed Frost Demon Gu had been his most valuable Gu; of those remaining, Blue Bird Ice Coffin Gu was now the most valuable.','FrostDemon was destroyed in139 and its loss confirmed140. It cannot remain an available target for Plunder in143. Preserve the explanation of its former value, but repair the present-tense inventory implication.','High continuity conflict; high minimal tense repair.','Do not resurrect FrostDemon or erase its history. Distinct from unresolvedIceMuscle inventory SR-060.')
prop(64,'Gu appearance',[(143,18),(143,19),(144,69)],'Preferred Ch143p18 一只灰色的甲虫→一个灰白色的方块. English: a grey-white block. Retain the detailed dice-like Stone Aperture description in144.','The same capturedGu is a grey beetle in143 but a hard square die-like object in144. No transformation or refinement upgrade is narrated between them.','High discrepancy; moderate preferred description.','Later StoneAperture descriptions or independent witness may determine direction; no invented metamorphosis.')
prop(65,'month regression',[(131,78),(132,63),(143,68),(144,72),(145,10),(145,28),(149,41)],'Working candidate Ch145p10 七月→八月; English: August. Do not implement until the late-August ThunderCrown forecast and the global season grid are reconciled.','LateJuly in131 is followed by8–9days in132, the funeral nextday143, and subsequent days before145’s upper-stage state;145 nevertheless labels the presentJuly again.','High accumulated-date tension; moderate exact replacement.','No exact day inlateJuly; combat and advancement transitions have open intervals. Separate month regression fromSR-054 seasonal naming. Preserve149prior-life lateAugust forecast.')
prop(66,'auction outcome continuity',[(111,46),(111,49),(123,74),(123,75),(147,83),(147,84)],'Candidate147p84: 之前，她曾想从商队处买一只酒虫给孙女药乐用，却没能如愿。她自己还需要一只酒虫作为材料… English: She had tried to buy one for her granddaughter from the caravan, but had failed; she also needed one herself.','111awards the auctioned LiquorWorm toFangYuan.123still seeks a worm for the granddaughter.147retrospectively claims a completed caravan purchase for her without supplying an alternative transaction.','High documented outcome mismatch; moderate correction.','An unreported second purchase could reconcile it. Do not invent that purchase or rewrite111’s award.')
prop(67,'cohort age',[(4,3),(110,16),(123,74),(124,57),(151,8),(159,96)],'Conditional candidate151p8 十五岁的少女→十六岁的少女. English: a sixteen-year-old girl.','The academy intake is15; YaoLe enters one cohort afterFangYuan. A January boundary passes before her death, whileFangYuan is explicitly17 in159. The nominal-age sequence therefore points to16, not15.','High cohort query; moderate exact-age repair.','Confirm universal intake age and nominal versus birthday age. Do not invent a birth date or use this to settle BaiNingBing’s age.')
prop(68,'attack range units',[(106,64),(142,87),(156,52),(157,10)],'Preferred local candidates: 十步→十米 for Moonglow’s range at142p87 and GoldenMoon’s unchanged range at156p52. Keep the enemy’s20-plus-pace distance in142 unchanged.','106explicitly fixes Moonglow’s unchanged range at10metres.142calls it10paces;156retains10paces for its successor;157again uses10metres for BloodMoon.','High inconsistent units; moderate repair direction.','No general conversion ofpaces,li,travel distances or allGu ranges. Distinguish effective range from maximum flight.')
prop(69,'range multiplier',[(106,65),(156,40)],'Ch156p40 是月光蛊的一倍→是月光蛊的两倍. English: twice Moonlight Gu’s range.','106states an increase ofonefold, to20metres from10.156’s “is one times” does not express that advantage.','High, with numerical corroboration.','Retain106’s增长一倍, which correctly means a100% increase. Do not globally replace一倍.')
prop(70,'ongoing versus future wave',[(152,69),(154,63)],'Ch154p63 狼潮在即→狼潮当前. English: with the wolf tide upon us / during the wolf tide.','ChiLian argues against changing leadership during a severe ongoing tide, after a three-day siege; imminent wording misplaces the overall crisis.','High context; moderate wording need.','If specifically forecasting a fresh attack, retain that narrower meaning explicitly. No new calendar interval inferred.')
prop(71,'first RevolutionThree Gu scope',[(152,35),(152,36),(152,42),(155,73),(155,78),(155,79),(155,80)],'Qualify155p78 as first lasting RevolutionThree equipmentGu: 第一只用于长期作战的三转蛊虫. English: His first RevolutionThree Gu for continued combat use had come directly from the clan.','The RevolutionThree consumable used in152 predates LightningWings. The intended equipment-order contrast is reasonable, but literal first-everRevolutionThree ownership is not.','High literal scope issue; moderate editorial need.','Keep consumables distinct from permanent equipment. No change to clan reward or actual fusion order.')
# Only locations and non-graphic descriptions of the appended notes are used in the editorial record.
prop(72,'paratext provenance',[(151,97),(151,98),(151,99)],'Separate the author-attributed note and compiler commentary from the narrative. Label their provenance unverified; do not incorporate the compiler’s conjectures into the story, chronology or fusion recipe.','The EPUB explicitly switches to attributed notes after paragraph96. The compiler speculates beyond the narrated events; an attribution inside this witness is not independent proof of authorial authenticity.','High narrative/paratext distinction; authenticity unresolved.','Original EPUB retained. Reader-facing sexualized passages involving a minor are omitted with paragraph markers, not silently rewritten. Another authenticated witness is needed.')
fixes=[(141,39,'他最终他将','他最终将'),(142,50,'熊力失去的地方','熊力死去的地方'),(143,91,'沟动了','勾动了'),(144,34,'望之脑后','忘之脑后'),(145,66,'征兆复出','征召复出'),(146,14,'这卖家','这卖价'),(146,19,'七十二快元石','七十二块元石'),(146,38,'冷不住发寒','忍不住发寒'),(147,56,'高阶买了','高价买了'),(149,98,'傍晚十分','傍晚时分'),(150,75,'关紧时刻','关键时刻'),(150,77,'乍缝剧变','乍逢剧变'),(150,84,'其实普通的野兽','岂是普通的野兽'),(150,89,'这是古月药姬','这时古月药姬'),(151,11,'因为艘向奶奶','因为我向奶奶'),(151,28,'张开欲言','张口欲言'),(151,79,'丁当跳出来','定当跳出来'),(151,81,'唱一声佛诺','唱一声佛号'),(152,7,'就得要蛊虫操纵驭兽蛊','就得要蛊师操纵驭兽蛊'),(152,34,'重现陷入','重新陷入'),(152,44,'气流从上而下','水流从上而下'),(152,63,'清楚掉','清除掉'),(154,22,'赤城惊疑','赤练惊疑'),(154,61,'控制不出','控制不住'),(154,67,'然而仰头','然后仰头'),(154,75,'操练这么多天','操劳这么多天'),(154,87,'方源勉强','方源面前'),(155,10,'找招呼','打招呼'),(155,20,'泛泛飞','泛泛'),(155,29,'那么秘密','那个秘密'),(155,32,'拍桌子的事后','拍桌子的时候'),(157,76,'自主主张','自作主张'),(158,34,'是家族屹立不倒','使家族屹立不倒'),(158,44,'再向药脉','在向药脉'),(158,67,'他们做了','他们坐了'),(159,24,'家老么','家老们'),(159,25,'倒并意外','倒并不意外'),(159,26,'也再所难免','也在所难免'),(159,67,'以他的能力','以她的能力'),(159,70,'若正要被','若真要被'),(159,79,'和方源一眼','和方源一样'),(159,103,'处于真心实意','出于真心实意'),(160,7,'具有身上寄居着','身上寄居着'),(160,58,'珍惜蛊虫','珍稀蛊虫')]
for c,p,a,b in fixes:assert a in src(c,p)['text'],(c,p,a)
v=prop(73,'source-copy bundle',[(c,p) for c,p,a,b in fixes],'Apply only the separately located minimal substitutions in copy_items after approval; no global replacements.','Local miskeys, duplicates, mistaken actor names, pronouns and malformed phrases. The full source remains unchanged.','High contextual confidence overall; individual items provisional.','Independent witnesses not checked. The quantity ofcopy items is separate fromproposal-record count.')
v['copy_items']=[dict(chapter=c,paragraph=p,original=a,replacement=b,status='Held; unapproved/unapplied') for c,p,a,b in fixes]
save('proposal-ledger.json',ledger)
tdoc=load('history/before-batch007-chronology-register.json');events=tdoc['entries'];te=len(events)
def event(timeline,kind,title,constraint,status,rr,dep='Full-corpus reconciliation outstanding.'):
 events.append(dict(id=f'T{len(events)+1:03d}',timeline=timeline,kind=kind,event=title,constraint=constraint,evidence_status=status,dependency=dep,sources=ss(rr)))
event('current life','continuous battle','QingShu death and next-day funeral','141–143continue samebattle;143p45deathconfirmed;funeralnextday;Baiawake10+hours.','No deathat140endpoint; nowconfirmed. Hoursnotdays.',[(142,42),(143,45),(143,68),(143,90)])
event('current life / recollection','age constraints','YaoLe and cohort','Standardintake15;YaoLeonecohortbehindFY;Januarypasses;151callsher15;159FY17.','Conditional16candidate;Bai143twenty-plus-yearsquestionhasnoverifiedexact-agecontradiction.',[(4,3),(123,74),(124,57),(151,8),(159,96)],'SR-067; no birthday invented.')
event('current life','month regression','Return to July','131lateJuly→1328–9days→143nextday→145manydays/upperstage, yet145July.','Monthsequencequery distinctfromseasonlabelquery.',[(131,78),(132,63),(143,68),(145,10),(145,28)],'SR-065 andSR-054.')
event('current life','natural promotion / relic use','Upper then peak','145uppercultivation/peakquality,Relicunused;146intendeduse;148peakcultivationconfirmed.','Actualusefallsininterveninginterval;dontadvanceat145knock.',[(145,25),(145,28),(145,41),(145,55),(146,50),(148,39)])
event('current life','quoted delivery / elapsed interval','Leaf contract and rescue','5leavesdeliverypromised3days;>halfmonthlaterrescue;Xiong6day/nightthirdsiege.','Promiseisnotshownfulfilment. Differentvillageseigecounts.',[(146,22),(146,23),(146,58),(147,28),(147,30)])
event('current life / prior-life forecast','return and raid','Reinforcements turn back','149returns15minafterdeparture;GuYuesecondraid;priorlifeThunderCrownlateAugust.','QingShu’scurrentdeathprecedesoldlifedeath; explicitbranchchange.',[(149,2),(149,30),(149,41),(149,42)])
event('current life','uncertain casualty observation','XiongJiaoMan struck','Thrownintobuilding,no furthermovementreported.','Notconfirmeddead. Do not manufacture a double-death contradiction from unconsciousness or off-screen status.',[(150,94)])
event('current life','captivity / imagined duration','YaoLe disappearance','Victimagines7–8dayscaptivity;actualintervalnotgiven;152three-dayabsenceandsearch.','Imaginedpunishmentnotelapsedtime; non-graphicrecordonly.',[(151,15),(152,69),(152,96)])
event('current life','fusion and breakthrough','RevolutionThree achieved','152halfshichen thenhalfke amidprocess;38pointsconsumed;repeatedpurges.','Approx1hourand7.5minutes underconventionalunits,notwholeprocesstotal.',[(152,1),(152,18),(152,48),(152,50),(152,60)])
event('current life / false testimony','alibi','ChiLian’s corroboration','153claims3daysretreatatChiLian;154revealslastnightblackmailalreadyRevolutionThree.','Intentionallyfalsealibi, notcontradictoryomniscienttimeline.',[(153,72),(153,79),(153,84),(154,12),(154,21)])
event('current life / mistaken inference','retrospective sponsorship','MoChen’s interpretations','Servantkilling2–3yearsago;imaginedChiinvestmentyear-plusago.','Murderrecapfitsopening→age17;investmentispoliticalmisreading.',[(153,16),(154,3),(159,96)])
event('current life','borrowing / fusion / access','Three days then next-day library','Halfshichennegotiation;3dayslaterSkyCanopy;recipevisit15minutes;nextday5+10minutes.','No claim allGu acquiredonsameday; dailyaccesscapnotrare-vaultpolicy.',[(155,63),(155,66),(155,74),(156,30),(156,68),(157,1),(157,27)])
event('current life','stipend and transfer','One day’s political transaction','300stipendreceived157→all300gifted158;NineLeaftradedforrewardtoken.','Samefunds,nottwoindependentpaymentsfromclan.',[(157,76),(158,38),(158,79),(158,90)])
event('current life','later raids / periodic mechanic','Near-ten attacks and BloodMoon cycle','159near10sieges,unspecifiedtimegap;BloodMoonmonthlyfewdaysat1/3attack.','Fourwolves159newraid,notfourdeadwolves152revived;low-powerdaysnotyetidentified.',[(152,69),(157,13),(159,2),(159,24),(160,24)])
event('current life','capacity consequence','Forty-four to forty-two','AfterbreakthroughresidueandCleanWateruse,159capacity42.','Explicit2percentage-pointloss,notretrospectiveerroror2%relative.',[(152,52),(159,39),(159,41),(159,43)],'SR-001/SR-007 updates.')
tdoc.update(scope='Sequential chronology throughChapter160; full-novel chronology not certified.',entries=events);save('chronology-register.json',tdoc)
rdoc=load('history/before-batch007-resource-register.json');resources=rdoc['entries'];re0=len(resources)
def res(kind,entry,calc,status,rr):resources.append(dict(id=f'R{len(resources)+1:03d}',kind=kind,entry=entry,calculation=calc,status=status,sources=ss(rr)))
res('loot and disposal','ManShigroupGu recoveredandsurrendered;Xiongcorpsesmissing.','Unknownnumber/totalmerit;do notbookBrownBearBase.','FiveBaiMasterskilledseparately;Guownershipnotautomaticallytransferredbydeath.',[(142,46),(142,49),(142,50),(142,55),(144,78)])
res('Plunder proceeds','RedIronRelic,StoneAperture,WaterShieldtakenfromBai.','Threecaptures;FrostDemonalreadydestroyednotavailable.','Firsttwoinstantlyrefined;allthreecountasacquired. StoneApertureNOTused.',[(143,10),(143,17),(143,18),(143,19),(143,39),(144,77)])
res('leaf sales','5at72paid;standard70laterreported80.','5×72=360;68/70earlierquotesnotamountpaid.','Deliverypromised3days;futureorderreceipt/wholecashbalanceunknown.',[(146,13),(146,17),(146,19),(146,22),(146,25),(146,42),(147,54)])
res('rescue price','OneBearEnslavementGu PLUS700stones.','Two components,notGuvalued700.','Acceptedterms;Guconveyed/refined;700alsoattestedbyrecipientreport.',[(147,13),(147,16),(147,23),(147,52),(147,78)])
res('breakthrough expense','RevolutionThreeconsumablecreatedandused;1000+stones.','44−38=6capacitypointsbeforeexpulsion;bodycapacitylater42.','OtherrecipeGuconsumed;YaoLeGuappropriatedbutcount/namesunknown. Noexactinventorytotalafterthis.',[(151,7),(152,32),(152,36),(152,42),(152,48),(152,51),(159,41)])
res('debt and stipend','3000loan+CleanWaterloan;twoIOUs.','Weekly100peace/300wartime;300actualfirstpaymentlatergiftedaway.','Loansnotearnings;repaymentnotnarrated.',[(155,27),(155,34),(155,41),(155,43),(155,63),(155,64),(157,76),(158,79)])
res('RevolutionThree equipment','LightningWingsfreeclanreward;SkyCanopysecondattempt.','WaterShielddiesonfailure;replacementwaterGu+WhiteJade→SkyCanopy.','Replacementidentityandpriceunspecified;firstlastingGu≠firstconsumable.',[(155,73),(155,74),(155,75),(155,76),(155,77),(155,78),(155,80)])
res('range and suppression','Moonglow10m;MoonScar20m;MoonShadowRevolution3/4/5 suppression60/30/15points.','MoonScar2×base;rank3capacity44whollysuppressedby60pointcap.','Do notcalculate44×40%remaining;rangeunitqueriesheld.',[(106,64),(106,65),(142,87),(156,40),(156,52),(156,60),(157,10)])
res('BloodMoon fusion','Moonglow+BloodQi→BloodMoon;secondattemptsuccess159.','AttackmonthlydropsTO1/3;ordinarybleedingusefulinlongbeastfights.','Firstfailedfusiondoesnotspecifywhichinputdied;noinventedreplacementpurchase.',[(157,4),(157,13),(159,31),(159,32),(159,33),(160,21),(160,24)])
res('ownership transfer and token','NineLeafactuallytransferredforrewardtoken;TushitaFlowerselected160.','NoNineLeafafter158;storagegapfilledat160.','Conditionalreverseexchangeunused;ThreeStepGrassstillnotbought.',[(148,9),(158,29),(158,38),(158,39),(160,66),(160,68),(160,69)])
res('capacity and cleanse','CleanWaterused;residueremovedbutaptitude42.','44−2=42;qualityRevolutionThreeinitialsilver.','FourFlavoursstillownedbutdoesnotrefineRevolutionThree;old44unchanged.',[(159,35),(159,39),(159,41),(159,46)])
rdoc.update(scope='Resources throughChapter160; exact total inventory withheld where capturedGu are unenumerated.',entries=resources);save('resource-register.json',rdoc)
np=sum(len(c['paragraphs']) for c in batch);nc=sum(sum(map(len,c['paragraphs'])) for c in batch);tp=sum(len(c['paragraphs']) for c in chapters);tc=sum(sum(map(len,c['paragraphs'])) for c in chapters)
p=load('history/before-batch007-reading-progress.json');p.update(sequential_chapters_completed=list(range(1,161)),paragraphs_read=tp,paragraph_characters_read=tc,next_chapter=161,remaining_chapters=2174,batch_report='Batch-007-Chapters-141-160.md',scope='Full sequential Chinese-source reading throughChapter160. Counts include EPUB paratext; not a fresh bilingual/live-site audit.',whole_novel_chronology_verified=False,chapters=[dict(chapter=c['corpus_chapter'],file=c['file'],title=c['title'],paragraphs=len(c['paragraphs']),first_paragraph=1,last_paragraph=len(c['paragraphs'])) for c in chapters]);p['batches'].append(dict(id='007',chapters=[141,160],paragraphs=np,characters=nc,reading_segments='Ten complete untruncated two-chapter outputs',note='Dense20-chapter batch. Ch151pp97–99 are separator and appended notes, not narrative.'));save('reading-progress.json',p)
save('batch-007-reading-log.json',dict(status='Complete',endpoint=160,next_unread=161,segments=[dict(chapters=[c,c+1],status='Full paragraph sequences printed and read; untruncated',paragraphs=sum(len(by[k]['paragraphs']) for k in [c,c+1])) for c in range(141,161,2)],note='Extraction alone not completion. No invented timestamps. Reading-copy omissions do not mean unread source.'))
lines=['# Numbered Chinese reading copy: Chapters141–160','Original paragraph numbering retained. Source text is unamended except explicitly marked reader-facing omissions in Ch151: sexualized passages involving a minor are not reproduced. The original EPUB is unchanged. Ch151pp97–99 are paratext, not narrative.','']
for c in batch:
 lines += [f"## Chapter {c['corpus_chapter']}: {c['title']}",f"Source: `{c['file']}`",'']
 for i,t in enumerate(c['paragraphs'],1):lines += [f'**{i:03d}.** {reader(c["corpus_chapter"],i,t)}','']
(r/'Batch-007-Numbered-Chinese-Source.md').write_text('\n'.join(lines))
for name,vs in [('Chronology',events[te:]),('Resources',resources[re0:])]:
 lines=[f'# Batch007: {name} jotter','Full source read throughChapter160; proposals remain held.','']
 for v in vs:
  lines += [f"## {v['id']}: {v.get('event',v['kind'])}"]
  for k in ['timeline','constraint','evidence_status','dependency','entry','calculation','status']:
   if k in v:lines += [f"**{k}:** {v[k]}"]
  for s in v['sources']:lines += [f"\n**Ch{s['chapter']} p{s['paragraph']}**",'> '+reader(s['chapter'],s['paragraph'],s['text']),'']
 (r/'jotters'/f'Batch-007-{name}.md').write_text('\n'.join(lines))
lines=['# Batch007: Continuity and held proposals',f'12new records; SR-073 contains{len(fixes)} individually located copy items. All unapproved/unapplied.','']
for v in new:
 lines += [f"## {v['id']}: {v['category']}"]
 for k in ['recommended_action','rationale','confidence','dependencies','publication_status']:lines += [f"**{k}:** {v[k]}"]
 for s in v['locations']:lines += [f"\n**Ch{s['chapter']} p{s['paragraph']}**",'> '+reader(s['chapter'],s['paragraph'],src(s['chapter'],s['paragraph'])['text'])]
 for z in v.get('copy_items',[]):lines += [f"- Ch{z['chapter']} p{z['paragraph']}: {z['original']} → {z['replacement']} (held)"]
 lines += ['']
lines += ['## Existing records updated']
for k,v in updates.items():lines += [f'**{k}:** {v}','']
lines += ['## Queries deliberately not upgraded to errors','BaiNingBing’s twenty-plus-year question lacks a contradictory exact age in the completed reading. Do not invent an18-year-old baseline.','XiongJiaoMan’s lack of movement after the impact is not an explicit death.','FangZheng’s middle→upper advancement fits interveningweeks; no fixedmonths requirement has been proved.','Ch151 philosophical and religious exposition is attributed narrative interpretation, not independently verified doctrine. Do not publish it as authoritative Buddhist teaching.','The FlowerWineMonk injury mystery is explicitly raised by FangYuan in156; do not repair it before its narrative resolution.']
(r/'jotters'/'Batch-007-Continuity.md').write_text('\n'.join(lines))
(r/'jotters'/'Batch-007-Preserve.md').write_text('''# Batch007: Preserve, do not repair

- Retain44% capacity before the new consequence,42% after159. CleanWater removes residue but does not restore the lost two percentage points.
- Preserve natural middle→upper cultivation, followed by the second Relic’s implied use to peak. Peak-quality essence does not itself mean peak cultivation.
- Preserve the false ChiLian alibi. The last-night blackmail explains his testimony; it does not relocate the breakthrough to his house.
- QingShu’s death is confirmed143, and the funeral is nextday. Bai’s wakeful period is10-plus hours.
- XiongJiaoMan is struck and unresponsive in150, not declared dead. An unnamed rescued classmate159 is not YaoLe returning to life.
- YaoLe’s imagined7–8day punishment is not an elapsed interval. Her stolenGu are not enumerated; do not invent an exact inventory total.
- Separate151’s narrative from its attributed and compiler notes. Compiler conjecture cannot establish additional events, births or recipe components. OriginalEPUB remains unchanged; restricted passages are marked omissions in the reading copy.
- A halfshichen is conventionally about an hour, unlike halfanordinaryhour; halfake is about7.5minutes. Local clocks are not the totalbreakthrough duration.
- Xiong’s thirdsiege and GuYue’s secondsiege are different counts. Threevisiblewolfkings plusa fourth concealed one is a reveal. Fournewwolves159 do not resurrect thefour killed152.
- The rescue costs oneBearEnslavementGu plus700stones, not a Gu pricedat700. Leaves5×72=360paid; three-daydelivery is a promise.
- 3000borrowedstones andCleanWaterare debts, not gifts. The300stipend is subsequently giftedto theuncle, not a second stipendpayment.
- WaterShield is destroyed in thefirstSkyCanopy attempt; anunidentified replacement andWhiteJade produceSkyCanopy onthesecond. FailedBloodMoon inputs are not specified.
- NineLeaf is retained until its actual158transfer. TushitaFlower selected160fillsstorage; ThreeStepGrass is stillnotowned, ChainsawCentipede stillnotcaptured.
- MonthlyBloodMoon attack dropsTOone-third, notBYone-third. Its success inalongbeastfight doesnotcontradict itslimitedvalue inshortGuMasterduels.
- Recipevault15-minuteaccess doesnotautomatically govern therare-Gu vault. Four/five-rank recipesremainrestricted; permittedborrowedGu operationdoesnotmean ownership.
- FlowerWineMonk’s injuries are adeliberately reopened mystery. ClanRevolutionFive holdings areFangYuan’s inference, not anobserved inventory.
- Do not convert everypace orli tometres. Local10pace/10metre attack-range repairs remainheld.
''')
report=f'''# Sequential close-reading review: Chapters 141–160
## Batch 007 | The Omniarch

**Completed: Chapters1–160 of2,334. Next unread:161. Remaining:2,174.**

This is a deliberately smaller, dense20-chapter batch, read in ten complete, untruncated two-chapter sequences: **{np:,} EPUB paragraphs / {nc:,} characters**. Cumulative coverage is **{tp:,} paragraphs / {tc:,} characters**. Extraction alone was not counted as reading.

**73 proposal records:72 active and held, one withdrawn.** Twelve records are new; SR-073 bundles{len(fixes)} individually located copy corrections. Proposal records and individual edits are therefore different counts. No source, site or repository changes were made. This is Chinese-source review, not a fresh bilingual or live-site audit; global chronology is not certified.

### Reading-copy and provenance note

Chapter151 includes96 narrative paragraphs, followed by a separator, an author-attributed note and compiler commentary. All were read, but the appended notes are not treated as narrated events. The compiler’s conjectures must not become invented story events or recipe requirements. The attribution to the author has not been independently authenticated.

The numbered reader-facing copy and workbook explicitly omit sexualized passages involving a minor, retaining paragraph numbers and omission markers. The original EPUB is unchanged. The report discusses those events non-graphically. Paragraph and character totals count the original EPUB text, including its paratext, not the shorter display copy.

## 1. Explicit aptitude change:44% becomes42%

This is the most important new accounting anchor:

> {src(159,41)['text']}

Chapter152’s breakthrough method leaves contamination. Repeated purges remove most of it; the borrowed Clean Water Gu subsequently removes the residual contamination. Nevertheless, Chapter159 explicitly confirms a permanent aptitude loss.

**Editorial treatment:** preserve44% in all earlier applicable scenes and42% after this consequence. The loss is **two percentage points**, not two percent of44. Cleaning away the residue does not restore the already-lost capacity.

This updates SR-001 and SR-007; it does not authorise retroactively changing the aperture, the earlier four/eight-shot calculations, or the GradeC anomaly in Chapter22.

The breakthrough itself consumes38 capacity points from the prior44, leaving6 before the remaining contaminated essence is expelled. RevolutionThree initial silver essence then replaces the RevolutionTwo essence. Four-Flavours remains owned in159, but its presence does not make it a RevolutionThree essence-refining Gu.

## 2. The calendar returns to July after advancing beyond late July

The relevant sequence is:

| Anchor | Evidence |
|---|---|
|Chapter131p78|Late July|
|Chapter132p63|Eight or nine days later|
|Chapter143p68|Funeral recalls yesterday’s battle|
|Chapter144p72|Fang Yuan is still middle-stage in the current example|
|Chapter145pp10,28|Present labelledJuly; natural upper-stage breakthrough occurred many days earlier|

> {src(145,10)['text']}

> {src(145,28)['text']}

**SR-065, held:** August is the working candidate for145’s month. The accumulated intervals make a return toJuly difficult; however, exact dates and later anchors must decide the final repair. I have not assigned an invented day.

This passage also explicitly callsJuly summer, strengthening the **internal** seasonal tension already logged under SR-054 against116’s “July, early autumn.” The earlier April/summer diagnosis SR-047 remains withdrawn. A new passage does not retrospectively justify that earlier assumption.

### Separate the former-life forecast

Chapter149 expects the Thunder Crown Wolf in **late August of the former-life sequence**. Qing Shu originally died helping defend the village at that stage. His current death in the Bai Ning Bing duel is explicitly earlier, a rebirth-branch change. Do not move his present death toAugust merely to match the memory.

### The three-day alibi is a lie, not a chronology error

Fang Yuan and Chi Lian claim that the missing three days were spent cultivating in Chi Lian’s residence. Chapter154 reveals what actually happened: Fang Yuan arrived **the previous night**, already RevolutionThree, and blackmailed Chi Lian into supporting him. That discrepancy is intentional and must be preserved.

Other clocks remain separate: the funeral is nextday; Bai lies awake for more than ten hours; a leaf delivery is promised in three days; the reinforcements turn back after fifteen minutes; and the later scene reaches nearly ten village assaults after an unspecified interval.

## 3. Destroyed Frost Demon Gu reappears in the inventory description

Chapter139 destroys Frost Demon Gu, and140 explicitly confirms the sacrifice. Yet143’s Plunder scene says:

> {src(143,35)['text']}

That wording treats the destroyed Gu as still among Bai Ning Bing’s possessions.

**SR-063, held:** retain the useful explanation but mark it as retrospective: “The destroyed Frost Demon Gu had been his most valuable Gu.” Then identify Blue Bird Ice Coffin Gu as the most valuable one remaining. Do not resurrect Frost Demon, or imply it remains a possible Plunder target.

Two further local identity findings:

- **SR-062:**142p76 attributes the giant tree spirit to Bai Ning Bing. It is **Gu Yue Qing Shu**; Bai is its prisoner.
- **SR-064:**143 describes the captured Stone Aperture Gu as a grey beetle, while144 describes it as a hard grey-white, dice-like block. No transformation is narrated. The working repair follows144’s fuller description, subject to later corroboration.

The earlier Ice Muscle ownership query, SR-060, remains open. Nothing here proves whether it was consumed, relinquished or still possessed.

## 4. Yao Ji’s auction recap and Yao Le’s age

### Failed auction versus claimed purchase

Chapter147 says:

> {src(147,84)['text']}

But111 awards the auctioned Liquor Worm to Fang Yuan, and123 still frames Yao Ji’s search as an attempt to obtain one for her granddaughter.

**SR-066, held:** describe the caravan transaction as an unsuccessful attempt, while retaining Yao Ji’s separate need for a refinement ingredient. An unreported alternative purchase could reconcile the passages, but none has been established. Do not invent one or reverse the witnessed auction result.

### One cohort behind, but still described as fifteen

Yao Le’s academy entry belongs to the cohort after Fang Yuan. The novel’s standard intake is fifteen; anotherJanuary passes before her death. Fang Yuan is explicitly seventeen by159, while151 calls her fifteen.

**SR-067, held:** sixteen is the conditional candidate under the established nominal-age sequence. The intake-age rule and age convention still need full-corpus confirmation; no birthday is invented.

By contrast, Bai Ning Bing’s “twenty-plus years” of questioning life in143 is **not** currently a proved age error. I found no contradictory exact age in the completed reading. His appearance and the three-year-old alliance recollection do not establish that he is eighteen.

## 5. Attack ranges: local repairs, not global rescaling

Chapter106 explicitly gives Moonglow a10-metre range, unchanged from Moonlight. Chapter142 gives it10paces, and156 calls Golden Moon’s unchanged range10paces. Chapter157 again gives Blood Moon a10-metre range.

**SR-068, held:** change only the two disputed range units,142p87 and156p52, frompaces to metres if later mechanics support106’s explicit baseline. Leave the enemy’s20-plus-pace distance in142 alone. Do not shrink geography, convert every pace into a metre, or collapse effective range into maximum flight distance.

Moon Scar has a better-supported numerical repair:

> {src(106,65)['text']}

Chapter156 nevertheless calls its range “one times” Moonlight’s range.

**SR-069, held:** “twice Moonlight Gu’s range.” Preserve106’s “increases by onefold,” which correctly describes10→20metres. This is not a global replacement rule for倍.

The Moon Shadow passage also strengthens the percentage-point convention: its suppression is60points againstRevolutionThree,30againstRevolutionFour and15againstRevolutionFive. The text says the60-point effect blocks all of a GradeC cultivator’s roughly40–50point capacity. It is not a60% relative reduction leaving40% of the available essence usable.

## 6. Resource and ownership reconciliation

| Event | What is established | What must not be invented |
|---|---|---|
|Man Shi group’s Gu|Recovered and surrendered for merit|An exact number or total merit payout|
|Xiong group’s corpses|Missing when Fang Yuan arrives|Acquisition of Brown Bear Innate Strength Gu|
|Plunder from Bai|Red Iron Relic, Stone Aperture, Water Shield|A surviving Frost Demon Gu|
|Leaf sale to Gu Yue Ye|Five at72stones, payment made:360|Fulfilled delivery merely because three days were promised|
|Rescue bargain|One Bear Enslavement Gu **plus700stones**|Treating700 as the Gu’s valuation|
|Chi Lian loans|3000stones andClean Water, twoIOUs|Income, gifts or completed repayment|
|First weekly stipend|300received, then all300given to the uncle|Two separate clan payments|
|Nine-Leaf Vitality Grass|Actually transferred in158for a reward token|Continuing leaf production from an already-surrenderedGu|
|Tushita Flower|Selected in160, filling the storage role|Earlier acquisition at the moment the token was offered|

### Advancement and fusion order

1. Natural middle→upper advancement is already complete in145. Four-Flavours supplies peak-quality essence, but he has not yet used the stolen relic.
2. Peak cultivation is confirmed in148, supporting use of the relic in the intervening interval.
3. The RevolutionThree breakthrough consumable is created and consumed in152.
4. Lightning Wings is the free clan reward.
5. The first Sky Canopy attempt fails and destroys Water Shield. An unidentified replacement water-defence Gu joins White Jade in the successful second attempt.
6. Blood Moon succeeds on its second attempt, confirmed159. The failed attempt does not identify a destroyed ingredient or a replacement purchase.

**SR-071** qualifies155’s “first RevolutionThree Gu” as the first lasting combat-equipment Gu. Otherwise it overlooks the RevolutionThree consumable used in152. Preserve the actual acquisition order rather than changing the reward.

I am **not giving a falsely exact total inventory** at this checkpoint. Chapter151 says Yao Le’s Gu were all taken and refined, without listing their number or identities. Temporary recipe inputs and recovered Gu also require separate accounting. The named equipment changes above are firm; a grand total is not.

Three-Step Fragrant Grass is still not acquired: Yao Ji blocked its purchase despite sufficient merit. The Chainsaw Golden Centipede is still an intended acquisition, not captured equipment.

## 7. Preserve apparent contradictions that the narration explains

- **Xiong Jiao Man’s status:**150 says she hits a building and shows no further movement. It does not explicitly declare her dead. No double-death finding is justified through160.
- **Three wolf kings become four:** the fourth is concealed and then revealed. The four killed by152 are not the same four attacking in159.
- **Different siege counts:** Xiong Village’s third encirclement and Gu Yue Village’s second are not competing counts of one village.
- **Qing Shu’s energy:** ambient qi is not infinite.142 explicitly exhausts the local supply. The RevolutionFour recovery comparison is hypothetical.
- **Blood Moon:** its periodic attack falls **to one-third**, not by one-third. Its usefulness in a prolonged beast fight does not contradict its limitations in short Gu Master duels.
- **Flower Wine Monk’s wounds:**156 deliberately reopens the mystery. Do not repair the discrepancy before the narrative explains it.
- **Fang Zheng’s promotion:** middle-stage145 and upper-stage158 can fit the intervening weeks. No fixed minimum duration has been established.
- **RevolutionFive clan Gu:**160 presents Fang Yuan’s inference, not a direct inspection of the clan’s holdings.

## 8. Copy corrections and checkpoint

SR-073 supplies{len(fixes)} individually located, minimal source-copy candidates, with the surrounding source quotations in the Continuity jotter. These include the grandfather/grandson slip in154p22, the female pronoun in159p67, and “price” or “before him” replacements where the source has obvious miskeys.

**SR-037 is extended to160p62:** the poisonGu was used by Wang Da, not Wang Er. Chapter160 also repeats爱别离, supporting the existing naming proposal SR-038.

At the endpoint, Fang Yuan has selected Tushita Flower and attends the former classmates’ banquet. No later events are counted. The numbered reading copy, four jotters, reading log, proposal ledger, **{len(events)} chronology entries**, **{len(resources)} resource entries**, and cumulative workbook are updated.

**Next actual unread:161. Next planned batch:161–180**, retaining the smaller batch size while this dense arc continues. All active changes remain held. Remote research backup remains unresolved; no upload or deletion occurred.
'''
report=re.sub(r'(?<=[A-Za-z])(?=\d)|(?<=\d)(?=[A-Za-z])',' ',report)
# Normalize common compact English compounds in the reader-facing report, not quoted Chinese.
for a,b in {'nextday':'next day','lateJuly':'late July','GuMaster':'Gu Master','RevolutionThree':'Revolution Three','RevolutionFour':'Revolution Four','RevolutionFive':'Revolution Five','GradeC':'Grade C','Four-Flavours':'Four-Flavours','Clean Water':'Clean Water'}.items():report=report.replace(a,b)
(r/'Batch-007-Chapters-141-160.md').write_text(report)
w=Workbook();s=w.active;s.title='Read me'
for row in [('Sequential review','Status'),('Completed','1–160 of2334'),('EPUB paragraphs',tp),('Source characters',tc),('Next unread',161),('Proposal records','73 total;72 activeheld;1withdrawn;noneapproved/applied'),('Copy bundle',f'SR-073 contains{len(fixes)}individually locateditems'),('Scope','Chinese-source review, not fresh bilingual/liveEnglish audit'),('Reading copy','Explicitly marked Ch151 omissions; originalEPUB unchanged; paratext not narrative'),('Chronology entries',len(events)),('Resource entries',len(resources)),('Global consistency','Not yet certified')]:s.append(row)
s=w.create_sheet('Chapter completion');s.append(['Chapter','File','Title','Paragraphs','Status'])
for c in chapters:s.append([c['corpus_chapter'],c['file'],c['title'],len(c['paragraphs']),'Full sequential read;151includes3paratextparagraphs' if c['corpus_chapter']==151 else 'Full sequential read'])
def ids(vs,k,c,p):return ', '.join(v['id'] for v in vs if any(z['chapter']==c and z['paragraph']==p for z in v[k]))
s=w.create_sheet('Paragraph register');s.append(['Chapter','Paragraph','Status','Proposal IDs','Chronology IDs','Resource IDs','Chinese reading copy'])
for c in chapters:
 for p,t in enumerate(c['paragraphs'],1):
  n=c['corpus_chapter'];s.append([n,p,'Read; not certified error-free',ids(ledger,'locations',n,p),ids(events,'sources',n,p),ids(resources,'sources',n,p),reader(n,p,t)])
s=w.create_sheet('Proposal ledger');s.append(['ID','Category','References','Action','Rationale','Confidence','Status','Dependencies','Publication','Approved','Applied','Latest update'])
for v in ledger:s.append([v['id'],v['category'],refs(v['locations']),v['recommended_action'],v.get('rationale','See originating report'),v.get('confidence','See status'),v['status'],v['dependencies'],v['publication_status'],'No','No',v.get('batch007_update',v.get('batch006_update',v.get('batch005_update',v.get('batch004_update',''))))])
s=w.create_sheet('Copy items');s.append(['Proposal','Chapter','Paragraph','Original','Candidate','Status'])
for v in ledger:
 for z in v.get('copy_items',[]):s.append([v['id'],z['chapter'],z['paragraph'],z['original'],z['replacement'],z['status']])
s=w.create_sheet('Chronology');s.append(['ID','Timeline','Kind','Event','Constraint','Evidence status','Dependencies','References','Exact permitted source'])
for v in events:s.append([v['id'],v['timeline'],v['kind'],v['event'],v['constraint'],v['evidence_status'],v['dependency'],refs(v['sources']),'\n'.join(reader(z['chapter'],z['paragraph'],z['text']) for z in v['sources'])])
s=w.create_sheet('Resources');s.append(['ID','Kind','Entry','Calculation','Status','References','Exact permitted source'])
for v in resources:s.append([v['id'],v['kind'],v['entry'],v['calculation'],v['status'],refs(v['sources']),'\n'.join(reader(z['chapter'],z['paragraph'],z['text']) for z in v['sources'])])
for s in w:
 s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
 for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='163B4B')
 for row in s.iter_rows(min_row=2):
  for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
 for col in s.columns:s.column_dimensions[col[0].column_letter].width=min(65,max(14,max(len(str(c.value or '')) for c in col[:30])+2))
w.save(r/'Sequential-Reading-and-Proposals.xlsx')
code=(r.parent/'make_whole_novel_audit_docx.py').read_text().replace("root=Path(__file__).parent","root=Path('/home/user/editorial/sequential-review')").replace('Whole-Novel-Aperture-and-Travel-Audit','Batch-007-Chapters-141-160').replace('WHOLE-NOVEL SOURCE AUDIT','SEQUENTIAL REVIEW / CHAPTERS141–160').replace('The Omniarch: Aperture and Travel Source Audit','The Omniarch: Sequential Review, Chapters141–160').replace('Full-corpus topical screening with selective contextual verification','Twenty complete chapters; active proposals held')
exec(compile(code,'docx-export','exec'),{})
(r/'CURRENT-STATUS.md').write_text(f'''# Sequential review checkpoint

Completed **Chapters1–160**. Nextunread **161**. Plannednextbatch **161–180**, dense20-chapterbatches. Remaining2174.

Batch007:{np}EPUBparagraphs/{nc}characters. Cumulative:{tp}paragraphs/{tc}characters.73records:72activeheld,1withdrawn. SR-073contains{len(fixes)}copyitems. Noneapproved/applied. FullnovelchronologyNOTcertified.

Latest: Batch-007-Chapters-141-160.md/.docx; Batch-007-Numbered-Chinese-Source.md; four007jotters; workbook/registers/readinglog. Reader-facingCh151sexualizedminorpassagesexplicitlyomitted,paragraphIDsretained;originalEPUBunchanged. Ch15197–99paratextnotnarrative;compilerconjecturesarenotcanonicalevents.

Carryforwards:
-159explicit44→42aptitudelossafterbreakthrough;CleanWaterused,doesnotrestorelostcapacity. Keepprior44. Noexactcurrentinventorytotal:151YaoLeGuappropriatedunnumbered.
-SR063destroyedFrostDemonpresenttense143needsretrospectiveframing. SR064StoneAperturebeetle143vsdice144. SR060IceMusclestatusstillunresolved.
-SR065145Julyregressionafter131lateJuly+intervals. SR054145JulySUMMERreinforcesinternalconflict116JulyAUTUMN. SR047withdrawalremains.
-SR066YaoJicompletedcaravanLiquorWormpurchase147vs111loss/123search. SR067YaoLe15vscohort-based16,conditional. Baiage143twenty-plusquestionnotprovederror.
-SR068142/15610paces→10metrescandidatesagainst106;SR069MoonScarrange2×,not1×. Nogeographicshrinkage.
-SR037adds160p62WangEr→WangDa;SR038reinforced160p61爱别离.
-XiongJiaoMan150unresponsiveNOTconfirmeddead;doubledeathNOTestablished. QingShudeath143confirmed; priorlife149lateAugustThunderCrownchangedbyrebirth.
-153three-dayChiLianretreatclaimisfalsealibi;154lastnightblackmailreveal,notchronologyrepair.
-ActualSkyCanopyreplacesWhiteJade+replacementwaterGu,WaterShielddiedfirstfailure;LightningWingsreward;BloodMoon2ndattemptsuccess;NineLeaftraded158;TushitaFlowerselected160;ThreeStepGrassandChainsawCentipedeNOTowned.
-3000loan+CleanWaterdebt;300stipendgivenuncle;leaf5×72=360paid,promiseddeliverynotconfirmed;rescueGuPLUS700.

No source/site/repositoryedits,uploadordeletions. Remoteresearchbackupunresolved. Do not rerun earliergenerators unchanged. Topical2334screeningNOTsequentialcompletion.
''')
assert len(chapters)==160 and len(ledger)==73 and len(new)==12
assert np==1878 and nc==63989 and tp==13167 and tc==480791
assert all(not v['approved'] and not v['applied'] for v in ledger)
assert all(z['text']==src(z['chapter'],z['paragraph'])['text'] for v in events+resources for z in v['sources'])
assert '—' not in report
texts={t for c in chapters for t in c['paragraphs']}
for l in report.splitlines():
 if l.startswith('> '):assert l[2:] in texts,l
print(dict(batch_paragraphs=np,batch_characters=nc,total_paragraphs=tp,total_characters=tc,proposal_records=len(ledger),active=72,withdrawn=1,copy_items=len(fixes),chronology=len(events),resources=len(resources),next=161))
