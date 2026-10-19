import json, copy, re
from pathlib import Path
from openpyxl import load_workbook
from docx import Document
from docx.shared import Pt
r = Path(__file__).parent
led = json.loads((r/'proposal-ledger.json').read_text())
by = {v['id']: v for v in led}
# automatic note for entries whose first location matches the witness exactly
auto = {}
for line in (r/'cand023.txt').read_text().splitlines():
    if not line.startswith('SR-'): continue
    parts = [p.strip() for p in line.split(' | ')]
    sid, loc, st = parts[0], parts[1], parts[4]
    if st.startswith('WITNESS EXACT'):
        auto[sid] = 'Hetushu copy has the same text as our source here (' + loc + '). Does not confirm the proposal.'
special = {
 'SR-195': 'Hetushu (ch437 p22) reads 远方 where our source reads 冤枉 (witness supports a change; the query still stands).',
 'SR-196': 'Hetushu (ch436 p89) reads 一旦 (supports the proposed 一 insertion). Held.',
 'SR-199': 'Hetushu (ch439 p4) reads 对方源 where our source reads 地方源. The witness supports a change but not the proposed deletion of 地. Decide at reconciliation.',
 'SR-215': 'Hetushu (ch453 p44) keeps 普通水浪 (the typo) and reads 三倍 where our source reads 十三倍 (new query SR-383).',
 'SR-217': 'Hetushu (ch456 p19) reads 儿女 where our source reads 耳女 (supports the proposed repair). Held.',
 'SR-228': 'Hetushu (ch458 p80) reads 三万两千头 where our source reads 三王两千头 (supports the 王→万 direction; query stands).',
 'SR-241': 'Hetushu (ch481 p74) reads 一百多只 where our source reads 一百都只 (supports the proposed repair). Held.',
 'SR-242': 'Hetushu (ch481 p64) reads 险地 where our source reads 贤弟 (witness gives a candidate; query stands).',
 'SR-244': 'Hetushu (ch485 p17) reads 一番设计 without 把 (supports deleting 把; 布设 vs 设计 remains open). Held.',
}
notes = {**auto, **special}
for sid, t in notes.items():
    if sid in by: by[sid]['witness_023_025'] = t
new = copy.deepcopy(led[-1]); new['id'] = 'SR-383'
new['category'] = 'Numeric (十三倍 vs 三倍), 453p44'
new['locations'] = [{'chapter': 453, 'paragraph': 44}]
new['recommended_action'] = 'Query. 453p44 says 普通水浪的十三倍. Hetushu reads 三倍. Confirm the ratio against 453p13 and the wolf-food record before any change. No text change.'
new['rationale'] = 'The Hetushu witness and our source differ on the multiplier at 453p44 (十三倍 vs 三倍). The text is the same otherwise. The difference is numeric and needs context to decide.'
new['confidence'] = 'Provisional (0.4).'
new['status'] = 'Provisional query; no approval'
new['dependencies'] = 'Witness collation for 453 (Hetushu). Whole-novel reconciliation of wolf-food ratios.'
new['approved'] = False; new['applied'] = False
new['publication_status'] = 'Held; not for publication'
new['witness_023_025'] = 'Witness (ch453 p44) reads 三倍 where our source reads 十三倍.'
for k in list(new.keys()):
    if k.startswith('batch') and k.endswith('_update'): new[k] = ''
if 'SR-383' not in by: led.append(new)
assert len(led) == 383 and led[-1]['id'] == 'SR-383'
(r/'proposal-ledger.json').write_text(json.dumps(led, ensure_ascii=False, indent=2))
by = {v['id']: v for v in led}
print('ledger', len(led), 'annotated', len(notes))

# Batch QA files: witness collated for 023-025
for b, lo, hi in [('batch-023',426,445),('batch-024',446,465),('batch-025',466,485)]:
    p = r/f'{b}-qa.json'
    q = json.loads(p.read_text()) if p.exists() else {}
    q['witness_collated'] = True
    q['witness_source'] = 'Hetushu book 7640 pages 5345664–5345723 (ch426–485); fact-check-023-025/'
    q['witness_summary'] = 'Mean exact-line agreement after watermark strip about 0.955. Witness confirms SR-196, SR-217, SR-241, SR-244 (partly). Witness supports queries SR-195, SR-228, SR-242. Witness reads differently from our source at 439p4 (SR-199) and 453p44 (new query SR-383). Most other entries: witness copies the source text (no confirmation).'
    p.write_text(json.dumps(q, ensure_ascii=False, indent=2))

# Batch reports: witness section and DOCX regeneration
rep = {
 'Batch-023-Chapters-426-445.md': (426,445,'Hetushu pages 5345664–5345683'),
 'Batch-024-Chapters-446-465.md': (446,465,'Hetushu pages 5345684–5345703'),
 'Batch-025-Chapters-466-485.md': (466,485,'Hetushu pages 5345704–5345723'),
}
for name, (lo, hi, pages) in rep.items():
    p = r/name; t = p.read_text(encoding='utf-8')
    t = re.sub(r'\n## 5\. Witness collation.*', '', t, flags=re.S)
    items = [(sid, notes_) for sid, notes_ in notes.items() if sid in by and any(lo <= l['chapter'] <= hi for l in by[sid]['locations'][:1])]
    lines = [f"\n## 5. Witness collation (Hetushu, ch{lo}–{hi})",
             f"- Collated against {pages} (fact-check-023-025/). Mean exact-line agreement after watermark strip is about 0.95.",
             "- Entries with a witness note below. Other held and query items in this chapter range match the witness text exactly, so the witness does not confirm them."]
    for sid, nt in items: lines.append(f"- {sid}: {nt}")
    lines.append("- New query SR-383 (453p44: 十三倍 vs 三倍) is in Batch 024.") if lo == 446 else None
    p.write_text(t.rstrip() + '\n' + '\n'.join(lines) + '\n', encoding='utf-8')
    d = Document(); st = d.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
    table = None; in_t = False
    clean = lambda s: re.sub(r'\*\*(.+?)\*\*', r'\1', s)
    for line in p.read_text(encoding='utf-8').splitlines():
        if not line.strip() or line.strip() == '---': in_t = False; continue
        if line.startswith('|'):
            if set(line.replace('|','').strip()) <= set('-: '): continue
            cells = [clean(c.strip()) for c in line.strip('|').split('|')]
            if not in_t: table = d.add_table(rows=0, cols=len(cells)); table.style = 'Light Grid Accent 1'; in_t = True
            row = table.add_row().cells
            for k, cv in enumerate(cells): row[k].text = cv
            continue
        in_t = False
        if line.startswith('### '): d.add_heading(line[4:], 2)
        elif line.startswith('## '): d.add_heading(line[3:], 1)
        elif line.startswith('# '): d.add_heading(line[2:], 0)
        elif line.startswith('- '): d.add_paragraph(clean(line[2:]), style='List Bullet')
        else: d.add_paragraph(clean(line))
    d.save(r/name.replace('.md', '.docx')); print('docx', name, len(items))

# Workbook: append SR-383 and annotate 023-025 rows
wb = load_workbook(r/'Sequential-Reading-and-Proposals.xlsx'); ws = wb['Proposal ledger']
HOLD = 'HOLD: full2334-chapter reading and global reconciliation outstanding'
for row in ws.iter_rows(min_row=2):
    sid = row[0].value
    if isinstance(sid, str) and sid.startswith('SR-') and sid in notes:
        row[11].value = (row[11].value or '') + ' | Witness (Hetushu 023–025): ' + notes[sid]
v = by['SR-383']
refs = '; '.join(f"Ch{l['chapter']} p{l['paragraph']}" for l in v['locations'])
ws.append([v['id'], v['category'], refs, v['recommended_action'], v['rationale'], v['confidence'], v['status'], v['dependencies'], HOLD, 'No', 'No', 'Batch 024 sequential reading; witness collated; ' + v['witness_023_025']])
for row in wb['Read me'].iter_rows(min_row=1, max_col=2):
    if row[0].value == 'Proposals':
        withdrawn = sum(1 for x in led if x['status'].startswith('Withdrawn')); resolved = sum(1 for x in led if x['status'].startswith('Resolved locally'))
        row[1].value = f"{len(led)} total; {len(led)-withdrawn-resolved} active (held or query); {withdrawn} withdrawn; {resolved} locally resolved"
wb.save(r/'Sequential-Reading-and-Proposals.xlsx'); print('workbook saved')
