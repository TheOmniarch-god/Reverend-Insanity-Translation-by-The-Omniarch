from pathlib import Path
import json, zipfile, re
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from docx import Document
r = Path(__file__).parent
load = lambda n: json.loads((r/n).read_text())
b = load('batch-028-source.json'); led = load('proposal-ledger.json'); p = load('reading-progress.json')
z = zipfile.ZipFile(r.parents[1]/'benchmark/sources/Gu-Zhen-Ren.epub')

# source vs EPUB
for c in b:
    actual = [n.get_text('', strip=True) for n in BeautifulSoup(z.read(c['file']), 'html.parser').find_all('p') if n.get_text('', strip=True)]
    assert c['paragraphs'] == actual, c['corpus_chapter']
assert [c['corpus_chapter'] for c in b] == list(range(526, 546))
N = sum(len(c['paragraphs']) for c in b); CH = sum(len(t) for c in b for t in c['paragraphs'])
assert N == 1761 and CH == 65397

# numbered source exact
src = (r/'Batch-028-Numbered-Chinese-Source.md').read_text(encoding='utf-8')
cur = None; found = []
for line in src.splitlines():
    m = re.match(r'## Chapter (\d+) ', line)
    if m: cur = int(m[1])
    m = re.match(r'(\d+)\.(\d+) (.*)', line)
    if m:
        assert int(m[1]) == cur
        assert m[3] == b[cur-526]['paragraphs'][int(m[2])-1]; found.append((cur, int(m[2])))
assert found == [(c['corpus_chapter'], i) for c in b for i in range(1, len(c['paragraphs'])+1)]
assert len(found) == 1761

# ledger
ids = [v['id'] for v in led]
assert len(led) == 306 and len(set(ids)) == 306 and ids[-1] == 'SR-306'
assert all(not v['approved'] and not v['applied'] for v in led)
for v in led:
    if int(v['id'][3:]) >= 286:
        for l in v['locations']:
            if 526 <= l['chapter'] <= 545:
                assert 1 <= l['paragraph'] <= len(b[l['chapter']-526]['paragraphs']), (v['id'], l)
        assert 'witness_collation' in v and v['witness_collation']
        assert v['status'].startswith(('Held repair', 'Provisional')), v['id']
assert sum(1 for v in led if int(v['id'][3:]) >= 286 and v['status'].startswith('Held repair')) == 12
assert sum(1 for v in led if int(v['id'][3:]) >= 286 and v['status'].startswith('Provisional')) == 9

# progress
assert p['next_chapter'] == 546 and p['sequential_chapters_completed'] == list(range(1, 546))
assert p['paragraphs_read'] == 46920 + N and p['paragraph_characters_read'] == 1645947 + CH
assert p['remaining_chapters'] == 1789

# workbook
w = load_workbook(r/'Sequential-Reading-and-Proposals.xlsx', read_only=True, data_only=True)
cc = [row for row in w['Chapter completion'].iter_rows(min_row=2, values_only=True) if row[0] and str(row[0]).isdigit() and int(row[0]) >= 526]
assert len(cc) == 20
pr = [row for row in w['Paragraph register'].iter_rows(min_row=2, values_only=True) if row[0] and str(row[0]).isdigit() and int(row[0]) >= 526]
assert len(pr) == N and all(row[6] == b[row[0]-526]['paragraphs'][row[1]-1] for row in pr)
pl = [row for row in w['Proposal ledger'].iter_rows(min_row=2, values_only=True) if row[0] and str(row[0]).startswith('SR-') and int(str(row[0])[3:]) >= 286]
assert len(pl) == 21
rm = dict(w['Read me'].iter_rows(min_row=2, max_col=2, values_only=True))
assert rm['Next unread'] == '546' and rm['Completed'] == '1–545 of 2334'

# docx
d = Document(r/'Batch-028-Chapters-526-545.docx')
txt = '\n'.join(x.text for x in d.paragraphs) + '\n' + '\n'.join(c.text for t in d.tables for row in t.rows for c in row.cells)
assert 'SR-286' in txt and 'SR-306' in txt and len(d.tables) == 2

# report hygiene
rep = (r/'Batch-028-Chapters-526-545.md').read_text(encoding='utf-8')
assert '—' not in rep
assert not any(line.startswith('> ') for line in rep.splitlines())

qa = {'status': 'passed', 'epub_paragraphs_rechecked': N, 'numbered_source_exact': True, 'numbered_source_lines': N,
      'workbook_chapter_rows': 20, 'workbook_paragraph_rows': N, 'workbook_paragraph_text_exact': True,
      'workbook_ledger_rows_added': 21, 'workbook_read_me_next_unread': 546, 'docx_proposal_ids_present': True,
      'docx_table_count': 2, 'ledger_total': 306, 'approved_or_applied': 0,
      'witness_collated_proposals': 21, 'witness_supports': ['SR-293'], 'em_dashes_in_report': 0,
      'progress_next': 546, 'progress_paragraphs': 46920 + N, 'progress_characters': 1645947 + CH}
(r/'batch-028-qa.json').write_text(json.dumps(qa, ensure_ascii=False, indent=2))
print(json.dumps(qa, ensure_ascii=False, indent=2))
