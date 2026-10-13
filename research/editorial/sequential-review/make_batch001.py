from pathlib import Path
import json,csv
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
r=Path(__file__).parent
p=r/'Batch-001-Chapters-1-10.md'
s=p.read_text()
insert='''## 10. Smaller textual and translation decisions

These are separate from the mathematical harmonisations. They are not newly verified defects in the live English.

### SR-008: poem title, Chapter 1 p57

The EPUB has **《将敬酒》**. Prefer the conventional poem title **《将进酒》** in the working-source apparatus. In English, retain an established translation of that title if the edition already has one; a workable rendering is **“Bring in the Wine”**. Do not turn an orthographic repair into a mandatory change of an otherwise sound English title. Classification: contextual source-title correction, with high confidence.

### SR-009: ceremony, Chapter 1 p63

The EPUB has **开窍大殿**, literally an awakening hall, in the instruction to rest for tomorrow's event. Prefer **开窍大典**, the awakening ceremony. Recommended English phrase: **“tomorrow’s Awakening Ceremony”**. The immediate event sequence strongly supports this repair. If the translation already says ceremony, leave it alone.

### SR-010: the uncle's question, Chapter 9 p17

The EPUB has **“方源，我的意向呢？”**, literally asking about the speaker's own intention while turning to Fang Yuan for his answer. Prefer **“方源，你的意向呢？”** in the working-source apparatus, and **“Fang Yuan, what do you think?”** in English. This restores the question's clear conversational function without adding a motive or a new offer. Classification: high-confidence contextual repair, not an independently attested variant.

### SR-011: room-size comparison, Chapter 9 p43

**大上两倍** has competing readings in ordinary usage: twice as large, or larger by twice the original amount. No dimensions in this scene adjudicate the distinction. My working rendering is **“The room was twice as large as his old one.”** Mark that as a translation choice of moderate confidence, not a proved numerical correction. Do not insist on “three times as large” purely from a mechanical parsing of the comparative. Reconsider only if another witness or a later dimension supplies firmer evidence.

### Other numbers to leave intact

- Chapter 1's 三个时辰 is six hours; Chapter 3's 一个时辰多丁点 is slightly over two hours. Night watches are not equivalent to that same number of elapsed hours.
- Rounded centuries and five-hundred-year life-history descriptions do not justify manufactured exact birthdays or a new lifespan total.
- Chapter 4's 三两步 can mean a few strides. It should not become a fixed measurement used to calculate the river's width.
- A general awakening rate of about sixty percent and a partial observed run of successes and failures do not constitute a statistical contradiction.
- The ancient migration in Chapter 4 states origins and settlement, not an exact transit time. No shortening of the world's geography is warranted by this passage.
- Chapter 9's accepted six-stone gift explains the six stones held in Chapter 10. The separate two-hundred-stone offer is refused; do not add it to Fang Yuan's funds.

'''
s=s.replace('## 10. Proposal status and continuation point',insert+'## 11. Proposal status and continuation point')
s=s.replace('| SR-007 | Define capacity points and preserve separate denominators. | Recommended house policy; no numerical recoding. |','| SR-007 | Define capacity points and preserve separate denominators. | Recommended house policy; no numerical recoding. |\n| SR-008–010 | Repair the poem title, ceremony character, and addressee pronoun. | Contextual source-copy proposals; retain already-sensible English. |\n| SR-011 | Use “twice as large” as the working room-size rendering. | Translation choice, not a demonstrated numerical repair. |')
p.write_text(s)
source=json.loads((r/'batch-001-source.json').read_text())
proposals=[
('SR-001','measurement',[(6,8),(10,19),(10,51)],'Express 44% as occupied aperture volume, retaining below-midpoint sea height.','Preferred editorial harmonisation','Later numerical height descriptions'),
('SR-002','chronology',[(8,48),(8,58)],'month → week; ten-plus days → first few days','Strong local harmonisation','Check Chapters 11–20 chronology'),
('SR-003','childhood',[(2,49),(2,50),(9,16)],'Fostering from three for twelve years, instead of five for ten.','Moderate-confidence conjecture','Later childhood references and witnesses'),
('SR-004','casualties',[(1,3),(2,11)],'Change Chapter 2 hundreds of thousands to ten million if harmonising.','Provisional preferred conjecture','V2 Ch180 chapter388.html p39 rechecked; sequential reading outstanding'),
('SR-005','source copy',[(8,42)],'有没有 → 又没有; preserve negative sense in English','High contextual confidence','No independent corrected witness verified'),
('SR-006','source copy',[(10,31)],'八方 → 八分','High confidence','Immediate numerical sequence'),
('SR-007','units',[(10,19),(10,32),(10,36),(10,39),(10,46),(10,51),(10,55)],'Define aperture-capacity percentages and percentage-point changes; preserve all source values.','Recommended policy','Do not invent universal gross/net recovery mechanism'),
('SR-008','source copy',[(1,57)],'将敬酒 → 将进酒; retain a sound existing English title.','High contextual confidence','Conventional poem title'),
('SR-009','source copy',[(1,63)],'开窍大殿 → 开窍大典; tomorrow’s Awakening Ceremony','High contextual confidence','Ceremony context'),
('SR-010','source copy',[(9,17)],'我的意向呢 → 你的意向呢; Fang Yuan, what do you think?','High contextual confidence','No independent corrected witness verified'),
('SR-011','translation choice',[(9,43)],'Working rendering: The room was twice as large as his old one.','Moderate confidence; not proven numerical repair','Revisit if dimensions or another witness settle comparative')]
ledger=[]
for id,kind,locs,action,status,dep in proposals:
 ledger.append(dict(id=id,category=kind,locations=[dict(chapter=c,paragraph=p,file=source[c-1]['file']) for c,p in locs],recommended_action=action,status=status,dependencies=dep,approved=False,applied=False))
(r/'proposal-ledger.json').write_text(json.dumps(ledger,ensure_ascii=False,indent=2))
progress=dict(source='/home/user/benchmark/sources/Gu-Zhen-Ren.epub',source_sha256='1ef1c29b6bddd7994f56696e2d091363aa8bbcee4889439465a9e2c51437dd5d',total_narrative_chapters=2334,sequential_chapters_completed=list(range(1,11)),paragraphs_read=715,paragraph_characters_read=26627,next_chapter=11,remaining_chapters=2324,batch_report=p.name,scope='Complete source paragraph sequences read. Not a fresh bilingual/live-site audit. Later topical excerpts do not count as sequential completion.',chapters=[dict(chapter=x['corpus_chapter'],file=x['file'],title=x['title'],paragraphs=len(x['paragraphs']),first_paragraph=1,last_paragraph=len(x['paragraphs'])) for x in source])
(r/'reading-progress.json').write_text(json.dumps(progress,ensure_ascii=False,indent=2))
lines=['# Source reading copy: Volume 1, Chapters 1–10','Nonempty paragraph numbers correspond to the review report. Chinese source text, not emended.','']
for ch in source:
 lines+=['## '+str(ch['corpus_chapter'])+'. '+ch['title'],'Source: `'+ch['file']+'`','']
 for i,t in enumerate(ch['paragraphs'],1):lines += [f'**{i:03d}.** {t}','']
(r/'Batch-001-Numbered-Chinese-Source.md').write_text('\n'.join(lines))
w=Workbook();a=w.active;a.title='Chapter completion';a.append(['Chapter','Source file','Title','Paragraphs','Status'])
for x in source:a.append([x['corpus_chapter'],x['file'],x['title'],len(x['paragraphs']),'Complete source sequence read'])
b=w.create_sheet('Paragraph register');b.append(['Chapter','Source paragraph','Status','Proposal IDs','Chinese source'])
for x in source:
 for i,t in enumerate(x['paragraphs'],1):
  ids=[v['id'] for v in ledger if any(l['chapter']==x['corpus_chapter'] and l['paragraph']==i for l in v['locations'])]
  b.append([x['corpus_chapter'],i,'Read; proposal recorded' if ids else 'Read; no proposal recorded in this batch',', '.join(ids),t])
c=w.create_sheet('Active proposals');c.append(['ID','Category','Locations','Recommended action','Status','Dependencies','Approved','Applied'])
for v in ledger:c.append([v['id'],v['category'],'; '.join(f"Ch{l['chapter']} p{l['paragraph']}" for l in v['locations']),v['recommended_action'],v['status'],v['dependencies'],'No','No'])
for sh in w:
 sh.freeze_panes='A2';sh.auto_filter.ref=sh.dimensions
 for cell in sh[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='163B4B')
 for row in sh.iter_rows(min_row=2):
  for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
 for col in sh.columns:
  sh.column_dimensions[col[0].column_letter].width=min(65,max(14,max(len(str(x.value or '')) for x in col[:30])+2))
w.save(r/'Sequential-Reading-and-Proposals.xlsx')
# Reuse the established Word layout while preserving this batch's own metadata.
code=(r.parent/'make_whole_novel_audit_docx.py').read_text()
code=code.replace("root=Path(__file__).parent", "root=Path('/home/user/editorial/sequential-review')")
code=code.replace('Whole-Novel-Aperture-and-Travel-Audit','Batch-001-Chapters-1-10').replace('WHOLE-NOVEL SOURCE AUDIT','SEQUENTIAL REVIEW: CHAPTERS 1–10').replace('The Omniarch: Aperture and Travel Source Audit','The Omniarch: Sequential Review, Chapters 1–10').replace('Full-corpus topical screening with selective contextual verification','Complete paragraph-by-paragraph source reading of Chapters 1–10; proposals not applied')
exec(compile(code,'batch-docx-export','exec'))
assert '—' not in s
print('Words:',len(s.split()),'Paragraphs:',sum(len(x['paragraphs']) for x in source),'Proposals:',len(ledger))
