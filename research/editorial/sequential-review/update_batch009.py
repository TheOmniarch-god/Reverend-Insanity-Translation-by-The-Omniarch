from pathlib import Path
import json, shutil
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
r=Path('/home/user/editorial/sequential-review')
def load(n):return json.loads((r/n).read_text())
def save(n,v):(r/n).write_text(json.dumps(v,ensure_ascii=False,indent=2))
for n in ['reading-progress.json','proposal-ledger.json','chronology-register.json','resource-register.json','working-standards-register.json']:
 d=r/'history'/('before-batch009-'+n)
 if not d.exists():shutil.copy2(r/n,d)
cs=sum([load(f'batch-{i:03d}-source.json') for i in range(1,10)],[]);by={c['corpus_chapter']:c for c in cs};batch=cs[180:200]
def src(c,p):return dict(chapter=c,paragraph=p,file=by[c]['file'],text=by[c]['paragraphs'][p-1])
def ss(rr):return [src(*z) for z in rr]
def loc(rr):return [{k:v for k,v in z.items() if k!='text'} for z in ss(rr)]
def refs(z):return '; '.join(f"Ch{x['chapter']} p{x['paragraph']}" for x in z)
def reader(c,p,t):
 if c==151 and (31<=p<=67 or p in [98,99]):return '[Restricted passage omitted from reader-facing copy. Original numbering retained; original EPUB unchanged. See earlier non-graphic editorial discussion.]'
 return t
ledger=load('history/before-batch009-proposal-ledger.json');new=[]
def pro(i,cat,rr,action,why,conf,dep):
 v=dict(id=f'SR-{i:03d}',category=cat,locations=loc(rr),recommended_action=action,rationale=why,confidence=conf,status='Active provisional proposal; no approval',dependencies=dep,approved=False,applied=False,publication_status='HOLD: full 2334-chapter reading and global reconciliation outstanding');ledger.append(v);new.append(v);return v
pro(85,'linked aptitude arithmetic',[(159,41),(162,56),(167,78),(198,25),(198,37),(198,55),(198,56)],'Preferred linked candidate: Ch198p25 四成三→四成二 and 五成三分→五成二分; p37 六成三分→六成二分; p56 四成三分→四成二分. Preserve the final 九成 (90%).','Three earlier explicit passages establish 42%, not 43%. Adding ten capacity points produces 52%, then62%. Later diminishing gains and nearly ten washes do not justify changing the final90 to89.','High conflict; strong linked candidate.','Check subsequent aptitude recaps and source variants. No retroactive change to the valid original44%; no universal decrement to every later figure.')
pro(86,'former-life inheritance date',[(177,28),(183,39),(183,72),(198,59)],'Conditional minimal wording at183p72: 四百年之后→数百年之后 and 提前了四百年的功夫→提前了数百年的功夫. Retain both original readings in evidence.','The same chapter first places finding Blood Sea inheritance after300-plus years, then after400 years. Exact preferred century is not established. RevolutionSix after400-plus years and earlier GradeB improvement are different milestones.','High discrepancy; moderate nonnumerical repair.','Later former-life milestones; do not conflate inheritance, founding a sect, RevolutionFive and RevolutionSix or silently pick one century.')
pro(87,'recalled acquisition identity',[(63,83),(63,85),(63,91),(79,30),(79,32),(128,22),(128,23),(188,81)],'Ch188p81 从地藏花中取出酒虫→从地藏花中取出白豕蛊.','The explicit first flower yields White Boar, the second Jade Skin, and the later one the Ear. The original Liquor Worm was not recovered from a flower. White Boar is the best locally witnessed dormant comparison.','High identity mismatch; strong White Boar candidate.','Check source variants; do not confuse the list of all inheritedGu in122 with the narrower set recovered from flowers.')
pro(88,'Mo family identity',[(172,62),(172,75),(176,59),(189,29)],'Ch189p29 漠尘也陨落→漠北也陨落.','The paragraph lists lost young talents after Qing Shu. Mo Chen is the surviving grandfather who negotiates with Fang Yuan; Mo Bei is the dead younger successor.','High contextual correction.','Confirm later cast recaps; keep Mo Chen’s concealed cultivation decline distinct from death.')
pro(89,'barrier disappearance sequence',[(198,89),(198,95),(198,96),(198,99),(198,103)],'Preferred minimal candidate at198p96: delete 否则等到这血罩消失 and retain 否则必让你们死无葬身之地！ Alternatively move actual disappearance to just before the attack, if a reliable witness supports that sequence.','p95 explicitly says the barrier has completely disappeared, but the immediately following threat waits for it to disappear. No second barrier is introduced.','High local ordering conflict; moderate choice of repair.','Source variants and intended threat timing. Do not change barrier mechanics to accommodate the slip.')
pro(90,'Stone Aperture recovery mechanism',[(187,56),(187,68),(188,93),(196,20),(198,47),(198,54)],'If Stone Aperture stops native recovery immediately, remove the native-aptitude contribution at188p93 and qualify187p56: the new snow-silver essence is supplied through Lotus-generated natural essence, not an independent aperture output. Do not remove187p68 automatic refinement of supplied essence.','187 and188 appear to credit innate recovery after the transformation;196 explicitly says Stone Aperture has no recovery ability, and198 says this ability returns when the stone effect is reversed. External essence conversion and native generation must be distinguished.','High mechanism query; conditional repair.','A progressive shutdown, different meanings of recovery, and later Stone Aperture explanations remain possible. Do not invent a hidden cutoff time.')
pro(91,'tournament horizon',[(179,79),(188,61)],'Conditional candidate188p61 今后百年→今后数年 if the line states the formal allocation term. Otherwise preserve the century as Xiong Lin’s claim about long-term consequences.','179 assigns profits for the next few years;188 describes a century of consequences. These need not be identical scopes, so a categorical numerical correction is premature.','Low-to-moderate issue; explicit conditional proposal.','Later tournament rules and whether this is rhetorical political impact rather than contract duration.')
pro(92,'blood river names',[(182,87),(184,11),(184,17),(188,74)],'Ch184p17 血海蟒→血河蟒; Ch188p74 血河墓地→血湖墓地.','The named snake is Blood River Python and the established burial site is Blood Lake Cemetery. These two local substitutions swap the neighbouring water terms.','High local consistency.','Later naming variants. Working English labels are identifiers, not an approved glossary.')
pro(93,'river propulsion and steering',[(200,6),(200,15),(200,54),(200,63),(200,72),(200,85)],'Conditional English at200p63: If they lost the sail, they would lose their means of steering toward shore. Chinese candidate 竹筏就无动力→竹筏就难以驶向岸边.','The current already drives the raft downstream. Sail damage threatens directed travel toward the bank; it does not remove all downstream propulsion. The original absolute no-power formulation conflates those functions.','Moderate physical-staging clarification.','Read subsequent raft and river travel. Wind and lateral approach conditions are unquantified; no invented speed or distance reduction.')
fixes=[
(181,43,'这是个迷','这是个谜'),(181,70,'和十绝体根本相比','和十绝体根本无法相比'),(181,77,'至始至终','自始至终'),(181,87,'无以伦比','无与伦比'),(181,93,'大炙','大炽'),(181,96,'无法在面对','无法再面对'),
(182,5,'微楞','微愣'),(182,62,'淬不及防','猝不及防'),(182,62,'被转入','被卷入'),(182,86,'怎么有会有','怎么会有'),
(183,38,'更准确的说','更准确地说'),(183,40,'恶名昭昭','恶名昭著'),(183,53,'临时反扑','临死反扑'),
(184,12,'眉眼出生长','眉眼处生长'),(184,12,'尽显无漏','尽显无遗'),(184,33,'身上中了蛊','身上种了蛊'),(184,47,'中下感动身受蛊','种下感同身受蛊'),(184,86,'山壁出','山壁处'),
(185,48,'血胡','血湖'),(185,66,'湖地','湖底'),(185,67,'正要被撞实了','真要被撞实了'),(185,79,'先前窜动','向前窜动'),(185,80,'蛇口长得巨大','蛇口张得巨大'),(185,94,'身上忧伤','身上有伤'),(185,102,'转身为血鬼飞僵','转生为血鬼飞僵'),
(186,17,'脑经','脑筋'),(187,58,'自取其辱','自取灭亡'),(187,66,'数千年的元莲仙尊','数千年前的元莲仙尊'),(187,75,'以无心算有心','以有心算无心'),
(188,10,'缓慢而决绝地回升','缓慢而坚定地回升'),(188,25,'用多少，就多少','用多少，就少多少'),(188,46,'树野','树叶'),(188,48,'收刮铭牌','搜刮铭牌'),(188,71,'电转','电钻'),(188,79,'个迷','个谜'),(188,88,'三对触角','三对触脚'),
(189,34,'悠然游然','悠然自得'),(189,37,'面现冲来的熊群','面向冲来的熊群'),(189,71,'却不留血','却不流血'),(189,85,'公认作弊','公然作弊'),
(190,15,'连忙微寒','脸庞微寒'),(190,34,'年纪相反','年纪相仿'),(190,82,'如影随形','如影随形'),(190,83,'听声辩位','听声辨位'),(190,93,'甩了一个颜色','甩了一个眼色'),(190,95,'不能让逃了','不能让他逃了'),(190,96,'就在这是','就在这时'),(190,98,'形式古朴','形制古朴'),(190,98,'脸色一罩','脸上一罩'),
(192,6,'我们的曾经的样子','我们曾经的样子'),(192,17,'尽夺蛊重','尽夺蛊虫'),(192,26,'狼潮上扬','浪潮上扬'),(192,29,'翅膀煽动','翅膀扇动'),(192,41,'古月一眼眶','古月一代眼眶'),
(193,21,'楞了一楞','愣了一愣'),(193,30,'转身为血鬼尸','转生为血鬼尸'),(193,33,'帮助他的身躯','绑住他的身躯'),(193,38,'张开怪叫','张口怪叫'),(193,85,'似的的蛊','似的蛊'),
(194,14,'承当了几次攻击','承受了几次攻击'),(194,47,'当他的血鬼尸躯','但他的血鬼尸躯'),
(195,22,'生机飘渺','生机渺茫'),(195,42,'这种留念','这种留恋'),(195,43,'竟然真的北冥冰魄体','竟然真的是北冥冰魄体'),(195,44,'在出现时','再出现时'),(195,71,'不能在控制','不能再控制'),
(196,8,'青矛山','青茅山'),(196,43,'绕你一命','饶你一命'),(196,53,'方源达到意识记忆','方源的意识记忆'),(196,53,'黄绿相见','黄绿相间'),(196,68,'显露起光团','显露出光团'),(196,88,'得意洋洋“','得意洋洋：“'),
(197,18,'顺势一体','顺势一提'),(197,27,'还在未方源','还在为方源'),(197,50,'对方源也来','对方源而言'),(197,61,'滔尽天下水','倾尽天下水'),(197,84,'砰然心动','怦然心动'),(197,86,'那我需要做？','那我需要做什么？'),
(198,8,'他们一起动手','我们一起动手'),(198,9,'再这下去','再这样下去'),(198,32,'绕过我们','饶过我们'),(198,84,'暂且绕过你','暂且饶过你'),
(199,11,'精力观望','凝神观望'),(199,13,'双臂尽是','双臂尽失'),(199,14,'一切都表面','一切都表明'),(199,25,'芊芊玉手','纤纤玉手'),(199,76,'睁开疲惫沉重的双眼','睁开疲惫沉重的双眼'),(199,79,'一直飞鹤','一只飞鹤'),
(200,11,'青矛竹','青茅竹'),(200,32,'青矛竹','青茅竹'),(200,33,'不在扩大','不再扩大'),(200,57,'青矛竹','青茅竹'),(200,79,'身上伤害数十处','身上伤口数十处')]
fixes=[z for z in fixes if z[2]!=z[3]]
for c,p,a,b in fixes:assert a in src(c,p)['text'],(c,p,a)
v=pro(94,'located copy-edit candidates',sorted(set((c,p) for c,p,a,b in fixes)),'Apply only the individually listed substitutions after approval; never as a blanket replacement.','Local typos, wrong characters, omissions and a few explicit diction candidates. Proper-name, arithmetic and continuity proposals are separately recorded.','High for obvious character slips; moderate for diction candidates.','Compare source witnesses and retain deliberate idiom/voice. Do not treat this mixed-confidence bundle as automatic approval.')
v['copy_items']=[dict(chapter=c,paragraph=p,original=a,replacement=b,status='HOLD; proposed only',confidence='Moderate contextual wording' if (c,p) in [(183,40),(183,53),(184,12),(185,102),(187,58),(187,75),(188,10),(189,34),(190,15),(193,30),(195,22),(197,61),(199,11)] else 'High local copy candidate') for c,p,a,b in fixes]
updates={'SR-001':'Ch198 introduces a conflicting43% retrospective baseline. New SR-085 holds the linked42→52→62 proposal while preserving the stated final90%.', 'SR-007':'199 Bai100→90 is10capacitypoints. FY199 retains oldRevolution3essence afterRevolution1demotion, then200 usesbronze. Recovery, capacity andquality remain distinct.', 'SR-053':'Current-life destruction in195/198–199 and anotherrewind do not settle conflicting dates for the original500-year former life.', 'SR-060':'200p58 says Bai’s carriedGu died in the self-explosion. This later loss does not determine whether IceMuscle remained in140.', 'SR-067':'181p9 age15 refers to the original recording.190p34 apparent相仿 candidate supplies no exact Bai age.', 'SR-081':'190p68 again states300paces while roots flutter, without stated ground contact. Mode/cultivation dependency remains unresolved; do not certify169200→300 unconditionally.'}
for v in ledger:
 if v['id'] in updates:v['batch009_update']=updates[v['id']]
events=load('history/before-batch009-chronology-register.json')['entries'];resources=load('history/before-batch009-resource-register.json')['entries'];ne=[];nr=[]
def ev(timeline,kind,event,constraint,rr,status='Explicit source anchors; interpretation remains revisable'):
 v=dict(id=f'T{len(events)+1:03d}',timeline=timeline,kind=kind,event=event,constraint=constraint,evidence_status=status,dependency='Whole-novel reconciliation outstanding',sources=ss(rr));events.append(v);ne.append(v)
def res(kind,entry,calc,rr,status='Held evidence register; no manuscript edit'):
 v=dict(id=f'R{len(resources)+1:03d}',kind=kind,entry=entry,calculation=calc,status=status,sources=ss(rr));resources.append(v);nr.append(v)
ev('Recorded past / current hypothesis','false inference','Fang Yuan misidentified as Ten Extreme','Age15 is recorded; narrator explicitly rejects the conclusion.',[(181,9),(181,35),(181,36)])
ev('Current shared pre-rewind sequence','early arrival','Letter accelerates Tie investigation','Original plan nextyear; craneletter triggers early arrival and separate bloodcalamity investigation.',[(182,6),(182,8),(184,40),(184,45),(192,6)])
ev('Current shared sequence / local flashback','scene offset','Lotus acquisition','182 cuts back15minutes to stone investment and acquisition.',[(182,35),(182,39),(182,55),(182,56)])
ev('Original500-year former life','conflicting milestone','Blood Sea inheritance date','183 says after300-plus and after400years; distinguish fromRevolution6 after400-plus andGradeB after200-plus.',[(183,39),(183,72),(177,28),(198,59)],'Unresolved discrepancy; SR-086')
ev('Ancient history / recollection','lifespan and founding','Founder and Sky Crane survive through unconventional methods','Founder hidden nearly1000years; former sectbrothers; decades flight/collection; suspendedlife is not normal uninterrupted lifespan.',[(184,49),(185,95),(192,4),(192,5),(192,11),(192,18),(192,21)])
ev('Current shared pre-rewind sequence','day boundary','Preliminary contest before main tournament','179 dayaftertomorrow;186dawn maincontest tomorrow. The actual fight is preliminary, not an erroneously advanced final.',[(179,79),(186,64),(186,66),(186,73)])
ev('Current shared sequence','elapsed travel','Underground escape','3–4hours through rubble; halfday wallwork (possibly idiomatic); additional30minutes to expose spider.189surface fight atsunset.',[(188,37),(188,67),(188,70),(189,1)])
ev('Current shared pre-rewind sequence','confirmed deaths','Xiong and clan leaders','Xiong JiaoMan alive189p21, explicitbeheadingp69. Three clanleaders die191p83. Earlier unresponsiveness does not prove an earlierdeath.',[(189,21),(189,24),(189,69),(191,83)])
ev('Current shared sequence / flashback','confirmed death','Tie Xue Leng and delayed Gu','193 recalls fatalbattle and finaltouch; contamination powersGu afterownerdeath.',[(190,98),(190,100),(193,25),(193,27),(193,61)])
ev('Overwritten first traversal after landing point','superseded outcome','Founder massacre and first ice explosion','194–196original: founder raisesaptitude, Bai explodes, founder survivesice, SkyCrane killsPython, FYselfexplodes. Not outcomes toduplicate inrevisedsequence.',[(194,75),(195,64),(195,68),(196,5),(196,7),(196,46)])
ev('Rewind transition','landing point','Second successful rebirth','196 returns just before BloodSkull extraction, matching193. Not another500yearrewind. Sameevents replay until intervention197.',[(193,83),(193,86),(196,54),(196,58),(196,67),(196,69)])
ev('Revised traversal after second rebirth','divergence','Founder thrown outside barrier','197changeschoice; Lotusoffer not transfer; founderlosesaccess, FYrefinessealedGu aftersealexpires.',[(197,3),(197,4),(197,21),(197,60),(197,87),(197,89)])
ev('Revised traversal','elapsed duration','Massacre and pursuit','1982–3shichen≈4–6hours, thenfurtherfight/flight, then2shichen≈4hour wait. No exactsunrise assigned.',[(198,53),(198,84),(198,85),(198,89),(198,92)])
ev('Revised traversal','death/revival','Bai explosion and Yin revival','Bai trapsSkyCrane, usesbotharms, revivedbeforeconsciousnessfullygone; femaleform, aptitude90, Revolution3retained.',[(199,4),(199,6),(199,14),(199,19),(199,23),(199,27),(199,33)])
ev('Revised traversal / parallel scenes','two-day cutaway','Sky Crane escapes and rescues Fang Zheng','199two days later, survivorsunderice; FangZhengwasmissing, notkilledinbarrier.3–5daydeathforecast notelapsed.',[(193,75),(199,35),(199,54),(199,65),(199,68),(199,80)])
ev('Revised traversal / VolumeTwo','five-day travel','River journey after Qing Mao','Corpus200=VolumeIICh1. Five-dayjourney overlaps the two-daycutaway, notseven days. WhiteBoneMountain notyetreached; attackforcesearlyshore.',[(200,1),(200,15),(200,16),(200,22),(200,38),(200,84),(200,86)])
res('Acquisition / cash','Lotus acquired and refined','182 roughly50000stones invested; not proof exact161gift was10000 or allcash fitinsideflower.',[(161,13),(182,39),(182,44),(182,55),(182,56)])
res('Control / contamination','BloodGu and ambient essence','Male governs eachbatgroup; BloodFrenzy enablesambientuptake but loss ofcontrol/death; BloodGuillotine doesnotneedowneressence.',[(183,43),(183,48),(183,49),(183,50),(184,64),(184,79)])
res('Acquisition / loss','Stone Aperture and Lightning Wings','Stoneconsumed187; Revolution3initial→peak. Wingsdiscarded/cutapart188, no laterinventorycredit.',[(187,40),(187,43),(187,45),(188,15),(188,16)])
res('Recovery / rank','Lotus output versus stone-aperture function','Sourceconflict heldSR090. Oldlight-silver notinstantlyallconverted; laterRevolution1 retainsold snow-silver toactivateYin.',[(187,56),(187,62),(187,68),(188,93),(196,20),(199,20)])
res('Acquisition / disposal','Thousand-Li Earth Wolf Spider','Refined188, contaminated189, abandoned190, explicitdead193. SourceRevolution5 retained.',[(188,80),(188,82),(189,56),(190,77),(193,65)])
res('Autonomous protection','White Form Snake','Unrefined, fights thenflees190.200saysno furthernews; excludesitfromBaiGu frozeninsidethebody.',[(190,62),(190,63),(200,12),(200,58)])
res('Sealing / barrier','Blood Curtain and Tie restraints','6mu barrier, Revolution5consumable, ownercanexitnotenter; Chains/Talisman contaminatedandexpire.197sealremoval precedesrefinement.',[(194,12),(194,13),(194,18),(194,26),(194,48),(194,51),(197,60),(197,87)])
res('Acquisition after rewind','Blood Skull and Yin–Yang pair','Two namedGu types comprise threeindividualGu; allRevolution4 refined197. Lotusonlyoffered, notdelivered.',[(194,3),(197,4),(197,49),(197,89),(197,92)])
res('Aptitude / rank','Blood Skull transforms Fang Yuan','43baseline conflicts42;53/63dependent candidates52/62; final90explicitretained. Rankfalls3→2→1initial asstoneeffectundone.',[(159,41),(198,25),(198,37),(198,41),(198,47),(198,54),(198,55),(198,56)])
res('Healing / ownership','Yin consumed, matching Yang retained','Bai100→90, Revolution3retained, femaleform. SpecificmatchingYang required; remainsFYproperty, notgiven toBai.',[(199,19),(199,26),(199,27),(199,30),(199,31),(199,32),(199,33),(199,49)])
res('Gu losses / preserved effects','Postexplosion inventory','Bai carriedGu frozen; escapedSnake distinct. FYdoubleboarstrength persistsRevolution1. Canopy/Ear/Lotus/Yang verified; nohealingGu.',[(200,12),(200,34),(200,44),(200,58),(200,65),(200,67),(200,88)])
res('Travel / conjecture','Crane King and raft','FYguessesCraneKing dead200;199showsSkyCrane ona giantcrane. Identitynotexplicit enough toassertnewcraneorresurrection. >8000kmriver, >1000li/dayclaim retained.',[(196,7),(199,85),(200,1),(200,15),(200,18)],'Crane death remains character conjecture, not confirmed inventory loss')
np=sum(len(c['paragraphs']) for c in batch);nc=sum(sum(map(len,c['paragraphs'])) for c in batch);tp=sum(len(c['paragraphs']) for c in cs);tc=sum(sum(map(len,c['paragraphs'])) for c in cs)
save('proposal-ledger.json',ledger);save('chronology-register.json',dict(scope='Sequential evidence through200; branch labels are overwritten/revised histories, not coexisting universes',entries=events));save('resource-register.json',dict(scope='Sequential evidence through200; distinguish holdings at rebirth landing point',entries=resources))
progress=load('history/before-batch009-reading-progress.json');progress.update(sequential_chapters_completed=list(range(1,201)),paragraphs_read=tp,paragraph_characters_read=tc,next_chapter=201,remaining_chapters=2134,batch_report='Batch-009-Chapters-181-200.md',scope='Full sequential Chinese-source reading through200; not fresh bilingual/live-site audit',whole_novel_chronology_verified=False)
progress['batches'].append(dict(id='009',chapters=[181,200],paragraphs=np,characters=nc,reading_segments='Ten complete untruncated two-chapter outputs'))
progress['chapters'] += [dict(chapter=c['corpus_chapter'],file=c['file'],title=c['title'],paragraphs=len(c['paragraphs']),first_paragraph=1,last_paragraph=len(c['paragraphs'])) for c in batch]
save('reading-progress.json',progress)
save('batch-009-reading-log.json',dict(status='Complete',endpoint=200,next_unread=201,segments=[dict(chapters=[a,a+1],status='Full paragraph sequences printed and read; untruncated',paragraphs=sum(len(by[n]['paragraphs']) for n in [a,a+1])) for a in range(181,201,2)],note='No invented timestamps; extraction alone is not completion'))
text=['# Numbered Chinese source: Chapters181–200','Original EPUB wording retained. Proposals are separate; corpus200 is VolumeIIChapter1.']
for c in batch:
 text += [f"\n## Corpus chapter {c['corpus_chapter']}: {c['title']}",f"Volume: {c['volume']} | EPUB: {c['file']}"]+[f"{i}. {t}" for i,t in enumerate(c['paragraphs'],1)]
(r/'Batch-009-Numbered-Chinese-Source.md').write_text('\n\n'.join(text))
for name,vs in [('Chronology',ne),('Resources',nr)]:
 out=[f'# Batch009 {name} jotter','All evidence provisional; no source edits.']
 for v in vs:
  out += [f"## {v['id']}: {v.get('event',v.get('entry'))}",v.get('constraint',v.get('calculation')),v.get('evidence_status',v.get('status'))]
  if 'timeline' in v:out.append('Timeline: '+v['timeline'])
  for z in v['sources']:out += [f"Ch{z['chapter']} p{z['paragraph']}",'> '+z['text']]
 (r/f'jotters/Batch-009-{name}.md').write_text('\n\n'.join(out))
out=['# Batch009 Continuity and proposed corrections','Every proposal remains unapproved and unapplied.']
for v in new:
 out += [f"## {v['id']}: {v['category']}",f"**Action:** {v['recommended_action']}",f"**Rationale:** {v['rationale']}",f"**Confidence:** {v['confidence']}",f"**Dependencies:** {v['dependencies']}",f"**Status:** {v['publication_status']}"]
 if 'copy_items' in v:
  for z in v['copy_items']:out += [f"### Ch{z['chapter']} p{z['paragraph']}: {z['original']} → {z['replacement']}",z['confidence'], '> '+src(z['chapter'],z['paragraph'])['text']]
 else:
  for z in v['locations']:out += [f"Ch{z['chapter']} p{z['paragraph']}",'> '+src(z['chapter'],z['paragraph'])['text']]
(r/'jotters/Batch-009-Continuity.md').write_text('\n\n'.join(out))
(r/'jotters/Batch-009-Preserve.md').write_text('''# Batch009 preservation cautions

- 181’s Ten Extreme hypothesis is explicitly false, not a new aptitude fact.
- Xiong Jiao Man was unresponsive earlier, not explicitly dead.189 supplies her actual death.
- Record the original500-year life separately from the overwritten and revised Qing Mao sequences. Do not infer parallel worlds from ledger branch labels.
- Bai’s two narrated explosions occur on different traversals of history. Do not add their consequences together.
- ThreeGu at197 means BloodSkull plus Yin and Yang. It is not a two-versus-three inventory error.
- A Gu’s aura can leave the aperture even when the Gu itself cannot. WeakCicada aura still suppressesRevolution4 after rebirth.
- Revolution changes do not automatically erase stored higher-quality essence or permanent physical strength.
- Lotus is acquired182. The197 offer is a ruse, not a transfer. Yin is consumed199; Yang remains with FangYuan.
- WhiteSnake leaves before Bai’s carriedGu freeze. Its survival does not contradict200’s internal inventory loss.
- FangZheng is missing before the massacre and survivesunderice; no resurrection. SkyCrane’s jadecoffin preserveslife, not confirmeddeath.
- 199’s two-day cutaway overlaps200’s five-day journey. Do not make seven days.
- Keep the final90% aptitude; do not reduce it to89 merely because the starting43 conflicts with42.
- Preserve large river dimensions and li units. No global geography shrinkage or invented conversion of paces to metres.
- No healingGu exists at the endpoint. WhiteBoneMountain remains ahead, not reached.
''')
stand=load('history/before-batch009-working-standards-register.json');stand['coverage_through']=200
su={'WS-01':'181 explicitly false Ten Extreme diagnosis;200 Crane King death is Fang Yuan’s guess. Neither overrides observed evidence.', 'WS-02':'198 has a real42/43 conflict: SR-085. Final90 retained.199/200 distinguish retained snow-silver reserves from new bronze essence.', 'WS-03':'187 StoneAperture raises stage;198 BloodSkull reverses it down toRevolution1 while aptitude reaches90.199 Bai keepsRevolution3 afterYin.', 'WS-04':'197 two namedGu types are three individuals. Snake unrefined and escapes; Yin consumed and matchingYang retained.186/187 Cicada cannotleaveaperture physically.', 'WS-05':'196 secondrewind requires separate overwritten/revised event labels.199two-day cutaway overlaps200five-dayjourney.198时辰 durations retained.', 'WS-06':'182Lotuspayment/acquisition completed;197offer notdelivered. Earlier approximategift rounding remains approximate.', 'WS-07':'190Ear300 withoutstatedgroundcontact expandsSR081.200river8000-pluskm andraft1000-plusli/day preservedpendingglobalcomparison.', 'WS-08':'181age15 isrecorded past;190similar-age wording cannotestablishBai exactage.', 'WS-09':'189MoChen/MoBei and198barrierordering areheldlocalcandidates. Interpretivecopyitems haveindividualconfidence.', 'WS-10':'Coverage200/2334 crossesintoVolumeII;2134unread. No whole-novel certification or manuscriptapproval.'}
for v in stand['entries']:v['batch009_update']=su[v['id']]
save('working-standards-register.json',stand)
out=['# Working editorial standards','## Evidence checkpoint: Chapters1–200','No story-wide factual standard is whole-novel-confirmed. These rules guide review, not publication approval. Earlier evidence retained; latest qualifications below.']
for v in stand['entries']:out += [f"## {v['id']}: {v['topic']}",f"**Status:** {v['status']}",f"**Working rule:** {v['rule']}",f"**Evidence:** {v['evidence']}",f"**Latest qualification:** {v['batch009_update']}",f"**Reopen when:** {v['reopen']}"]
(r/'WORKING-STANDARDS.md').write_text('\n\n'.join(out))
print('REGISTERS',np,nc,tp,tc,'proposals',len(ledger),'copy',len(fixes),'chronology',len(events),'resources',len(resources))
def quote(c,p):return f"**Ch{c} p{p}:**\n\n> {src(c,p)['text']}"
report=f'''# Sequential close-reading review: Chapters 181–200
## Batch 009 | The Omniarch

**Completed: Chapters 1–200 of 2,334. Next unread: 201. Remaining: 2,134.**

This batch contains **{np:,} paragraphs / {nc:,} source characters**, read in ten complete, untruncated two-chapter segments. Cumulative coverage is **{tp:,} paragraphs / {tc:,} characters**. Corpus Chapter 200 is **Volume Two, Chapter One**, not a duplicate Chapter One in the corpus numbering.

The ledger contains **94 proposal records: 93 active and held, one withdrawn**. Ten records are new. SR-094 contains **{len(fixes)} individually located copy candidates**, with lower confidence marked for interpretive wording rather than obvious character slips. No proposals are approved or applied. This is Chinese-source review, not a fresh audit of the live English translation.

## 1. The principal numerical finding: 42% becomes 43% without an intervening change

The earlier loss is explicit:

{quote(159,41)}

Chapters 162p56 and 167p78 independently retain 42%. But Chapter 198 changes the retrospective baseline:

{quote(198,25)}

The first ten-point gain is then followed by another:

{quote(198,37)}

**SR-085, preferred linked candidate, held:**

| Location | Current source | Candidate |
|---|---|---|
|198p25, retrospective baseline|43%|42%|
|198p25, first gain|53%|52%|
|198p37, second gain|63%|62%|
|198p56, retrospective starting value|43%|42%|
|198p56, final aptitude|90%|**Keep 90%**|

This is a linked correction, not a blanket numerical reduction. The final aptitude follows nearly ten applications with progressively smaller gains. The text does not supply enough fixed increments to derive 89%. Likewise, the original 44% before the explicitly narrated loss remains valid.

Later recaps and source witnesses still need checking before publication. The contradiction is strongly supported locally; that does not make the remedy approved.

## 2. Another rebirth: the chronology now needs separate traversals

The events are not one uninterrupted sequence. Chapter 196 returns Fang Yuan to the moment immediately before the Blood Skull Gu is extracted, already narrated at the end of Chapter 193.

{quote(196,67)}

{quote(196,69)}

### The chronology map

| Sequence | What belongs to it |
|---|---|
|Original 500-year former life|His old cultivation, inheritance and sect history; separate from the current Qing Mao disaster|
|Shared current-life sequence through the landing point|Lotus acquisition, Stone Aperture use, Wings and Spider losses, three clan leaders killed, Tie Xue Leng’s death and delayed restraints|
|Overwritten first traversal after the landing point|Founder raises his aptitude, Bai explodes, founder survives the ice, Sky Crane kills the python, Fang Yuan self-explodes|
|Revised traversal beginning in196|Events replay until197; Fang Yuan throws the founder outside, takes the sealed Gu, raises his own aptitude; Sky Crane kills the founder; Bai traps Sky Crane and is revived|

These labels are editorial bookkeeping. Chapter 187 describes one world and one river of time, not several simultaneously existing universes.

**Do not double-count** the Blood Skull acquisition, both narrated explosions, both sets of casualties, or both durations. Knowledge travels back with Fang Yuan; the later physical outcomes of the overwritten sequence do not become possessions or events in the revised one.

## 3. The volume ending is not an excuse to stop

This checkpoint crosses into Volume Two, and immediately adds evidence that qualifies earlier assumptions:

- Bai’s carried Gu died in the explosion. The snake had already left, so it is not part of that loss.
- Fang Yuan’s physical strength persists despite his rank falling to Revolution One.
- His old snow-silver essence briefly survives that demotion, but by Chapter 200 he is using bronze essence.
- The final aptitude is 90%, yet operating Revolution Three Sky Canopy remains extremely costly at Revolution One.
- The party has no healing Gu at the endpoint.

Working standards can stabilise as methods. Factual claims about all cultivation, all inventories or all travel still cannot be certified from one volume.

## 4. Time anchors: forecasts, flashbacks and overlapping cutaways

### The investigator’s early arrival now has a cause

{quote(182,6)}

The white crane’s letter causes Tie Xue Leng to abandon his earlier schedule. Chapter 184 explains that it reports a blood calamity; Chapter 192 confirms that Sky Crane sent it. This supports keeping Chapter 113’s later date as an earlier plan, not altering it to match the eventual arrival.

### The tournament date is not contradicted by the preliminary fighting

Chapter 179 places the main tournament on the day after tomorrow. At dawn in186, the main event is tomorrow and the **preliminary** event begins. It lasts until sunset and requires thirty nameplates for qualification. It is not a prematurely held main tournament.

### The escape and the revised barrier sequence

- Chapter 182p35 looks back fifteen minutes to the Lotus attempt.
- Chapter 188 gives three or four hours of tunnel travel, then “half a day” of wall work and a further half-hour. The half-day expression may be idiomatic; it is not automatically twelve precisely measured hours.
- Chapter 198’s two or three 时辰 amount conventionally to about four to six hours, followed later by another two 时辰, about four hours, while Sky Crane pursues the founder.
- Those durations must not be silently halved to keep every event in daylight. No exact sunrise timestamp has yet been established for the later fighting.
- Chapter 199 cuts to Sky Crane **two days later**. Chapter 200 resumes the fugitives after **five days of river travel**. These are overlapping viewpoints, not a seven-day total.

### A genuine former-life date query

Chapter 183p39 places the Blood Sea inheritance after three hundred-plus years; p72 says after four hundred years. **SR-086** holds a possible “centuries later” formulation without pretending that it settles the exact year. Reaching Revolution Six after four hundred-plus years and gaining Grade B after two hundred-plus years are separate milestones, not automatic fixes for the inheritance date.

## 5. Revolution, stored essence and Stone Aperture must be kept separate

Stone Aperture Gu is consumed in187 and raises Fang Yuan from Revolution Three initial stage to peak stage. His existing light-silver reserves are not instantly replaced wholesale with snow-silver essence.

Blood Skull Gu later reverses the stone transformation. His aptitude rises while his cultivation falls, eventually reaching Revolution One initial stage. This is an explicit special interaction, not a general rule that better aptitude lowers rank.

Chapter 199 then says:

{quote(199,20)}

That explains the otherwise puzzling activation of a Revolution Four healing Gu by a newly demoted Revolution One cultivator. Stored essence quality and current rank are not identical variables.

### Unresolved native-recovery mechanism: SR-090

Earlier passages seem to credit the transformed aperture’s own recovery, particularly188p93. But196 says:

{quote(196,20)}

Chapter 198p54 says native recovery is available again after the aperture returns to a light membrane. The question is whether187–188 loosely refer to refining Lotus-supplied natural essence, describe residual recovery before a later shutdown, or contain a genuine mechanism slip.

The conditional repair removes an independent native-recovery contribution where necessary. It **does not** remove the ability to refine supplied natural essence. No hidden progressive shutdown is being invented as a settled answer.

## 6. Acquisition, disposal and control ledger

| Item | Confirmed sequence through200 |
|---|---|
|Heavenly Essence Treasure Lotus|Actually captured and refined182; the197 offer is a ruse, not a transfer|
|Stone Aperture Gu|Consumed187; effect later reversed by Blood Skull|
|Lightning Wings Gu|Discarded and cut apart188; not simply injured forever|
|Thousand-Li Earth Wolf Spider|Refined188, old contamination reactivates189, abandoned190, explicitly dead193|
|White Form Immortal Snake|Unrefined autonomous protector; flees190, no further news by200|
|Blood Skull Gu|Taken sealed197, refined only after the seal expires|
|Yin–Yang Rotation Gu|A pair: two individuals. Together with Blood Skull, three Gu, not a counting error|
|Yin Gu|Consumed to revive Bai199, restoring the body in female form and lowering aptitude from100% to90%|
|Matching Yang Gu|Retained by Fang Yuan; not handed over to Bai|
|Bai’s remaining carried Gu|Declared frozen to death200; later loss does not resolve the old140 inventory query|

The Lotus’s approximately50000-stone investment is now completed. Chapter 182 rounds the earlier near10000 gift to10000, but this does not create a new receipt or prove that all the money was packed into Tushita Flower.

The Earth Treasury Flower comparison in188p81 names the Liquor Worm. **SR-087** proposes White Boar Gu instead, supported by the actual flower extraction in63. Jade Skin and the Ear are other flower acquisitions; the list of all inherited Gu must not be mistaken for a list of flower contents.

## 7. Deaths and apparent deaths: preserve the distinctions

**Xiong Jiao Man:** Chapter 189 first shows her leading the assault, with her new Revolution Three cultivation explained. Her beheading is explicit atp69. The earlier unresponsive condition in150 was not an explicit death. There is no established double death to correct.

**Tie Xue Leng:** The mask and iron hand protect his daughter before the text explicitly confirms his death through the recollection in193. His contaminated Gu continuing to act is explained by their absorption of ambient essence. It is not evidence that he remains alive.

**Fang Zheng:** He was separated from the retreat and never placed inside the final massacre. Sky Crane finds him alive under the ice two days later. That is survival, not resurrection.

**Sky Crane:** Jade burial preserves him inside the ice. The narrative never confirms his death there.

**The Crane King:** Fang Yuan guesses in200 that it may have died because Sky Crane returned alone. Chapter 199 nevertheless shows Sky Crane riding a giant crane. The exact mount identity is not stated there. Keep the death claim as Fang Yuan’s inference; neither invent a replacement mount nor assert that a dead king resurrected.

## 8. Other held continuity proposals

| Record | Issue | Proposed handling |
|---|---|---|
|SR-088|189p29 says Mo Chen died among the lost young talents|Prefer Mo Bei; the surviving grandfather negotiated with Fang Yuan shortly before|
|SR-089|198p95 says the barrier is gone; p96 threatens what happens when it disappears|Prefer a present threat without the already-expired condition; witness-based scene reordering is an alternative|
|SR-091|188’s century of tournament consequences versus179’s few-year profit allocation|Conditional query only: political impact may be broader than the formal allocation period|
|SR-092|Blood Sea Python and Blood River Cemetery slips|Prefer established Blood River Python and Blood Lake Cemetery respectively|
|SR-093|200p63 says losing the sail removes all propulsion|Prefer loss of the ability to steer toward shore; the river current still drives downstream movement|
|SR-094|{len(fixes)} local copy candidates|Individual source locations and confidence in the Continuity jotter and workbook; no blanket substitution|

Examples of clear copy candidates include 忧伤→有伤 in185p94, 蛊重→蛊虫 in192p17, 帮助→绑住 in193p33, and 双臂尽是→双臂尽失 in199p13. Interpretive diction candidates are deliberately marked more cautiously.

## 9. Range and geography are still open

Chapter 190 again gives a300-pace Ear range without explicitly describing ground contact. This expands SR-081; it does not prove that the200-pace value in169 should always be300. Mode, cultivation and effective operating range still need reconciliation.

Chapter 200 describes the Huang Long River as over8000kilometres long and the raft as travelling over1000li per day. These are preserved. The total river length, daily downstream travel and remaining distance to White Bone Mountain are different measurements. No global reduction of geography is authorised.

Likewise, a hundred-li competition area is not a statement that every contestant travels a hundred li, and the vast ice imagery does not automatically supply a surveyed mountain diameter.

## 10. Endpoint and saved work

Fang Yuan and Bai Ning Bing have been forced ashore **before reaching White Bone Mountain**. Their raft is wrecked, both are injured, Fang Yuan has Revolution One initial cultivation and90% aptitude, Bai retains Revolution Three cultivation and90% aptitude, and neither has a healing Gu available.

Saved with this report:

- Full numbered, unemended Chinese source for181–200.
- Four jotters: Chronology, Resources, Continuity and Preserve.
- The close-reading working notes and ten-segment reading log.
- Updated94-record proposal ledger, **{len(events)} chronology entries**, **{len(resources)} resource entries**, workbook and ten-standard register.

**Next actual unread:201.** Whole-novel verification remains incomplete. No source, site or repository edits, uploads or deletions occurred; remote research backup remains unresolved.
'''
(r/'Batch-009-Chapters-181-200.md').write_text(report)
# Workbook: cumulative evidence, not merely this batch's summary.
w=Workbook();s=w.active;s.title='Read me'
for row in [('Sequential reading','Status'),('Completed','1–200 of2334'),('Next unread',201),('Paragraphs',tp),('Characters',tc),('Proposals','94total;93activeheld;1withdrawn'),('New copy candidates',len(fixes)),('Whole-novel certified','No'),('Approved/applied changes','None'),('Scope','Chinese-source review, not a fresh live-site audit'),('Timeline','Overwritten and revised histories must not be counted twice'),('Source note','Original EPUB unchanged; earlier151 restricted excerpts retain omission markers')]:s.append(row)
s=w.create_sheet('Chapter completion');s.append(['Chapter','File','Title','Paragraphs','Status'])
for c in cs:s.append([c['corpus_chapter'],c['file'],c['title'],len(c['paragraphs']),'Full sequential read'])
def ids(vs,k,c,p):return ', '.join(v['id'] for v in vs if any(z['chapter']==c and z['paragraph']==p for z in v[k]))
s=w.create_sheet('Paragraph register');s.append(['Chapter','Paragraph','Status','Proposal IDs','Chronology IDs','Resource IDs','Chinese reading copy'])
for c in cs:
 for p,t in enumerate(c['paragraphs'],1):
  n=c['corpus_chapter'];s.append([n,p,'Read; not certified error-free',ids(ledger,'locations',n,p),ids(events,'sources',n,p),ids(resources,'sources',n,p),reader(n,p,t)])
s=w.create_sheet('Proposal ledger');s.append(['ID','Category','References','Action','Rationale','Confidence','Status','Dependencies','Publication','Approved','Applied','Latest update'])
for v in ledger:
 latest=next((v[f'batch{i:03d}_update'] for i in range(9,0,-1) if f'batch{i:03d}_update' in v),'')
 s.append([v['id'],v['category'],refs(v['locations']),v['recommended_action'],v.get('rationale','See original report'),v.get('confidence','See originating report'),v['status'],v['dependencies'],v['publication_status'],'No','No',latest])
s=w.create_sheet('Copy items');s.append(['Proposal','Chapter','Paragraph','Original','Candidate','Confidence','Status'])
for v in ledger:
 for z in v.get('copy_items',[]):s.append([v['id'],z['chapter'],z['paragraph'],z['original'],z['replacement'],z.get('confidence','See originating proposal'),z['status']])
s=w.create_sheet('Chronology');s.append(['ID','Timeline','Kind','Event','Constraint','Evidence status','Dependencies','References','Exact source'])
for v in events:s.append([v['id'],v['timeline'],v['kind'],v['event'],v['constraint'],v['evidence_status'],v['dependency'],refs(v['sources']),'\n'.join(reader(z['chapter'],z['paragraph'],z['text']) for z in v['sources'])])
s=w.create_sheet('Resources');s.append(['ID','Kind','Entry','Calculation','Status','References','Exact source'])
for v in resources:s.append([v['id'],v['kind'],v['entry'],v['calculation'],v['status'],refs(v['sources']),'\n'.join(reader(z['chapter'],z['paragraph'],z['text']) for z in v['sources'])])
s=w.create_sheet('Working standards');s.append(['ID','Topic','Status','Rule','Evidence','Reopening trigger','Latest qualification'])
for v in stand['entries']:s.append([v['id'],v['topic'],v['status'],v['rule'],v['evidence'],v['reopen'],v['batch009_update']])
for s in w:
 s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
 for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='163B4B')
 for row in s.iter_rows(min_row=2):
  for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
 for col in s.columns:s.column_dimensions[col[0].column_letter].width=min(65,max(14,max(len(str(c.value or '')) for c in col[:30])+2))
w.save(r/'Sequential-Reading-and-Proposals.xlsx')
code=(r.parent/'make_whole_novel_audit_docx.py').read_text().replace("root=Path(__file__).parent","root=Path('/home/user/editorial/sequential-review')").replace('Whole-Novel-Aperture-and-Travel-Audit','Batch-009-Chapters-181-200').replace('WHOLE-NOVEL SOURCE AUDIT','SEQUENTIAL REVIEW / CHAPTERS181–200').replace('The Omniarch: Aperture and Travel Source Audit','The Omniarch: Sequential Review, Chapters181–200').replace('Full-corpus topical screening with selective contextual verification','Twenty complete chapters; rewritten history requires separate event tracking')
exec(compile(code,'docx-export','exec'),{})
(r/'CURRENT-STATUS.md').write_text(f'''# Sequential review checkpoint

Completed **1–200 of2334**. Next unread **201**. Remaining **2134**. Suggested next batch201–220, adaptive to density.

Batch009: {np}paragraphs/{nc}characters. Cumulative: {tp}paragraphs/{tc}characters.94proposalrecords:93activeheld,1withdrawn. SR094contains{len(fixes)}copycandidates.145chronology entries,84resources. No approved/applied changes. Whole-novel chronology NOTcertified.

Latest deliverable: Batch-009-Chapters-181-200.md/.docx. Source: Batch-009-Numbered-Chinese-Source.md. Four009jotters + complete batch-009-working.md, readinglog, registers, workbook and workingstandards updated.200=VolumeIICh1.

Carry forward:
- SR085 linked19843→42,53→52,63→62,final90PRESERVED. Actual159/162/16742 explicit. No globalaptituderecalculation.
- SR086183formerlifeinheritance300plusvs400 unresolved; proposed nonnumericalhundreds. Revolution6after400plus andGradeBafter200plusdifferentmilestones.
- SR087188EarthTreasuryLiquorWorm→WhiteBoar witnessed63. SR088189MoChen→MoBei. SR089198barrieralreadygonebeforefuturethreat. SR090stoneaperturenative recovery187/188vs196/198 unresolved. SR091tournamentcenturyvsfewyears conditionalrhetoricalscope. SR092bloodwaternames. SR093sailsteeringvsdownstreampropulsion. SR09491copyitems mixedconfidence.
-196rewind landsjustbefore193BloodSkullextraction. Original500yearlife, sharedcurrenthistory, overwrittenpostlanding194–196, revisedpostlanding196–200 DISTINCT. No duplicate deaths/Gu/payments/time. One time river, notcoexistingworlds.
- Lotusacquired182;197offerNOTtransfer. StoneApertureconsumed187;FYRevolution3initial→peak then198BloodSkullreversal→Revolution1initial/90aptitude.199usesretainedRevolution3essence,200bronze. Strengthpersists. Wingsdiscarded/destroyed188;Spideracquired188contaminated189abandoned190dead193. Snakeunrefinedflees190. Bai carriedGu frozen200 doesnotresolve140inventory.
-197BloodSkull+YinYangpair=3Gu, refinedafterTalismanexpiry.199Yinconsumed:Bairevivedfemale,Revolution3/90aptitude;matchingYangstillFYowned. CicadaweakafterrewindcanstillrepressRevolution4aura, cannotphysicallyleaveaperture186/187.
- XiongJiaoManalive189thenexplicitbeheading;no150doubledeath. Threeclanleadersdead191. TieXueLengdeathconfirmed193flashback. RevisedSkyCrane killsfounder thenfrozenbutalive;199escapes2dayslater/rescuesmissingFangZheng. FY200CraneKingdeathconjectureNOTfact;199giantcranerideexactidentityunconfirmed.
-1992daycutaway overlaps2005dayjourney, not7days.186prelimsnotmaincompetition.18215minflashback.1982–3shichen≈4–6h thenlater2shichen≈4h; noforceddaylightclock.
-190Ear300withoutstatedrootcontactexpandsSR081;stillconditional.200river8000+km/raft1000+liadaypreserved. WhiteBoneMountainNOTreached:forcedashore,injumed,nohealingGu,raftwrecked. FullreadingcontinuesbeyondVolumeI.

No source/site/repository edits, uploads or deletions. Remote backup unresolved. OriginalEPUBunchanged. Do not rerun earliergenerators unchanged.009generator safely reruns frompre009history.
'''.replace('injumed','injured'))
assert len(cs)==200 and len(ledger)==94 and len(new)==10
assert np==1825 and nc==63205 and tp==16853 and tc==606662
assert all(not v['approved'] and not v['applied'] for v in ledger)
assert all(z['text']==src(z['chapter'],z['paragraph'])['text'] for v in events+resources for z in v['sources'])
texts={t for c in cs for t in c['paragraphs']}
for line in report.splitlines():
 if line.startswith('> '):assert line[2:] in texts
print('QA PASSED',dict(chapters=200,paragraphs=tp,characters=tc,proposals=len(ledger),new_records=len(new),copy_items=len(fixes),chronology=len(events),resources=len(resources),next=201))
