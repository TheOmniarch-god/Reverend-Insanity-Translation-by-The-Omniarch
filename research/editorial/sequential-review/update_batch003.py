from pathlib import Path
import json,shutil
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
r=Path(__file__).parent
load=lambda n:json.loads((r/n).read_text())
save=lambda n,x:(r/n).write_text(json.dumps(x,ensure_ascii=False,indent=2))
progress=load('reading-progress.json')
assert progress['next_chapter']<=51,'Do not regress a later completed batch.'
chapters=[]
for b in ['001','002','003']:chapters+=load(f'batch-{b}-source.json')
by={x['corpus_chapter']:x for x in chapters};batch=chapters[20:]
assert list(by)==list(range(1,51))
h=r/'history';h.mkdir(exist_ok=True)
for n in ['reading-progress.json','proposal-ledger.json','chronology-register.json','Sequential-Reading-and-Proposals.xlsx']:
 dest=h/('before-batch003-'+n)
 if not dest.exists():shutil.copy2(r/n,dest)
def locs(pairs):return [dict(chapter=c,paragraph=p,file=by[c]['file'],text=by[c]['paragraphs'][p-1]) for c,p in pairs]
ledger=load('proposal-ledger.json')
for v in ledger:
 if v['id']=='SR-012':
  v['recommended_action']='Current preferred investigation direction: Ch18 twelve → six, rather than raising Ch20 six. Ch23 states6.5 total early refinement cost. No change approved; precise cost scope and overall ledger still unresolved.'
  v['batch003_update']='Ch23 p49 supports lower expenditure, but needs separate reconciliation with the three-stone advance. Supersedes the tentative direction of raising Ch20 to at least twelve.'
  if not any(l['chapter']==23 for l in v['locations']):v['locations'] += [dict(chapter=23,paragraph=49,file=by[23]['file'])]
  v['dependencies']='SR-019 full opening cash reconciliation; independent source witnesses; distinguish Liquor Worm cost from all refinement.'
new=[
('SR-017','chronology',[(16,9),(23,48),(30,20),(30,36),(39,58),(44,45),(46,32)],'Investigate nearly100-day replay trigger versus roughly46 days since discovery. Candidate more than a month only if later evidence confirms local trigger repair; no exact replacement approved.','Substantial conflict under natural activation model; initial activation inferred','Later legacy/recording explanations and source witnesses; do not invent a hidden earlier activation'),
('SR-018','chronology',[(44,45),(45,34),(47,13),(47,63),(47,73),(47,86)],'Working candidate: change Ch47 murder/dispute/cleanup labels from last night to the night before last as a linked set.','Locally strong second-night versus departure mismatch','Ch51 onward investigation chronology and stay/departure evidence'),
('SR-019','resource accounting',[(23,42),(23,49),(23,52),(23,64),(26,6),(28,71),(28,72),(29,3),(29,4)],'Retain 20 as working downstream balance. Candidate explicitly include the already-paid 3-stone advance (Ch23 p42): subtotal23 minus advance3 equals20, if not included in cost categories. Reconcile44.5 peak separately.','Displayed subtotal differs from cash balance; narrated advance may explain difference. Cost scope unresolved.','SR-012; partial stones; stipends; inventories and unitemised costs'),
('SR-020','resource accounting',[(28,71),(39,58)],'Candidate clarify118 as112 from classmates plus6 in personal stipends across two rounds.','Strong arithmetic explanation, source wording needs qualification','Confirm collection rate and recipients remain unchanged'),
('SR-021','aptitude and staging',[(22,39),(22,40),(22,78)],'Hold linked GradeC/38%/three-blade conflict. Do not change grade alone without addressing non-Moonlight GradeD pupils.','Source mismatch; no safe isolated one-word correction','Later aptitude thresholds and scene evidence; earlier whole-corpus finding'),
('SR-022','capacity',[(29,12),(33,48),(34,7),(43,9),(50,65)],'Preferred local candidate: Ch43 four-and-a-half tenths →44%; preserve all normal44% references.','Strong isolated-number harmonisation','Global aptitude sequence remains under review'),
('SR-023','action continuity',[(27,34),(29,48),(34,63)],'Preserve observed two-finger strike; qualify retrospective claims to mainly palm strikes rather than denying all finger use.','Direct action conflicts with later categorical summaries','Fit each affected sentence minimally; no invented combat actions'),
('SR-024','source names',[(26,9),(31,37),(48,18)],'Ch26 andCh48 MoChen→MoBei; Ch31 MoBei→MoChen. Follow participant identity, not global name replacement.','High contextual confidence','No independent corrected witness verified'),
('SR-025','source names',[(45,62)],'JiaFu→JiaJinSheng as the interlocutor in this sentence.','High contextual confidence','Keep previous and later JiaFu references unchanged'),
('SR-026','source names',[(50,32)],'ChiLian→ChiCheng as grandson thanking grandfather.','High contextual confidence','No global replacement'),
('SR-027','source copy',[(33,12)],'First你→我 in bed comparison: I have a bed to sleep in; what about you?','High contextual confidence','Preserve surrounding taunt'),
('SR-028','source copy',[(24,42)],'根本就能形成有效打击→根本就不能形成有效打击','High contextual confidence','Negative supported by inexperienced users and paragraph reasoning')]
for id,cat,ls,action,status,dep in new:
 v=dict(id=id,category=cat,locations=[{k:v for k,v in x.items() if k!='text'} for x in locs(ls)],recommended_action=action,status=status,dependencies=dep,approved=False,applied=False,publication_status='HOLD: full-corpus reading and global reconciliation outstanding')
 ledger=[x for x in ledger if x['id']!=id]+[v]
save('proposal-ledger.json',ledger)
chron=load('chronology-register.json');events=[v for v in chron['entries'] if int(v['id'][1:])<23]
def ev(id,timeline,kind,ls,title,constraint,status,dep=''):
 events.append(dict(id=id,timeline=timeline,kind=kind,event=title,constraint=constraint,evidence_status=status,dependency=dep,sources=locs(ls)))
ev('T023','current life','next morning',[(21,1),(21,14)],'Fang Zheng arrives after Fang Yuan','Dawn; Fang Yuan collected first prize at midnight','Direct next-morning link')
ev('T024','current life','elapsed interval',[(22,5),(23,48)],'Moonlight lesson after refinement','Ten-odd days in Ch22; more precise16days inCh23 same lesson day','Use16 as explicit anchor; loose十来天 not automatically a separate error')
ev('T025','current life','scheduled then elapsed',[(23,55),(24,1),(26,1),(26,3)],'Assessment and first extortion','Three days to test; subsequent same-day evening stipend','Confirmed local chain')
ev('T026','current life','forecast',[(23,58),(39,21),(39,39)],'Caravan expected then arrives late','Expected in under a month; normallyApril, actualMay; delays named','Preserve forecast and stated travel delay')
ev('T027','current life','schedule',[(26,14),(49,67)],'Class officers and later squads','Next3months for first appointments; later one-year squad horizon','Future schedule, not elapsed time')
ev('T028','current life','weekly interval',[(30,20),(30,36),(31,1),(31,29)],'First to second extortion','3days+4days=7; retrospectively confirmed','Multiple consistent anchors')
ev('T029','current life','mission and forecast',[(31,5),(31,16),(31,23),(31,24)],'Mo Yan returns; next wolf tide forecast','Mission ten-plus/more than halfmonth; tide nextyear, usual3year recurrence','Approximate durations and forecast; not exact future event date')
ev('T030','current life','idiom',[(33,41)],'First/fifteenth-day threat','Saying about inability to evade forever','Not a calendar appointment')
ev('T031','current life','overnight sequence',[(33,52),(34,20),(34,28),(34,31),(35,62),(36,33),(36,36)],'Gao Wan killed following second-extortion night','Late night→pre-dawn→sunrise; morning aftermath','Direct scene sequence')
ev('T032','current life','ordered duration',[(37,54),(37,55)],'Mo Yan confined','Seven-day sentence','Order, not seven days elapsed')
ev('T033','Ren Zu myth','lifespan and physical age',[(38,44),(38,47),(38,53),(38,54),(38,56)],'Lifespan Gu and rejuvenation','1day then1hour remaining;80years lifespan; existing age100; youthful body20','Separate lifespan, calendar age and physical condition')
ev('T034','current life','calendar and interval',[(39,1),(39,13),(39,44),(39,58)],'May caravan arrival','Halfmonth afterGaoWan death; morning region arrival/noon camp/evening market','Explicit month and coarse interval')
ev('T035','current life','market sequence',[(44,1),(44,42),(44,45),(45,34),(45,71)],'Second market day and night','Next noon; second evening; departure morning after tomorrow','Explicit relative dates','SR-018')
ev('T036','previous life and current family schedule','retrospective and current start',[(45,67),(45,68),(45,69)],'Jia inheritance competition','Prior-life struggle6–7years; present secondyear after rules announcedlastyear','Separate remembered outcome from present trajectory')
ev('T037','current life / recording timer','trigger conflict',[(16,9),(46,32)],'Image-wall inheritance trigger','Nearly100days replay versus roughly46days reconstructed since discovery','Activation at discovery inferred, not conclusively dated','SR-017; natural-model mismatch, no invented prior activation')
ev('T038','current life','date conflict',[(44,45),(47,13),(47,63),(47,73),(47,86)],'Murder night versus departure','Second-night murder; departure after3daystay; last-night retrospectives','Locally inconsistent labels under explicit market sequence','SR-018; check investigation witnesses')
ev('T039','current life','elapsed weather interval',[(47,1),(48,1),(48,2)],'Rain after departure','Rain lasts4days; seasonal transition described','Explicit broad interval; not a new absolute June date')
ev('T040','current life','same day and overnight',[(48,11),(49,4),(50,4),(50,63),(50,87),(50,88)],'Combat class to first middle-stage breakthrough','Same-day injury recalled that evening; overnight advance completednextmorning','Direct relative chain')
ev('T041','current life','recurring assistance and forecast',[(50,25),(50,55)],'Other pupils cultivation schedules','ChiLian assistance every3days; FangZheng forecasts3–4days','Do not record forecast as completed breakthrough')
chron.update(scope='Sequential source evidence through Chapter50. Unresolved conflicts retained; no global chronology certification.',entries=events)
save('chronology-register.json',chron)
resources=[]
def rr(id,ls,kind,entry,calculation,status):resources.append(dict(id=id,sources=locs(ls),kind=kind,entry=entry,calculation=calculation,status=status))
rr('R001',[(11,53),(11,65),(12,77),(13,18),(16,25),(16,62)],'nominal opening cash','Six held; inn2; wine2; legacy15','6−2−2+15=17 nominally','Partial earlier stone depletion and settlement timing not fully itemised')
rr('R002',[(18,34),(18,36),(18,64)],'transfer and separate inventory','FY receives5.5;FZlaterreceives6','Keep physicalpieces versus equivalents distinct','Actual5.5 transfer; second6belongs toFZ')
rr('R003',[(18,84),(20,29),(23,49)],'conflicting expenditure','12 in subperiod;6LiquorWorm total;6.5earlyrefinement','Cannot all be literal totals on identicalbasis','SR-012: current preferred investigation direction12→6;held')
rr('R004',[(21,14),(23,49)],'prize and cumulative account','20prize received byCh21;44.5peak recalledCh23','6+15+5.5+20=46.5 gross listed majorreceipts before expenses;42.5 after firstinn/wine4','Not a certifiedcashpeak: timing,earlierdepletion,stipends and expenses remain open')
rr('R005',[(23,35),(23,46),(23,48)],'food consumption versus purchasing','4petals/day;jar2stoneslasts4days','4/10+2/4=0.9/day;16×0.9=14.4','Reported14.5 not automaticallywrong due inventory/rounding')
rr('R006',[(23,42),(23,49),(23,52)],'subtotal and advance query','44.5peak minus6.5refining minus14.5food minus0.5living','44.5−6.5−14.5−0.5=23; subtract the already-narrated3 advance separately and20 remains.','SR-019: determine whether itemised costs exclude the current advance. Do not double-debit it; earlier44.5 peak remains unverified.')
rr('R007',[(23,42),(23,64),(23,66)],'prepayment and inventory','Currentwine/mealadvance3;stockfivejars10;petalsperish5days','Do not subtractprepaymenttwice or countstockconsumptionasnewcashpurchase','Separate cash,creditandfoodinventory')
rr('R008',[(26,6),(28,71),(28,72),(29,3),(29,4)],'linked actual balances','Stipend3;prize10;56levies;then20jarsat2','20−10+3+10+56=79;79−40=39','Localarithmeticconsistent;inheritsunresolved20startingbalance')
rr('R009',[(29,18),(29,23),(29,24),(29,29),(34,13),(34,15),(34,16)],'essence processing','10activation+10processed→5refined','20/5=4overallinput/output;44×4=176initialcapacitypoints','Cumulativethroughput,notsimultaneous176%storage')
rr('R010',[(33,48),(34,7),(34,28),(35,52)],'cultivation and combat','44→12→44;about30minrefill;3stonesovernight;under10afterfight','Do not inferexactbladecountfromonlyexplicitlydescribedattacks','Statedstates;grossflowsnotfullyspecified')
rr('R011',[(36,55),(37,54),(39,58)],'fine and compensation','Fine30;compensation30toFY;50toGaoWanfamily','FYnetfine/compensation0;family50notFYincome','Orderedearlier;Ch39confirmsFYpayment/receipt')
rr('R012',[(28,71),(39,58)],'collection scope','Tworoundsreported118','2×56=112;2×(56+3)=118','SR-020candidateexplicitlyincludeownstipends')
rr('R013',[(39,58),(40,42),(40,46),(41,20)],'reported cash and buying limit','98cash;9affordable;7prudentlimit;6purchased','floor(98/10)=9;actualcost60','Differentquantities,notcontradictorycounts;98notfullyderivedfrompriorledger')
rr('R014',[(42,69),(42,70),(43,17)],'sale and subsequentdepletion','Toadsold500;remaining538;partstoneusedafterreturn','98−60+500=538','Exacttransactionarithmetic;538notpermanentbalance')
rr('R015',[(43,9),(50,65)],'capacity conflict','45%inCh43;44%againCh50','Local45versusrepeated44','SR-022preferred44held')
rr('R016',[(45,11),(45,12),(45,13),(45,80),(45,85),(48,51),(48,54)],'other recipients and unrealisedincome','Fraudvictimgets500;imagesale2000agreednotcompleted;sparprize20notclaimed','Do notcreditanyofthesetoFY','Keeprecipientandtransactionstatus')
rr('R017',[(50,37),(50,47),(50,48),(50,49),(50,54)],'Fang Zheng separate account','RecoverytwiceFY;3remaining;academy3minuslevy1;family3weeklyplusvariableextras','3stonesat1/daylast3days','DonomergewithFYbalance;recoverymatches8vers4earlier')
rr('R018',[(50,65),(50,85)],'breakthrough reserve','44middle-stagepoints→20remaining','Netnarrateddrop24points','Overnightgrosscostnotfullyspecified')
save('resource-register.json',dict(scope='Opening throughChapter50. Explicit local arithmetic distinguished from uncertified cumulative accounts.',entries=resources))
# Progress reflects completed reading, never mere extraction.
progress.update(sequential_chapters_completed=list(range(1,51)),paragraphs_read=3605,paragraph_characters_read=138379,next_chapter=51,remaining_chapters=2284,batch_report='Batch-003-Chapters-21-50.md',whole_novel_chronology_verified=False,publication_policy='All emendations held for complete2334-chapter review and global reconciliation.',scope='Full sequential Chinese source reading throughChapter50; not a fresh bilingual/live-site audit; later topical excerpts excluded from completion.',chapters=[dict(chapter=x['corpus_chapter'],file=x['file'],title=x['title'],paragraphs=len(x['paragraphs']),first_paragraph=1,last_paragraph=len(x['paragraphs'])) for x in chapters])
progress['batches']=[b for b in progress['batches'] if b['id']!='003']+[dict(id='003',chapters=[21,50],paragraphs=2140,characters=83815,reading_segments='Fifteen full untruncated two-chapter outputs')]
save('reading-progress.json',progress)
# Readable exact source copy.
lines=['# Numbered Chinese source: Chapters21–50','Unemended nonempty EPUBbody paragraphs. Source numbering matches Batch003.','']
for x in batch:
 lines += [f"## Chapter {x['corpus_chapter']}: {x['title']}",f"Source: `{x['file']}`",'']
 for i,t in enumerate(x['paragraphs'],1):lines += [f'**{i:03d}.** {t}','']
(r/'Batch-003-Numbered-Chinese-Source.md').write_text('\n'.join(lines))
# Jotters: evidence is copied verbatim, not silently emended.
j=r/'jotters';j.mkdir(exist_ok=True)
def refs(ss):return '; '.join(f"Ch{s['chapter']} p{s['paragraph']}" for s in ss)
ls=['# Chronology jotter: Batch003','All entries provisional for global reconciliation. Day labels in the main report are analytical, not canonical dates.','']
for v in events[22:]:
 ls += [f"## {v['id']}: {v['event']}",f"Timeline: {v['timeline']}. Type: {v['kind']}.",v['constraint'],v['evidence_status'],f"Dependency: {v['dependency'] or 'Continue sequential cross-checking.'}",f"References: {refs(v['sources'])}",'']
 for s in v['sources']:ls += [f"**Ch{s['chapter']} p{s['paragraph']}**",'> '+s['text'],'']
(j/'Batch-003-Chronology.md').write_text('\n'.join(ls))
ls=['# Resource jotter: throughChapter50','Local arithmetic is not equivalent to a fully reconstructed cumulative balance. Do not apply repairs yet.','']
for v in resources:
 ls += [f"## {v['id']}: {v['kind']}",v['entry'],f"Calculation: {v['calculation']}",f"Status: {v['status']}",f"References: {refs(v['sources'])}",'']
 for s in v['sources']:ls += [f"**Ch{s['chapter']} p{s['paragraph']}**",'> '+s['text'],'']
(j/'Batch-003-Resources.md').write_text('\n'.join(ls))
ls=['# Continuity and textual-query jotter: Batch003','All proposals held. Source-copy candidates are not live-English defect claims.','']
for v in ledger:
 if int(v['id'][3:])>=17 or v['id']=='SR-012':
  ls += [f"## {v['id']}: {v['category']}",v['recommended_action'],f"Evidence status: {v['status']}",f"Dependencies: {v['dependencies']}",f"References: {refs(v['locations'])}",'']
ls += ['## Q-001: academy disciplinary rules','Ch28 p31 threatens expulsion for Gu combat; Ch36 p55 gives thirty stones for the first offence, and Ch39 p58 confirms the fine. Preserve as an open credibility/rule-application query. Do not automatically present the initial threat as authoritative omniscient law.','', '## Q-002: middle-stage essence colour','Ch50 p70 uses 翠绿 for splashes amid repeated 苍绿 middle-stage descriptions. Treat as a lexical/lore watch item, not a mandatory correction of every green adjective.','']
(j/'Batch-003-Continuity.md').write_text('\n'.join(ls))
(j/'Batch-003-Preserve.md').write_text('''# Decisions to preserve the text: Batch 003

- Three years without a Grade A talent: verified at Ch1 p50, Ch6 p71, and Ch25 p3. Do not replace years with generations.
- Revolution Three academy elder and Revolution Two martial instructor are different people (Ch22 p16; Ch24 p14; Ch36 p34).
- No opening Fang Yuan Revolution Five claim was located in the reread first-twenty-chapter evidence. Do not manufacture a Revolution Five/Revolution Six conflict from an assumption.
- Effective moonblade range and distance before complete disappearance are not automatically the same (Ch22 pp28–30; Ch24 p38; Ch25 p32).
- The May caravan arrival is explicitly late; April is the usual arrival month (Ch39 pp21,39).
- The wall's initial apparent lack of further treasure is Fang Yuan's limited conclusion. Ch46 reveals more information. Preserve the surprise rather than revising earlier beliefs.
- A large tent can house a small gambling business (Ch40 pp10,24,59).
- Nine stones affordable, seven the prudent ceiling, six selected: three different quantities (Ch40 pp42,46; Ch41 p20).
- The toad's quick refinement is explained by its dormant weakened state (Ch42 pp40–42); it does not invalidate earlier refinement difficulties.
- Three Gu requiring expenditure in Ch42 p66 can exclude the Cicada's self-supplied food. Do not change the count without preserving that context.
- Thirty-three combat victories do not make Fang Yuan physically inexhaustible. Ch48 explicitly distinguishes experience from bodily capacity.
- Injury healed does not eliminate the need for convalescence: Mo Bei's treatment in Ch49 and advice to rest in Ch50 can coexist.
- Forecasts about cultivation, wolf tides, appointments, and political succession are not completed events.
- The proposed image-wall sale and abandoned sparring prize are not received income. Compensation paid to another character is not Fang Yuan's money.
- Ren Zu's mythic rejuvenation distinguishes physical youth from elapsed life and additional lifespan. Do not enter it into Fang Yuan's present-life calendar.
- Preserve character error, bias, bluff, incomplete knowledge, and deliberate misdirection. Contradiction detection is not permission to make every speaker omniscient.
''')
# Cumulative workbook.
w=Workbook();sh=w.active;sh.title='Read me'
for row in [['Sequential review','Current status'],['Completed chapters','1–50 of2334'],['Complete paragraphs',3605],['Source characters',138379],['Next chapter',51],['Publication policy','HOLD all emendations for full-corpus reconciliation'],['Read does not mean','Certified error-free or globally harmonised'],['Scope','Chinese source reading; no fresh live-English validation'],['Chronology','41 evidence entries; unresolved conflicts retained'],['Resources','18 entries; local calculations distinguished from uncertified cumulative balances']]:sh.append(row)
sh=w.create_sheet('Chapter completion');sh.append(['Chapter','File','Title','Paragraphs','Status'])
for x in chapters:sh.append([x['corpus_chapter'],x['file'],x['title'],len(x['paragraphs']),'Full sequential source read'])
sh=w.create_sheet('Paragraph register');sh.append(['Chapter','Paragraph','Reading status','Proposal IDs','Chronology IDs','Resource IDs','Chinese source'])
for x in chapters:
 for p,t in enumerate(x['paragraphs'],1):
  c=x['corpus_chapter'];ids=[v['id'] for v in ledger if any(s['chapter']==c and s['paragraph']==p for s in v['locations'])];ts=[v['id'] for v in events if any(s['chapter']==c and s['paragraph']==p for s in v['sources'])];rs=[v['id'] for v in resources if any(s['chapter']==c and s['paragraph']==p for s in v['sources'])]
  sh.append([c,p,'Read; no error-free certification',', '.join(ids),', '.join(ts),', '.join(rs),t])
sh=w.create_sheet('Held proposals');sh.append(['ID','Category','Locations','Working action','Evidence status','Dependencies','Publication status','Approved','Applied'])
for v in ledger:sh.append([v['id'],v['category'],refs(v['locations']),v['recommended_action'],v['status'],v['dependencies'],v['publication_status'],'No','No'])
sh=w.create_sheet('Chronology');sh.append(['ID','Timeline','Kind','Event','Constraint','Evidence status','Dependencies','References','Exact source'])
for v in events:sh.append([v['id'],v['timeline'],v['kind'],v['event'],v['constraint'],v['evidence_status'],v['dependency'],refs(v['sources']),'\n'.join(f"Ch{s['chapter']} p{s['paragraph']}: {s['text']}" for s in v['sources'])])
sh=w.create_sheet('Resources');sh.append(['ID','Kind','Entry','Calculation','Status','References','Exact source'])
for v in resources:sh.append([v['id'],v['kind'],v['entry'],v['calculation'],v['status'],refs(v['sources']),'\n'.join(f"Ch{s['chapter']} p{s['paragraph']}: {s['text']}" for s in v['sources'])])
for sh in w:
 sh.freeze_panes='A2';sh.auto_filter.ref=sh.dimensions
 for cell in sh[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='163B4B')
 for row in sh.iter_rows(min_row=2):
  for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
 for col in sh.columns:sh.column_dimensions[col[0].column_letter].width=min(65,max(14,max(len(str(x.value or '')) for x in col[:30])+2))
w.save(r/'Sequential-Reading-and-Proposals.xlsx')
# Main report only; no time spent creating separate Word files for every jotter.
code=(r.parent/'make_whole_novel_audit_docx.py').read_text()
code=code.replace("root=Path(__file__).parent", "root=Path('/home/user/editorial/sequential-review')").replace('Whole-Novel-Aperture-and-Travel-Audit','Batch-003-Chapters-21-50').replace('WHOLE-NOVEL SOURCE AUDIT','SEQUENTIAL REVIEW / CHAPTERS21–50').replace('The Omniarch: Aperture and Travel Source Audit','The Omniarch: Sequential Review, Chapters21–50').replace('Full-corpus topical screening with selective contextual verification','Thirty complete source chapters; all emendations held for global reconciliation')
exec(compile(code,'docx-export','exec'),{})
assert sum(len(x['paragraphs']) for x in batch)==2140
assert sum(sum(map(len,x['paragraphs'])) for x in batch)==83815
assert sum(len(x['paragraphs']) for x in chapters)==3605
assert sum(sum(map(len,x['paragraphs'])) for x in chapters)==138379
assert len(ledger)==28 and len(events)==41 and len(resources)==18
assert all(not x['approved'] and not x['applied'] for x in ledger)
assert '—' not in (r/'Batch-003-Chapters-21-50.md').read_text()
print('Verified complete through50. Batch2140paragraphs/83815characters. Total3605/138379. Next51. All28proposalsheld.')
