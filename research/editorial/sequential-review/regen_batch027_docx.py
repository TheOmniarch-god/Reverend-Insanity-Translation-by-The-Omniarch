from pathlib import Path
import re
from docx import Document
from docx.shared import Pt
r=Path('/home/user/editorial/sequential-review')
# DOCX from markdown
md=(r/'Batch-027-Chapters-506-525.md').read_text().splitlines()
doc=Document()
st=doc.styles['Normal'];st.font.name='Calibri';st.font.size=Pt(10.5)
for line in md:
    if not line.strip() or line.strip()=='---': continue
    if line.startswith('# '): doc.add_heading(line[2:],0)
    elif line.startswith('## '): doc.add_heading(line[3:],1)
    elif line.startswith('### '): doc.add_heading(line[4:],2)
    elif line.startswith('- '): doc.add_paragraph(re.sub(r'\*\*(.+?)\*\*',r'\1',line[2:]),style='List Bullet')
    elif line.startswith('|'):
        if set(line.replace('|','').strip())<=set('-: '): continue
        cells=[c.strip() for c in line.strip('|').split('|')]
        if not hasattr(doc,'_t') or doc._t is None or doc._t_done:
            doc._t=doc.add_table(rows=0,cols=len(cells)); doc._t.style='Light Grid Accent 1'; doc._t_done=False
        row=doc._t.add_row().cells
        for k,cv in enumerate(cells): row[k].text=re.sub(r'\*\*(.+?)\*\*',r'\1',cv)
    else:
        doc._t_done=True
        doc.add_paragraph(re.sub(r'\*\*(.+?)\*\*',r'\1',line))
doc.save(r/'Batch-027-Chapters-506-525.docx')
print('docx saved', len(doc.paragraphs), len(doc.tables))
