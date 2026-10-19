import json
from pathlib import Path
from openpyxl import load_workbook
r = Path(__file__).parent
WB = r/'Sequential-Reading-and-Proposals.xlsx'
wb = load_workbook(WB)
led = {v['id']: v for v in json.loads((r/'proposal-ledger.json').read_text())}
src = json.loads((r/'batch-033-source.json').read_text())
ids_new = [f"SR-{i}" for i in range(384, 395)]

# 1. Proposal ledger: append SR-384..SR-394 if absent
pl = wb['Proposal ledger']
present = {pl.cell(rr,1).value for rr in range(2, pl.max_row+1)}
added_ledger = 0
for sid in ids_new:
    if sid in present: continue
    e = led[sid]
    loc = e['locations'][0]
    ref = f"Ch{loc['chapter']} p{loc['paragraph']}"
    pl.append([sid, e['category'], ref, e['recommended_action'], e['rationale'], e['confidence'], e['status'],
               e['dependencies'], e['publication_status'], 'No', 'No', e.get('witness_033', '')])
    added_ledger += 1

# 2. Chapter completion: append 626..645 if absent
cc = wb['Chapter completion']
have = {cc.cell(rr,1).value for rr in range(2, cc.max_row+1)}
added_ch = 0
for c in src:
    if c['corpus_chapter'] in have: continue
    cc.append([c['corpus_chapter'], c['file'], c['title'], len(c['paragraphs']), 'Full sequential read'])
    added_ch += 1

# 3. Paragraph register: append 626..645 paragraphs if absent
pr = wb['Paragraph register']
have_p = {pr.cell(rr,1).value for rr in range(2, pr.max_row+1)}
added_p = 0
# map proposals to paragraphs
pmap = {}
for sid in ids_new:
    loc = led[sid]['locations'][0]
    pmap.setdefault((loc['chapter'], loc['paragraph']), []).append(sid)
for c in src:
    if c['corpus_chapter'] in have_p: continue
    for i, t in enumerate(c['paragraphs'], 1):
        pids = ', '.join(pmap.get((c['corpus_chapter'], i), [])) or None
        pr.append([c['corpus_chapter'], i, 'Read; not certified error-free', pids, None, None, t])
        added_p += 1

# 4. Read me
rm = wb['Read me']
total_active = sum(1 for v in led.values() if v['status'].startswith(('Held','Provisional','Query','HOLD')) or 'query' in v['status'].lower() or 'held' in v['status'].lower())
upd = {
 'Completed': '1–645 of 2334',
 'Next unread': '646',
 'Paragraphs': '58460',
 'Characters': '2044665',
 'Proposals': f"{len(led)} total; {len(led)-5} active (held or query); 4 withdrawn; 1 locally resolved",
}
for rr in range(1, rm.max_row+1):
    k = rm.cell(rr,1).value
    if k in upd: rm.cell(rr,2).value = upd[k]
wb.save(WB)
print('added ledger', added_ledger, 'chapters', added_ch, 'paragraphs', added_p)
print('ledger rows', pl.max_row-1, 'chapter rows', cc.max_row-1, 'paragraph rows', pr.max_row-1)
