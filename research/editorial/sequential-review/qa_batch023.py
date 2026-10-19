from pathlib import Path
import json,zipfile,hashlib,re
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from docx import Document
r=Path(__file__).parent
load=lambda n:json.loads((r/n).read_text())
b=load('batch-023-source.json');led=load('proposal-ledger.json');p=load('reading-progress.json')
z=zipfile.ZipFile(r.parents[1]/'benchmark/sources/Gu-Zhen-Ren.epub')
for c in b:
    actual=[n.get_text('',strip=True) for n in BeautifulSoup(z.read(c['file']),'html.parser').find_all('p') if n.get_text('',strip=True)]
    assert c['paragraphs']==actual,c['corpus_chapter']
assert [c['corpus_chapter'] for c in b]==list(range(426,446))
src=(r/'Batch-023-Numbered-Chinese-Source.md').read_text();cur=None;found=[]
for line in src.splitlines():
    m=re.match(r'## Chapter (\d+) ',line)
    if m: cur=int(m[1])
    m=re.match(r'(\d+)\.(\d+) (.*)',line)
    if m:
        assert int(m[1])==cur
        assert m[3]==b[cur-426]['paragraphs'][int(m[2])-1];found.append((cur,int(m[2])))
assert found==[(c['corpus_chapter'],i) for c in b for i in range(1,len(c['paragraphs'])+1)]
assert len(found)==1744
ids=[v['id'] for v in led];assert len(led)==208 and len(set(ids))==208
assert all(not v.get('approved') and not v.get('applied') for v in led)
for v in led:
    if v['id'] in [f'SR-{i}' for i in range(194,209)]:
        for l in v['locations']: assert 1<=l['paragraph']<=len(b[l['chapter']-426]['paragraphs'])
assert p["next_chapter"]==446 and p['sequential_chapters_completed'][-1]==445 and len(p['sequential_chapters_completed'])==445
assert p['paragraphs_read']==38245+1744 and p["paragraph_characters_read"]==1328051+sum(len(t) for c in b for t in c['paragraphs'])
assert p["remaining_chapters"]==1889
w=load_workbook(r/'Sequential-Reading-and-Proposals.xlsx',read_only=True,data_only=True)
cc=[row for row in w['Chapter completion'].iter_rows(min_row=2,values_only=True) if row[0] and int(row[0])>=426]
assert len(cc)==20
pr=[row for row in w['Paragraph register'].iter_rows(min_row=2,values_only=True) if row[0] and int(row[0])>=426]
assert len(pr)==1744 and all(row[6]==b[row[0]-426]['paragraphs'][row[1]-1] for row in pr)
pl=[row for row in w['Proposal ledger'].iter_rows(min_row=2,values_only=True) if row[0] and str(row[0]).startswith('SR-18') or (row[0] and str(row[0]) in [f'SR-{i}' for i in range(194,209)])]
assert len([x for x in pl if x[0] in [f'SR-{i}' for i in range(194,209)]])==15
rm=dict(w['Read me'].iter_rows(min_row=2,max_col=2,values_only=True));assert rm["Next unread"]=="446"
d=Document(r/'Batch-023-Chapters-426-445.docx');txt='\n'.join(x.text for x in d.paragraphs)
assert 'SR-194' in txt and 'SR-208' in txt and len(d.tables)==1
rep=(r/'Batch-023-Chapters-426-445.md').read_text();assert '—' not in rep
for line in rep.splitlines():
    if line.startswith('> '): assert line[2:] in {t for c in b for t in c['paragraphs']}
for f in (r/'fact-check-022').glob('ch*-hetushu.html'): pass
qa={'status':'passed','epub_paragraphs_rechecked':1744,'numbered_source_exact':True,'numbered_source_lines':1744,'workbook_chapter_rows':20,'workbook_paragraph_rows':1744,'workbook_paragraph_text_exact':True,'workbook_ledger_rows_added':15,'workbook_read_me_next_unread':446,'docx_proposal_ids_present':True,'docx_table_count':1,'ledger_total':208,'approved_or_applied':0,'witness_hashes_verified':0,'report_blockquotes_exact':True,'em_dashes_in_report':0,'progress_next':446,'progress_paragraphs':39989,'progress_characters':1391431}
(r/'batch-023-qa.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2));print(json.dumps(qa,ensure_ascii=False,indent=2))
