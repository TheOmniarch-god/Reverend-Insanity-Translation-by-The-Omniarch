from pathlib import Path
import json,shutil
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
r=Path('/home/user/editorial/sequential-review')
def load(n):return json.loads((r/n).read_text())
def save(n,v):(r/n).write_text(json.dumps(v,ensure_ascii=False,indent=2))
for n in ['reading-progress.json','proposal-ledger.json','chronology-register.json','resource-register.json']:
 d=r/'history'/('before-batch006-'+n)
 if not d.exists():shutil.copy2(r/n,d)
chapters=sum([load(f'batch-{i:03d}-source.json') for i in range(1,7)],[]);by={c['corpus_chapter']:c for c in chapters};batch=chapters[110:140]
def src(c,p):return dict(chapter=c,paragraph=p,file=by[c]['file'],text=by[c]['paragraphs'][p-1])
def ss(rr):return [src(*x) for x in rr]
def loc(rr):return [{k:v for k,v in z.items() if k!='text'} for z in ss(rr)]
def refs(z):return '; '.join(f"Ch{x['chapter']} p{x['paragraph']}" for x in z)
ledger=load('history/before-batch006-proposal-ledger.json');new=[]
for v in ledger:
 if v['id']=='SR-047':
  v['superseded_action']=v['recommended_action'];v['recommended_action']='WITHDRAW the proposed usual-April versus usual-summer harmonisation. Retain the source pending the separate internal season-label query SR-054.';v['status']='Withdrawn: earlier diagnosis too categorical';v['publication_status']='WITHDRAWN; do not apply';v['batch006_update']='Ch116 calls July early autumn; Gregorian season assumptions cannot establish April versus summer as a contradiction. Ch139 creates a separate internal seasonal tension, not proof that SR-047 was sound.';v['rationale']='The previous diagnosis assumed a month-to-season mapping that the text does not establish. This withdrawal does not certify the whole calendar.';v['confidence']='High confidence in withdrawing the categorical diagnosis.'
 if v['id'] in ['SR-001','SR-007','SR-021']:
  v['batch006_update']='Ch115p30 and135p72 retain44%;135p40 distinguishes GradeA maximum99% from TenExtreme100%. No blanket aptitude recalculation. SeeSR-052 for local shot arithmetic.'
def prop(i,cat,rr,action,reason,conf,dep):
 v=dict(id=f'SR-{i:03d}',category=cat,locations=loc(rr),recommended_action=action,rationale=reason,confidence=conf,status='Active provisional proposal; no approval',dependencies=dep,approved=False,applied=False,publication_status='HOLD: full2334-chapter reading and global reconciliation outstanding');ledger.append(v);new.append(v);return v
prop(49,'source negation',[(111,72)],'一个都能少→一个都不能少. English: He could afford to neglect neither his cultivation nor his Gu.','Preparation requires both; missing不 reverses the proposition.','High contextual confidence.','Independent witness not yet checked; source-copy proposal, not a verified live-English error.')
prop(50,'awakening schedule',[(109,56),(110,16),(111,80),(113,11),(123,74)],'If the speaker is meant to give current information, 今年要参加开窍大典→今年参加了开窍大典. Otherwise preserve the prospective wording as stale reported information and annotate the discrepancy.','The ceremony is announced for day after tomorrow, then tomorrow, but the seventh-day caravan scene still describes it prospectively. Winter later confirms YaoLe is this year’s GradeB newcomer.','High timing query; moderate need to alter dialogue.','Check exact day boundaries, speaker knowledge and later ceremony references. Do not postpone the whole ceremony automatically.')
prop(51,'essence remainder',[(115,30),(115,47),(116,1),(116,19),(116,33)],'If exact accounting is intended, Ch116p19只剩下另一半→还剩下六分. English: six percentage points remained.','Eleven points minus five equals six; the following five-point moonblade leaves one. Half is defensible as approximate prose, not exact arithmetic.','High arithmetic; moderate editorial need.','Preserve aperture-capacity units and possible recovery during action. Do not change eleven to ten.')
prop(52,'shot capacity',[(114,24),(115,30),(115,47),(115,48),(122,9),(122,10),(135,72)],'If Ch122p10 states the maximum number of full-cost shots, 十六记→十七记. English: seventeen shots. Alternatively explicitly label sixteen as a conservative firing budget.','With44% capacity and10/5/2.5-point cost, floor totals are4/8/17. Doubling the rounded-down8 incorrectly loses one available shot. Six high-stage shots leave29points, consistent with more than half remaining.','High conditional arithmetic; held pending global mechanics.','Confirm cost scaling, full-strength shots, other costs and whether sixteen is intended as a maximum. Preserve44%; do not introduce an unstated reserve as fact.')
prop(53,'previous-life destruction date',[(124,57),(126,30),(124,38),(135,51),(135,53)],'Preferred interim candidate at Ch135p51: 狼潮过去之后的三年后→狼潮过去之后. English: After the wolf tide had passed. Retain both original numerical claims in the evidence register until the global chronology determines which one to restore.','January alliance recollection places the disaster two years later; the later recollection puts it three years after the wolf tide has passed. Both describe the former life. Current-life butterfly effects do not reconcile two accounts of that same history.','High conflict; moderate nonnumerical repair.','Full previous-life timeline, wolf-tide duration and later destruction recaps. Do not automatically change two to four or pretend the proposed deletion is approved.')
prop(54,'internal seasonal labels',[(116,90),(117,27),(124,57),(127,22),(127,84),(129,20),(131,78),(132,63),(139,27)],'Provisional minimal candidate: Ch139p27在这炎热的夏季→在这炎热的时节; English: in this hot weather. Hold direction until all month/season anchors are reconciled.','Ch116 explicitly labels July early autumn, whereas Ch131 lateJuly plus8–9days inCh132 leads into Ch139’s explicit hot summer. This is an internal seasonal tension, not merely a Gregorian-calendar objection.','High evidence of differing labels; moderate need/direction of repair.','Weather can differ by year; exact calendar unidentified. SR-047 withdrawn separately. No automatic conversion to lunar or Gregorian dates.')
prop(55,'dialogue roles',[(125,40),(125,58),(125,59)],'你防守你进攻→我防守你进攻. English: I defend; you attack.','XiongJiang proposes defending while BaiBingYi attacks; subsequent action confirms those roles.','High.','No independent corrected witness checked.')
prop(56,'Gu name',[(132,74),(132,76),(132,78)],'Ch132p76两步芳草蛊→三步芳草蛊. Retain one Three-Step Fragrant Grass Gu name.','Chapter title and adjacent references identify oneGu, not a second two-step variety.','High.','Later nomenclature check; no purchase yet established.')
prop(57,'reported killing versus narrated death',[(132,24),(132,33),(132,54),(133,96),(135,11),(138,16)],'If literal self-report is intended, revise Ch133p96 to: 前些时候，杀了一个白家的蛊师，有些趣味，又从一个熊家的蛊师尸体上夺了这只蛊虫，你看看。 Otherwise preserve his claim and mark it as an inaccurate attribution.','A wolf kills XiongZhan; BaiNingBing kills a BaiGuMaster and loots XiongZhan. Ch135 accuses him of withholding aid and138 explicitly calls his responsibility indirect.','High account discrepancy; moderate emendation.','A boast may claim indirect responsibility. Do not rewrite the witnessed deaths to fit dialogue or treat FangYuan’s later self-serving account as neutral.')
prop(58,'uncertain casualty quantity',[(131,67)],'Working candidate 二五个小组→二十五个小组, twenty-five groups, ONLY after checking another witness. Until then label the number corrupt/uncertain, not an established25.','二五 is malformed in this ordinary numerical context. Inserting十 is plausible but not the only possible corruption.','Moderate candidate; high need for witness verification.','Independent source edition required before settling25 versus a range/other reading. Group removals are not identical to number of individual deaths.')
prop(59,'probability wording',[(138,71)],'Working English: Two of the three clans would favour him; only one would favour Bai Ning Bing. Preserve the tactical comparison without claiming measured two-thirds odds.','Counting two favourable clan labels out of three does not establish equal encounter probabilities or winning probability.','High mathematical qualification; moderate adaptation preference.','Could be an intentionally rough tactical heuristic. Check patrol distribution if exact probabilities are to be retained.')
prop(60,'BaiNingBing RevolutionThree inventory',[(137,12),(137,90),(139,24),(140,23),(140,24),(140,25)],'Keep two as the current source reading. If later evidence confirms IceMuscle is no longer held or is non-active, clarify that only two active RevolutionThree combatGu remained before the sacrifice; otherwise revisit the total.','IceMuscle is explicitly RevolutionThree and permanently changed his body. Ch140 counts only FrostDemon and BlueBirdIceCoffin. Prior use does not prove continued ownership, and passive permanence does not itself prove consumption.','High query; low confidence in any numerical correction.','Determine IceMuscle consumption/disposal mechanism. Do not silently change two→three or one→two.')
fixes=[(114,11,'从从','从'),(114,16,'后面的时候','后面的石猴'),(116,3,'不能能','不能'),(116,69,'四五天猴头','四五颗猴头'),(117,79,'不敢兴趣','不感兴趣'),(118,9,'这只是不是蟾死了','这只蟾是不是死了'),(118,36,'河岸出','河岸处'),(118,84,'躯干出来','驱赶出来'),(118,87,'蛊虫气息的会','蛊虫的气息会'),(119,21,'无妨惩处','无法惩处'),(120,12,'接着一股股','借着一股股'),(121,51,'四位酒虫','四味酒虫'),(121,67,'每次合练蛊师','每次合炼蛊虫'),(121,72,'隐去本。','隐去本体。'),(122,65,'有了残狼','有的残狼'),(122,72,'不属于青书','不输于青书'),(123,70,'前世后两百年后','前世两百年后'),(125,102,'你一次你是','这一次你是'),(127,31,'显露了狼群','陷入了狼群'),(128,5,'在外门绕','在外面绕'),(128,8,'再次抵挡石林','再次抵达石林'),(129,69,'无比辟易','无不辟易'),(130,4,'独自一人到陷入','独自一人陷入'),(130,36,'伤势了斗志','丧失了斗志'),(130,69,'已经入傍晚','已经如傍晚'),(130,71,'有何可俱','有何可惧'),(130,72,'就在之时','就在这时'),(130,96,'就这古月鹏','就着古月鹏'),(131,18,'人生入戏','人生如戏'),(131,86,'平死一搏','拼死一搏'),(132,11,'但若此次必须','但此次必须'),(135,43,'海水解口','海水解渴'),(136,28,'发现了目光','发现了目标'),(137,38,'的确让他的真元的确','的确让他的真元'),(138,47,'被被狼群','被狼群'),(139,6,'他人对的真正实力','他人对他的真正实力'),(139,27,'周围了空气','周围的空气'),(139,40,'月刃站在','月刃斩在'),(139,85,'青书扔想','青书仍想')]
for c,p,a,b in fixes:assert a in src(c,p)['text'],(c,p,a)
v=prop(61,'source-copy bundle',[(c,p) for c,p,a,b in fixes],'Apply only the individually listed minimal substitutions after approval; see copy_items. No global substitutions.','Contextual miskeys, duplicated characters, missing word endings and malformed phrases. Bundle retains a separate location and replacement for every item.','High contextual confidence overall; each item remains provisional.','Independent witnesses not checked. Preserve the unmodified numbered source. Not an assertion that live English reproduces these errors.')
v['copy_items']=[dict(chapter=c,paragraph=p,original=a,replacement=b,status='Held; unapproved/unapplied') for c,p,a,b in fixes]
save('proposal-ledger.json',ledger)
tdoc=load('history/before-batch006-chronology-register.json');events=tdoc['entries'];te=len(events)
for v in events:
 if v['id']=='T078':v['dependency']='SR-050 now carries the ceremony query. SR-047 withdrawn; SR-054 separately tracks seasonal labels.'
def event(timeline,kind,title,constraint,status,rr,dep='Full-corpus reconciliation outstanding.'):
 events.append(dict(id=f'T{len(events)+1:03d}',timeline=timeline,kind=kind,event=title,constraint=constraint,evidence_status=status,dependency=dep,sources=ss(rr)))
event('current life','auction clocks','Caravan sale timing','Revolution1halfday,Revolution2oneday,Revolution3twodays fromfirstbid; caravan8days; seventhdayRelic.','Display/bid/award/use are distinct.',[(111,13),(111,14),(111,73),(111,79),(111,80),(113,67)])
event('current life','ceremony query','YaoLe prospective wording','109dayaftertomorrow/110tomorrow versus113seventhdayprospective awakening;123currentyearnewcomer.','Speakerinformationmaybestale; exactscene-boundaryopen.',[(109,56),(110,16),(113,11),(123,74)],'SR-050.')
event('current life','duration and negotiation','Return home and inheritance sale','About1yearaway; negotiation>2hours;2–3yearpaybackisbuyerforecast.','Ch90hasnoaunt/FYface-to-facemeeting. Earlierworkingfalsealarmwithdrawn.',[(90,6),(90,18),(112,23),(112,42)])
event('forecast / prior comparison','future dates','Wolf tide and investigator','Marchcaravan forecastsnextyearwave; TieXueLengexpectedyearafter next.','Relativeyearsanchoredtospeakerdate,notalltofirstJiamurder.',[(111,69),(112,5),(113,54),(113,60)])
event('current life','advancement','Relic purchase precedes use','Caravan+3daysstillinitial;115promotesmiddle,additional~15minutesbenefit;30minuteslaterforest.','Do notadvancecultivationatpurchase.',[(114,1),(114,90),(115,16),(115,20),(115,21),(115,30),(115,31)])
event('current life','season and retrospective event','Spring to July','116spring/summertransitionthenJulyearlyautumn; rankdiscovered1dayaftermonkeykill isretrospective.','Do notplacepromotiondiscoveryafterallofsummer.',[(116,81),(116,82),(116,84),(116,90),(117,27)],'SR-054; SR-047withdrawal.')
event('current life','same-day crisis','River-swallowing toad','Appeared3daysbeforemeeting; ChiShanfails;FYdeploys;toaddeparts;500rewardordered.','Toadnotcaptured. AdjacentRenZubranchstorynotelapsedpresenttime.',[(117,3),(119,33),(120,14),(120,45),(120,58)])
event('previous-life future / nested narrative','JiangFan','Future legend recalled','JiangFanarises200+yearslater; quota/monthand3months’fishbelongtohisstory.','Noapertureopened; do notadd200yearstopresentlife.',[(118,4),(118,12),(118,41),(118,57),(118,65)])
event('current life','forecast versus delay','Stone forest progress','2days/3pillars,1.5monthexpected;>2monthsactualslowprogress,firstwintersnow.','Delayexplicitlycausedbywolfactivityandpatrols.',[(122,13),(122,21),(122,22),(122,25),(122,31)])
event('current life','winter calendar','December to January alliance','Winterintensifies;Decembercasualties;Januarywinterendalliance.','NoexactdayorGregorianconversion.',[(124,43),(124,48),(124,57),(124,58)])
event('previous-life recollection','lifespan accounting','Five lifespanGu','5×100extension+natural100=potential600;actual500+beforeforcedself-destruction.','Potentiallifespanisnotelapsedage.',[(126,78),(126,80),(126,82)])
event('previous-life forecast','destruction conflict','QingMao remembered destruction','126two yearslater versus135three yearsafterwolfwave.','Samepriorbranchconflict;notcurrentbutterflyeffect.',[(126,30),(135,51),(135,53)],'SR-053; nofinalnumberchosen.')
event('current life / retrospective','fusion date','HiddenScale completion','127sayscompletedBEFOREJanuaryalliance,despitefirstshownuseafter>1month.','Noautomaticfusionatfirstshownuse.',[(123,49),(123,68),(127,22),(127,37),(127,38)])
event('current life','intervals and bodily adaptation','EarthHearingGrass','Spring>1month+>halfmonth; cavegap>1month;feedseveraldays;implantnight→nextnoonadapted.','Right-earlateralityconsistent; poemwinterimagerynotcalendarchange.',[(127,84),(128,2),(128,49),(128,75),(129,1),(129,7)])
event('current life','wave onset and time skip','Main wolf tide','129summernear;130actualonset;stormovernight;131ten-plusdaysworsening,thenlateJuly.','LateJulynotmerely10daysafterstorm;unquantifiedpassageexists.',[(129,20),(130,51),(130,100),(131,62),(131,69),(131,78)])
event('myth / nested narrative','RenZu','Loneliness heart','Oneovernightvigil,eyesbecomechildren.','NotFYcalendarorbiological-ageclock.',[(131,42),(131,54),(131,57),(131,58)])
event('current life','post-July date and season','Eight or nine days later','132opens8–9daysafterlateJulyaction;139describeshotsummer.','Internaltensionwith116Julyearlyautumn;weatheralonecannotsettlecalendar.',[(131,78),(132,63),(139,27)],'SR-054; donotreinstateSR-047.')
event('current life','continuous battle / feint','Arm injury and sacrifice','135apparentrightfracture→137revealedfeint→139actualright-arm/FrostDemonexplosion.','Leftarmfrozen136doesnotchangewhicharmissacrificed.',[(135,65),(136,92),(137,12),(139,24),(140,23)])
event('current life','endpoint and casualties','QingShu fight begins','Two teammates+YaoHong+elderdead;FangZhengobserves;QingShustartsWoodCharm.','QingShuNOTconfirmeddeadby140. No141+outcomecounted.',[(140,3),(140,55),(140,61),(140,76),(140,87)])
tdoc.update(scope='Sequential chronology throughChapter140; complete-novel chronology not certified.',entries=events);save('chronology-register.json',tdoc)
rdoc=load('history/before-batch006-resource-register.json');resources=rdoc['entries'];re=len(resources)
def res(kind,entry,calc,status,rr):resources.append(dict(id=f'R{len(resources)+1:03d}',kind=kind,entry=entry,calculation=calc,status=status,sources=ss(rr)))
res('auction awards','SecondLiquorWorm,BlackBoar,RedIronRelic purchased; CleanWaternotbid.','LiquorWormwinningmargin20,notfinalprice. RedIron3000display/5000–8000predictionnotawardprice.','Onlyactualpurchasesbooked; exactpricesunknown.',[(111,46),(111,49),(111,64),(111,68),(111,76),(113,48),(113,49),(113,50)])
res('assets and production','Land/houses/wineshop/8servantssold;NineLeafretained.','Threechestsproceedsnotnumericbalance; nineleavesnotproduceddaily.','123leafrevenue>400, notautomatic365-dayyield.',[(112,2),(112,71),(112,76),(121,69),(123,44),(123,45),(123,46)])
res('Moonglow essence','44capacity;10initial/5middle/2.5upperpoints.','Initial4,middle8,upper17maximumconditionally;11−5−5=1.','SR-051/052held; actualsixupperstagecastsleave29.',[(114,24),(115,30),(115,47),(115,48),(116,1),(116,19),(116,33),(122,9),(122,10)])
res('bitter wine and reward','Sixcenturyclams;500reward.','2broken+4whole=6; rewardordered120confirmed121.','Toadnotowned; clamsnotallGu.',[(120,42),(120,58),(121,68)])
res('FourFlavours fusion','TwoLiquorWorms+4winejars→oneFourFlavours;firstattemptsuccess.','>400stones, notnecessarilyexact400;doubleessencequalitynotrankadvancement.','Futurecostexceedsleaf-onlyincome.',[(121,32),(121,38),(121,44),(121,52),(121,61),(121,65),(121,67)])
res('trade and payment','200fromChiShan;BlackBoartradedforFishScale.','Cashdifferencepromisedbutnotquantified; BlackBoarstrengthretained.','HiddenStone+FishScale→HiddenScale completedbeforealliance.',[(123,20),(123,67),(123,71),(127,37),(127,38)])
res('leaf prices','Previousclan55→60;FYasks65;priorlifepeak100.','65−60=5premium; pricequote≠sale.','QingShudeclines; donotrecord65receipts.',[(124,27),(124,29),(124,30),(124,36),(124,38),(124,41)])
res('war currency','10points/eye;5points/jinrice;25points/stone.','Oneeye=0.4stoneatstatedexchange;onewolf2eyes=20points,not1stone.','Futureinflationseparatefromcurrentrates.',[(127,14),(132,68),(132,89)])
res('prospectiveloot','RockGu3900points;2–3operationsintendedforThreeStepGrass.','3900/25=156stoneexchangeequivalent,notactualcashreceived.','NoRockGurecoveryorThreeStepGrasspurchaseestablishedby140.',[(132,77),(133,28),(133,37)])
res('aptitude and lifespan','GradeAupto99%;TenExtreme100%;FY44%.','Five100-yearLifespanGu+100natural=600potential.','Aperturecapacity,quality,recoveryandagearedifferentquantities.',[(126,82),(135,40),(135,72),(140,15)])
res('inventory checkpoint','137explicitseven;138Plundertakenandrefined→eight.','5through110,+Liquor,+Black,+Relic=8;Relicconsumed7;+HiddenStone8;Liquorfusion7;Black/Fishswapnet0;Hiddenfusion6;+EarGrass7;+Plunder8.','Eight:Moonglow,WhiteJade,HiddenScale,FourFlavours,Cicada,EarGrass,NineLeaf,Plunder.',[(115,21),(116,64),(121,48),(127,37),(128,37),(137,46),(138,11)])
res('BaiGu ownership/use','IceMuscleusedpreviously;FrostDemonconsumed;BlueBirdretained.','140twoRevolution3beforeexplosion/oneafterneedsIceMusclestatus,notautomatic3/2.','PlunderheldunrefinedbyBai,instantlyrefinedbyFY; distinguishpossession/control.',[(137,12),(137,90),(138,16),(140,24),(140,25)])
rdoc.update(scope='Resources throughChapter140; balances incomplete where prices or transactions are unspecified.',entries=resources);save('resource-register.json',rdoc)
np=sum(len(c['paragraphs']) for c in batch);nc=sum(sum(map(len,c['paragraphs'])) for c in batch);tp=sum(len(c['paragraphs']) for c in chapters);tc=sum(sum(map(len,c['paragraphs'])) for c in chapters)
p=load('history/before-batch006-reading-progress.json');p.update(sequential_chapters_completed=list(range(1,141)),paragraphs_read=tp,paragraph_characters_read=tc,next_chapter=141,remaining_chapters=2194,batch_report='Batch-006-Chapters-111-140.md',scope='Full sequential Chinese source reading throughChapter140. Not a fresh bilingual/live-site audit.',whole_novel_chronology_verified=False,chapters=[dict(chapter=c['corpus_chapter'],file=c['file'],title=c['title'],paragraphs=len(c['paragraphs']),first_paragraph=1,last_paragraph=len(c['paragraphs'])) for c in chapters]);p['batches'].append(dict(id='006',chapters=[111,140],paragraphs=np,characters=nc,reading_segments='Fifteen complete untruncated two-chapter outputs'));save('reading-progress.json',p)
save('batch-006-reading-log.json',dict(status='Complete',endpoint=140,next_unread=141,segments=[dict(chapters=[c,c+1],status='Full paragraph sequences printed and read; untruncated',paragraphs=sum(len(by[k]['paragraphs']) for k in [c,c+1])) for c in range(111,141,2)],note='Extraction alone not completion. No invented timestamps.'))
lines=['# Numbered Chinese source: Chapters111–140','Unemended nonempty EPUB body paragraphs, numbered to match the review.','']
for c in batch:
 lines += [f"## Chapter {c['corpus_chapter']}: {c['title']}",f"Source: `{c['file']}`",'']
 for i,t in enumerate(c['paragraphs'],1):lines += [f'**{i:03d}.** {t}','']
(r/'Batch-006-Numbered-Chinese-Source.md').write_text('\n'.join(lines))
for name,vs in [('Chronology',events[te:]),('Resources',resources[re:])]:
 lines=[f'# Batch006: {name} jotter','Complete sequential reading throughChapter140. No emendations applied.','']
 for v in vs:
  lines += [f"## {v['id']}: {v.get('event',v['kind'])}"]
  for k in ['timeline','constraint','evidence_status','dependency','entry','calculation','status']:
   if k in v:lines += [f"**{k}:** {v[k]}"]
  for s in v['sources']:lines += [f"\n**Ch{s['chapter']} p{s['paragraph']}**",'> '+s['text'],'']
 (r/'jotters'/f'Batch-006-{name}.md').write_text('\n'.join(lines))
lines=['# Batch006: Continuity and held proposals','13 new ledger records; SR-061 contains39 individually located copy items. SR-047 withdrawn. No source/site changes.','']
for v in new:
 lines += [f"## {v['id']}: {v['category']}"]
 for k in ['recommended_action','rationale','confidence','dependencies','publication_status']:lines += [f"**{k}:** {v[k]}"]
 for s in v['locations']:lines += [f"\n**Ch{s['chapter']} p{s['paragraph']}**",'> '+src(s['chapter'],s['paragraph'])['text']]
 for z in v.get('copy_items',[]):lines += [f"- Ch{z['chapter']} p{z['paragraph']}: {z['original']} → {z['replacement']} (held)"]
 lines += ['']
lines += ['## Withdrawal and restraint','SR-047’s prior April/summer harmonisation is withdrawn. SR-054 is a distinct internal season-label query.','The Ch112 aunt-contact suspicion was discarded after rereading Ch90: FangYuan was not present during the aunt/uncle conversation.','Ch133 signal visibility is100li, while ear reconnaissance is300paces. Neither authorises a global geographical reduction.','Ch129 sound propagation explanation is a narrative analogy, not a measured acoustics model; range also depends on attenuation and ground contact. No Gu range alteration proposed.']
(r/'jotters'/'Batch-006-Continuity.md').write_text('\n'.join(lines))
preserve='''# Batch006: Preserve, do not repair

- Preserve auction display, predictions, bid, award and use as separate events. Clean Water Gu was not bought. Second Liquor Worm, Black Boar and Red Iron Relic were bought; exact prices remain unspecified.
- Nine leaves are not produced every day. Asking65 per leaf is not a completed sale. A previous-life100 peak is not the current price.
- The inherited grass is retained when the other property is sold. Three chests do not yield a numerical cash balance.
- Red Iron Relic is unused at Chapter114 and consumed in115. A purchase does not advance cultivation by itself.
- Preserve44% capacity and separate essence quality from bodily cultivation. Four-Flavours does not turn middle-stage cultivation into upper-stage cultivation.
- A hundred-beast king may command500+ monkeys. These are classes, not exact headcounts.
- The15-minute fight and subsequent repeated expeditions need not be one continuous15-minute event.
- The stone-monkey king is finished by a rival. The chapter title need not specify the final blow literally.
- A predicted1.5-month route clearance can fail; the ensuing delay is explicitly explained.
- Hidden Scale is completed before the alliance but first shown later. Black Boar’s permanent strength survives its trade.
- Earth Hearing Grass replaces the RIGHT ear. Ground contact enables its full300-pace range; moving without rooted contact reduces it below half that range.
- JiangFan’s200-plus-year future story, recipe discoveries, RenZu’s myth and previous-life destruction forecasts do not run on the current-life clock.
- Five100-year lifespan extensions give600 potential years including the natural100, not an actual600-year life before rebirth.
- BaiNingBing’s apparent right-arm fracture is a feint. The later right-arm sacrifice is real. These are not inconsistent healing events.
- RevolutionThree cultivation with diluted RevolutionTwo essence is not an actual rank demotion. RevolutionThree Gu can still be used at excessive cost.
- TenExtreme recovery is fast but not unlimited. Healing while fleeing is explicitly supplied by regeneration.
- A possessed Gu need not be refined: BaiNingBing fails to refine Plunder; FangYuan immediately succeeds.
- FangYuan’s claim that Bai killed the whole Xiong group suppresses his own killing of XiongJiang. XiongLin survives at this checkpoint.
- Chapter140 starts QingShu’s WoodCharm use; it does not yet narrate his death. Never import the next chapter into coverage.
- Do not restore the withdrawn April/summer diagnosis. The separateJuly/autumn versus later-summer query remains unresolved; neither Gregorian nor lunar conversion is established.
- Ch112’s roughly-year-long absence is not contradicted by Ch90: the earlier invitation was rejected through Shen, not a face-to-face visit.
'''
(r/'jotters'/'Batch-006-Preserve.md').write_text(preserve)
report=f'''# Sequential close-reading review: Chapters 111–140
## Batch 006 | The Omniarch

**Completed: Chapters 1–140 of 2,334. Next unread: 141. Remaining: 2,194.**

This batch was read in fifteen complete, untruncated two-chapter sequences: **{np:,} paragraphs / {nc:,} source characters**. Cumulative coverage: **{tp:,} paragraphs / {tc:,} characters**. Extraction was not counted as reading.

This is Chinese-source review, not a fresh bilingual or live-site audit. There are **61 proposal records: 60 active and held, one withdrawn**. Thirteen records are new; one bundles39 individually located copy corrections, so records are not synonymous with individual edits. Nothing is approved or applied. Whole-novel chronology remains uncertified.

## 1. Main chronological finding: the remembered destruction date conflicts

Chapter126, during the January alliance sequence, recalls destruction two years later:

> {src(126,30)['text']}

Chapter135 gives a different anchor:

> {src(135,51)['text']}

These are accounts of the **same previous life**, not a comparison between the old and reborn branches. The wolf tide is still ahead during the alliance sequence, and Chapter124 forecasts at least a year of it. “Two years later” and “three years after the wolf tide” therefore cannot simply be interchanged.

**SR-053, held:** the least committal interim wording is “After the wolf tide had passed,” deleting the disputed three-year interval from135 rather than inventing a replacement date. That is a proposal only. The original figures remain in the register, and the whole-novel evidence must determine which numerical account, if either, should ultimately stand.

The future investigation is a separate clock: Jia Fu says Tie Xue Leng is expected **the year after next** in Chapter113. Jiang Fan’s200-plus-year future history, the future discovery of recipes, and Ren Zu’s overnight vigil are also kept outside the present-life calendar.

## 2. A correction to my earlier diagnosis, plus a separate seasonal query

**SR-047 is withdrawn.** Batch005 treated the usual April arrival and the usual summer arrival as a firm conflict. That was too categorical. The text itself calls July “early autumn” in116; a Gregorian month-to-season assumption cannot settle the earlier passage.

> {src(116,90)['text']}

This does **not** mean the calendar is now consistent. A new, internal tension appears later: Chapter131 reaches late July, Chapter132 advances eight or nine days, and Chapter139 explicitly calls the resulting scene summer:

> {src(131,78)['text']}

> {src(132,63)['text']}

> {src(139,27)['text']}

**SR-054, held:** “in this hot weather” is a minimal candidate for the last phrase. It preserves the heat without committing to a disputed season. The direction remains provisional. Different yearly weather is possible, and no exact lunar or Gregorian calendar has been established. This new query does not reinstate the withdrawn April/summer diagnosis.

### Awakening ceremony: prospective wording lingers into caravan day seven

Chapter109 announces the ceremony for the day after tomorrow; Chapter110 says tomorrow. The caravan’s seventh-day sale sequence nevertheless contains:

> {src(113,11)['text']}

Chapter123 confirms Yao Le is that year’s GradeB academy newcomer. **SR-050:** if the speaker is meant to be current, change “will attend this year” to “attended this year.” If this is stale gossip, preserve it as such. Do not move the entire ceremony to rescue one speaker’s wording.

### An aunt-contact false alarm was rejected

Chapter112’s “not seen for a year” initially warranted checking against90. Rereading90 shows Fang Yuan was **not present** during the aunt and uncle’s conversation: he refused their invitation through Shen. There is no demonstrated recent face-to-face encounter to contradict112. No correction is proposed.

## 3. Essence arithmetic: preserve capacity, fix only what is warranted

| Situation | Full-cost expenditure | Capacity-based result |
|---|---|---|
| RevolutionTwo initial essence |10 percentage points per Moonglow shot|4 shots from44%|
| RevolutionTwo middle essence |5 points per shot|8 shots from44%|
| Upper-quality essence via Four-Flavours |2.5 points per shot under the stated scaling|17 shots from44%, not16|
| Chapter116 counterattack |11 minus5 defence minus5 attack|1 point remaining|

Chapter122 says:

> {src(122,10)['text']}

**SR-052, held:** if these are maximum full-cost shot counts, seventeen is the local mathematical candidate. Doubling an already rounded-down eight loses a usable shot. Six upper-quality shots leave29 points, consistent with122’s “more than half.” Other simultaneous costs or an intentionally conservative firing budget could justify fewer shots, but the text does not state that reserve here.

**SR-051, held:** Chapter116’s “only the other half remained” can be made exact as “six percentage points remained.” It is also defensible as approximate prose. Do not change eleven to ten merely to make “half” literal.

There is **no proposal to alter Fang Yuan’s44% aptitude**. Chapter135 separately confirms GradeA can reach99%, while a TenExtreme physique reaches100%. Capacity, essence quality, recovery rate and cultivation stage remain different variables.

## 4. Purchases, fusion and the eight-Gu checkpoint

### Confirmed acquisitions, not auction estimates

Chapter113 confirms three purchases: the second Liquor Worm, Black Boar Gu and Red Iron Relic Gu. The second Liquor Worm wins by20 stones, but its final price is not given. Red Iron’s display price and predicted winning prices are not completed transaction amounts. Fang Yuan explicitly does not bid for Clean Water Gu in111.

The inherited property sale retains Nine-Leaf Vitality Grass. Three chests of stones do not supply an exact cash balance. Chapter112 explicitly says he is **not producing leaves every day**, so the earlier daily-earnings passage cannot support uninterrupted annual revenue.

The relic remains unused in114 and is consumed in115. Four-Flavours fusion succeeds on the first attempt in121, consumes the two Liquor Worms and four wines, and costs **more than400 stones**. His500-stone toad reward is confirmed, but it is not an exact net-profit calculation.

Black Boar is traded for Fish Scale; the promised price difference is not numerically reported. Hidden Scale is retrospectively confirmed to have been fused **before the alliance**, not at its first shown use in127.

### Inventory reconciliation

Chapter137 explicitly lists seven:

> {src(137,46)['text']}

Chapter138 adds the captured and immediately refined Plunder Gu. The endpoint therefore supports **eight permanent held Gu**:

1. Moonglow Gu
2. White Jade Gu
3. Hidden Scale Gu
4. Four-Flavours Liquor Worm
5. Spring Autumn Cicada
6. Earth Hearing Flesh Ear Grass
7. Nine-Leaf Vitality Grass
8. Plunder Gu

Detached sale leaves are separate stock. Three-Step Fragrant Grass remains a desired purchase, not an acquisition. The Chainsaw Golden Centipede is identified but not captured. Man Shi’s Rock Gu is a prospective3900-point recovery, not booked income.

### Prices and war currency

The clan’s leaf price moves from55 to60; Fang Yuan asks65 and Qing Shu declines. The remembered100-stone peak belongs to the former life. Current exchange is10 merit points per wolf eye,5 per jin of rice, and25 per primeval stone. Two eyes yield20 points, not one whole stone.

## 5. Bai Ning Bing: preserve the reveals, query the actual discrepancies

- **Apparent fracture versus real amputation:** the right-arm fracture is described tentatively in135 and exposed as a feint in137. Chapter139’s right-arm sacrifice is real. The frozen left arm in136 does not change that sequence.
- **Suppression is not demotion:** RevolutionThree cultivation remains; diluted RevolutionTwo essence restrains its expression. RevolutionThree Gu are expensive to use, not absolutely unusable. Full RevolutionThree expression returns in140.
- **Possession is not refinement:** Bai held Plunder for days without completing refinement. Fang Yuan’s immediate success is explicitly exceptional.
- **Ice Muscle inventory, SR-060:**137 says RevolutionThree Ice Muscle has permanently changed his body.140 counts only two RevolutionThree Gu before Frost Demon is sacrificed. We do not yet know whether Ice Muscle was consumed or otherwise relinquished. Hold the query; do not mechanically change two to three.

Chapter133’s claim to have killed a Xiong clansman also conflicts with the narrated scene: the wolf kills Xiong Zhan, while Bai kills a Bai clansman and loots Xiong Zhan. Chapters135 and138 explicitly support indirect responsibility. **SR-057** offers a literal corrected self-report but preserves the possibility of a boast claiming indirect responsibility.

Fang Yuan’s later accusation that Bai killed the Xiong group is different: it conceals Fang Yuan’s own killing of Xiong Jiang. Do not harmonise the action to that self-serving accusation.

## 6. Copy corrections and numerical restraint

The detailed Continuity jotter supplies exact source paragraphs, replacements, rationales, dependencies and held status for every new record. Especially clear examples are:

| Reference | Source | Held candidate |
|---|---|---|
|111p72|一个都能少|一个都不能少|
|125p40|你防守你进攻|我防守你进攻|
|132p76|两步芳草蛊|三步芳草蛊|
|116p69|四五天猴头|四五颗猴头|
|121p51|四位酒虫|四味酒虫|
|139p40|月刃站在|月刃斩在|

The malformed casualty number **二五个小组** in131 is not silently treated as25. **SR-058** records 二十五 as a plausible candidate requiring another textual witness.

Chapter138’s two favourable clans out of three also does not establish a measured two-thirds encounter probability, much less winning probability. **SR-059** proposes stating the favourable categories rather than presenting exact odds. The original may be a rough tactical heuristic; no patrol distribution has been established.

## 7. Endpoint and saved records

By140, two male teammates, Yao Hong and the protecting elder have died. Fang Zheng observes from a distance. Qing Shu begins using Wood Charm Gu in the final paragraph. **His death has not yet occurred in the completed reading.**

The numbered Chinese source, four clean jotters, reading log,61-record proposal ledger, {len(events)}-entry chronology register, {len(resources)}-entry resource register and cumulative workbook are updated. Remote repository backup remains unresolved; no upload or deletion occurred.

**Next planned batch: Chapters141–170. Next actual unread:141.** All active proposals remain held for full2334-chapter reconciliation.
'''
# Readable typography in deliverables; leave source quotations unchanged.
import re
report=re.sub(r'(?<=[A-Za-z])(?=\d)|(?<=\d)(?=[A-Za-z])',' ',report)
(r/'Batch-006-Chapters-111-140.md').write_text(report)
w=Workbook();s=w.active;s.title='Read me'
for row in [('Sequential review','Status'),('Completed','1–140 of2334'),('Paragraphs',tp),('Source characters',tc),('Next unread',141),('Proposal records','61 total;60 held;1 withdrawn;none approved/applied'),('Copy bundle','SR-061 contains39 individually located items'),('Scope','Chinese source; not fresh liveEnglish audit'),('Chronology entries',len(events)),('Resource entries',len(resources)),('Global consistency','Not yet certified')]:s.append(row)
s=w.create_sheet('Chapter completion');s.append(['Chapter','File','Title','Paragraphs','Status'])
for c in chapters:s.append([c['corpus_chapter'],c['file'],c['title'],len(c['paragraphs']),'Full sequential read'])
def ids(vs,k,c,p):return ', '.join(v['id'] for v in vs if any(z['chapter']==c and z['paragraph']==p for z in v[k]))
s=w.create_sheet('Paragraph register');s.append(['Chapter','Paragraph','Status','Proposal IDs','Chronology IDs','Resource IDs','Chinese source'])
for c in chapters:
 for p,t in enumerate(c['paragraphs'],1):
  n=c['corpus_chapter'];s.append([n,p,'Read; not certified error-free',ids(ledger,'locations',n,p),ids(events,'sources',n,p),ids(resources,'sources',n,p),t])
s=w.create_sheet('Proposal ledger');s.append(['ID','Category','References','Action','Rationale','Confidence','Status','Dependencies','Publication','Approved','Applied','Latest update'])
for v in ledger:s.append([v['id'],v['category'],refs(v['locations']),v['recommended_action'],v.get('rationale','See originating report'),v.get('confidence','See status'),v['status'],v['dependencies'],v['publication_status'],'No','No',v.get('batch006_update',v.get('batch005_update',v.get('batch004_update','')))])
s=w.create_sheet('Copy items');s.append(['Proposal','Chapter','Paragraph','Original','Candidate','Status'])
for z in ledger[-1]['copy_items']:s.append(['SR-061',z['chapter'],z['paragraph'],z['original'],z['replacement'],z['status']])
s=w.create_sheet('Chronology');s.append(['ID','Timeline','Kind','Event','Constraint','Evidence status','Dependencies','References','Exact source'])
for v in events:s.append([v['id'],v['timeline'],v['kind'],v['event'],v['constraint'],v['evidence_status'],v['dependency'],refs(v['sources']),'\n'.join(z['text'] for z in v['sources'])])
s=w.create_sheet('Resources');s.append(['ID','Kind','Entry','Calculation','Status','References','Exact source'])
for v in resources:s.append([v['id'],v['kind'],v['entry'],v['calculation'],v['status'],refs(v['sources']),'\n'.join(z['text'] for z in v['sources'])])
for s in w:
 s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
 for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='163B4B')
 for row in s.iter_rows(min_row=2):
  for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
 for col in s.columns:s.column_dimensions[col[0].column_letter].width=min(65,max(14,max(len(str(c.value or '')) for c in col[:30])+2))
w.save(r/'Sequential-Reading-and-Proposals.xlsx')
code=(r.parent/'make_whole_novel_audit_docx.py').read_text().replace("root=Path(__file__).parent","root=Path('/home/user/editorial/sequential-review')").replace('Whole-Novel-Aperture-and-Travel-Audit','Batch-006-Chapters-111-140').replace('WHOLE-NOVEL SOURCE AUDIT','SEQUENTIAL REVIEW / CHAPTERS111–140').replace('The Omniarch: Aperture and Travel Source Audit','The Omniarch: Sequential Review, Chapters111–140').replace('Full-corpus topical screening with selective contextual verification','Thirty complete chapters; active proposals held')
exec(compile(code,'docx-export','exec'),{})
(r/'CURRENT-STATUS.md').write_text(f'''# Sequential review checkpoint

Completed **Chapters1–140**. Next unread **141**. Planned next batch **141–170**. Remaining **2194**.

Batch006:{np}paragraphs/{nc}characters. Cumulative:{tp}paragraphs/{tc}characters.61proposalrecords:60activeheld,1withdrawn;noneapproved/applied. SR-061contains39separatecopyitems. Whole-novelchronology NOTcertified.

Latest: Batch-006-Chapters-111-140.md/.docx; Batch-006-Numbered-Chinese-Source.md; four006jotters; cumulativeworkbook/registers/logupdated. No source/site/repositoryedit,remoteuploadordeletions. Remote research backup unresolved. Neverstorecredentialsinfiles.

Carry forward:
- SR-047withdrawn: April/summercategoricaldiagnosisunsupported. SR-054separateinternalJulyearlyautumnversuspost-lateJulyhotsummerquery; calendarNOTresolved.
- SR-053samepreviouslifedestructiontwo yearslater126 versus3yearsafterwolfwave135. Minimalnonnumericcandidateonly; fullnovelneeded.
- SR-050YaoLeprospectiveawakeninginday7gossip despite109/110nearceremony; speakerknowledgeunresolved.
- SR-05244capacityat2.5pointsallows17shots,not16ifmaximum; SR-05111−5=6,thenblade5leaves1. Noaptitudechange.
- EightFYGuat140:Moonglow,WhiteJade,HiddenScale,FourFlavours,Cicada,EarGrass,NineLeaf,Plunder. ThreeStepGrassandChainsawCentipedeNOTowned.
- SR-060IceMuscleprevioususevs140twoRevolution3Guquery; noforcedcountchange.
- Ch131二五个casualtynumberuncertain;25candidateONLY. Ch133Xiongdeathclaimcontradictsnarratedwolfkill; laterindirectresponsibilityexplicit.
- Ch140QingShustartsWoodCharm; hisdeathNOTyetread. Rightfracturefeintdistinctfromrealrightarmsacrifice.
- Auntface-to-facefalsealarmdiscardedafter90reread. EarlierCh23advanceandCh55JiaJinShengbuyerrestraintremain.

Do not rerun previousbatchgenerators unchanged. Topical2334screeningisNOTsequentialcompletion.
''')
assert len(chapters)==140 and len(ledger)==61 and len(fixes)==39
assert all(not v['approved'] and not v['applied'] for v in ledger)
assert all(s['text']==src(s['chapter'],s['paragraph'])['text'] for v in events+resources for s in v['sources'])
assert '—' not in report
print(dict(batch_paragraphs=np,batch_characters=nc,total_paragraphs=tp,total_characters=tc,proposal_records=len(ledger),active=60,withdrawn=1,chronology=len(events),resources=len(resources),next=141))
