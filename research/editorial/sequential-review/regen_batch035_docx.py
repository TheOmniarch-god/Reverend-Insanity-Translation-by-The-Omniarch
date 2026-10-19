from pathlib import Path
from docx import Document
from docx.shared import Pt
r = Path(__file__).parent
lines = (r/'Batch-035-Chapters-666-685.md').read_text(encoding='utf-8').split('\n')
doc = Document()
doc.styles['Normal'].font.size = Pt(10)
i = 0
while i < len(lines):
    ln = lines[i]
    if ln.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].startswith('|'):
            cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
            if not all(set(c) <= set('-: ') for c in cells): rows.append(cells)
            i += 1
        t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = 'Table Grid'
        for a, row in enumerate(rows):
            for b, c in enumerate(row):
                t.cell(a, b).text = c
                for p in t.cell(a, b).paragraphs:
                    for run in p.runs: run.font.size = Pt(8)
        continue
    if ln.startswith('# '): doc.add_heading(ln[2:], 0)
    elif ln.startswith('## '): doc.add_heading(ln[3:], 1)
    elif ln.startswith('- '): doc.add_paragraph(ln[2:], style='List Bullet')
    elif ln.strip(): doc.add_paragraph(ln)
    i += 1
doc.save(r/'Batch-035-Chapters-666-685.docx')
print('docx saved', len(doc.paragraphs), len(doc.tables))
