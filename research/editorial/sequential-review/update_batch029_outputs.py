from pathlib import Path
import json, re
from openpyxl import load_workbook
from docx import Document
from docx.shared import Pt
r = Path(__file__).parent
s28 = json.loads((r/'batch-028-source.json').read_text())
s29 = json.loads((r/'batch-029-source.json').read_text())
src = s28 + s29
led = json.loads((r/'proposal-ledger.json').read_text())
HOLD = 'HOLD: full2334-chapter reading and global reconciliation outstanding'
assert len(led) == 337 and led[-1]['id'] == 'SR-337'
pmap = {}
for v in led:
    if int(v['id'][3:]) >= 286:
        for l in v['locations']:
            pmap.setdefault((l['chapter'], l['paragraph']), []).append(v['id'])
wb = load_workbook(r/'Sequential-Reading-and-Proposals.xlsx')
def drop(ws, pred):
    for idx in range(ws.max_row, 1, -1):
        if pred(ws.cell(idx, 1).value): ws.delete_rows(idx)
drop(wb['Chapter completion'], lambda v: v is not None and str(v).isdigit() and int(v) >= 526)
drop(wb['Paragraph register'], lambda v: v is not None and str(v).isdigit() and int(v) >= 526)
drop(wb['Proposal ledger'], lambda v: v is not None and str(v).startswith('SR-') and int(str(v)[3:]) >= 286)
ws = wb['Chapter completion']
for c in src:
    ws.append([c['corpus_chapter'], c['file'], c['title'], len(c['paragraphs']), 'Full sequential read'])
ws = wb['Paragraph register']
for c in src:
    for i, t in enumerate(c['paragraphs'], 1):
        ids = pmap.get((c['corpus_chapter'], i))
        ws.append([c['corpus_chapter'], i, 'Read; not certified error-free', ', '.join(ids) if ids else 'None', 'None', 'None', t])
ws = wb['Proposal ledger']
for v in led:
    if int(v['id'][3:]) >= 286:
        refs = '; '.join(f"Ch{l['chapter']} p{l['paragraph']}" for l in v['locations'])
        batch = 'Batch 028' if int(v['id'][3:]) <= 306 else 'Batch 029'
        ws.append([v['id'], v['category'], refs, v['recommended_action'], v['rationale'], v['confidence'], v['status'], v['dependencies'], HOLD, 'No', 'No', f'{batch} sequential reading; source-internal; witness collated'])
N = sum(len(c['paragraphs']) for c in src); CH = sum(len(t) for c in src for t in c['paragraphs'])
vals = {'Completed': '1–565 of 2334', 'Next unread': '566', 'Paragraphs': str(48681 + 1800), 'Characters': str(1711344 + 63491),
        'Proposals': '337 total; 332 active (held or query); 4 withdrawn; 1 locally resolved'}
for row in wb['Read me'].iter_rows(min_row=1, max_col=2):
    k = row[0].value
    if k in vals: row[1].value = vals[k]
wb.save(r/'Sequential-Reading-and-Proposals.xlsx')
print('workbook saved', N, CH)

def build_docx(md_name, out_name):
    md = (r/md_name).read_text(encoding='utf-8').splitlines()
    doc = Document()
    st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
    doc._t = None; doc._t_done = True
    for line in md:
        if not line.strip() or line.strip() == '---': continue
        if line.startswith(('# ','## ','### ')): doc._t_done = True
        if line.startswith('# '): doc.add_heading(line[2:], 0)
        elif line.startswith('## '): doc.add_heading(line[3:], 1)
        elif line.startswith('### '): doc.add_heading(line[4:], 2)
        elif line.startswith('- '): doc.add_paragraph(re.sub(r'\*\*(.+?)\*\*', r'\1', line[2:]), style='List Bullet')
        elif line.startswith('|'):
            if set(line.replace('|', '').strip()) <= set('-: '): continue
            cells = [c.strip() for c in line.strip('|').split('|')]
            if doc._t is None or doc._t_done:
                doc._t = doc.add_table(rows=0, cols=len(cells)); doc._t.style = 'Light Grid Accent 1'; doc._t_done = False
            row = doc._t.add_row().cells
            for k, cv in enumerate(cells): row[k].text = re.sub(r'\*\*(.+?)\*\*', r'\1', cv)
        else:
            doc._t_done = True
            doc.add_paragraph(re.sub(r'\*\*(.+?)\*\*', r'\1', line))
    doc.save(r/out_name)
    return len(doc.tables)
print('docx 028 tables', build_docx('Batch-028-Chapters-526-545.md', 'Batch-028-Chapters-526-545.docx'))
print('docx 029 tables', build_docx('Batch-029-Chapters-546-565.md', 'Batch-029-Chapters-546-565.docx'))
