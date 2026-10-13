from pathlib import Path
import json, shutil
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
r=Path('/home/user/editorial/sequential-review')
def load(n):return json.loads((r/n).read_text())
def save(n,x):(r/n).write_text(json.dumps(x,ensure_ascii=False,indent=2))
chapters=sum([load(f'batch-{i:03d}-source.json') for i in range(1,6)],[]);by={c['corpus_chapter']:c for c in chapters};batch=[by[c] for c in range(81,111)]
def src(c,p):return dict(chapter=c,paragraph=p,file=by[c]['file'],text=by[c]['paragraphs'][p-1])
def sources(rr):return [src(*x) for x in rr]
def loc(rr):return [{k:v for k,v in s.items() if k!='text'} for s in sources(rr)]
def refs(ss):return '; '.join(f"Ch{s['chapter']} p{s['paragraph']}" for s in ss)
for name in ['reading-progress.json','proposal-ledger.json','chronology-register.json','resource-register.json']:
 dst=r/'history'/('before-batch005-'+name)
 if not dst.exists():shutil.copy2(r/name,dst)
ledger=[v for v in load('proposal-ledger.json') if int(v['id'][3:])<=38]
updates={
'SR-024':'Add Ch82p79 and Ch83p1: the pupil is Mo Bei, not elder Mo Chen. Preserve the actual elder at Ch83p51.',
'SR-030':'Ch81p39 explicitly confirms57 pupils, strengthening the56-paying-peer basis for67 levy-only /70 including own stipend. No basis for silently adding a57th payer.',
'SR-032':'Ch109 March awakening announcement strengthens the March-entry/April-first-week calendar. Distinguish academic-year labels from literal elapsed months; private teaching has its own start date.',
'SR-021':'Ch91p41 explicitly sets GradeC at40–59%. This strengthens the Ch22 GradeC38% mismatch, but changing its grade alone still does not resolve the earlier scene staging.',
'SR-001':'Ch99p5 retains44% aperture occupancy after RevolutionTwo advancement. No aptitude or aperture-volume recalculation authorised.'}
for v in ledger:
 if v['id'] in updates:v['batch005_update']=updates[v['id']]
 if v['id']=='SR-024':
  for z in loc([(82,79),(83,1)]):
   if z not in v['locations']:v['locations'].append(z)
new=[]
def prop(i,cat,rr,action,reason,confidence,deps):
 v=dict(id=f'SR-{i:03d}',category=cat,locations=loc(rr),recommended_action=action,rationale=reason,confidence=confidence,status='Active provisional proposal; no approval',dependencies=deps,approved=False,applied=False,publication_status='HOLD: full2334-chapter reading and global reconciliation outstanding');ledger.append(v);new.append(v)
prop(39,'academy duration',[(81,13),(81,16),(85,3),(87,52),(97,55),(99,6),(99,9)],'Where elapsed duration is explicit, prefer academy year / this school year for schooling and since spring or through spring and autumn for Cicada recovery. Do not replace every occurrence of year.','The spring entry and winter graduation do not establish12 elapsed months. Academic-year and seasonal-cycle usages can be legitimate shorthand; 整整一年 and 满一年 need especially careful treatment.','High need to distinguish meanings; moderate need for emendation.','SR-002, SR-017, SR-032; later school calendars and Cicada recovery mechanics. Do not impose a365-day maturation requirement.')
prop(40,'private instruction duration',[(58,54),(63,7),(81,93),(84,39),(86,50)],'Ch84p39 被一个四转蛊师教导一年→被一个四转蛊师教导了半年多; working English: taught by a Revolution Four Gu Master for a little over half a year.','Private lessons begin only after the Jia inquiry. The JadeSkin gift and midyear tusk scheme are each recalled as roughly half a year ago at graduation, not a full year of tuition.','High local conflict; moderate exact replacement.','Reconcile all winter dates. A less numerical several months remains available if later dates oppose half a year.')
prop(41,'prize amount',[(81,31),(81,41),(86,68),(89,14),(89,48)],'Preferred linked repair:150→100 in Ch81pp31,41; retain100 in Ch86/89 and the25-per-month ×4-month comparison.','Two initial150 announcements conflict with the later100 announcement, repeated100 recollection and matching rent arithmetic. The later coherent chain favours reducing the initial number, not inflating all downstream amounts.','High inconsistency; moderate-to-high preferred direction.','Later payout/cash evidence and possible unannounced prize reduction. Initial reading preference100→150 superseded before delivery.')
prop(42,'age and calendar',[(86,30),(90,7),(90,20),(101,4),(104,13)],'Ch90p7 如今已经十六岁了→过了年就十六岁了; English: He will be sixteen after the New Year.','The aunt says already16, but the same chapter has Fang Yuan explicitly say he will be16 after New Year; Ch86 still calls him15. Later inheritance scenes consistently use16.','High immediate mismatch; high conditional repair.','Preserve possible nominal/New-Year age convention rather than invent a birthday or change all winter ages.')
prop(43,'Gu name',[(62,88),(82,61),(83,38),(84,18),(93,48)],'赤丸蛐蛐蛊→龙丸蛐蛐蛊 in Ch82p61 and Ch93p48, retaining one English cricketGu name.','The same pupil’s selected movementGu is 龙丸 in Ch62,83,84; no replacement or evolution is narrated.','High local identity confidence.','Later nomenclature and independent textual witnesses; not a fresh liveEnglish error claim.')
prop(44,'clan branch',[(93,35),(93,57)],'Ch93p57 方源虽然不是漠脉中人→方源虽然不是赤脉中人.','The leader is Chi Shan, head of the Chi-branch team. The explanation concerns protecting a newcomer outside his own branch, not the rival Mo branch.','High contextual confidence.','Preserve elsewhere the actual Mo/Chi distinction; source-copy repair only.')
prop(45,'Gu continuity',[(99,50),(99,52),(100,1),(100,57)],'Ch100p57 玉皮蛊→白玉蛊; English: White Jade Gu.','JadeSkin has already been consumed in the successful WhiteJade fusion. More than10days then pass; honey was collected a few days before the current scene. Chapter98 specifically set up the strongerGu for this task.','High contextual confidence.','No second JadeSkin or alternative collection date established; continue full-corpus cross-check.')
prop(46,'physical staging',[(106,26),(106,27),(106,33),(106,45)],'Ch106p26 摊开右手→摊开左手.','The grass is on the left palm, the right fingers pluck it, and it retracts into the left palm. One opening right-hand reference conflicts with the continuous action.','High contextual confidence.','No hand-transfer action narrated; preserve the correctly used right plucking hand.')
prop(47,'caravan calendar',[(39,21),(39,39),(109,34),(109,45)],'If harmonising, distinguish usualApril arrival from this year’s prior-life summer arrival: Normally the caravan came in April; in my previous life it arrived that summer. This time it was here in March, two or three months earlier than I remembered.','Ch39 establishes usualApril and an explicitly lateMay visit. Ch109 now calls summer the usual annual schedule, separately describing a deliberately changed March arrival. The butterfly effect explains the current change, not conflicting claims about the normal baseline.','High baseline tension; moderate linked wording.','Check later caravan histories and calendars. Retain the2–3month comparison only against prior-life summer, not usualApril.')
prop(48,'source copy',[(106,19)],'一心己用→一心多用; English: divide his attention among several tasks.','Context describes multitasking during fusion; the earlier explanation and current action support 多, not 己.','High contextual confidence.','No independent corrected witness checked.')
save('proposal-ledger.json',ledger)
tdoc=load('chronology-register.json');events=tdoc['entries'][:59]
def event(timeline,kind,title,constraint,status,rr,dep='Full-corpus reconciliation outstanding.'):
 events.append(dict(id=f'T{len(events)+1:03d}',timeline=timeline,kind=kind,event=title,constraint=constraint,evidence_status=status,dependency=dep,sources=sources(rr)))
event('current life','school-year / elapsed duration','Graduation chronology','Nearwinter→winter final cultivation test→nextday tournament; academic year repeatedly described as a year.','Do not equate academic-year labels with12elapsed months; explicit fullness remains queried.',[(81,16),(81,34),(81,40),(82,1),(85,3),(87,52)],'SR-039 and SR-032.')
event('current life','private teaching retrospect','Tuition duration','Ch84 says1year; lessons began afterJia inquiry; winter gift/tusk recaps sayhalfyear.','Separate privateinstruction start from academyentry.',[(58,54),(81,93),(84,39),(86,50)],'SR-040.')
event('current life','age / New Year','Fifteen to sixteen','Fifteen at exam; aunt already16 versus FYafterNewYear16; later16 at inheritance.','Do not invent exactbirthday. Samechapter mismatch retained.',[(86,30),(90,7),(90,20),(101,4),(104,13)],'SR-042.')
event('current life','healing / recovery','Fist injuries','Wounds close immediately; weakness persists7–8days.','Function recovery differs from closure; first mission3dayslater does not prove fullstrength restored.',[(86,22),(86,33),(87,36)])
event('current life','clock / travel','First mission and lodging','Exam+3days;30minute outwardrun;1hour collection; returnafternoon; housesearchuntilnight;30minute rental agreement.','Narrated local clocks, not whole-day travel assumptions.',[(87,36),(88,19),(88,31),(88,58),(88,70),(89,46),(90,42)])
event('current life','breakthrough mechanism','Resource-fed RevolutionTwo breakthrough','4–5days expected; interruption of15minutes permits crystalwall fullrecovery.','Capacity44 does not become55; sustained replenishment is explicit exception to ordinary burst requirement.',[(90,70),(91,52),(91,55),(91,59)])
event('current life','day count / retrospective wording','Four days cultivation then sleep','Firstmorningshopping; nextdaynoonsearch; third-daysearch; fourth-daybreakthrough; wake nextafternoon; fifthdayexit but five-days-ago recaps.','Distinguish inclusive day labels from elapsed days; low-level possible off-by-one remains open.',[(91,1),(91,66),(92,26),(92,33),(92,36),(92,47),(92,52),(92,54),(92,60),(92,70)],'Do not add a sixthday solely to force every phrase into elapsed24hour units.')
event('current life','location / misdirection','Decoy room and actual retreat','Rental furnishings and note mislead searchers; actual retreat atJiangHe house in village.','Room-setting reveal is intentional, not teleportation or double occupancy.',[(92,20),(92,54),(92,56),(92,60),(92,61)])
event('current life / forecast','cycle and expected event','Wolf tide timing','Largewave every3years; nextyear forecast; Ch97 says1yearlater. Electricwolves appear early in minorwave.','Do not substitute minorwave for mainwolf catastrophe; exactfutureinterval awaiting occurrence.',[(93,22),(93,23),(95,87),(97,32),(97,36)])
event('current life','battle clocks','Beast wave evening','Poisonfight5minutes; reinforcements10minutes; JiaoSan first presumeddead, then seenmoving, then killed.','Premature inference corrected within scene, not resurrection.',[(95,61),(96,4),(96,51),(96,55),(96,84)])
event('current life','administrative intervals','Inquiry and inheritance mission','Three days afterwave: hourlong questioning; toldreturn3dayslater for mission.','Bribe explains imposed delay. Missiondelivery follows transition; exact intervening clock not supplied.',[(97,7),(97,19),(97,30),(97,48),(98,2),(98,20)])
event('current life / Cicada','seasonal cycle','Cicada recovery','Near1yearrest; onecycle ofspring/autumn invoked.','Seasonal interpretation is Fang Yuan’s emerging understanding, not established fixed365daycooldown.',[(99,6),(99,9),(99,11)],'SR-039; later recovery and time-flow evidence.')
event('prior-life future','future discovery','Fusion knowledge dates','Recipe improved150years later; understanding principle broadlyrecognised300years later.','Knowledge from oldtimeline, not current waiting periods.',[(99,45),(99,61)])
event('current life','rent / loose transition','Inheritance cash-crunch date','>10days afterfusion; honeycollectedfewdaysago; rentalmonthprepaid butnow2daysarrears.','Rent gives additional elapsedconstraint; unquantified preparation/transitions remain open. No invented exact day for fusion.',[(92,20),(100,1),(100,57),(101,33)])
event('uncle biography','ages / life intervals','Dong Tu history','15deployment;16Revolution2;17upper;18peak;19attemptRevolution3;20injury;8yearsdespondency;29birthdayreappraisal;30restart;45retirement.','Age20injury not fixed to20thbirthday; no automatic8→9repair.',[(101,11),(101,12),(101,13),(101,14),(101,15)])
event('current life','transfer / relative poverty','Inheritance handed over','Submission→3dayslatertransfer; later sayspoor daybeforeyesterday.','Can remainpoor duringprocessing; not automatically3days→2days mismatch.',[(102,15),(102,47),(103,59)])
event('prior life / conditional future','forecasts and proposed contract','Qing Shu and LiquorWorm offer','QingShu diedinpriorlifewave1yearlater; offered5yearuse ofSevenFragranceGu conditionalonfuturefusion; neveraccepted. PriorRevolution2plateau>100years.','Separate oldbranch death, presentforecast and unrealisedcontract.',[(105,6),(105,26),(105,30)])
event('current life','fusion recovery interval','Moonglow attempts','One day aftermeeting firstattemptfails; ≥3daysrecoveryestimated;3dayslater succeeds.','Do not call secondattempt immediate or firsttry success.',[(106,1),(106,13),(106,24),(106,56),(106,60)])
event('current life','season / calendar','Spring and March caravan','Winterends→spring; firstmeetingJiangYa recalled1yearearlier; Marchcaravan; ceremonydayaftertomorrow then tomorrowinshopscene.','Shopping daynotexplicitly anchored: couldbearrival+1. Preserveopenboundaryquery ratherthan force eitherrelativeword.',[(107,2),(107,25),(109,34),(109,56),(110,16)],'SR-047; later ceremony timing.')
event('current life historical recollection','flashback / convalescence','Man Shi and Bai Ning Bing','Encounter2yearsago: BaiNingBingRevolution2 butcapablekillingRevolution3. Currentdefeat followedby3daysrecoveryperrumour.','Historicalexception and reportedrecovery distinctfrompresentcombatclock.',[(108,108),(108,117),(108,119),(109,11)])
tdoc.update(scope='Sequential evidence throughChapter110; global chronology not certified.',entries=events);save('chronology-register.json',tdoc)
rdoc=load('resource-register.json');resources=rdoc['entries'][:27]
def resource(kind,entry,calc,status,rr):resources.append(dict(id=f'R{len(resources)+1:03d}',kind=kind,entry=entry,calculation=calc,status=status,sources=sources(rr)))
resource('cohort and prize','57pupils;150announced then100recapped with25/month×4months.','56payingpeers confirmed;100rent comparison exact.','SR-030 strengthened; SR-041 preferred initial150→100, held.',[(81,39),(81,41),(86,68),(89,14),(89,48)])
resource('first task proceeds','Teamreceives6;leader2,otherfour1each. Base task2.','2+4×1=6; increaseby2×base gives3×base. FangYuanactualincome1.','Do not credit all6 to FangYuan or translate 多了两倍 as totaldoubling.',[(88,71),(88,74)])
resource('capacity and replenishment','C44%; ordinaryRevolution2breakthroughthreshold55%; GradeC40–59%.','Ch92 remainingreserve goesfrom44 downbyabout20 toabout24, then13 beforebreakthrough. Naturalregenerationandstones mean grosscostnot equalnetdrop.','Explicit exception; no44→55 aptitude rewrite.',[(90,70),(91,41),(91,42),(92,37),(92,38),(99,5)])
resource('pre-inheritance cash','Once>1000held;2monthsrunwayforecast beforebreakthrough; later onlyonefusionaffordable; final2.5cash.','Oneboarking tusk receivedfree and consumedinfusion. Exactstonecountspent notgiven.','No certified cumulativebalance; forecasts not outlays.',[(91,29),(91,35),(99,19),(99,21),(99,22),(100,58)])
resource('inventory ledger','SevenGu aftersecondLittleLight;WhiteJadefusion reduces7→6; inheritedNineLeaf restores6→7.','OneLittleLightdies→6; replacementimplied→7;3inputs→oneMoonglow means7→5 permanentGu, excludingdetachedleafstock.','Replacementpriceunreported; secondLiquorWormnotyetpurchasedby110.',[(91,31),(98,32),(99,50),(102,57),(102,60),(106,15),(106,17),(106,61)])
resource('feeding rates','JadeSkin2liang/10days;WhiteJade8liang/20days.','0.2→0.4liang/day:2×averagejade consumption. Pigfeedingceased.','Combinedmarketcost explicitlyhigher; selfgatheredjade reducescashneed. Do not infer4×average.',[(100,50),(100,51),(100,52),(100,54),(100,55)])
resource('inheritance and refused settlement','Uncleoffers1000→2000→2500, rejected. Actualassets>10mu,8servants,3houses,1wine-shop andNineLeafgrass.','Unclewithdrawscurrentmonthshopreceipts; workerswages+20%,manager+40%; basewagesnotgiven.','Noofferedcashbookedasincome; no invented assetvaluation.',[(101,33),(101,55),(102,50),(102,60),(103,25),(103,33)])
resource('leaf production and commerce','20essencepoints perleaf;9leavesregrown inhalfday;50retail each; >400reportednet; later9/day and agentmargin.','9×20=180points;44reserve allows2beforetop-up.9×50=450retailgross; clanretail55 vs FangYuanretail50.','Exactproductioncostandagentmarginunreported;450not proven exactwholesalereceipt.',[(106,40),(106,41),(106,44),(106,47),(106,48),(107,22),(107,37),(107,38),(107,41)])
resource('combat ratios and range','TwoLittleLight notstackable;Moonglow3×Moonlightattack, still10mrange;MoonScar20m.','Defence16monkeyremark inCh100 notsamecultivationcontext as12limitCh79. No linear conversionofsurvivalcounttodefencemultiplier.','Do not silentlycorrect16→12 or add thetwoLittleLightbonuses.',[(79,81),(100,8),(100,9),(106,61),(106,62),(106,64),(106,65)])
resource('extortion and auction','ManShidemands500; LiquorWormadvertised500 versusnormal580.','580−500=80. Sealedbids determinefinalprice; noFangYuanbidyetby110.','Extortiondemandnotpaidexpense; displayprice not completedpurchase.',[(108,11),(110,10),(110,11),(110,27),(110,55)])
rdoc.update(scope='Resources throughChapter110; local arithmetic distinguished from incomplete running balances.',entries=resources);save('resource-register.json',rdoc)
np=sum(len(c['paragraphs']) for c in batch);nc=sum(sum(map(len,c['paragraphs'])) for c in batch);tp=sum(len(c['paragraphs']) for c in chapters);tc=sum(sum(map(len,c['paragraphs'])) for c in chapters)
p=load('reading-progress.json');p.update(sequential_chapters_completed=list(range(1,111)),paragraphs_read=tp,paragraph_characters_read=tc,next_chapter=111,remaining_chapters=2224,batch_report='Batch-005-Chapters-81-110.md',scope='Full sequential Chinese source reading throughChapter110. Not a fresh liveEnglish/bilingual audit.',whole_novel_chronology_verified=False,chapters=[dict(chapter=c['corpus_chapter'],file=c['file'],title=c['title'],paragraphs=len(c['paragraphs']),first_paragraph=1,last_paragraph=len(c['paragraphs'])) for c in chapters]);p['batches']=[v for v in p['batches'] if v['id']!='005']+[dict(id='005',chapters=[81,110],paragraphs=np,characters=nc,reading_segments='Fifteen complete untruncated two-chapter outputs')];save('reading-progress.json',p)
save('batch-005-reading-log.json',dict(status='Complete',endpoint=110,next_unread=111,segments=[dict(chapters=[c,c+1],status='Full paragraph sequences printed and read; untruncated',paragraphs=sum(len(by[k]['paragraphs']) for k in [c,c+1])) for c in range(81,111,2)],note='Extraction alone not completion. No invented timestamps.'))
lines=['# Numbered Chinese source: Chapters81–110','Unemended nonempty EPUB body paragraphs, numbered to match the review.','']
for c in batch:
 lines += [f"## Chapter {c['corpus_chapter']}: {c['title']}",f"Source: `{c['file']}`",'']
 for i,t in enumerate(c['paragraphs'],1):lines += [f'**{i:03d}.** {t}','']
(r/'Batch-005-Numbered-Chinese-Source.md').write_text('\n'.join(lines))
for name,vs in [('Chronology',events[59:]),('Resources',resources[27:])]:
 lines=[f'# Batch005: {name} jotter','Complete source reading throughChapter110. Changes remain held.','']
 for v in vs:
  lines += [f"## {v['id']}: {v.get('event',v['kind'])}"]
  for k in ['timeline','kind','constraint','evidence_status','dependency','entry','calculation','status']:
   if k in v:lines += [f"**{k.replace('_',' ').title()}:** {v[k]}"]
  for s in v['sources']:lines += [f"\n**Ch{s['chapter']} p{s['paragraph']}**",'> '+s['text'],'']
 (r/'jotters'/f'Batch-005-{name}.md').write_text('\n'.join(lines))
lines=['# Batch005: Continuity and active proposals','All proposals unapproved and unapplied.','']
for v in new:
 lines += [f"## {v['id']}: {v['category']}",f"**Action:** {v['recommended_action']}",f"**Rationale:** {v['rationale']}",f"**Confidence:** {v['confidence']}",f"**Dependencies:** {v['dependencies']}",f"**Status:** {v['publication_status']}",'']
 for s in v['locations']:lines += [f"**Ch{s['chapter']} p{s['paragraph']}**",'> '+src(s['chapter'],s['paragraph'])['text'],'']
lines += ['## Existing proposals updated']
for k,v in updates.items():lines += [f'**{k}:** {v}','']
lines += ['## Unsettled queries and active restraint','- Ch83p14 says a moonblade dissipates beyond10m, more categorical than the earlier distinction between effective10m and physicalflight15/16m. Working clarification: loses effective striking power beyond10m, not automatic shrinking of trajectories. Hold pending later mechanics.','- Ch83p42 比你耐用三倍: retain a ratio query rather than silently interpret every 比…倍 as the same multiplier. No independent numerical example here decides3× versus4×.','- Ch92 fifth-day exit versus five-days-ago recaps: if a strict elapsed correction is eventually necessary, consider fourdaysago or severaldaysago in both recaps. For now retain inclusive-count uncertainty.','- Ch101 rent arrears: treat the prepaidmonth+2days as a constraint on unspecified intervals, not proof that the ten-plus-day transition must be rewritten.','- Ch94p51 inheritance delay past16: do not decide whether this means the sixteenth birthday or the whole year at16 without the later legal context.','- Ch109 dayaftertomorrow / Ch110 tomorrow: if scenes occur on different days, an explicit nextday scene cue could clarify; no mandatory change until later anchors resolve shopping date.','- Ch99p64 success rate提高两成: relative improvement versus percentage-point addition is not explicitly settled. Do not encode a base success rate or certainty.','- Ch105p17 oneGu/oneability and Ch106p80 vitalGu protection are general claims to cross-check against later exceptions. No correction approved now.','']
(r/'jotters'/'Batch-005-Continuity.md').write_text('\n'.join(lines))
preserve='''# Batch005: Preserve, do not repair

- A school year is not automatically twelve elapsed months. Differentiate school/calendar labels from explicit duration claims and from the separate period of private instruction.
- The final exam precedes the NewYear age transition. Do not turn every fifteen into sixteen or invent a birthday.
- The four-day training description and fifth-day exit use inclusive day labels. Preserve the ambiguity of the five-days-ago recaps pending tighter anchors.
- Renting by the day is an available pricing arrangement, not a prohibition on prepaying a month. The monthly payment is explicit.
- Fang Yuan trains in Jiang He's house while misleading pursuers with the furnished rented room. This is a location reveal, not a continuity mistake.
- RevolutionTwo advancement uses repeated replenishment. A normal55% burst threshold does not mandate a55% maximum aperture for every successful cultivator.
- Chi Cheng has separately advanced toRevolutionTwo. His assumption that Fang Yuan remainsRevolutionOne is mistaken character knowledge.
- Apparent deaths followed by survival in these scenes are corrected observations. Jiao San's actual final death is distinct from Fang Yuan's first assessment.
- The female team member did not voluntarily compete for the hiding place. Fang Yuan's later testimony is a lie; do not edit the assault to match it.
- The hidden WhiteJadeGu explains survival and honey collection. Its existence is deliberately concealed from other characters.
- JadeSkin's16-monkey defensive estimate follows advancement beyond the earlier12-monkey scene. No identical conditions or linear defensive multiplier are established.
- Seven ownedGu after the inheritance is consistent:7→6 throughWhiteJadefusion, then6→7 withNineLeafgrass. Detached leafGu are separately produced sale inventory.
- Fusion failure kills oneLittleLight. A replacement is implied by the later successful recipe, but no exact purchase price is given.
- WhiteBoar's permanent strength remains after theGu is consumed. WhiteJade itself does not grant new strength.
- TwoLittleLightGu are recipe inputs, not two simultaneously stacking attack bonuses.
- Business revenue, retail price, dealer margin, stored assets and spendable cash are different quantities. Do not turn nineleaves×50 into an exact wholesale payment without the margin.
- The uncle's1000/2000/2500 offers and ManShi's500 demand are not completed payments.
- Fang Yuan's threatened deliberate defeat and homelessness are bargaining moves, not actual events.
- QingShu's old-life death, the proposed5yearGu use and the150/300year future discoveries stay separate from the present calendar.
- Ch107 spring and Ch109 March are real calendar advances. The earlycaravan is an explicit rebirth consequence; only the conflicting usual-arrival baseline is queried.
- At the Chapter110 endpoint the secondLiquorWorm is displayed, not purchased. Its500 price is an auction starting indication, not the final transaction.
- Ch96's Earth plank/necessity analogy has not been verified as a legal or historical account; preserve its narrative role without presenting it as universal legal advice.
'''
(r/'jotters'/'Batch-005-Preserve.md').write_text(preserve)
report=f'''# Sequential close-reading review: Chapters81–110
## Batch005 | The Omniarch

**Actual completed coverage: Chapters1–110 of2,334. Next unread:111. Remaining:2,224.**

All30chapters in this batch were read in fifteen complete, untruncated two-chapter sequences: **{np:,} paragraphs / {nc:,} source characters**. Cumulative coverage is **{tp:,} paragraphs / {tc:,} characters**. Extracting the chapters was not counted as reading.

This is Chinese-source review, not a new bilingual or live-site audit. **All{len(ledger)} proposals remain held, unapproved and unapplied.** Whole-novel temporal consistency is not yet verified. No source, site or repository changes were made.

## 1. Main chronology findings

### School year, elapsed year and private tuition must be separated

Chapter109 locates the next awakening ceremony in March. Together with Chapter7's April lesson, seven days after awakening, this strengthens the spring-entry calendar. It does not support moving the original entry to January just to rescue Chapter72's half-year claim.

Several new passages describe an academy year, almost a full year, or a full year of extortion. Some may be ordinary school-year shorthand. They should not all be mechanically replaced. The stronger independent problem is Chapter84's claim that Fang Zheng has been privately taught for a year:

> {src(84,39)['text']}

Those lessons began after the Jia inquiry in Chapter58. At graduation, the later JadeSkin gift and the June tusk scheme are each recalled as roughly half a year earlier. The preferred held wording for private tuition is **a little over half a year**, with **several months** available if later dates require less precision.

Chapter99's Cicada passage also distinguishes a seasonal cycle of spring and autumn from a measured twelve-month recovery period. Do not invent a365-day minimum for its partial recovery.

### Age fifteen versus sixteen: a local, repairable mismatch

In Chapter90, the aunt says Fang Yuan is already sixteen. In the same chapter Fang Yuan says he will be sixteen after NewYear. Chapter86 still calls him fifteen; the later inheritance scene calls both twins sixteen.

**Working correction:** change the aunt's statement to **He will be sixteen after the NewYear.** This preserves the sequence without inventing a birthday or imposing a Western age convention.

### The new caravan contradiction is not the butterfly effect itself

Chapter39 distinguished the caravan's usual April arrival from that year's late May arrival. Chapter109 now describes summer as the usual arrival season:

> {src(109,45)['text']}

Its current March arrival is explicitly attributed to changes caused by Fang Yuan. Preserve that. The issue is the conflicting description of the ordinary historical schedule.

The held clarification distinguishes **usually April** from **that year's prior-life summer arrival**. The two-or-three-month advance should then compare March with the remembered summer arrival, not with April. Later caravan evidence may revise this proposal.

## 2. A proposal changed during this reading

### Year-end prize: preferred150→100, not100→150

| Evidence | Amount |
|---|---:|
| Ch81p31, anticipated prize |150|
| Ch81p41, formal announcement |150|
| Ch86p68, chief's award announcement |100|
| Ch89p14, teammate's recollection |100|
| Ch89p48, prize covers four months at25/month |100 implied|

After Chapter86 alone, raising100to150 was a plausible direction. Chapter89 supplies a coherent100-stone chain, including **25×4=100**. The preferred proposal now changes the two initial150 figures to100.

This is still an emendation, not an approved correction. An unannounced reduced payout or later accounting evidence could change the interpretation. No fictitious50-stone fee has been introduced to reconcile the figures.

## 3. Active proposals added

| ID | Subject | Working action |
|---|---|---|
|SR-039|Academy/Cicada duration wording|Distinguish school year and seasonal cycle from literal elapsed year; clarify only the explicit problem passages.|
|SR-040|Private teaching|Ch84 one year→a little over half a year, provisionally.|
|SR-041|Year-end prize|Ch81pp31,41:150→100; retain coherent later100 chain.|
|SR-042|Aunt's age statement|Ch90 already sixteen→will be sixteen after NewYear.|
|SR-043|MovementGu name|Ch82p61 andCh93p48:赤丸→龙丸, matching the pupil's other appearances.|
|SR-044|Branch name|Ch93p57:Mo branch→Chi branch in the ChiShan-team explanation.|
|SR-045|Honey collectionGu|Ch100p57:JadeSkin→WhiteJade; the former was already consumed in fusion.|
|SR-046|Hand continuity|Ch106p26:right→left; right fingers pluck leaves from the left-palm grass.|
|SR-047|Caravan baseline|Separate usualApril arrival from the prior-life summer arrival; preserve the deliberate earlyMarch change.|
|SR-048|Multitasking copy error|Ch106p19:一心己用→一心多用.|

The continuity jotter and ledger give exact source quotations, rationale, confidence and dependencies for each. High-confidence copy repairs remain held under the same full-corpus policy as temporal proposals.

### Earlier proposals strengthened or extended

- **SR-024:** add Mo Chen→Mo Bei atCh82p79 andCh83p1. Preserve the actual elderMoChen atCh83p51.
- **SR-030:** Chapter81 explicitly counts57pupils, supporting56paying peers. This strengthens the earlier67/70-versus68 accounting query.
- **SR-021:** Chapter91 explicitly gives GradeC as40–59%. This strengthens the GradeC38% mismatch, but does not alone resolve the original scene's staging.
- **SR-001:** Chapter99 retains44% aperture occupancy after RevolutionTwo advancement. No blanket aptitude recalculation is justified.

## 4. Chronology chain and unresolved boundaries

| Sequence | Evidence and caution |
|---|---|
|Final exam→first mission|Three days. Hand wounds healed, but seven-to-eight-day weakness need not have ended.|
|First mission|Thirty-minute outward run, one-hour collection, return by afternoon; househunting continues into evening.|
|Rental→retreat|Next morning shopping and departure; the rented room is a decoy. Actual cultivation is at JiangHe's village house.|
|Breakthrough|Fourth-day success after four days/three nights; sleep until next afternoon, fifth-day exit. Five-days-ago recaps may use inclusive counting.|
|Return journey|Evening beastwave; five-minute poisonfight, later ten-minute wait for reinforcements.|
|Investigation|Three days after the wave; one-hour interrogation. Another three-day administrative delay announced for the inheritance mission.|
|Fusion→honey turn-in|More than ten days; exact preparation intervals are not all stated.|
|Rent arrears|One month was prepaid; uncle later reports two days overdue. This constrains the open timeline, not an excuse to invent an exact missing interval.|
|Inheritance|Three days from submission to actual transfer. The day-before-yesterday poverty remark can refer to the waiting period.|
|Moonglow refinement|One day after the discussion: failure. Three days later: success after recovery.|
|Spring caravan|Winter gives way to spring; March arrival. Ceremony is dayaftertomorrow at reception, tomorrow in the shop conversation. Shopping may occur the nextday, but the scene boundary is not explicit.|

The last two-day/one-day expressions are **not yet a compulsory correction**. If later evidence confirms next-day shopping, a scene cue may suffice. Likewise, a fifth calendar day need not mean five complete24-hour periods have elapsed.

## 5. Resources and mechanics: what now reconciles

### Breakthrough does not change aptitude

The ordinary55% peak-essence threshold is explicitly qualified by exceptions. Fang Yuan has44%, and succeeds through prolonged replenishment while maintaining pressure on the aperture wall. The text supplies the mechanism. Do not change44to55 or erase the threshold.

### Inventory is consistent

**7Gu→6 after WhiteJade fusion→7 after inheriting NineLeaf Vitality Grass.** OneLittleLight then dies in failed Moonglow refinement. A replacement is implied before the successful three-input recipe, but its purchase price is not supplied. Successful Moonglow fusion leaves five permanentGu, apart from produced leafGu inventory.

Chapter110 has not yet given Fang Yuan a secondLiquorWorm. It is still an auction prospect.

### Leaf production and sale are not the same account

Each leaf requires20essence points to regrow, so nine require180points of cumulative expenditure. A44-point reserve supports two before replenishment, not nine stored at once. This is throughput, not a180% aperture.

Nine leaves at50retail give450gross. The text reports more than400in earnings, but does not itemise production cost or JiangYa's margin. Chapter107 distinguishes the clan's55retail price from Fang Yuan's50retail price. Do not assume the agent hands over exactly450.

WhiteJade consumes eightliang of jade every20days, compared with JadeSkin's twoliang every10days. That doubles average jade consumption, not quadruples it. Pigfeeding ceases, but the narrator still says the combined market feeding cost rises. Self-sourced jade makes cash expenditure a separate question.

### Cash and assets

The pre-inheritance cash low is explicitly **2.5stones**. The uncle's1000/2000/2500settlement offers are refused. Inherited property includes over10mu of farmland, eightservants, threebamboo houses, a wine-shop and NineLeaf Vitality Grass. Its cash valuation is not given. The uncle has already withdrawn the shop's current-month receipts; those cannot be credited to Fang Yuan.

The displayed LiquorWorm price is500against a stated normal580. The80difference is real, but sealed bids determine the actual price. No purchase or bid outcome is claimed throughChapter110.

## 6. Safeguards against false corrections

- Fang Yuan's accounts of the female team member entering the pig voluntarily, and of a bear providing the honey opportunity, are concealment, not neutral facts to harmonise with the action.
- JiaoSan is initially presumed dead, then observed moving, then killed. Preserve that sequence.
- PermanentWhiteBoar strength survives consumption of theGu in fusion. WhiteJade does not itself grant new strength.
- The later sixteen-monkey estimate does not share identical cultivation conditions with the earlier twelve-monkey estimate. No automatic16→12 correction.
- TwoLittleLight recipe inputs do not stack two amplification bonuses before fusion.
- QingShu's death is recalled from the former life. The proposed five-year right to use a higherLiquorWorm is an unaccepted conditional offer.
- Onehundred-plus years of prior trading and onehundred-plus years stuck atRevolutionTwo need not be disjoint periods. Do not add them as consecutive blocks.
- Ch105's distant108,000li expression may be idiomatic. No geographical shrinkage or precise travel-time model is authorised by it alone.
- The threatening500-stone demand in the tavern is not a paid bill.

## 7. Saved checkpoint

The numbered source, four clean jotters, reading log, proposal ledger, chronology register, resource register and cumulative workbook are updated through110. Only small metadata checkpoints were archived; no duplicate workspace backup was created.

**Next planned batch: Chapters111–140. Next actual unread chapter:111.** No later topical knowledge has been counted as sequential completion.
'''
# Keep the delivered report readable; compact register strings remain internal.
import re
for a,b in {'NewYear':'New Year','JadeSkin':'Jade Skin','WhiteJade':'White Jade','LittleLight':'Little Light','WhiteBoar':'White Boar','NineLeaf':'Nine-Leaf','LiquorWorm':'Liquor Worm','Moonglow':'Moonglow','ChiShan':'Chi Shan','JiangHe':'Jiang He','JiangYa':'Jiang Ya','JiaoSan':'Jiao San','ManShi':'Man Shi','QingShu':'Qing Shu','MoChen':'Mo Chen','RevolutionTwo':'Revolution Two','GradeC':'Grade C','dayaftertomorrow':'day after tomorrow','nextday':'next day','onehundred':'one hundred','eights ervants':'eight servants','eightservants':'eight servants','threebamboo':'three bamboo','twoliang':'two liang','eightliang':'eight liang','five-year right':'five-year right'}.items():
 report=report.replace(a,b)
report=re.sub(r'(?<=[A-Za-z])(?=\d)|(?<=\d)(?=[A-Za-z])',' ',report)
report=report.replace('of2,334','of 2,334').replace('unread:111','unread: 111').replace('Remaining:2,224','Remaining: 2,224').replace('not the final transaction','not the final transaction')
(r/'Batch-005-Chapters-81-110.md').write_text(report)
w=Workbook();s=w.active;s.title='Read me'
for row in [('Sequential review','Status'),('Completed','1–110 of2334'),('Paragraphs',tp),('Source characters',tc),('Next unread',111),('Proposals',f'{len(ledger)}held; none approved/applied'),('Scope','Chinese-source reading, not fresh liveEnglish audit'),('Chronology entries',len(events)),('Resource entries',len(resources)),('Global consistency','Not yet certified')]:s.append(row)
s=w.create_sheet('Chapter completion');s.append(['Chapter','File','Title','Paragraphs','Status'])
for c in chapters:s.append([c['corpus_chapter'],c['file'],c['title'],len(c['paragraphs']),'Full sequential read'])
s=w.create_sheet('Paragraph register');s.append(['Chapter','Paragraph','Status','Proposal IDs','Chronology IDs','Resource IDs','Chinese source'])
def ids(vs,k,c,p):return ', '.join(v['id'] for v in vs if any(z['chapter']==c and z['paragraph']==p for z in v[k]))
for c in chapters:
 for p,t in enumerate(c['paragraphs'],1):
  n=c['corpus_chapter'];s.append([n,p,'Read, not certified error-free',ids(ledger,'locations',n,p),ids(events,'sources',n,p),ids(resources,'sources',n,p),t])
s=w.create_sheet('Held proposals');s.append(['ID','Category','Locations','Working action','Rationale','Confidence / status','Dependencies','Publication','Approved','Applied','Latest update'])
for v in ledger:s.append([v['id'],v['category'],refs(v['locations']),v['recommended_action'],v.get('rationale','See originating report'),v.get('confidence',v['status']),v['dependencies'],v['publication_status'],'No','No',v.get('batch005_update',v.get('batch004_update',''))])
s=w.create_sheet('Chronology');s.append(['ID','Timeline','Kind','Event','Constraint','Evidence status','Dependencies','References','Exact source'])
for v in events:s.append([v['id'],v['timeline'],v['kind'],v['event'],v['constraint'],v['evidence_status'],v['dependency'],refs(v['sources']),'\n'.join(z['text'] for z in v['sources'])])
s=w.create_sheet('Resources');s.append(['ID','Kind','Entry','Calculation','Status','References','Exact source'])
for v in resources:s.append([v['id'],v['kind'],v['entry'],v['calculation'],v['status'],refs(v['sources']),'\n'.join(z['text'] for z in v['sources'])])
for s in w:
 s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
 for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='163B4B')
 for row in s.iter_rows(min_row=2):
  for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
 for col in s.columns:s.column_dimensions[col[0].column_letter].width=min(65,max(14,max(len(str(c.value or '')) for c in col[:30])+2))
w.save(r/'Sequential-Reading-and-Proposals.xlsx')
code=(r.parent/'make_whole_novel_audit_docx.py').read_text().replace("root=Path(__file__).parent","root=Path('/home/user/editorial/sequential-review')").replace('Whole-Novel-Aperture-and-Travel-Audit','Batch-005-Chapters-81-110').replace('WHOLE-NOVEL SOURCE AUDIT','SEQUENTIAL REVIEW / CHAPTERS81–110').replace('The Omniarch: Aperture and Travel Source Audit','The Omniarch: Sequential Review, Chapters81–110').replace('Full-corpus topical screening with selective contextual verification','Thirty complete chapters; all emendations held')
exec(compile(code,'docx-export','exec'),{})
(r/'CURRENT-STATUS.md').write_text(f'''# Sequential review checkpoint

Full sequential reading completed: **Chapters1–110**. Next unread: **111**. Planned next batch: **111–140**.

Batch005: {np}paragraphs/{nc}characters. Cumulative: {tp}paragraphs/{tc}characters. Remaining2224chapters. All{len(ledger)}proposals held, none approved or applied. Whole-novel chronology NOT verified.

Latest main report: Batch-005-Chapters-81-110.md/.docx. Actualsource: Batch-005-Numbered-Chinese-Source.md. Fourcleanjotters, readinglog and cumulativeworkbook updated. No source/site/repository changes. No remoteupload or backup created; do not store credentials in projectfiles. Workspace originals retained.

Important carry-forwards:
- SR-041 preferred prize repair is Ch81 150→100, supportedCh86/89 and25×4rent comparison. Initial reverse-direction speculation withdrawn before delivery.
- SR-040 Ch84one-yearprivateinstruction is distinct from academic-year labels; candidatehalfyearplus.
- SR-042 Ch90auntalready16 versusafterNewYear16; preserveageconvention.
- SR-047 Ch109usualsummer versusCh39usualApril; preserveintentionalearlyMarchbutterflyeffect andpriorlifesummer separately.
- Ch110auctionunresolved: secondLiquorWorm NOT yetowned,500displayprice NOT purchaseprice. Continue111withoutpeeking-at-outcome counted asread.
- Ch101rentmonth+2days constrains openintervals; Ch92inclusivecount andCh109/110ceremonyboundary unresolved,notforced.
- Ch23three-stoneadvance remains viableexplanationof20cash. Ch55toadbuyerJiaJinShengconfirmed; noJiaFurepair. Falsealibis, rumours andmistakenidentities mustremainseparatelylabelled.

Do not rerun earlierbatchgenerators unchanged: they wouldresetcoverage. Oldtopicalscreeningthrough2334 isNOT sequentialreading.
''')
assert len(chapters)==110 and len(ledger)==48
assert all(not v['approved'] and not v['applied'] for v in ledger)
assert np==2534 and nc==92351 and tp==8538 and tc==319874
assert '—' not in report
print(dict(batch_paragraphs=np,batch_characters=nc,total_paragraphs=tp,total_characters=tc,proposals=len(ledger),chronology=len(events),resources=len(resources),next=111))
