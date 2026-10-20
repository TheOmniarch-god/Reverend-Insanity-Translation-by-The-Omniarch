import openpyxl, json, glob
SR='/home/user/editorial/sequential-review'
srcs=[]
for b in ['036','037','038','039','040','041']:
    srcs+=json.load(open(f'{SR}/batch-{b}-source.json'))
srcs=sorted(srcs,key=lambda c:c['corpus_chapter'])
assert [c['corpus_chapter'] for c in srcs]==list(range(686,956)), 'range mismatch'
ledger=json.load(open(f'{SR}/proposal-ledger.json'))
pmap={}
for e in ledger:
    n=int(e['id'].split('-')[1])
    pass
    for loc in e['locations']:
        if loc.get('chapter') and loc.get('paragraph'):
            pmap.setdefault((loc['chapter'],loc['paragraph']),[]).append(e['id'])
wb=openpyxl.load_workbook(f'{SR}/Sequential-Reading-and-Proposals.xlsx')
cc=wb['Chapter completion']; pr=wb['Paragraph register']; rm=wb['Read me']
existing={r[0] for r in cc.iter_rows(min_row=2,values_only=True)}
assert 685 in existing and 686 not in existing
newchars=0; newparas=0
for c in srcs:
    ch=c['corpus_chapter']
    cc.append([ch,c['file'],c['title'],len(c['paragraphs']),'Full sequential read'])
    for i,p in enumerate(c['paragraphs'],1):
        ids=pmap.get((ch,i))
        pr.append([ch,i,'Read; not certified error-free',', '.join(sorted(set(ids))) if ids else None,None,None,p])
        newchars+=len(p); newparas+=1
prev_chars=2171266; prev_paras=61680
for row in rm.iter_rows(min_row=1,max_row=rm.max_row):
    k=row[0].value
    if k=='Completed': row[1].value='1–955 of 2334'
    elif k=='Next unread': row[1].value='956'
    elif k=='Paragraphs': row[1].value=str(prev_paras+newparas)
    elif k=='Characters': row[1].value=str(prev_chars+newchars)
    elif k=='Proposals': row[1].value='710 total (SR-001 to SR-710); batches 037–041 and 036 added; nothing approved or applied'
wb.save(f'{SR}/Sequential-Reading-and-Proposals.xlsx')
print('chapters added',len(srcs),'paragraphs added',newparas,'chars added',newchars,'linked paragraphs',sum(1 for k in pmap if 686<=k[0]<=955))
