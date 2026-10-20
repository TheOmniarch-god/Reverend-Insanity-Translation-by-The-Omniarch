import sys, re
from docx import Document
from docx.shared import Pt
from docx.enum.section import WD_ORIENT
def convert(md_path, docx_path):
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    st = doc.styles['Normal']; st.font.size = Pt(9)
    lines = open(md_path, encoding='utf-8').read().splitlines()
    i = 0
    def add_runs(par, text):
        parts = re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text)
        for p in parts:
            if p.startswith('**') and p.endswith('**'): r = par.add_run(p[2:-2]); r.bold = True
            elif p.startswith('`') and p.endswith('`'): r = par.add_run(p[1:-1]); r.font.name = 'Consolas'
            else: par.add_run(p)
    while i < len(lines):
        ln = lines[i]
        if ln.startswith('|'):
            block = []
            while i < len(lines) and lines[i].startswith('|'):
                block.append(lines[i]); i += 1
            rows = [[c.strip() for c in b.strip().strip('|').split('|')] for b in block if not re.match(r'^\|\s*-', b)]
            ncol = max(len(r) for r in rows)
            t = doc.add_table(rows=len(rows), cols=ncol); t.style = 'Table Grid'
            for ri, r in enumerate(rows):
                for ci in range(ncol):
                    cell = t.cell(ri, ci); cell.text = ''
                    add_runs(cell.paragraphs[0], r[ci] if ci < len(r) else '')
                    for run in cell.paragraphs[0].runs:
                        run.font.size = Pt(8)
                        if ri == 0: run.bold = True
            doc.add_paragraph()
            continue
        if ln.startswith('# '): doc.add_heading(ln[2:], level=0)
        elif ln.startswith('## '): doc.add_heading(ln[3:], level=1)
        elif ln.startswith('- '): add_runs(doc.add_paragraph(style='List Bullet'), ln[2:])
        elif ln.strip(): add_runs(doc.add_paragraph(), ln)
        i += 1
    doc.save(docx_path)
if __name__ == '__main__':
    convert(sys.argv[1], sys.argv[2]); print('ok', sys.argv[2])
